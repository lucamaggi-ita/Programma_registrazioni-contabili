from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fantozzi.relatorios_gerenciais_ok import (
    IDENTIFIER_HEADER,
    OUTPUT_LANCAMENTO_HEADERS,
    _derive_identifier,
    _format_currency,
    _lancamento_matches_special,
    _read_lancamentos,
    _read_operacoes_especiais,
    _resolve_excel_paths,
    _special_keys,
    _sum_values,
    _text,
)


AUDITORIA_GERAL_SHEET = "Auditoria_Geral"
RECONCILIACAO_STATUS_SHEET = "Reconciliação_Status"
RECONCILIACAO_OPERACOES_SHEET = "Reconcil_Operacoes_Especiais"
SOBREPOSICOES_SHEET = "Sobreposicoes"
TOTAIS_ORIGEM_SHEET = "Totais_por_Origem"
INCONSISTENCIAS_SHEET = "Inconsistencias"
AMOSTRA_INCLUIDOS_SHEET = "Amostra_Incluidos_OK"
AMOSTRA_EXCLUIDOS_SHEET = "Amostra_Excluidos"
INSTRUCOES_SHEET = "Instrucoes"

SOURCE_HEADER = "arquivo_excel_origem"

AUDITORIA_HEADERS = ["indicador", "valor", "status", "observacao"]
STATUS_HEADERS = [
    "status_revisao",
    "quantidade",
    "valor_total",
    "percentual_quantidade",
    "percentual_valor",
]
OPERACOES_HEADERS = [
    "tipo_operacao_especial",
    "quantidade_total",
    "quantidade_status_ok",
    "quantidade_status_revisar",
    "valor_total",
    "valor_status_ok",
    "valor_status_revisar",
    "tratamento",
]
SOBREPOSICOES_HEADERS = ["categoria", "quantidade", "valor_total", "observacao"]
TOTAIS_ORIGEM_HEADERS = [
    "identificador_operacional",
    "total_origem",
    "status_ok",
    "status_revisar",
    "operacoes_especiais",
    "incluidos_ok_nao_especiais",
    "valor_incluido_ok_nao_especial",
]
INCONSISTENCIAS_HEADERS = ["tipo_inconsistencia", "quantidade", "detalhe", "gravidade"]
AMOSTRA_EXCLUIDOS_HEADERS = ["motivo_exclusao", *OUTPUT_LANCAMENTO_HEADERS]


class AuditoriaRelatoriosGerenciaisOkError(Exception):
    pass


@dataclass(frozen=True)
class AuditoriaRelatoriosGerenciaisOkResult:
    entradas: tuple[Path, ...]
    excel_paths: tuple[Path, ...]
    pastas_operacoes_especiais: tuple[Path, ...]
    xlsx_path: Path
    report_path: Path | None
    total_lancamentos_origem: int
    total_status_ok: int
    total_status_revisar: int
    total_operacoes_especiais: int
    total_incluidos_ok_nao_especiais: int
    total_ok_excluidos_por_operacao_especial: int
    total_revisar_com_operacao_especial: int
    valor_total_origem: Decimal
    valor_total_status_ok: Decimal
    valor_total_status_revisar: Decimal
    valor_total_operacoes_especiais: Decimal
    valor_total_incluido_ok_nao_especial: Decimal
    resultado_auditoria: str
    inconsistencias: tuple[dict[str, object], ...]


def auditar_relatorios_gerenciais_ok_provisorios(
    entradas: Iterable[Path | str],
    xlsx_path: Path | str,
    report_path: Path | str | None = None,
    pastas_operacoes_especiais: Iterable[Path | str] | None = None,
) -> AuditoriaRelatoriosGerenciaisOkResult:
    entradas = tuple(Path(path) for path in entradas)
    xlsx_path = Path(xlsx_path)
    report_path = Path(report_path) if report_path else None
    pastas_operacoes = tuple(Path(path) for path in (pastas_operacoes_especiais or ()))
    excel_paths = _resolve_excel_paths(entradas)

    all_rows: list[dict[str, object]] = []
    for excel_path in excel_paths:
        identifier = _derive_identifier(excel_path)
        for row in _read_lancamentos(excel_path):
            all_rows.append({IDENTIFIER_HEADER: identifier, SOURCE_HEADER: str(excel_path), **row})

    operacoes = _read_operacoes_especiais(pastas_operacoes)
    special_keys = _special_keys(operacoes)
    for row in all_rows:
        row["eh_operacao_especial"] = _lancamento_matches_special(row, special_keys)

    included_ok = [
        row
        for row in all_rows
        if _status(row) == "OK" and not row["eh_operacao_especial"]
    ]
    ok_special = [
        row for row in all_rows if _status(row) == "OK" and row["eh_operacao_especial"]
    ]
    revisar_plain = [
        row
        for row in all_rows
        if _status(row) == "REVISAR" and not row["eh_operacao_especial"]
    ]
    revisar_special = [
        row
        for row in all_rows
        if _status(row) == "REVISAR" and row["eh_operacao_especial"]
    ]
    invalid_status = [row for row in all_rows if _status(row) not in {"OK", "REVISAR"}]
    inconsistencias = _inconsistencias(all_rows, invalid_status)
    resultado = _resultado_auditoria(inconsistencias, revisar_plain, revisar_special, operacoes)

    result = AuditoriaRelatoriosGerenciaisOkResult(
        entradas=entradas,
        excel_paths=excel_paths,
        pastas_operacoes_especiais=pastas_operacoes,
        xlsx_path=xlsx_path,
        report_path=report_path,
        total_lancamentos_origem=len(all_rows),
        total_status_ok=len(included_ok) + len(ok_special),
        total_status_revisar=len(revisar_plain) + len(revisar_special),
        total_operacoes_especiais=len(operacoes),
        total_incluidos_ok_nao_especiais=len(included_ok),
        total_ok_excluidos_por_operacao_especial=len(ok_special),
        total_revisar_com_operacao_especial=len(revisar_special),
        valor_total_origem=_sum_values(all_rows, "valor"),
        valor_total_status_ok=_sum_values([*included_ok, *ok_special], "valor"),
        valor_total_status_revisar=_sum_values([*revisar_plain, *revisar_special], "valor"),
        valor_total_operacoes_especiais=_sum_values(operacoes, "valor_total_nfe"),
        valor_total_incluido_ok_nao_especial=_sum_values(included_ok, "valor"),
        resultado_auditoria=resultado,
        inconsistencias=tuple(inconsistencias),
    )

    workbook = Workbook()
    _write_table(workbook.active, AUDITORIA_GERAL_SHEET, AUDITORIA_HEADERS, _auditoria_rows(result))
    _write_table(workbook.create_sheet(RECONCILIACAO_STATUS_SHEET), RECONCILIACAO_STATUS_SHEET, STATUS_HEADERS, _status_rows(all_rows))
    _write_table(workbook.create_sheet(RECONCILIACAO_OPERACOES_SHEET), RECONCILIACAO_OPERACOES_SHEET, OPERACOES_HEADERS, _operacoes_rows(operacoes, all_rows))
    _write_table(workbook.create_sheet(SOBREPOSICOES_SHEET), SOBREPOSICOES_SHEET, SOBREPOSICOES_HEADERS, _sobreposicoes_rows(included_ok, ok_special, revisar_plain, revisar_special))
    _write_table(workbook.create_sheet(TOTAIS_ORIGEM_SHEET), TOTAIS_ORIGEM_SHEET, TOTAIS_ORIGEM_HEADERS, _totais_origem_rows(all_rows, operacoes))
    _write_table(workbook.create_sheet(INCONSISTENCIAS_SHEET), INCONSISTENCIAS_SHEET, INCONSISTENCIAS_HEADERS, inconsistencias or [_no_inconsistencies_row()])
    _write_table(workbook.create_sheet(AMOSTRA_INCLUIDOS_SHEET), AMOSTRA_INCLUIDOS_SHEET, OUTPUT_LANCAMENTO_HEADERS, included_ok[:50])
    _write_table(workbook.create_sheet(AMOSTRA_EXCLUIDOS_SHEET), AMOSTRA_EXCLUIDOS_SHEET, AMOSTRA_EXCLUIDOS_HEADERS, _amostra_excluidos(revisar_plain, ok_special, revisar_special))
    _write_instrucoes(workbook.create_sheet(INSTRUCOES_SHEET))

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)
    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(relatorio_texto(result, included_ok, ok_special, revisar_plain, revisar_special), encoding="utf-8", newline="\n")
    return result


def console_lines(result: AuditoriaRelatoriosGerenciaisOkResult) -> list[str]:
    lines = [
        "[OK] Auditoria dos relatorios gerenciais OK provisorios gerada.",
        f"     Total origem: {result.total_lancamentos_origem}",
        f"     Status OK: {result.total_status_ok}",
        f"     Status REVISAR: {result.total_status_revisar}",
        f"     OK nao especiais incluidos: {result.total_incluidos_ok_nao_especiais}",
        f"     OK excluidos por operacao especial: {result.total_ok_excluidos_por_operacao_especial}",
        f"     Resultado auditoria: {result.resultado_auditoria}",
        f"     XLSX gerado: {result.xlsx_path}",
    ]
    if result.report_path:
        lines.append(f"     Relatorio: {result.report_path}")
    return lines


def relatorio_texto(
    result: AuditoriaRelatoriosGerenciaisOkResult,
    included_ok: list[dict[str, object]],
    ok_special: list[dict[str, object]],
    revisar_plain: list[dict[str, object]],
    revisar_special: list[dict[str, object]],
) -> str:
    lines = [
        "AUDITORIA DOS RELATORIOS GERENCIAIS OK PROVISORIOS",
        "==================================================",
        "",
        "Status do relatorio: PROVISORIO",
        "",
        "Criterios auditados:",
        "- Inclui somente status_revisao = OK",
        "- Exclui REVISAR",
        "- Exclui operacoes especiais",
        "",
        "Reconciliacao:",
        f"- Total origem: {result.total_lancamentos_origem}",
        f"- Status OK: {result.total_status_ok}",
        f"- Status REVISAR: {result.total_status_revisar}",
        f"- Operacoes especiais: {result.total_operacoes_especiais}",
        f"- OK nao especiais incluidos: {result.total_incluidos_ok_nao_especiais}",
        f"- OK excluidos por operacao especial: {result.total_ok_excluidos_por_operacao_especial}",
        f"- REVISAR com operacao especial: {result.total_revisar_com_operacao_especial}",
        "",
        "Valores:",
        f"- Valor total origem: {_format_currency(result.valor_total_origem)}",
        f"- Valor total OK: {_format_currency(result.valor_total_status_ok)}",
        f"- Valor total REVISAR: {_format_currency(result.valor_total_status_revisar)}",
        f"- Valor total operacoes especiais: {_format_currency(result.valor_total_operacoes_especiais)}",
        f"- Valor total incluido OK nao especial: {_format_currency(result.valor_total_incluido_ok_nao_especial)}",
        "",
        "Sobreposicoes:",
        f"- OK nao especial incluido: {len(included_ok)} | {_format_currency(_sum_values(included_ok, 'valor'))}",
        f"- OK com operacao especial excluido: {len(ok_special)} | {_format_currency(_sum_values(ok_special, 'valor'))}",
        f"- REVISAR nao especial excluido: {len(revisar_plain)} | {_format_currency(_sum_values(revisar_plain, 'valor'))}",
        f"- REVISAR com operacao especial excluido: {len(revisar_special)} | {_format_currency(_sum_values(revisar_special, 'valor'))}",
        "",
        "Inconsistencias:",
    ]
    if result.inconsistencias:
        lines.extend(f"- {row['tipo_inconsistencia']}: {row['quantidade']} | {row['gravidade']}" for row in result.inconsistencias)
    else:
        lines.append("- Nenhuma inconsistencia critica encontrada")
    lines.extend(
        [
            "",
            "Resultado auditoria:",
            result.resultado_auditoria,
            "",
            "Alertas:",
            "- Revisao contabil real pendente.",
            "- Nao usar como DRE definitiva.",
            "- Nao usar como livro diario definitivo.",
        ]
    )
    return "\n".join(lines) + "\n"


def _auditoria_rows(result: AuditoriaRelatoriosGerenciaisOkResult) -> list[dict[str, object]]:
    values = [
        ("status_relatorio", "PROVISORIO", "OK", "Auditoria provisoria."),
        ("total_lancamentos_origem", result.total_lancamentos_origem, "OK", "Total lido dos Excel."),
        ("total_status_ok", result.total_status_ok, "OK", "Lancamentos com status OK."),
        ("total_status_revisar", result.total_status_revisar, "ATENCAO" if result.total_status_revisar else "OK", "Lancamentos excluidos por revisao."),
        ("total_operacoes_especiais", result.total_operacoes_especiais, "ATENCAO" if result.total_operacoes_especiais else "OK", "NF-e especiais auditadas."),
        ("total_incluidos_ok_nao_especiais", result.total_incluidos_ok_nao_especiais, "OK", "Base incluida nos relatorios OK."),
        ("total_ok_excluidos_por_operacao_especial", result.total_ok_excluidos_por_operacao_especial, "ATENCAO" if result.total_ok_excluidos_por_operacao_especial else "OK", "OK retirados por operacao especial."),
        ("total_revisar_com_operacao_especial", result.total_revisar_com_operacao_especial, "ATENCAO" if result.total_revisar_com_operacao_especial else "OK", "Intersecao REVISAR e especial."),
        ("valor_total_origem", result.valor_total_origem, "OK", "Valor total dos lancamentos lidos."),
        ("valor_total_status_ok", result.valor_total_status_ok, "OK", "Valor dos lancamentos OK antes da exclusao especial."),
        ("valor_total_status_revisar", result.valor_total_status_revisar, "ATENCAO" if result.total_status_revisar else "OK", "Valor em revisao."),
        ("valor_total_operacoes_especiais", result.valor_total_operacoes_especiais, "ATENCAO" if result.total_operacoes_especiais else "OK", "Valor total das NF-e especiais."),
        ("valor_total_incluido_ok_nao_especial", result.valor_total_incluido_ok_nao_especial, "OK", "Valor incluido nos relatorios OK provisorios."),
        ("resultado_auditoria", result.resultado_auditoria, result.resultado_auditoria, "Resultado consolidado."),
    ]
    return [{"indicador": item[0], "valor": item[1], "status": item[2], "observacao": item[3]} for item in values]


def _status_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    total_count = len(rows) or 1
    total_value = _sum_values(rows, "valor") or Decimal("1.00")
    output = []
    for status in sorted({_status(row) for row in rows}):
        selected = [row for row in rows if _status(row) == status]
        value = _sum_values(selected, "valor")
        output.append(
            {
                "status_revisao": status or "SEM_STATUS",
                "quantidade": len(selected),
                "valor_total": value,
                "percentual_quantidade": (Decimal(len(selected)) / Decimal(total_count) * Decimal("100")).quantize(Decimal("0.01")),
                "percentual_valor": (value / total_value * Decimal("100")).quantize(Decimal("0.01")),
            }
        )
    return output


def _operacoes_rows(operacoes: list[dict[str, object]], rows: list[dict[str, object]]) -> list[dict[str, object]]:
    output = []
    for tipo in sorted({_text(row.get("tipo_operacao_especial")) for row in operacoes}):
        ops = [row for row in operacoes if _text(row.get("tipo_operacao_especial")) == tipo]
        keys = _special_keys(ops)
        matched = [row for row in rows if _lancamento_matches_special(row, keys)]
        matched_ok = [row for row in matched if _status(row) == "OK"]
        matched_revisar = [row for row in matched if _status(row) == "REVISAR"]
        output.append(
            {
                "tipo_operacao_especial": tipo,
                "quantidade_total": len(ops),
                "quantidade_status_ok": len(matched_ok),
                "quantidade_status_revisar": len(matched_revisar),
                "valor_total": _sum_values(ops, "valor_total_nfe"),
                "valor_status_ok": _sum_values(matched_ok, "valor"),
                "valor_status_revisar": _sum_values(matched_revisar, "valor"),
                "tratamento": "Excluido dos relatorios OK provisorios",
            }
        )
    return output


def _sobreposicoes_rows(*groups: list[dict[str, object]]) -> list[dict[str, object]]:
    labels = [
        ("OK nao especial incluido", "Base dos relatorios gerenciais OK provisorios."),
        ("OK com operacao especial excluido", "Status OK, mas fora por operacao especial."),
        ("REVISAR nao especial excluido", "Pendente de revisao contabil humana."),
        ("REVISAR com operacao especial excluido", "Intersecao de revisao e operacao especial."),
    ]
    return [
        {"categoria": label, "quantidade": len(rows), "valor_total": _sum_values(rows, "valor"), "observacao": obs}
        for (label, obs), rows in zip(labels, groups)
    ]


def _totais_origem_rows(rows: list[dict[str, object]], operacoes: list[dict[str, object]]) -> list[dict[str, object]]:
    op_counts = Counter(_text(row.get("identificador_operacional")) for row in operacoes)
    output = []
    for identifier in sorted({_text(row.get(IDENTIFIER_HEADER)) for row in rows}):
        selected = [row for row in rows if _text(row.get(IDENTIFIER_HEADER)) == identifier]
        included = [row for row in selected if _status(row) == "OK" and not row.get("eh_operacao_especial")]
        output.append(
            {
                "identificador_operacional": identifier,
                "total_origem": len(selected),
                "status_ok": sum(1 for row in selected if _status(row) == "OK"),
                "status_revisar": sum(1 for row in selected if _status(row) == "REVISAR"),
                "operacoes_especiais": op_counts.get(identifier, 0),
                "incluidos_ok_nao_especiais": len(included),
                "valor_incluido_ok_nao_especial": _sum_values(included, "valor"),
            }
        )
    return output


def _inconsistencias(rows: list[dict[str, object]], invalid_status: list[dict[str, object]]) -> list[dict[str, object]]:
    output = []
    if invalid_status:
        output.append({"tipo_inconsistencia": "STATUS_INVALIDO", "quantidade": len(invalid_status), "detalhe": "Lancamentos sem status OK ou REVISAR.", "gravidade": "ERRO"})
    missing_supplier = [row for row in rows if not _text(row.get("fornecedor_ou_favorecido"))]
    if missing_supplier:
        output.append({"tipo_inconsistencia": "FORNECEDOR_AUSENTE", "quantidade": len(missing_supplier), "detalhe": "Fornecedor ausente, mas calculo continua.", "gravidade": "ATENCAO"})
    return output


def _no_inconsistencies_row() -> dict[str, object]:
    return {"tipo_inconsistencia": "Nenhuma inconsistencia critica encontrada", "quantidade": 0, "detalhe": "-", "gravidade": "OK"}


def _amostra_excluidos(
    revisar_plain: list[dict[str, object]],
    ok_special: list[dict[str, object]],
    revisar_special: list[dict[str, object]],
) -> list[dict[str, object]]:
    rows = []
    rows.extend({"motivo_exclusao": "REVISAR", **row} for row in revisar_plain[:25])
    rows.extend({"motivo_exclusao": "OPERACAO_ESPECIAL", **row} for row in ok_special[:25])
    rows.extend({"motivo_exclusao": "REVISAR_E_OPERACAO_ESPECIAL", **row} for row in revisar_special[:25])
    return rows


def _resultado_auditoria(
    inconsistencias: list[dict[str, object]],
    revisar_plain: list[dict[str, object]],
    revisar_special: list[dict[str, object]],
    operacoes: list[dict[str, object]],
) -> str:
    if any(row.get("gravidade") == "ERRO" for row in inconsistencias):
        return "ERRO"
    if inconsistencias or revisar_plain or revisar_special or operacoes:
        return "ATENCAO"
    return "OK"


def _status(row: dict[str, object]) -> str:
    return _text(row.get("status_revisao")).upper()


def _write_table(sheet, title: str, headers: list[str], rows: list[dict[str, object]]) -> None:
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header)) for header in headers])
    _style_table(sheet)
    _set_widths(sheet, headers)


def _write_instrucoes(sheet) -> None:
    lines = [
        "AUDITORIA DOS RELATORIOS GERENCIAIS OK PROVISORIOS",
        "",
        "Esta auditoria verifica os relatorios gerenciais OK provisorios.",
        "Ela nao transforma dados em definitivos.",
        "Ela nao altera classificacoes.",
        "Ela explica sobreposicoes entre OK, REVISAR e operacoes especiais.",
        "Os relatorios continuam provisorios.",
        "Nao usar como DRE definitiva.",
        "Nao usar como livro diario definitivo.",
    ]
    for line in lines:
        sheet.append([line])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 120
    for row in sheet.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _style_table(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, (int, float)) and _money_header(sheet, cell.column):
                cell.number_format = '#,##0.00'


def _money_header(sheet, column: int) -> bool:
    header = _text(sheet.cell(row=1, column=column).value)
    return "valor" in header or "percentual" in header


def _set_widths(sheet, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, min(len(header) + 8, 52))


def _cell_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "SIM" if value else "NAO"
    return value
