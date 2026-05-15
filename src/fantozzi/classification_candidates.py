from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from fantozzi.excel_writer import LANCAMENTOS_SHEET, REVISAO_SHEET
from fantozzi.normalization import normalize


VALUE_HEADERS = ("valor",)
DESCRIPTION_HEADERS = (
    "descricao",
    "descricao_produto",
    "produto",
    "descricao_itens",
    "observacoes",
    "documento",
    "chave_origem",
)


class ClassificationCandidatesError(Exception):
    pass


@dataclass(frozen=True)
class CandidateSourceRow:
    description: str
    normalized_description: str
    value: Decimal
    supplier: str
    document: str
    status_revisao: str
    alertas: str
    conta_id: str
    conta_nome: str
    cfop: str
    natureza_operacao: str
    tipo_operacao_especial: str
    confianca_operacao_especial: str
    tratamento_operacao_especial: str


@dataclass(frozen=True)
class ClassificationCandidate:
    description: str
    normalized_description: str
    occurrences: int
    total_value: Decimal
    suppliers: tuple[str, ...]
    documents: tuple[str, ...]
    statuses: tuple[str, ...]
    alerts: tuple[str, ...]
    accounts: tuple[str, ...]
    cfops: tuple[str, ...]
    naturezas_operacao: tuple[str, ...]
    tipos_operacao_especial: tuple[str, ...]
    confiancas_operacao_especial: tuple[str, ...]
    tratamentos_operacao_especial: tuple[str, ...]


@dataclass(frozen=True)
class ClassificationCandidatesResult:
    excel_path: Path
    source_sheet: str
    total_candidates: int
    source_rows: int
    total_review_value: Decimal
    candidates: tuple[ClassificationCandidate, ...]


def generate_classification_candidates_report(
    excel_path: Path | str,
    *,
    generated_at: datetime | None = None,
) -> tuple[ClassificationCandidatesResult, str]:
    result = analyze_classification_candidates(excel_path)
    return result, report_text(result, generated_at=generated_at)


def analyze_classification_candidates(
    excel_path: Path | str,
) -> ClassificationCandidatesResult:
    excel_path = Path(excel_path)
    _validate_file(excel_path)

    try:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ClassificationCandidatesError(
            f"Arquivo Excel nao pode ser lido: {excel_path}"
        ) from exc

    source_sheet_name, source_rows = _source_rows(workbook)
    candidates = _aggregate_candidates(source_rows)
    total_review_value = sum(
        (candidate.total_value for candidate in candidates), Decimal("0.00")
    )

    return ClassificationCandidatesResult(
        excel_path=excel_path,
        source_sheet=source_sheet_name,
        total_candidates=len(candidates),
        source_rows=sum(candidate.occurrences for candidate in candidates),
        total_review_value=total_review_value.quantize(Decimal("0.01")),
        candidates=tuple(candidates),
    )


def report_text(
    result: ClassificationCandidatesResult,
    *,
    generated_at: datetime | None = None,
) -> str:
    generated_at = generated_at or datetime.now()
    lines = [
        "FANTOZZI - RELATORIO DE CANDIDATOS PARA CLASSIFICACAO",
        "======================================================",
        f"Arquivo Excel: {result.excel_path}",
        f"Aba analisada: {result.source_sheet}",
        f"Data/hora relatorio: {generated_at:%Y-%m-%d %H:%M:%S}",
        "",
        "Resumo",
        "----------------------------------------",
        f"Total de candidatos: {result.total_candidates}",
        f"Linhas de origem em revisao: {result.source_rows}",
        f"Valor total em revisao: {_format_currency_br(result.total_review_value)}",
    ]

    if not result.candidates:
        lines.extend(
            [
                "",
                "Nenhum candidato para classificacao encontrado.",
            ]
        )
    else:
        lines.extend(["", "Candidatos", "----------------------------------------"])
        for index, candidate in enumerate(result.candidates, start=1):
            lines.extend(_candidate_lines(index, candidate))

    lines.extend(
        [
            "",
            "Observacoes",
            "----------------------------------------",
            "- Este relatorio e somente leitura.",
            "- Nenhuma regra YAML foi alterada.",
            "- Nenhuma classificacao foi aplicada automaticamente.",
            "- Use as sugestoes apenas como apoio para revisao manual.",
        ]
    )
    return "\n".join(lines) + "\n"


def console_lines(result: ClassificationCandidatesResult) -> list[str]:
    return [
        "[OK] Relatorio de candidatos de classificacao concluido",
        f"Total de candidatos: {result.total_candidates}",
        f"Linhas de origem em revisao: {result.source_rows}",
        f"Valor total em revisao: {_format_currency_br(result.total_review_value)}",
    ]


def save_report(report_path: Path | str, text: str) -> None:
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(text, encoding="utf-8", newline="\n")


def _source_rows(workbook: object) -> tuple[str, tuple[CandidateSourceRow, ...]]:
    if LANCAMENTOS_SHEET in workbook.sheetnames:
        rows = _rows_from_sheet(workbook[LANCAMENTOS_SHEET], only_revision=True)
        if rows or REVISAO_SHEET not in workbook.sheetnames:
            return LANCAMENTOS_SHEET, rows

    if REVISAO_SHEET in workbook.sheetnames:
        rows = _rows_from_sheet(workbook[REVISAO_SHEET], only_revision=False)
        if rows:
            return REVISAO_SHEET, rows
        if LANCAMENTOS_SHEET in workbook.sheetnames:
            return LANCAMENTOS_SHEET, ()
        return REVISAO_SHEET, rows

    raise ClassificationCandidatesError(
        f"Aba obrigatoria ausente: {LANCAMENTOS_SHEET} ou {REVISAO_SHEET}"
    )


def _rows_from_sheet(
    sheet: Worksheet,
    *,
    only_revision: bool,
) -> tuple[CandidateSourceRow, ...]:
    headers = _header_indexes(sheet)
    _validate_headers(headers, sheet.title)

    rows = []
    for row in _data_rows(sheet):
        status_revisao = _text(_cell(row, headers, "status_revisao"))
        if only_revision and status_revisao.upper() == "OK":
            continue

        description = _text(_first_cell(row, headers, DESCRIPTION_HEADERS))
        normalized_description = normalize(description) if description else ""
        if not normalized_description:
            continue

        rows.append(
            CandidateSourceRow(
                description=description,
                normalized_description=normalized_description,
                value=_decimal_value(_cell(row, headers, "valor")),
                supplier=_text(
                    _first_cell(
                        row,
                        headers,
                        ("fornecedor_nome", "fornecedor", "fornecedor_ou_favorecido"),
                    )
                ),
                document=_text(
                    _first_cell(row, headers, ("documento", "numero_nfe", "chave_origem"))
                ),
                status_revisao=status_revisao or "(vazio)",
                alertas=_text(_first_cell(row, headers, ("alertas", "alerta"))),
                conta_id=_text(_cell(row, headers, "conta_id")),
                conta_nome=_text(_cell(row, headers, "conta_nome")),
                cfop=_text(_cell(row, headers, "cfop")),
                natureza_operacao=_text(_cell(row, headers, "natureza_operacao")),
                tipo_operacao_especial=_text(_cell(row, headers, "tipo_operacao_especial")),
                confianca_operacao_especial=_text(
                    _cell(row, headers, "confianca_operacao_especial")
                ),
                tratamento_operacao_especial=_text(
                    _cell(row, headers, "tratamento_operacao_especial")
                ),
            )
        )

    return tuple(rows)


def _aggregate_candidates(
    rows: tuple[CandidateSourceRow, ...]
) -> tuple[ClassificationCandidate, ...]:
    grouped: dict[str, list[CandidateSourceRow]] = defaultdict(list)
    for row in rows:
        grouped[row.normalized_description].append(row)

    candidates = []
    for normalized_description, group in grouped.items():
        descriptions = Counter(row.description for row in group)
        description = descriptions.most_common(1)[0][0]
        total_value = sum((row.value for row in group), Decimal("0.00"))
        candidates.append(
            ClassificationCandidate(
                description=description,
                normalized_description=normalized_description,
                occurrences=len(group),
                total_value=total_value.quantize(Decimal("0.01")),
                suppliers=_sorted_unique(row.supplier for row in group),
                documents=_sorted_unique(row.document for row in group),
                statuses=_sorted_unique(row.status_revisao for row in group),
                alerts=_sorted_unique(row.alertas for row in group),
                accounts=_sorted_unique(
                    _join_code_name(row.conta_id, row.conta_nome) for row in group
                ),
                cfops=_sorted_unique(row.cfop for row in group),
                naturezas_operacao=_sorted_unique(row.natureza_operacao for row in group),
                tipos_operacao_especial=_sorted_unique(
                    row.tipo_operacao_especial for row in group
                ),
                confiancas_operacao_especial=_sorted_unique(
                    row.confianca_operacao_especial for row in group
                ),
                tratamentos_operacao_especial=_sorted_unique(
                    row.tratamento_operacao_especial for row in group
                ),
            )
        )

    return tuple(
        sorted(
            candidates,
            key=lambda candidate: (
                -candidate.total_value,
                -candidate.occurrences,
                candidate.normalized_description,
            ),
        )
    )


def _candidate_lines(index: int, candidate: ClassificationCandidate) -> list[str]:
    return [
        f"{index}) {candidate.description}",
        "----------------------------------------",
        f"Ocorrencias: {candidate.occurrences}",
        f"Valor total: {_format_currency_br(candidate.total_value)}",
        f"Fornecedores: {_join_values(candidate.suppliers)}",
        f"Documentos: {_join_values(candidate.documents)}",
        f"Status: {_join_values(candidate.statuses)}",
        f"Alerta: {_join_values(candidate.alerts)}",
        f"Contas atuais: {_join_values(candidate.accounts)}",
        f"CFOP: {_join_values(candidate.cfops)}",
        f"Natureza operacao: {_join_values(candidate.naturezas_operacao)}",
        f"Tipo operacao especial: {_join_values(candidate.tipos_operacao_especial)}",
        f"Confianca operacao especial: {_join_values(candidate.confiancas_operacao_especial)}",
        f"Tratamento operacao especial: {_join_values(candidate.tratamentos_operacao_especial)}",
        "",
        "Texto normalizado:",
        candidate.normalized_description,
        "",
        "Sugestao manual para analise:",
        "- verificar a conta gerencial correta antes de alterar regras",
        "- se confirmado, adicionar alias manualmente em config/regras_classificacao_produtos.yaml",
        "",
        "Bloco YAML sugerido somente como rascunho:",
        "- conta_id: \"PREENCHER_MANUALMENTE\"",
        "  conta_nome: \"PREENCHER_MANUALMENTE\"",
        "  aliases:",
        f"    - {candidate.normalized_description}",
        "",
    ]


def _validate_file(path: Path) -> None:
    if not path.exists():
        raise ClassificationCandidatesError(f"Arquivo Excel nao encontrado: {path}")
    if not path.is_file():
        raise ClassificationCandidatesError(f"Arquivo Excel nao e um arquivo: {path}")


def _validate_headers(headers: dict[str, int], sheet_name: str) -> None:
    missing_value = [header for header in VALUE_HEADERS if header not in headers]
    if missing_value:
        raise ClassificationCandidatesError(
            f"Colunas obrigatorias ausentes em {sheet_name}: "
            + ", ".join(missing_value)
        )

    if not any(header in headers for header in DESCRIPTION_HEADERS):
        raise ClassificationCandidatesError(
            f"Colunas de descricao ausentes em {sheet_name}: "
            + ", ".join(DESCRIPTION_HEADERS)
        )


def _header_indexes(sheet: Worksheet) -> dict[str, int]:
    return {
        str(cell.value): index
        for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)))
        if cell.value is not None
    }


def _data_rows(sheet: Worksheet) -> Iterable[tuple[object, ...]]:
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(_has_value(value) for value in row):
            yield row


def _cell(row: tuple[object, ...], indexes: dict[str, int], header: str) -> object:
    index = indexes.get(header)
    if index is None or index >= len(row):
        return None
    return row[index]


def _first_cell(
    row: tuple[object, ...], indexes: dict[str, int], headers: tuple[str, ...]
) -> object:
    for header in headers:
        value = _cell(row, indexes, header)
        if _has_value(value):
            return value
    return None


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _has_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def _decimal_value(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0.00")
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _sorted_unique(values: Iterable[str]) -> tuple[str, ...]:
    return tuple(sorted({value for value in values if value}))


def _join_values(values: tuple[str, ...]) -> str:
    return ", ".join(values) if values else "-"


def _join_code_name(code: str, name: str) -> str:
    if code and name:
        return f"{code} - {name}"
    return code or name


def _format_currency_br(value: Decimal) -> str:
    text = f"{value.quantize(Decimal('0.01')):.2f}"
    integer_part, decimal_part = text.split(".")
    groups = []
    while integer_part:
        groups.append(integer_part[-3:])
        integer_part = integer_part[:-3]
    return "R$ " + ".".join(reversed(groups)) + "," + decimal_part
