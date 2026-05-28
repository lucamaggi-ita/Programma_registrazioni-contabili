from __future__ import annotations

import csv
import hashlib
import json
import shutil
import time
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Mapping

from fantozzi.dre_gerencial_provisoria import gerar_dre_gerencial_ok_provisoria
from fantozzi.excel_writer import LANCAMENTOS_SHEET
from fantozzi.lancamentos_manuais_gerenciais import (
    CSV_OPTIONAL_HEADERS,
    CSV_REQUIRED_HEADERS,
    processar_lancamentos_manuais_gerenciais,
)
from fantozzi.pacote_contabil_gerencial import gerar_pacote_contabil_gerencial_provisorio
from fantozzi.pipeline_paths import resolver_arquivos_lancamentos_xlsx
from fantozzi.pre_livro_diario import gerar_pre_livro_diario_teste
from fantozzi.relatorios_gerenciais_ok import gerar_relatorios_gerenciais_ok_provisorios
from openpyxl import load_workbook


STATUS_BASELINE_APROVADA = "APROVADA"
TIPO_RUN_INCREMENTAL = "INCREMENTAL"

BLOCCO_BASELINE_AUSENTE = "BLOCCO_INCREMENTAL_BASELINE_AUSENTE"
BLOCCO_BASELINE_NAO_APROVADA = "BLOCCO_INCREMENTAL_BASELINE_NAO_APROVADA"
BLOCCO_BASELINE_INSUFICIENTE = "BLOCCO_INCREMENTAL_BASELINE_INSUFICIENTE"
BLOCCO_STRUCTURAL_CHANGE = "BLOCCO_INCREMENTAL_STRUCTURAL_CHANGE"
BLOCCO_MANUALI_NAO_APROVADO = "BLOCCO_INCREMENTAL_MANUALI_NAO_APROVADO"
BLOCCO_DELTA_JA_APLICADO = "BLOCCO_INCREMENTAL_DELTA_JA_APLICADO"


class RunsIncrementaisError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(f"{code}: {message}")
        self.code = code
        self.message = message


@dataclass(frozen=True)
class BaselineAprovadaResult:
    baseline_path: Path
    baseline_id: str
    run_id: str
    source_run_path: Path
    reports_run_path: Path


@dataclass(frozen=True)
class RunIncrementalResult:
    manifest_path: Path
    incremental_id: str
    baseline_id: str
    status: str
    output_run_dir: Path
    reports_run_dir: Path
    elapsed_seconds: float


def registrar_baseline_aprovada(
    manifest_run_path: Path | str,
    saida_path: Path | str,
    project_root: Path | str | None = None,
) -> BaselineAprovadaResult:
    project_root = Path(project_root) if project_root else Path.cwd()
    manifest_run_path = Path(manifest_run_path)
    saida_path = Path(saida_path)

    if not manifest_run_path.exists():
        raise RunsIncrementaisError(
            BLOCCO_BASELINE_AUSENTE,
            f"manifest_run.json nao encontrado: {manifest_run_path}",
        )

    manifest = _load_json(manifest_run_path)
    run_id = _text(manifest.get("run_id")) or manifest_run_path.parent.name
    baseline_id = f"baseline_{run_id}"
    reports_run_path = manifest_run_path.parent
    source_run_path = _source_run_path_from_manifest(manifest, run_id)

    required_paths = {
        "source_run_path": source_run_path,
        "pre_livro_path": reports_run_path / "pre_livro_diario_gerencial_teste.xlsx",
        "dre_path": reports_run_path / "dre_gerencial_ok_provisoria.xlsx",
        "pacote_path": reports_run_path / "pacote_contabil_gerencial_provisorio",
    }
    missing = [name for name, path in required_paths.items() if not path.exists()]
    if missing:
        raise RunsIncrementaisError(
            BLOCCO_BASELINE_INSUFICIENTE,
            "baseline sem outputs obrigatorios: " + ", ".join(missing),
        )

    structural_hashes = _structural_hashes(project_root)
    banca_manifest = manifest.get("lancamentos_banca_aprovados") or {}
    manuais_manifest = manifest.get("lancamentos_manuais_gerenciais") or {}
    manuali_hashes: list[str] = []
    manuali_input = _text(manuais_manifest.get("input_csv"))
    if manuais_manifest.get("executado") and manuali_input:
        manuali_path = _resolve_project_path(project_root, manuali_input)
        if manuali_path.exists():
            manuali_hashes.append(_sha256_file(manuali_path))

    periodo = _periodo_baseline_from_manifest(manifest, source_run_path)
    baseline = {
        "baseline_id": baseline_id,
        "run_id": run_id,
        "status": STATUS_BASELINE_APROVADA,
        "periodo_ini": periodo[0],
        "periodo_fim": periodo[1],
        "created_at": _now_iso(),
        "source_run_path": str(source_run_path),
        "reports_run_path": str(reports_run_path),
        "pre_livro_path": str(required_paths["pre_livro_path"]),
        "dre_path": str(required_paths["dre_path"]),
        "pacote_path": str(required_paths["pacote_path"]),
        "inputs": {
            "xml": {
                "path": _text((manifest.get("parametros") or {}).get("xml_base")),
                "hash": _xml_fingerprint(project_root, manifest),
                "count": (manifest.get("totais") or {}).get("xml_processados", 0),
            },
            "banca": {
                "path": _text(banca_manifest.get("input_xlsx")),
                "hash": _optional_file_hash(project_root, banca_manifest.get("input_xlsx")),
                "count": banca_manifest.get("linhas_lidas", 0),
            },
            "manuali": {
                "path": manuali_input,
                "hash": manuali_hashes[0] if manuali_hashes else None,
                "count": manuais_manifest.get("lancamentos_aceitos", 0),
            },
        },
        "structural_hashes": structural_hashes,
        "structural_hash": _hash_jsonable(structural_hashes),
        "applied_manuali_hashes": manuali_hashes,
        "incrementais_aplicados": [],
    }
    _write_json(saida_path, baseline)
    return BaselineAprovadaResult(
        baseline_path=saida_path,
        baseline_id=baseline_id,
        run_id=run_id,
        source_run_path=source_run_path,
        reports_run_path=reports_run_path,
    )


def executar_run_incremental_manuais(
    baseline_path: Path | str,
    manuali_csv_path: Path | str,
    manuali_manifest_path: Path | str,
    saida_dir: Path | str,
    project_root: Path | str | None = None,
) -> RunIncrementalResult:
    started = time.perf_counter()
    project_root = Path(project_root) if project_root else Path.cwd()
    baseline_path = Path(baseline_path)
    manuali_csv_path = Path(manuali_csv_path)
    manuali_manifest_path = Path(manuali_manifest_path)
    saida_dir = Path(saida_dir)

    if not baseline_path.exists():
        raise RunsIncrementaisError(
            BLOCCO_BASELINE_AUSENTE,
            f"baseline aprovada nao encontrada: {baseline_path}",
        )

    baseline = _load_json(baseline_path)
    if _text(baseline.get("status")).upper() != STATUS_BASELINE_APROVADA:
        raise RunsIncrementaisError(
            BLOCCO_BASELINE_NAO_APROVADA,
            f"status baseline invalido: {baseline.get('status')}",
        )

    current_structural_hashes = _structural_hashes(project_root)
    current_structural_hash = _hash_jsonable(current_structural_hashes)
    if _text(baseline.get("structural_hash")) != current_structural_hash:
        raise RunsIncrementaisError(
            BLOCCO_STRUCTURAL_CHANGE,
            "hash estrutural mudou; execute uma full run explicita.",
        )

    manuali_manifest = _load_json(manuali_manifest_path)
    manuali_status = _text(manuali_manifest.get("status_validacao")).upper()
    if manuali_status != "APROVADO":
        raise RunsIncrementaisError(
            BLOCCO_MANUALI_NAO_APROVADO,
            f"manifest manuali nao aprovado: {manuali_status or 'AUSENTE'}",
        )

    manuali_fp = _manuali_fingerprint(manuali_csv_path, manuali_manifest_path)
    _validar_periodo_manuali(baseline, manuali_fp)
    applied_hashes = set(baseline.get("applied_manuali_hashes") or [])
    if manuali_fp["hash"] in applied_hashes:
        raise RunsIncrementaisError(
            BLOCCO_DELTA_JA_APLICADO,
            f"delta manuali ja aplicado: {manuali_fp['hash']}",
        )

    source_run_path = Path(_text(baseline.get("source_run_path")))
    if not source_run_path.exists():
        raise RunsIncrementaisError(
            BLOCCO_BASELINE_INSUFICIENTE,
            f"source_run_path da baseline nao encontrado: {source_run_path}",
        )

    incremental_id = _incremental_id(saida_dir)
    output_run_dir = saida_dir / "reprocessamento_output" / "runs" / incremental_id
    reports_run_dir = saida_dir / "reprocessamento_reports" / "runs" / incremental_id
    reports_run_dir.mkdir(parents=True, exist_ok=True)
    output_run_dir.parent.mkdir(parents=True, exist_ok=True)
    shutil.copytree(source_run_path, output_run_dir, dirs_exist_ok=True)

    manuali_dir = output_run_dir / f"MANUAL_GERENCIAL_INCREMENTAL_{incremental_id}"
    manuali_dir.mkdir(parents=True, exist_ok=True)
    manuali_result = processar_lancamentos_manuais_gerenciais(
        manuali_csv_path,
        project_root / "config" / "plano_contas_gerencial.yaml",
        project_root / "config" / "mapeamento_dre.yaml",
        manuali_dir / "lancamentos_manuais_gerenciais.xlsx",
        manuali_dir / "lancamentos_manuais_gerenciais_aplicados.csv",
        reports_run_dir / "auditoria_lancamentos_manuais_gerenciais.txt",
    )

    pre_livro_path = reports_run_dir / "pre_livro_diario_gerencial_teste.xlsx"
    pre_livro_report = reports_run_dir / "auditoria_pre_livro_diario_gerencial_teste.txt"
    lancamentos_xlsx = resolver_arquivos_lancamentos_xlsx(output_run_dir)
    gerar_pre_livro_diario_teste(
        lancamentos_xlsx,
        pre_livro_path,
        relatorio_auditoria_path=pre_livro_report,
    )

    relatorios_path = reports_run_dir / "relatorios_gerenciais_ok_provisorios.xlsx"
    relatorios_report = reports_run_dir / "relatorios_gerenciais_ok_provisorios.txt"
    gerar_relatorios_gerenciais_ok_provisorios(
        [output_run_dir],
        relatorios_path,
        report_path=relatorios_report,
    )

    dre_path = reports_run_dir / "dre_gerencial_ok_provisoria.xlsx"
    dre_report = reports_run_dir / "dre_gerencial_ok_provisoria.txt"
    dre_result = gerar_dre_gerencial_ok_provisoria(
        [output_run_dir],
        project_root / "config" / "mapeamento_dre.yaml",
        dre_path,
        report_path=dre_report,
    )

    manifest_incremental = _incremental_manifest(
        baseline,
        incremental_id,
        manuali_manifest,
        manuali_fp,
        manuali_result,
        dre_result.status_dre,
        output_run_dir,
        reports_run_dir,
        pre_livro_path,
        dre_path,
        reports_run_dir / "pacote_contabil_gerencial_provisorio",
        started,
    )
    _write_json(reports_run_dir / "manifest_run.json", manifest_incremental)
    pacote_result = gerar_pacote_contabil_gerencial_provisorio(reports_run_dir)
    manifest_incremental["outputs"]["pacote"] = str(pacote_result.pacote_dir)
    manifest_incremental["pacote_contabil_gerencial"] = {
        "executado": True,
        "status": pacote_result.status,
        "diretorio": str(pacote_result.pacote_dir),
        "manifest": str(pacote_result.manifest_path),
    }
    manifest_incremental["elapsed_seconds"] = round(time.perf_counter() - started, 3)

    manifest_path = reports_run_dir / "manifest_incremental.json"
    _write_json(manifest_path, manifest_incremental)
    _write_json(reports_run_dir / "manifest_run.json", manifest_incremental)
    _registrar_delta_aplicado(baseline_path, baseline, manifest_incremental, manuali_fp)

    return RunIncrementalResult(
        manifest_path=manifest_path,
        incremental_id=incremental_id,
        baseline_id=_text(baseline.get("baseline_id")),
        status=_text(manifest_incremental.get("status")),
        output_run_dir=output_run_dir,
        reports_run_dir=reports_run_dir,
        elapsed_seconds=float(manifest_incremental["elapsed_seconds"]),
    )


def _incremental_manifest(
    baseline: Mapping[str, object],
    incremental_id: str,
    manuali_manifest: Mapping[str, object],
    manuali_fp: Mapping[str, object],
    manuali_result: object,
    dre_status: str,
    output_run_dir: Path,
    reports_run_dir: Path,
    pre_livro_path: Path,
    dre_path: Path,
    pacote_path: Path,
    started: float,
) -> dict[str, object]:
    return {
        "tipo_run": TIPO_RUN_INCREMENTAL,
        "baseline_id": baseline.get("baseline_id"),
        "baseline_run_id": baseline.get("run_id"),
        "periodo_ini": baseline.get("periodo_ini"),
        "periodo_fim": baseline.get("periodo_fim"),
        "baseline_periodo_ini": baseline.get("periodo_ini"),
        "baseline_periodo_fim": baseline.get("periodo_fim"),
        "incremental_id": incremental_id,
        "run_id": incremental_id,
        "data_execucao": _now_iso(),
        "status_run": "PROVISORIO_CONCLUIDO",
        "status": "APROVADO" if dre_status == "OK_PROVISORIO" else "APROVADO_CON_RISERVA",
        "fontes_reprocessadas": {
            "xml": False,
            "banca": False,
            "comprovantes": False,
            "conta_corrente": False,
            "manuali": True,
        },
        "fontes_reutilizadas": {
            "xml": True,
            "banca": True,
            "comprovantes": True,
            "conta_corrente": True,
        },
        "lancamentos_manuais_gerenciais": {
            "executado": True,
            "input_csv": str(manuali_result.input_csv),
            "output_xlsx_run": str(manuali_result.output_xlsx),
            "linhas_lidas": manuali_result.linhas_lidas,
            "lancamentos_aceitos": manuali_result.lancamentos_aceitos,
            "linhas_tecnicas_geradas": manuali_result.linhas_tecnicas_geradas,
            "total_manual": str(manuali_result.total_manual),
        },
        "manuali": {
            "manifest_status": manuali_manifest.get("status_validacao"),
            "linhas": manuali_fp["linhas"],
            "valor_total": str(manuali_fp["valor_total"]),
            "hash": manuali_fp["hash"],
            "periodo_ini": manuali_fp["periodo_ini"],
            "periodo_fim": manuali_fp["periodo_fim"],
            "competencias_afetadas": manuali_fp["competencias"],
        },
        "structural_change_detected": False,
        "full_run_executed": False,
        "dre_status": dre_status,
        "outputs": {
            "output_run_dir": str(output_run_dir),
            "reports_run_dir": str(reports_run_dir),
            "pre_livro": str(pre_livro_path),
            "dre": str(dre_path),
            "pacote": str(pacote_path),
        },
        "relatorios_gerados": [
            str(pre_livro_path),
            str(reports_run_dir / "relatorios_gerenciais_ok_provisorios.xlsx"),
            str(dre_path),
        ],
        "erros": [],
        "avisos": [],
        "elapsed_seconds": round(time.perf_counter() - started, 3),
    }


def _registrar_delta_aplicado(
    baseline_path: Path,
    baseline: dict[str, object],
    manifest_incremental: Mapping[str, object],
    manuali_fp: Mapping[str, object],
) -> None:
    applied_hashes = list(baseline.get("applied_manuali_hashes") or [])
    applied_hashes.append(str(manuali_fp["hash"]))
    baseline["applied_manuali_hashes"] = sorted(set(applied_hashes))
    incrementais = list(baseline.get("incrementais_aplicados") or [])
    incrementais.append(
        {
            "incremental_id": manifest_incremental.get("incremental_id"),
            "tipo": "manuali",
            "hash": manuali_fp["hash"],
            "status": manifest_incremental.get("status"),
            "manifest": str(
                Path(str(manifest_incremental["outputs"]["reports_run_dir"]))
                / "manifest_incremental.json"
            ),
        }
    )
    baseline["incrementais_aplicados"] = incrementais
    baseline["run_id"] = manifest_incremental.get("incremental_id")
    baseline["source_run_path"] = manifest_incremental["outputs"]["output_run_dir"]
    baseline["reports_run_path"] = manifest_incremental["outputs"]["reports_run_dir"]
    baseline["pre_livro_path"] = manifest_incremental["outputs"]["pre_livro"]
    baseline["dre_path"] = manifest_incremental["outputs"]["dre"]
    baseline["pacote_path"] = manifest_incremental["outputs"]["pacote"]
    baseline["ultimo_incremental_id"] = manifest_incremental.get("incremental_id")
    baseline["ultima_atualizacao_tipo"] = TIPO_RUN_INCREMENTAL
    baseline_periodo_ini = _text(baseline.get("periodo_ini"))
    baseline_periodo_fim = _text(baseline.get("periodo_fim"))
    manuali_periodo = manifest_incremental.get("manuali") or {}
    baseline["periodo_ini"] = baseline_periodo_ini or _text(
        manuali_periodo.get("periodo_ini")
    )
    baseline["periodo_fim"] = baseline_periodo_fim or _text(
        manuali_periodo.get("periodo_fim")
    )
    baseline["updated_at"] = _now_iso()
    _write_json(baseline_path, baseline)


def _manuali_fingerprint(
    csv_path: Path,
    manifest_path: Path,
) -> dict[str, object]:
    if not csv_path.exists():
        raise RunsIncrementaisError(
            BLOCCO_BASELINE_INSUFICIENTE,
            f"CSV manuali nao encontrado: {csv_path}",
        )
    headers = [*CSV_REQUIRED_HEADERS, *CSV_OPTIONAL_HEADERS]
    linhas = 0
    valor_total = Decimal("0.00")
    row_hashes: list[str] = []
    datas: list[date] = []
    competencias: set[str] = set()
    with csv_path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file, delimiter=";")
        missing = [header for header in CSV_REQUIRED_HEADERS if header not in (reader.fieldnames or [])]
        if missing:
            raise RunsIncrementaisError(
                BLOCCO_BASELINE_INSUFICIENTE,
                "CSV manuali sem colunas obrigatorias: " + ", ".join(missing),
            )
        for row in reader:
            linhas += 1
            valor_total += Decimal(_text(row.get("valor")) or "0")
            data_lancamento = _parse_iso_date(_text(row.get("data")))
            datas.append(data_lancamento)
            competencias.add(data_lancamento.strftime("%Y-%m"))
            normalized = {header: _text(row.get(header)) for header in headers}
            row_hashes.append(_hash_jsonable(normalized))
    manifest_hash = _sha256_file(manifest_path)
    combined = _hash_jsonable({"rows": sorted(row_hashes), "manifest": manifest_hash})
    return {
        "hash": combined,
        "linhas": linhas,
        "valor_total": valor_total.quantize(Decimal("0.01")),
        "row_hashes": row_hashes,
        "manifest_hash": manifest_hash,
        "periodo_ini": min(datas).isoformat() if datas else None,
        "periodo_fim": max(datas).isoformat() if datas else None,
        "competencias": sorted(competencias),
    }


def _validar_periodo_manuali(
    baseline: Mapping[str, object],
    manuali_fp: Mapping[str, object],
) -> None:
    periodo_ini = _text(baseline.get("periodo_ini"))
    periodo_fim = _text(baseline.get("periodo_fim"))
    manual_ini = _text(manuali_fp.get("periodo_ini"))
    manual_fim = _text(manuali_fp.get("periodo_fim"))
    if not periodo_ini or not periodo_fim or not manual_ini or not manual_fim:
        return
    if manual_ini < periodo_ini or manual_fim > periodo_fim:
        raise RunsIncrementaisError(
            BLOCCO_BASELINE_INSUFICIENTE,
            (
                "a baseline corrente nao cobre estas datas; manuali fora do "
                "periodo coberto pela baseline: "
                f"{manual_ini}..{manual_fim} fora de {periodo_ini}..{periodo_fim}. "
                "As datas do Excel podem estar corretas; registre/promova uma "
                "baseline aprovada que cubra o periodo. Nao sera feita full run "
                "automatica."
            ),
        )


def _periodo_lancamentos(source_run_path: Path) -> tuple[str | None, str | None]:
    datas: list[date] = []
    for excel_path in resolver_arquivos_lancamentos_xlsx(source_run_path):
        try:
            workbook = load_workbook(excel_path, data_only=True, read_only=True)
        except Exception:
            continue
        if LANCAMENTOS_SHEET not in workbook.sheetnames:
            continue
        sheet = workbook[LANCAMENTOS_SHEET]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        if "data_documento" not in headers:
            continue
        data_idx = headers.index("data_documento")
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or data_idx >= len(row):
                continue
            parsed = _coerce_date(row[data_idx])
            if parsed:
                datas.append(parsed)
    if not datas:
        return None, None
    return min(datas).isoformat(), max(datas).isoformat()


def _periodo_baseline_from_manifest(
    manifest: Mapping[str, object],
    source_run_path: Path,
) -> tuple[str | None, str | None]:
    periodo_ini = _text(manifest.get("periodo_ini")) or _text(
        manifest.get("baseline_periodo_ini")
    )
    periodo_fim = _text(manifest.get("periodo_fim")) or _text(
        manifest.get("baseline_periodo_fim")
    )
    if periodo_ini or periodo_fim:
        return periodo_ini or None, periodo_fim or None
    return _periodo_lancamentos(source_run_path)


def _parse_iso_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise RunsIncrementaisError(
            BLOCCO_BASELINE_INSUFICIENTE,
            f"data manuali invalida para incremental: {value}",
        ) from exc


def _coerce_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _min_date_text(left: str, right: str) -> str | None:
    values = [value for value in (left, right) if value]
    return min(values) if values else None


def _max_date_text(left: str, right: str) -> str | None:
    values = [value for value in (left, right) if value]
    return max(values) if values else None


def _source_run_path_from_manifest(manifest: Mapping[str, object], run_id: str) -> Path:
    parametros = manifest.get("parametros") or {}
    saida_output = _text(parametros.get("saida_output"))
    if saida_output:
        return Path(saida_output) / "runs" / run_id
    output_run_dir = _text((manifest.get("outputs") or {}).get("output_run_dir"))
    if output_run_dir:
        return Path(output_run_dir)
    raise RunsIncrementaisError(
        BLOCCO_BASELINE_INSUFICIENTE,
        "manifest sem parametros.saida_output ou outputs.output_run_dir",
    )


def _structural_hashes(project_root: Path) -> dict[str, str | None]:
    rel_paths = (
        "config/plano_contas_gerencial.yaml",
        "config/mapeamento_dre.yaml",
        "config/regras_classificacao_produtos.yaml",
        "config/regras_classificacao_comprovantes.yaml",
        "src/fantozzi/reprocessamento_pipeline.py",
        "src/fantozzi/pre_livro_diario.py",
        "src/fantozzi/dre_gerencial_provisoria.py",
        "src/fantozzi/relatorios_gerenciais_ok.py",
        "src/fantozzi/lancamentos_manuais_gerenciais.py",
        "src/fantozzi/pacote_contabil_gerencial.py",
    )
    hashes: dict[str, str | None] = {}
    for rel_path in rel_paths:
        path = project_root / rel_path
        hashes[rel_path] = _sha256_file(path) if path.exists() else None
    return hashes


def _xml_fingerprint(project_root: Path, manifest: Mapping[str, object]) -> str | None:
    xml_base = _text((manifest.get("parametros") or {}).get("xml_base"))
    if not xml_base:
        return None
    xml_path = _resolve_project_path(project_root, xml_base)
    if not xml_path.exists():
        return None
    if xml_path.is_file():
        return _sha256_file(xml_path)
    rows = []
    for path in sorted(xml_path.rglob("*.xml")):
        rows.append(
            {
                "path": str(path.relative_to(xml_path)),
                "size": path.stat().st_size,
                "sha256": _sha256_file(path),
            }
        )
    return _hash_jsonable(rows)


def _optional_file_hash(project_root: Path, value: object) -> str | None:
    path_text = _text(value)
    if not path_text:
        return None
    path = _resolve_project_path(project_root, path_text)
    return _sha256_file(path) if path.exists() and path.is_file() else None


def _resolve_project_path(project_root: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else project_root / path


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _hash_jsonable(value: object) -> str:
    payload = json.dumps(
        _json_ready(value),
        sort_keys=True,
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _load_json(path: Path) -> dict[str, object]:
    with path.open("r", encoding="utf-8") as file:
        return json.load(file)


def _write_json(path: Path, payload: Mapping[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(_json_ready(payload), indent=2, ensure_ascii=False),
        encoding="utf-8",
        newline="\n",
    )


def _json_ready(value: object) -> object:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, Mapping):
        return {str(key): _json_ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_ready(item) for item in value]
    return value


def _incremental_id(saida_dir: Path) -> str:
    prefix = datetime.now().strftime("%Y-%m-%d")
    runs_root = saida_dir / "reprocessamento_reports" / "runs"
    existing = sorted(runs_root.glob(f"{prefix}_*")) if runs_root.exists() else []
    return f"{prefix}_{len(existing) + 1:03d}"


def _now_iso() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()
