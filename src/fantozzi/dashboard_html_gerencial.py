from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from fantozzi.dre_pdf_gerencial import (
    DrePdfGerencialError,
    _find_dre_xlsx,
    _fmt_brl,
    _load_dre_geral,
    _load_manifest,
    _load_periodos,
    _to_decimal,
)


class DashboardHtmlGerencialError(Exception):
    pass


@dataclass(frozen=True)
class DashboardHtmlResult:
    html_path: Path
    dre_xlsx_source: Path
    run_id: str | None
    periodos: tuple[str, ...]
    total_linhas: int
    status: str


def gerar_dashboard_html(
    run_dir: Path | str,
    output_dir: Path | str | None = None,
) -> DashboardHtmlResult:
    run_dir = Path(run_dir)
    if not run_dir.exists() or not run_dir.is_dir():
        raise DashboardHtmlGerencialError(f"run-dir nao encontrado: {run_dir}")

    try:
        dre_xlsx = _find_dre_xlsx(run_dir)
        dre_rows = _load_dre_geral(dre_xlsx)
    except DrePdfGerencialError as exc:
        raise DashboardHtmlGerencialError(str(exc)) from exc
    if not dre_rows:
        raise DashboardHtmlGerencialError(f"Nenhuma linha DRE encontrada em: {dre_xlsx}")

    periodos = _load_periodos(dre_xlsx)
    mensal_rows = _load_dre_mensal(dre_xlsx)
    manifest = _load_manifest(run_dir)
    run_id = manifest.get("run_id") or None

    if output_dir is None:
        output_dir = run_dir / "output_visual"
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    html_path = output_dir / "02_DASHBOARD_GERENCIAL.html"
    _render_html(html_path, dre_rows, mensal_rows, periodos, run_id)

    return DashboardHtmlResult(
        html_path=html_path,
        dre_xlsx_source=dre_xlsx,
        run_id=run_id,
        periodos=tuple(periodos),
        total_linhas=len(dre_rows),
        status="TESTE_TECNICO_NAO_OFICIAL",
    )


def console_lines(result: DashboardHtmlResult) -> list[str]:
    lines = [
        "[OK] Dashboard HTML gerencial gerado.",
        f"     Status: {result.status}",
        f"     HTML: {result.html_path}",
        f"     Fonte DRE: {result.dre_xlsx_source}",
        f"     Linhas DRE: {result.total_linhas}",
    ]
    if result.run_id:
        lines.append(f"     Run: {result.run_id}")
    if result.periodos:
        lines.append(f"     Periodos: {' | '.join(result.periodos)}")
    return lines


def _load_dre_mensal(xlsx_path: Path) -> list[dict]:
    try:
        wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
        if "DRE_Mensal" not in wb.sheetnames:
            wb.close()
            return []
        ws = wb["DRE_Mensal"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return []
    if len(rows) < 2:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        d = {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
        if any(v is not None for v in d.values()):
            result.append(d)
    return result


_PALETTE = [
    "#1F4E78",
    "#2E75B6",
    "#5B9BD5",
    "#A8C4E0",
    "#F4A261",
    "#E76F51",
    "#2A9D8F",
    "#E9C46A",
]

_CSS = """\
* { box-sizing: border-box; margin: 0; padding: 0; }
body { font-family: Arial, sans-serif; background: #F4F6FA; color: #222; font-size: 13px; }
.header { background: #1F4E78; color: white; padding: 18px 32px 14px; }
.header h1 { font-size: 20px; font-weight: bold; margin-bottom: 4px; }
.header-meta { font-size: 11px; opacity: 0.85; margin-top: 6px; display: flex; flex-wrap: wrap; gap: 12px; align-items: center; }
.status-badge { background: #E04B2A; color: white; font-size: 10px; font-weight: bold; padding: 2px 8px; border-radius: 3px; letter-spacing: 0.5px; }
.meta-badge { background: rgba(255,255,255,0.15); padding: 2px 8px; border-radius: 3px; font-size: 10px; }
.main { padding: 24px 32px; }
.cards { display: flex; flex-wrap: wrap; gap: 16px; margin-bottom: 24px; }
.card { background: white; border-radius: 8px; padding: 16px 20px; flex: 1; min-width: 180px; box-shadow: 0 1px 4px rgba(0,0,0,0.08); border-left: 4px solid #2E75B6; }
.card-label { font-size: 11px; color: #666; text-transform: uppercase; letter-spacing: 0.5px; margin-bottom: 6px; }
.card-value { font-size: 18px; font-weight: bold; color: #1F4E78; }
.card-sub { font-size: 10px; color: #888; margin-top: 4px; }
.card-total { border-left-color: #1F4E78; }
.charts-row { display: flex; gap: 16px; margin-bottom: 24px; }
.chart-half { flex: 1; background: white; border-radius: 8px; box-shadow: 0 1px 4px rgba(0,0,0,0.08); padding: 8px; overflow: hidden; min-width: 0; }
.chart-full { background: white; border-radius: 8px; box-shadow: 0 1px 4px rgba(0,0,0,0.08); padding: 8px; margin-bottom: 24px; overflow: hidden; }
.section-title { font-size: 14px; font-weight: bold; color: #1F4E78; margin-bottom: 12px; }
.table-section { background: white; border-radius: 8px; box-shadow: 0 1px 4px rgba(0,0,0,0.08); padding: 20px; margin-bottom: 24px; }
table { width: 100%; border-collapse: collapse; font-size: 12px; }
thead th { background: #1F4E78; color: white; padding: 8px 12px; text-align: left; font-weight: bold; }
thead th.align-right { text-align: right; }
tbody tr:nth-child(even) { background: #F4F6FA; }
tbody tr.group-row { background: #D9E1F2; font-weight: bold; }
tbody tr.subtotal-row { background: #EEF2F9; font-weight: bold; border-top: 1px solid #8EA9C1; }
tbody tr.total-row { background: #1F4E78; color: white; font-weight: bold; font-size: 13px; }
td { padding: 6px 12px; }
td.align-right { text-align: right; font-variant-numeric: tabular-nums; }
.footer { font-size: 10px; color: #888; padding: 16px 32px 24px; border-top: 1px solid #DDE2EA; }
@media (max-width: 700px) { .charts-row { flex-direction: column; } .main { padding: 16px; } }
"""


def _render_html(
    html_path: Path,
    dre_rows: list[dict],
    mensal_rows: list[dict],
    periodos: list[str],
    run_id: str | None,
) -> None:
    try:
        import plotly.graph_objects as go
    except ImportError as exc:
        raise DashboardHtmlGerencialError(
            "Plotly nao esta instalado. Execute: pip install 'plotly>=5.0'"
        ) from exc

    grupos: dict[str, Decimal] = {}
    for row in dre_rows:
        grupo = str(row.get("grupo_dre") or "—")
        valor = _to_decimal(row.get("valor"))
        grupos[grupo] = grupos.get(grupo, Decimal("0")) + valor

    grand_total = sum(grupos.values(), Decimal("0"))
    grupo_names = list(grupos.keys())
    grupo_values = [float(v) for v in grupos.values()]
    bar_colors = [_PALETTE[i % len(_PALETTE)] for i in range(len(grupo_names))]

    fig_bar = go.Figure(go.Bar(
        x=grupo_names,
        y=grupo_values,
        marker_color=bar_colors,
        text=[_fmt_brl(v) for v in grupo_values],
        textposition="outside",
        hovertemplate="%{x}<br>%{text}<extra></extra>",
    ))
    fig_bar.update_layout(
        title=dict(text="Subtotal por Grupo DRE", font=dict(size=14, color="#1F4E78")),
        yaxis_title="R$",
        plot_bgcolor="#F8FAFC",
        paper_bgcolor="#FFFFFF",
        font=dict(family="Arial, sans-serif", size=11),
        margin=dict(l=70, r=30, t=50, b=80),
        height=360,
        showlegend=False,
    )

    fig_pie = go.Figure(go.Pie(
        labels=grupo_names,
        values=grupo_values,
        marker=dict(colors=bar_colors),
        hovertemplate="%{label}<br>%{percent}<br>%{text}<extra></extra>",
        text=[_fmt_brl(v) for v in grupo_values],
        textinfo="percent+label",
        hole=0.3,
    ))
    fig_pie.update_layout(
        title=dict(text="Composicao por Grupo", font=dict(size=14, color="#1F4E78")),
        paper_bgcolor="#FFFFFF",
        font=dict(family="Arial, sans-serif", size=11),
        margin=dict(l=20, r=20, t=50, b=20),
        height=360,
        showlegend=False,
    )

    line_div = ""
    if len(periodos) >= 2 and mensal_rows:
        mensal_agg: dict[tuple[str, str], Decimal] = {}
        for row in mensal_rows:
            ano_mes = str(row.get("ano_mes") or "")
            grupo = str(row.get("grupo_dre") or "—")
            valor = _to_decimal(row.get("valor"))
            key = (ano_mes, grupo)
            mensal_agg[key] = mensal_agg.get(key, Decimal("0")) + valor

        fig_line = go.Figure()
        for i, grupo in enumerate(grupo_names):
            y_vals = [
                float(mensal_agg.get((mes, grupo), Decimal("0")))
                for mes in periodos
            ]
            fig_line.add_trace(go.Scatter(
                x=periodos,
                y=y_vals,
                mode="lines+markers",
                name=grupo,
                line=dict(color=_PALETTE[i % len(_PALETTE)], width=2),
                marker=dict(size=6),
                hovertemplate="%{x}<br>" + grupo + "<br>R$ %{y:,.2f}<extra></extra>",
            ))
        fig_line.update_layout(
            title=dict(text="Andamento Mensal por Grupo", font=dict(size=14, color="#1F4E78")),
            yaxis_title="R$",
            xaxis_title="Mês",
            plot_bgcolor="#F8FAFC",
            paper_bgcolor="#FFFFFF",
            font=dict(family="Arial, sans-serif", size=11),
            margin=dict(l=70, r=30, t=50, b=60),
            height=360,
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
        )
        line_div = (
            '<div class="chart-full">'
            + fig_line.to_html(full_html=False, include_plotlyjs=False)
            + "</div>"
        )

    # Bar chart carries the embedded plotly.js (True = inline, offline-safe)
    bar_html = fig_bar.to_html(full_html=False, include_plotlyjs=True)
    pie_html = fig_pie.to_html(full_html=False, include_plotlyjs=False)

    periodo_str = " · ".join(periodos) if periodos else "—"
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    run_badge = (
        '<span class="meta-badge">Run: ' + run_id + "</span>" if run_id else ""
    )

    cards_html = _build_cards_html(grupos, grand_total, len(periodos), periodo_str)
    table_html = _build_table_html(dre_rows)

    html = (
        "<!DOCTYPE html>\n"
        '<html lang="pt-BR">\n'
        "<head>\n"
        '<meta charset="UTF-8">\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1.0">\n'
        "<title>DRE Gerencial — Dashboard " + periodo_str + "</title>\n"
        "<style>\n" + _CSS + "</style>\n"
        "</head>\n"
        "<body>\n"
        '<div class="header">\n'
        "  <h1>DRE Gerencial Provisória — Dashboard</h1>\n"
        '  <div class="header-meta">\n'
        "    <span>Período: " + periodo_str + "</span>\n"
        "    <span>Gerado em: " + now_str + "</span>\n"
        "    " + run_badge + "\n"
        '    <span class="status-badge">TESTE TÉCNICO / NÃO OFICIAL</span>\n'
        "  </div>\n"
        "</div>\n"
        '<div class="main">\n'
        + cards_html + "\n"
        '  <div class="charts-row">\n'
        '    <div class="chart-half">' + bar_html + "</div>\n"
        '    <div class="chart-half">' + pie_html + "</div>\n"
        "  </div>\n"
        + line_div + "\n"
        '  <div class="table-section">\n'
        '    <div class="section-title">Detalhamento DRE</div>\n'
        + table_html + "\n"
        "  </div>\n"
        "</div>\n"
        '<div class="footer">\n'
        "  Relatório gerencial provisório gerado automaticamente a partir da full run Fantozzi. "
        "Não substitui escrituração contábil oficial sem revisão do contador. "
        "<strong>TESTE TÉCNICO / NÃO OFICIAL</strong>\n"
        "</div>\n"
        "</body>\n"
        "</html>\n"
    )

    html_path.write_text(html, encoding="utf-8")


def _build_cards_html(
    grupos: dict[str, Decimal],
    grand_total: Decimal,
    n_periodos: int,
    periodo_str: str,
) -> str:
    parts = ['  <div class="cards">']
    for grupo, total in grupos.items():
        label = (grupo[:35] + "…") if len(grupo) > 35 else grupo
        parts.append(
            '    <div class="card">'
            '<div class="card-label">' + label + "</div>"
            '<div class="card-value">' + _fmt_brl(total) + "</div>"
            "</div>"
        )
    parts.append(
        '    <div class="card card-total">'
        '<div class="card-label">Total Geral</div>'
        '<div class="card-value">' + _fmt_brl(grand_total) + "</div>"
        '<div class="card-sub">'
        + str(n_periodos)
        + " período(s): "
        + periodo_str
        + "</div>"
        "</div>"
    )
    parts.append("  </div>")
    return "\n".join(parts)


def _build_table_html(dre_rows: list[dict]) -> str:
    grouped: dict[str, list[dict]] = {}
    for row in dre_rows:
        grupo = str(row.get("grupo_dre") or "—")
        if grupo not in grouped:
            grouped[grupo] = []
        grouped[grupo].append(row)

    rows_html: list[str] = []
    grand_total = Decimal("0")

    for grupo, rows in grouped.items():
        rows_html.append(
            '<tr class="group-row">'
            "<td>" + grupo + "</td>"
            '<td class="align-right"></td>'
            "</tr>"
        )
        subtotal = Decimal("0")
        for row in rows:
            linha = str(row.get("linha_dre") or "—")
            valor = _to_decimal(row.get("valor"))
            subtotal += valor
            grand_total += valor
            rows_html.append(
                "<tr>"
                "<td>&nbsp;&nbsp;&nbsp;" + linha + "</td>"
                '<td class="align-right">' + _fmt_brl(valor) + "</td>"
                "</tr>"
            )
        rows_html.append(
            '<tr class="subtotal-row">'
            "<td>&nbsp;&nbsp;&nbsp;Sub-total " + grupo + "</td>"
            '<td class="align-right">' + _fmt_brl(subtotal) + "</td>"
            "</tr>"
        )

    rows_html.append(
        '<tr class="total-row">'
        "<td>TOTAL GERAL</td>"
        '<td class="align-right">' + _fmt_brl(grand_total) + "</td>"
        "</tr>"
    )

    return (
        "    <table>\n"
        '      <thead><tr>'
        "<th>Grupo / Linha DRE</th>"
        '<th class="align-right">Valor (R$)</th>'
        "</tr></thead>\n"
        "      <tbody>"
        + "".join(rows_html)
        + "</tbody>\n"
        "    </table>"
    )
