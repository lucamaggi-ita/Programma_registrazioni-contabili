from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fantozzi.classification_review_export import CSV_HEADERS


DEFAULT_PLANO_CONTAS = (
    Path(__file__).resolve().parents[2] / "config" / "plano_contas_gerencial.yaml"
)

REVISAO_SHEET = "Revisao"
PLANO_CONTAS_SHEET = "Plano_de_Contas"
ACOES_VALIDAS_SHEET = "Acoes_Validas"
GUIA_ALIAS_REGRAS_SHEET = "Guia_Alias_e_Regras"
EXEMPLOS_SHEET = "Exemplos"
INSTRUCOES_SHEET = "Instrucoes"

FRIENDLY_COLUMNS = [
    "conta_escolhida",
    "explicacao_acao",
    "alerta_preenchimento",
    "nivel_risco_regra",
    "exemplo_decisao",
]

REVISAO_COLUMN_ORDER = [
    "descricao_normalizada",
    "descricao_exemplo",
    "ocorrencias",
    "fornecedores",
    "valor_total",
    "conta_atual_nome",
    "alertas",
    "acao_manual",
    "conta_escolhida",
    "palavras_chave_sugeridas",
    "nivel_risco_regra",
    "alerta_preenchimento",
    "observacao_manual",
    "conta_id_sugerida",
    "conta_nome_sugerida",
    "tipo_candidato",
    "documentos",
    "status_revisao",
    "conta_atual_id",
    "explicacao_acao",
    "exemplo_decisao",
]

EDITABLE_COLUMNS = {
    "acao_manual",
    "conta_escolhida",
    "palavras_chave_sugeridas",
    "observacao_manual",
}

ALLOWED_ACTIONS_DESCRIPTIONS = {
    "ANALISAR_COM_CONTADOR": (
        "Usar quando precisa de decisao do contador. Nao gera regra automatica."
    ),
    "CLASSIFICAR": (
        "Classificar este candidato agora, sem necessariamente criar alias ou regra futura."
    ),
    "CRIAR_ALIAS": (
        "Usar para produto especifico e recorrente. Opcao mais segura que regra ampla."
    ),
    "CRIAR_REGRA": (
        "Usar somente quando palavras-chave identificam claramente uma categoria."
    ),
    "IGNORAR": "Usar quando o candidato nao deve gerar classificacao ou regra.",
    "MANTER_REVISAO": "Usar quando ainda nao ha seguranca para classificar.",
}

ACTION_EXAMPLES = {
    "ANALISAR_COM_CONTADOR": "Caso fiscal/contabil duvidoso ou operacao especial.",
    "CLASSIFICAR": "Produto revisado manualmente para uma conta existente.",
    "CRIAR_ALIAS": "COCA COLA LATA 350ML -> Bebidas / Refrigerantes.",
    "CRIAR_REGRA": "REFRIGERANTE; COCA COLA; FANTA; SPRITE; GUARANA.",
    "IGNORAR": "Candidato nao deve gerar proposta.",
    "MANTER_REVISAO": "Manter pendente ate existir decisao segura.",
}

GENERIC_KEYWORDS = {
    "UN",
    "UND",
    "UNIDADE",
    "KG",
    "G",
    "GR",
    "ML",
    "L",
    "LT",
    "CX",
    "CAIXA",
    "PCT",
    "PACOTE",
    "PC",
    "LATA",
    "PET",
    "VIDRO",
    "SACO",
    "BRANCO",
    "PRETO",
    "AZUL",
    "VERDE",
    "EXTRA",
    "ESPECIAL",
    "TRADICIONAL",
    "NEUTRO",
    "ZERO",
    "LIGHT",
    "DIET",
}


class ClassificationReviewWorkbookError(Exception):
    pass


def gerar_planilha_revisao_classificacao(
    csv_path: Path | str,
    output_path: Path | str,
    plano_contas_path: Path | str | None = None,
) -> Path:
    csv_path = Path(csv_path)
    output_path = Path(output_path)
    plano_contas_path = Path(plano_contas_path or DEFAULT_PLANO_CONTAS)

    headers, rows = ler_csv_revisao(csv_path)
    contas = carregar_plano_contas(plano_contas_path)

    workbook = Workbook()
    workbook.security.lockStructure = False
    criar_aba_revisao(workbook, headers, rows, contas)
    criar_aba_plano_de_contas(workbook, contas)
    criar_aba_acoes_validas(workbook)
    criar_aba_guia_alias_e_regras(workbook)
    criar_aba_exemplos(workbook)
    criar_aba_instrucoes(workbook)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def ler_csv_revisao(csv_path: Path | str) -> tuple[list[str], list[dict[str, str]]]:
    csv_path = Path(csv_path)
    if not csv_path.exists():
        raise ClassificationReviewWorkbookError(
            f"CSV de revisao nao encontrado: {csv_path}"
        )
    if not csv_path.is_file():
        raise ClassificationReviewWorkbookError(
            f"CSV de revisao nao e um arquivo: {csv_path}"
        )
    if csv_path.stat().st_size == 0:
        raise ClassificationReviewWorkbookError("CSV de revisao esta vazio.")

    try:
        with csv_path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.DictReader(file, delimiter=";")
            if not reader.fieldnames:
                raise ClassificationReviewWorkbookError(
                    "CSV de revisao nao possui cabecalho valido."
                )
            headers = [header.strip() for header in reader.fieldnames if header]
            missing_headers = [header for header in CSV_HEADERS if header not in headers]
            if missing_headers:
                raise ClassificationReviewWorkbookError(
                    "Colunas obrigatorias ausentes: " + ", ".join(missing_headers)
                )
            return headers, [dict(row) for row in reader]
    except UnicodeDecodeError as exc:
        raise ClassificationReviewWorkbookError(
            f"CSV de revisao nao pode ser lido em UTF-8: {csv_path}"
        ) from exc


def carregar_plano_contas(plano_path: Path | str) -> list[dict[str, str]]:
    plano_path = Path(plano_path)
    if not plano_path.exists():
        raise ClassificationReviewWorkbookError(
            f"Plano de contas nao encontrado: {plano_path}"
        )
    if not plano_path.is_file():
        raise ClassificationReviewWorkbookError(
            f"Plano de contas nao e um arquivo: {plano_path}"
        )

    try:
        with plano_path.open("r", encoding="utf-8") as file:
            data = yaml.safe_load(file) or {}
    except Exception as exc:
        raise ClassificationReviewWorkbookError(
            f"Plano de contas nao pode ser lido: {plano_path}"
        ) from exc

    contas = data.get("contas", [])
    if not isinstance(contas, list):
        raise ClassificationReviewWorkbookError(
            "Plano de contas nao possui lista valida em 'contas'."
        )

    rows: list[dict[str, str]] = []
    for item in contas:
        if not isinstance(item, dict):
            continue
        conta_id = _text(item.get("id") or item.get("conta_id"))
        conta_nome = _text(item.get("nome") or item.get("conta_nome"))
        grupo = _text(item.get("grupo"))
        rows.append(
            {
                "conta_id": conta_id,
                "conta_nome": conta_nome,
                "grupo": grupo,
                "nivel": _account_level(conta_id),
                "conta_formatada": _format_account_choice(conta_id, conta_nome),
                "ativo": _text(item.get("ativo", "sim")) or "sim",
                "observacao": _text(item.get("observacao")),
            }
        )
    return rows


def criar_aba_revisao(
    workbook: Workbook,
    headers: list[str],
    rows: list[dict[str, str]],
    contas: list[dict[str, str]],
) -> None:
    sheet = workbook.active
    sheet.title = REVISAO_SHEET
    output_headers = _review_headers(headers)
    sheet.append(output_headers)

    for row in rows:
        enriched_row = _enrich_review_row(row, contas)
        sheet.append([enriched_row.get(header, "") for header in output_headers])

    _style_header(sheet)
    _apply_revisao_styles(sheet, output_headers)
    _set_filter_and_freeze(sheet)
    _apply_action_validation(sheet, output_headers, len(rows))
    _apply_account_validation(sheet, output_headers, len(rows), len(contas))
    _protect_revisao_sheet(sheet, output_headers)


def criar_aba_plano_de_contas(
    workbook: Workbook, contas: Iterable[dict[str, str]]
) -> None:
    sheet = workbook.create_sheet(PLANO_CONTAS_SHEET)
    headers = [
        "conta_id",
        "conta_nome",
        "grupo",
        "nivel",
        "conta_formatada",
        "ativo",
        "observacao",
    ]
    sheet.append(headers)
    for conta in contas:
        sheet.append([conta.get(header, "") for header in headers])

    _style_header(sheet)
    _set_filter_and_freeze(sheet)
    _set_widths(sheet, {"A": 16, "B": 42, "C": 38, "D": 10, "E": 62, "F": 12, "G": 50})


def criar_aba_acoes_validas(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(ACOES_VALIDAS_SHEET)
    sheet.append(["acao_manual", "descricao", "exige_conta", "exige_palavras_chave"])
    for action, description in ALLOWED_ACTIONS_DESCRIPTIONS.items():
        sheet.append(
            [
                action,
                description,
                "sim" if action in {"CLASSIFICAR", "CRIAR_ALIAS", "CRIAR_REGRA"} else "nao",
                "sim" if action == "CRIAR_REGRA" else "nao",
            ]
        )

    _style_header(sheet)
    _set_filter_and_freeze(sheet)
    _set_widths(sheet, {"A": 28, "B": 92, "C": 18, "D": 24})
    for row in sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


def criar_aba_guia_alias_e_regras(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(GUIA_ALIAS_REGRAS_SHEET)
    rows = [
        ["Tema", "Orientacao"],
        ["ALIAS", "Descricao especifica de um produto conhecido e recorrente."],
        ["REGRA", "Palavras-chave para uma familia clara de produtos."],
        ["Preferencia", "Se estiver em duvida entre CRIAR_ALIAS e CRIAR_REGRA, prefira CRIAR_ALIAS."],
        ["Risco", "CRIAR_REGRA pode classificar muitos produtos futuros se as palavras forem genericas."],
        [],
        ["Situacao", "Melhor acao", "Por que"],
        ["Produto especifico recorrente", "CRIAR_ALIAS", "Mais seguro e pontual"],
        ["Familia clara de produtos", "CRIAR_REGRA", "Ajuda a automatizar casos futuros"],
        ["Duvida contabil", "ANALISAR_COM_CONTADOR", "Evita classificacao errada"],
        ["Duvida operacional", "MANTER_REVISAO", "Mantem pendente"],
        ["Produto irrelevante para regra", "IGNORAR", "Nao gera proposta"],
    ]
    for row in rows:
        sheet.append(row)

    _style_header(sheet)
    _set_widths(sheet, {"A": 34, "B": 34, "C": 78})
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def criar_aba_exemplos(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(EXEMPLOS_SHEET)
    rows = [
        ["Tipo", "Descricao exemplo", "Acao", "Conta", "Palavras-chave", "Motivo"],
        ["Alias", "COCA COLA LATA 350ML", "CRIAR_ALIAS", "Bebidas / Refrigerantes", "", "Produto especifico e recorrente."],
        ["Alias", "DETERGENTE YPE NEUTRO 500ML", "CRIAR_ALIAS", "Material de Limpeza", "", "Descricao especifica. Evita regra generica com NEUTRO ou 500ML."],
        ["Alias", "FARINHA DE TRIGO DONA BENTA 1KG", "CRIAR_ALIAS", "Insumos / Alimentos", "", "Produto especifico. Classificacao pontual e segura."],
        ["Regra", "COCA COLA LATA 350ML", "CRIAR_REGRA", "Bebidas / Refrigerantes", "REFRIGERANTE; COCA COLA; FANTA; SPRITE; GUARANA", "As palavras identificam refrigerantes."],
        ["Regra", "DETERGENTE YPE NEUTRO 500ML", "CRIAR_REGRA", "Material de Limpeza", "DETERGENTE; DESINFETANTE; AGUA SANITARIA; LIMPADOR", "As palavras identificam produtos de limpeza."],
        ["Regra", "COPO DESCARTAVEL 300ML", "CRIAR_REGRA", "Embalagens / Descartaveis", "COPO DESCARTAVEL; PRATO DESCARTAVEL; TALHER DESCARTAVEL; GUARDANAPO", "As palavras identificam descartaveis de uso operacional."],
        [],
        ["Evite palavras genericas", "; ".join(sorted(GENERIC_KEYWORDS))],
        ["Explicacao", "Essas palavras aparecem em muitos produtos diferentes e podem criar regras perigosas."],
    ]
    for row in rows:
        sheet.append(row)

    _style_header(sheet)
    _set_widths(sheet, {"A": 24, "B": 36, "C": 22, "D": 36, "E": 58, "F": 70})
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def criar_aba_instrucoes(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(INSTRUCOES_SHEET)
    instructions = [
        "Planilha assistida de revisao de classificacao",
        "",
        "1. Trabalhe somente na aba Revisao.",
        "2. Nao altere as abas tecnicas.",
        "3. Escolha acao_manual na lista.",
        "4. Escolha conta_escolhida na lista, quando necessario.",
        "5. Nao digite codigo de conta manualmente.",
        "6. Use CRIAR_ALIAS para produto especifico.",
        "7. Use CRIAR_REGRA somente quando as palavras-chave forem claras.",
        "8. Se tiver duvida, use MANTER_REVISAO ou ANALISAR_COM_CONTADOR.",
        "9. Salve o arquivo.",
        "10. Rode o comando de importacao da planilha.",
        "",
        "Nunca faca:",
        "- apagar linhas;",
        "- mudar descricoes originais;",
        "- digitar conta inventada;",
        "- criar regra com palavra generica;",
        "- usar CRIAR_REGRA se nao tiver certeza.",
        "",
        "Nesta fase, nenhuma regra YAML e modificada automaticamente.",
    ]
    for line in instructions:
        sheet.append([line])

    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 112
    for row in sheet.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _review_headers(headers: list[str]) -> list[str]:
    all_headers = [*headers, *[header for header in FRIENDLY_COLUMNS if header not in headers]]
    ordered = [header for header in REVISAO_COLUMN_ORDER if header in all_headers]
    ordered.extend(header for header in all_headers if header not in ordered)
    return ordered


def _enrich_review_row(
    row: dict[str, str], contas: list[dict[str, str]]
) -> dict[str, str]:
    enriched = {header: _text(value) for header, value in row.items()}
    conta_id = enriched.get("conta_id_sugerida", "")
    conta_nome = enriched.get("conta_nome_sugerida", "")
    if conta_id and not conta_nome:
        conta_nome = _account_name_by_id(contas, conta_id)
    enriched.setdefault("conta_escolhida", _format_account_choice(conta_id, conta_nome))

    acao = _text(enriched.get("acao_manual")).upper()
    enriched["explicacao_acao"] = ALLOWED_ACTIONS_DESCRIPTIONS.get(acao, "")
    enriched["alerta_preenchimento"] = _fill_alert(enriched)
    enriched["nivel_risco_regra"] = _risk_level(acao)
    enriched["exemplo_decisao"] = ACTION_EXAMPLES.get(acao, "")
    return enriched


def _fill_alert(row: dict[str, str]) -> str:
    acao = _text(row.get("acao_manual")).upper()
    conta = _text(row.get("conta_escolhida")) or _text(row.get("conta_id_sugerida"))
    palavras = _text(row.get("palavras_chave_sugeridas"))
    if acao in {"CLASSIFICAR", "CRIAR_ALIAS", "CRIAR_REGRA"} and not conta:
        return "Escolha uma conta para esta acao."
    if acao == "CRIAR_REGRA" and not palavras:
        return "Informe palavras-chave especificas para criar regra."
    return ""


def _risk_level(acao: str) -> str:
    if acao == "CRIAR_REGRA":
        return "ALTO"
    if acao == "CRIAR_ALIAS":
        return "BAIXO"
    if acao == "CLASSIFICAR":
        return "MEDIO"
    return "BAIXO"


def _style_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _apply_revisao_styles(sheet, headers: list[str]) -> None:
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    generated_fill = PatternFill("solid", fgColor="D9EAF7")
    technical_fill = PatternFill("solid", fgColor="E7E6E6")
    valor_total_index = _column_index(headers, "valor_total")

    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = _suggested_width(header)
        if header in EDITABLE_COLUMNS:
            fill = editable_fill
        elif header in CSV_HEADERS:
            fill = technical_fill
        else:
            fill = generated_fill
        sheet.cell(row=1, column=index).fill = fill
        sheet.cell(row=1, column=index).font = Font(bold=True, color="000000")

        for cell in sheet.iter_cols(
            min_col=index,
            max_col=index,
            min_row=2,
            max_row=max(sheet.max_row, 2),
        ):
            for item in cell:
                item.alignment = Alignment(wrap_text=True, vertical="top")
                item.fill = fill
                if index == valor_total_index:
                    item.alignment = Alignment(
                        horizontal="right", wrap_text=True, vertical="top"
                    )


def _apply_action_validation(sheet, headers: list[str], row_count: int) -> None:
    if "acao_manual" not in headers:
        return
    column_letter = get_column_letter(_column_index(headers, "acao_manual"))
    last_row = max(row_count + 1, 2)
    validation = DataValidation(
        type="list",
        formula1=f"={ACOES_VALIDAS_SHEET}!$A$2:$A${len(ALLOWED_ACTIONS_DESCRIPTIONS) + 1}",
        allow_blank=True,
    )
    validation.error = "Escolha uma acao manual valida."
    validation.errorTitle = "Acao manual invalida"
    validation.prompt = "Selecione uma acao manual valida."
    validation.promptTitle = "acao_manual"
    sheet.add_data_validation(validation)
    validation.add(f"{column_letter}2:{column_letter}{last_row}")


def _apply_account_validation(
    sheet, headers: list[str], row_count: int, account_count: int
) -> None:
    if "conta_escolhida" not in headers or account_count == 0:
        return
    column_letter = get_column_letter(_column_index(headers, "conta_escolhida"))
    last_row = max(row_count + 1, 2)
    validation = DataValidation(
        type="list",
        formula1=f"={PLANO_CONTAS_SHEET}!$E$2:$E${account_count + 1}",
        allow_blank=True,
    )
    validation.error = "Escolha uma conta existente na aba Plano_de_Contas."
    validation.errorTitle = "Conta invalida"
    validation.prompt = "Selecione uma conta do plano de contas."
    validation.promptTitle = "conta_escolhida"
    sheet.add_data_validation(validation)
    validation.add(f"{column_letter}2:{column_letter}{last_row}")


def _protect_revisao_sheet(sheet, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        locked = header not in EDITABLE_COLUMNS
        for cell in sheet.iter_cols(
            min_col=index,
            max_col=index,
            min_row=2,
            max_row=max(sheet.max_row, 2),
        ):
            for item in cell:
                item.protection = Protection(locked=locked)
    sheet.protection.sheet = True
    sheet.protection.enable()


def _set_filter_and_freeze(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _set_widths(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _suggested_width(header: str) -> int:
    widths = {
        "tipo_candidato": 24,
        "descricao_normalizada": 34,
        "descricao_exemplo": 42,
        "ocorrencias": 12,
        "fornecedores": 38,
        "documentos": 26,
        "valor_total": 16,
        "conta_atual_id": 16,
        "conta_atual_nome": 34,
        "status_revisao": 18,
        "alertas": 34,
        "acao_manual": 28,
        "conta_escolhida": 58,
        "conta_id_sugerida": 20,
        "conta_nome_sugerida": 34,
        "palavras_chave_sugeridas": 42,
        "observacao_manual": 44,
        "explicacao_acao": 54,
        "alerta_preenchimento": 38,
        "nivel_risco_regra": 18,
        "exemplo_decisao": 48,
    }
    return widths.get(header, max(14, min(len(header) + 4, 40)))


def _column_index(headers: list[str], name: str) -> int:
    return headers.index(name) + 1


def _format_account_choice(conta_id: str, conta_nome: str) -> str:
    conta_id = _text(conta_id)
    conta_nome = _text(conta_nome)
    if conta_id and conta_nome:
        return f"{conta_id} - {conta_nome}"
    if conta_id:
        return conta_id
    return ""


def _account_name_by_id(contas: list[dict[str, str]], conta_id: str) -> str:
    normalized_id = _text(conta_id)
    for conta in contas:
        if conta.get("conta_id") == normalized_id:
            return conta.get("conta_nome", "")
    return ""


def _account_level(conta_id: str) -> str:
    conta_id = _text(conta_id)
    if not conta_id:
        return ""
    return str(conta_id.count(".") + 1)


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()
