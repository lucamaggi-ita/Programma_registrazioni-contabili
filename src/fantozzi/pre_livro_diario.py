from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fantozzi.excel_writer import HEADERS, LANCAMENTOS_SHEET
from fantozzi.operation_classifier import (
    OperationAuditError,
    OperacaoEspecialAuditItem,
    auditar_operacoes_especiais,
)
from fantozzi.relatorios_gerenciais_ok import _is_bonificacao_recebida_decidida


LANCAMENTOS_OK_SHEET = "Lancamentos_OK"
LANCAMENTOS_REVISAR_SHEET = "Lancamentos_REVISAR"
OPERACOES_ESPECIAIS_SHEET = "Operacoes_Especiais"
RESUMO_CONTA_SHEET = "Resumo_por_Conta"
RESUMO_MENSAL_SHEET = "Resumo_Mensal"
AUDITORIA_SHEET = "Auditoria"
INSTRUCOES_SHEET = "Instrucoes"

SOURCE_HEADER = "arquivo_excel_origem"
OUTPUT_LANCAMENTO_HEADERS = [SOURCE_HEADER, *HEADERS]

OPERACOES_HEADERS = [
    "pasta_xml_origem",
    "tipo_operacao_especial",
    "confianca",
    "destino_gerencial_sugerido",
    "status_revisao_sugerido",
    "cnpj_destinatario",
    "fornecedor",
    "cnpj_fornecedor",
    "documento",
    "serie",
    "chave_nfe",
    "data_emissao",
    "natureza_operacao",
    "cfops_encontrados",
    "valor_itens_relevantes",
    "valor_total_nfe",
    "motivos",
    "alertas",
    "produtos_envolvidos",
]

RESUMO_CONTA_HEADERS = [
    "conta_id",
    "conta_nome",
    "grupo_conta",
    "qtd_lancamentos_ok",
    "valor_ok",
    "qtd_lancamentos_revisar",
    "valor_revisar",
    "qtd_lancamentos_total",
    "valor_total",
]

RESUMO_MENSAL_HEADERS = [
    "competencia",
    "qtd_lancamentos_ok",
    "valor_ok",
    "qtd_lancamentos_revisar",
    "valor_revisar",
    "qtd_lancamentos_total",
    "valor_total",
]

AUDITORIA_HEADERS = ["Indicador", "Valor", "Status", "Observacao"]


class PreLivroDiarioError(Exception):
    pass


@dataclass(frozen=True)
class AuditoriaIndicador:
    indicador: str
    valor: object
    status: str
    observacao: str


@dataclass(frozen=True)
class PreLivroDiarioResult:
    output_path: Path
    excel_paths: tuple[Path, ...]
    pastas_operacoes_especiais: tuple[Path, ...]
    lancamentos_ok: int
    lancamentos_revisar: int
    operacoes_especiais: int
    total_por_conta: int
    total_meses: int
    valor_total_lancamentos_ok: Decimal
    valor_total_lancamentos_revisar: Decimal
    valor_total_lancamentos_geral: Decimal
    valor_total_operacoes_especiais: Decimal
    status_geral_auditoria: str
    auditoria: tuple[AuditoriaIndicador, ...]


def gerar_pre_livro_diario_teste(
    excel_paths: Iterable[Path | str],
    output_path: Path | str,
    pastas_operacoes_especiais: Iterable[Path | str] | None = None,
    relatorio_auditoria_path: Path | str | None = None,
) -> PreLivroDiarioResult:
    excel_paths = tuple(Path(path) for path in excel_paths)
    output_path = Path(output_path)
    pastas_operacoes = tuple(Path(path) for path in (pastas_operacoes_especiais or ()))
    relatorio_auditoria_path = (
        Path(relatorio_auditoria_path) if relatorio_auditoria_path else None
    )

    if not excel_paths:
        raise PreLivroDiarioError("Informe ao menos um Excel de lancamentos.")

    lancamentos_ok: list[dict[str, object]] = []
    lancamentos_revisar: list[dict[str, object]] = []
    resumo_por_conta: dict[tuple[str, str, str], dict[str, object]] = {}
    resumo_mensal: dict[str, dict[str, object]] = {}

    for excel_path in excel_paths:
        for row in _read_lancamentos(excel_path):
            status = _text(row.get("status_revisao")).upper()
            is_bonificacao_decidida = _is_bonificacao_recebida_decidida(row)
            enriched = {SOURCE_HEADER: str(excel_path), **row}
            if status == "OK" or is_bonificacao_decidida:
                lancamentos_ok.append(enriched)
            else:
                lancamentos_revisar.append(enriched)
            _accumulate_resumo_por_conta(
                resumo_por_conta, row, force_ok=is_bonificacao_decidida
            )
            _accumulate_resumo_mensal(
                resumo_mensal, row, force_ok=is_bonificacao_decidida
            )

    operacoes_especiais = _read_operacoes_especiais(pastas_operacoes)
    resumo_conta_rows = _resumo_conta_rows(resumo_por_conta)
    resumo_mensal_rows = _resumo_mensal_rows(resumo_mensal)
    auditoria = _build_auditoria(
        lancamentos_ok,
        lancamentos_revisar,
        operacoes_especiais,
        resumo_conta_rows,
        resumo_mensal_rows,
    )

    workbook = Workbook()
    _write_lancamentos_sheet(
        workbook.active,
        LANCAMENTOS_OK_SHEET,
        lancamentos_ok,
    )
    _write_lancamentos_sheet(
        workbook.create_sheet(LANCAMENTOS_REVISAR_SHEET),
        LANCAMENTOS_REVISAR_SHEET,
        lancamentos_revisar,
    )
    _write_operacoes_sheet(
        workbook.create_sheet(OPERACOES_ESPECIAIS_SHEET),
        operacoes_especiais,
    )
    _write_resumo_sheet(
        workbook.create_sheet(RESUMO_CONTA_SHEET),
        RESUMO_CONTA_HEADERS,
        resumo_conta_rows,
    )
    _write_resumo_sheet(
        workbook.create_sheet(RESUMO_MENSAL_SHEET),
        RESUMO_MENSAL_HEADERS,
        resumo_mensal_rows,
    )
    _write_auditoria_sheet(
        workbook.create_sheet(AUDITORIA_SHEET),
        auditoria,
    )
    _write_instrucoes_sheet(
        workbook.create_sheet(INSTRUCOES_SHEET),
        excel_paths,
        pastas_operacoes,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    if relatorio_auditoria_path:
        salvar_relatorio_auditoria(relatorio_auditoria_path, auditoria)

    return PreLivroDiarioResult(
        output_path=output_path,
        excel_paths=excel_paths,
        pastas_operacoes_especiais=pastas_operacoes,
        lancamentos_ok=len(lancamentos_ok),
        lancamentos_revisar=len(lancamentos_revisar),
        operacoes_especiais=len(operacoes_especiais),
        total_por_conta=len(resumo_por_conta),
        total_meses=len(resumo_mensal),
        valor_total_lancamentos_ok=_sum_values(lancamentos_ok, "valor"),
        valor_total_lancamentos_revisar=_sum_values(lancamentos_revisar, "valor"),
        valor_total_lancamentos_geral=_sum_values(
            [*lancamentos_ok, *lancamentos_revisar], "valor"
        ),
        valor_total_operacoes_especiais=_sum_values(
            operacoes_especiais, "valor_total_nfe"
        ),
        status_geral_auditoria=_status_geral_auditoria(auditoria),
        auditoria=tuple(auditoria),
    )


def console_lines(result: PreLivroDiarioResult) -> list[str]:
    return [
        "[OK] Pre-livro diario gerencial de teste gerado.",
        f"     Excel de origem: {len(result.excel_paths)}",
        f"     Lancamentos OK: {result.lancamentos_ok}",
        f"     Lancamentos REVISAR: {result.lancamentos_revisar}",
        f"     Operacoes especiais: {result.operacoes_especiais}",
        f"     Resumo por conta: {result.total_por_conta}",
        f"     Resumo mensal: {result.total_meses}",
        f"     Status auditoria: {result.status_geral_auditoria}",
        f"     Arquivo gerado: {result.output_path}",
    ]


def salvar_relatorio_auditoria(
    report_path: Path | str,
    auditoria: Iterable[AuditoriaIndicador],
) -> None:
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        relatorio_auditoria_text(auditoria),
        encoding="utf-8",
        newline="\n",
    )


def relatorio_auditoria_text(auditoria: Iterable[AuditoriaIndicador]) -> str:
    lines = [
        "FANTOZZI - AUDITORIA DO PRE-LIVRO DIARIO GERENCIAL DE TESTE",
        "===========================================================",
        "",
        "Este relatorio e um controle preliminar. Nao e livro diario definitivo.",
        "",
        "Indicadores",
        "-----------",
    ]
    for item in auditoria:
        lines.append(
            f"{item.indicador};{_audit_value_text(item.valor)};{item.status};{item.observacao}"
        )
    return "\n".join(lines) + "\n"


def _read_lancamentos(excel_path: Path) -> list[dict[str, object]]:
    _validate_file(excel_path, "Excel de lancamentos")
    try:
        workbook = load_workbook(excel_path, data_only=True)
    except Exception as exc:
        raise PreLivroDiarioError(
            f"Excel de lancamentos nao pode ser lido: {excel_path}"
        ) from exc

    if LANCAMENTOS_SHEET not in workbook.sheetnames:
        raise PreLivroDiarioError(f"Aba obrigatoria ausente: {LANCAMENTOS_SHEET}")

    sheet = workbook[LANCAMENTOS_SHEET]
    headers = [cell.value for cell in sheet[1]]
    missing = [header for header in HEADERS if header not in headers]
    if missing:
        raise PreLivroDiarioError(
            "Colunas obrigatorias ausentes em Lancamentos: " + ", ".join(missing)
        )

    rows: list[dict[str, object]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(_has_value(value) for value in values):
            continue
        raw = dict(zip(headers, values))
        rows.append({header: raw.get(header) for header in HEADERS})
    return rows


def _read_operacoes_especiais(
    pastas: tuple[Path, ...],
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for pasta in pastas:
        try:
            result = auditar_operacoes_especiais(pasta)
        except OperationAuditError as exc:
            raise PreLivroDiarioError(str(exc)) from exc
        for item in result.itens:
            rows.append(_operacao_row(pasta, item))
    return rows


def _operacao_row(pasta: Path, item: OperacaoEspecialAuditItem) -> dict[str, object]:
    nfe = item.nfe
    classificacao = item.classificacao
    cfops = ", ".join(sorted({produto.cfop for produto in nfe.produtos if produto.cfop}))
    valor_itens = sum((produto.valor for produto in item.produtos), Decimal("0.00"))
    produtos = "; ".join(
        f"{produto.descricao} | CFOP {produto.cfop} | R$ {_money_text(produto.valor)}"
        for produto in item.produtos
    )
    return {
        "pasta_xml_origem": str(pasta),
        "tipo_operacao_especial": classificacao.tipo,
        "confianca": classificacao.confianca,
        "destino_gerencial_sugerido": classificacao.destino_gerencial,
        "status_revisao_sugerido": classificacao.status_revisao,
        "cnpj_destinatario": nfe.cnpj_destinatario,
        "fornecedor": nfe.nome_emitente,
        "cnpj_fornecedor": nfe.cnpj_emitente,
        "documento": nfe.numero,
        "serie": nfe.serie,
        "chave_nfe": nfe.chave,
        "data_emissao": nfe.data_emissao,
        "natureza_operacao": nfe.natureza_operacao,
        "cfops_encontrados": cfops,
        "valor_itens_relevantes": valor_itens.quantize(Decimal("0.01")),
        "valor_total_nfe": nfe.valor_total.quantize(Decimal("0.01")),
        "motivos": "; ".join(classificacao.motivos),
        "alertas": "; ".join(classificacao.alertas),
        "produtos_envolvidos": produtos,
    }


def _accumulate_resumo_por_conta(
    resumo: dict[tuple[str, str, str], dict[str, object]],
    row: dict[str, object],
    *,
    force_ok: bool = False,
) -> None:
    key = (
        _text(row.get("conta_id")),
        _text(row.get("conta_nome")),
        _text(row.get("grupo_conta")),
    )
    bucket = resumo.setdefault(
        key,
        {
            "conta_id": key[0],
            "conta_nome": key[1],
            "grupo_conta": key[2],
            "qtd_lancamentos_ok": 0,
            "valor_ok": Decimal("0.00"),
            "qtd_lancamentos_revisar": 0,
            "valor_revisar": Decimal("0.00"),
        },
    )
    _accumulate_status_bucket(bucket, row, force_ok=force_ok)


def _accumulate_resumo_mensal(
    resumo: dict[str, dict[str, object]],
    row: dict[str, object],
    *,
    force_ok: bool = False,
) -> None:
    competencia = _text(row.get("competencia")) or "SEM_COMPETENCIA"
    bucket = resumo.setdefault(
        competencia,
        {
            "competencia": competencia,
            "qtd_lancamentos_ok": 0,
            "valor_ok": Decimal("0.00"),
            "qtd_lancamentos_revisar": 0,
            "valor_revisar": Decimal("0.00"),
        },
    )
    _accumulate_status_bucket(bucket, row, force_ok=force_ok)


def _accumulate_status_bucket(
    bucket: dict[str, object],
    row: dict[str, object],
    *,
    force_ok: bool = False,
) -> None:
    value = _decimal_value(row.get("valor"))
    if force_ok or _text(row.get("status_revisao")).upper() == "OK":
        bucket["qtd_lancamentos_ok"] = int(bucket["qtd_lancamentos_ok"]) + 1
        bucket["valor_ok"] = bucket["valor_ok"] + value
    else:
        bucket["qtd_lancamentos_revisar"] = int(bucket["qtd_lancamentos_revisar"]) + 1
        bucket["valor_revisar"] = bucket["valor_revisar"] + value


def _resumo_conta_rows(
    resumo: dict[tuple[str, str, str], dict[str, object]]
) -> list[dict[str, object]]:
    rows = []
    for row in resumo.values():
        rows.append(_with_totals(row))
    return sorted(rows, key=lambda row: (_text(row["conta_id"]), _text(row["conta_nome"])))


def _resumo_mensal_rows(resumo: dict[str, dict[str, object]]) -> list[dict[str, object]]:
    rows = []
    for row in resumo.values():
        rows.append(_with_totals(row))
    return sorted(rows, key=lambda row: _text(row["competencia"]))


def _with_totals(row: dict[str, object]) -> dict[str, object]:
    qtd_ok = int(row["qtd_lancamentos_ok"])
    qtd_revisar = int(row["qtd_lancamentos_revisar"])
    valor_ok = row["valor_ok"]
    valor_revisar = row["valor_revisar"]
    return {
        **row,
        "qtd_lancamentos_total": qtd_ok + qtd_revisar,
        "valor_total": valor_ok + valor_revisar,
    }


def _build_auditoria(
    lancamentos_ok: list[dict[str, object]],
    lancamentos_revisar: list[dict[str, object]],
    operacoes_especiais: list[dict[str, object]],
    resumo_conta_rows: list[dict[str, object]],
    resumo_mensal_rows: list[dict[str, object]],
) -> list[AuditoriaIndicador]:
    total_ok = len(lancamentos_ok)
    total_revisar = len(lancamentos_revisar)
    total_geral = total_ok + total_revisar
    total_operacoes = len(operacoes_especiais)
    valor_ok = _sum_values(lancamentos_ok, "valor")
    valor_revisar = _sum_values(lancamentos_revisar, "valor")
    valor_geral = valor_ok + valor_revisar
    valor_operacoes = _sum_values(operacoes_especiais, "valor_total_nfe")
    resumo_conta_total = _sum_int_values(resumo_conta_rows, "qtd_lancamentos_total")
    resumo_conta_valor = _sum_values(resumo_conta_rows, "valor_total")
    resumo_mensal_total = _sum_int_values(resumo_mensal_rows, "qtd_lancamentos_total")
    resumo_mensal_valor = _sum_values(resumo_mensal_rows, "valor_total")
    count_status = "OK" if total_ok + total_revisar == total_geral else "ERRO"
    conta_status = (
        "OK"
        if resumo_conta_total == total_geral and resumo_conta_valor == valor_geral
        else "ERRO"
    )
    mensal_status = (
        "OK"
        if resumo_mensal_total == total_geral and resumo_mensal_valor == valor_geral
        else "ERRO"
    )
    operacoes_status = "ATENCAO_CONTABIL" if total_operacoes else "NAO_APLICAVEL"
    revisao_status = "ATENCAO" if total_revisar else "OK"
    status_geral = _derive_status_geral(
        total_revisar=total_revisar,
        total_operacoes=total_operacoes,
        has_error=count_status == "ERRO" or conta_status == "ERRO" or mensal_status == "ERRO",
    )

    return [
        AuditoriaIndicador(
            "total_lancamentos_ok",
            total_ok,
            "OK",
            "Lancamentos com status_revisao OK e BONIFICACAO_RECEBIDA decidida em 2.05.003.",
        ),
        AuditoriaIndicador(
            "total_lancamentos_revisar",
            total_revisar,
            revisao_status,
            "Lancamentos que exigem revisao manual antes de qualquer livro definitivo.",
        ),
        AuditoriaIndicador(
            "total_lancamentos_geral",
            total_geral,
            count_status,
            "Total calculado como OK + REVISAR.",
        ),
        AuditoriaIndicador(
            "total_operacoes_especiais",
            total_operacoes,
            operacoes_status,
            "Operacoes especiais ficam separadas dos lancamentos normais.",
        ),
        AuditoriaIndicador(
            "total_resumo_por_conta",
            len(resumo_conta_rows),
            conta_status,
            "Resumo_por_Conta confere quantidade e valor total dos lancamentos.",
        ),
        AuditoriaIndicador(
            "total_resumo_mensal",
            len(resumo_mensal_rows),
            mensal_status,
            "Resumo_Mensal confere quantidade e valor total dos lancamentos.",
        ),
        AuditoriaIndicador(
            "valor_total_lancamentos_ok",
            valor_ok,
            "OK",
            "Soma dos valores dos lancamentos OK.",
        ),
        AuditoriaIndicador(
            "valor_total_lancamentos_revisar",
            valor_revisar,
            revisao_status,
            "Soma dos valores dos lancamentos em revisao.",
        ),
        AuditoriaIndicador(
            "valor_total_lancamentos_geral",
            valor_geral,
            "OK" if conta_status == "OK" and mensal_status == "OK" else "ERRO",
            "Soma dos lancamentos OK e REVISAR usados nos resumos.",
        ),
        AuditoriaIndicador(
            "valor_total_operacoes_especiais",
            valor_operacoes,
            operacoes_status,
            "Soma do valor total das NF-e classificadas como operacoes especiais.",
        ),
        AuditoriaIndicador(
            "status_geral_auditoria",
            status_geral,
            status_geral,
            "Status consolidado da auditoria do pre-livro diario de teste.",
        ),
    ]


def _derive_status_geral(
    *,
    total_revisar: int,
    total_operacoes: int,
    has_error: bool,
) -> str:
    if has_error:
        return "ERRO"
    if total_operacoes:
        return "ATENCAO_CONTABIL"
    if total_revisar:
        return "ATENCAO"
    return "OK"


def _status_geral_auditoria(auditoria: Iterable[AuditoriaIndicador]) -> str:
    for item in auditoria:
        if item.indicador == "status_geral_auditoria":
            return _text(item.valor)
    return "ERRO"


def _sum_values(rows: Iterable[dict[str, object]], field: str) -> Decimal:
    total = Decimal("0.00")
    for row in rows:
        total += _decimal_value(row.get(field))
    return total.quantize(Decimal("0.01"))


def _sum_int_values(rows: Iterable[dict[str, object]], field: str) -> int:
    total = 0
    for row in rows:
        try:
            total += int(row.get(field) or 0)
        except (TypeError, ValueError):
            continue
    return total


def _write_lancamentos_sheet(
    sheet,
    title: str,
    rows: list[dict[str, object]],
) -> None:
    sheet.title = title
    sheet.append(OUTPUT_LANCAMENTO_HEADERS)
    for row in rows:
        sheet.append([row.get(header) for header in OUTPUT_LANCAMENTO_HEADERS])
    _style_table(sheet)
    _set_widths_for_headers(sheet, OUTPUT_LANCAMENTO_HEADERS)


def _write_operacoes_sheet(sheet, rows: list[dict[str, object]]) -> None:
    sheet.append(OPERACOES_HEADERS)
    for row in rows:
        sheet.append([row.get(header) for header in OPERACOES_HEADERS])
    _style_table(sheet)
    _set_widths_for_headers(sheet, OPERACOES_HEADERS)


def _write_resumo_sheet(
    sheet,
    headers: list[str],
    rows: list[dict[str, object]],
) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    _style_table(sheet)
    _set_widths_for_headers(sheet, headers)


def _write_auditoria_sheet(
    sheet,
    auditoria: list[AuditoriaIndicador],
) -> None:
    sheet.append(AUDITORIA_HEADERS)
    for item in auditoria:
        sheet.append(
            [
                item.indicador,
                item.valor,
                item.status,
                item.observacao,
            ]
        )
    _style_table(sheet)
    _set_widths_for_headers(sheet, AUDITORIA_HEADERS)


def _write_instrucoes_sheet(
    sheet,
    excel_paths: tuple[Path, ...],
    pastas_operacoes: tuple[Path, ...],
) -> None:
    lines = [
        "PRE-LIVRO DIARIO GERENCIAL DE TESTE",
        "",
        "Este arquivo e um relatorio de controle, nao um livro diario definitivo.",
        "Nenhuma regra nova foi aplicada.",
        "Nenhum config/*.yaml foi modificado.",
        "Nenhum CSV, XLSX assistido ou Excel original foi modificado.",
        "",
        "Abas:",
        "- Lancamentos_OK: lancamentos com status_revisao OK e BONIFICACAO_RECEBIDA decidida em 2.05.003.",
        "- Lancamentos_REVISAR: lancamentos que exigem revisao humana.",
        "- Operacoes_Especiais: comodato, bonificacao e outros casos auditados a partir dos XML informados.",
        "- Resumo_por_Conta: totais por conta, separados entre OK e REVISAR.",
        "- Resumo_Mensal: totais por competencia, separados entre OK e REVISAR.",
        "- Auditoria: indicadores de coerencia calculados em Python para controle preliminar.",
        "",
        "Observacao importante:",
        "- Este arquivo e um pre-livro diario gerencial de teste.",
        "- Nao e um livro diario definitivo.",
        "- Lancamentos_REVISAR exigem revisao manual.",
        "- Operacoes_Especiais nao devem ser tratadas automaticamente como compra normal.",
        "- COMODATO deve permanecer fora da DRE ate validacao contabil.",
        "- BONIFICACAO_RECEBIDA decidida em 2.05.003 entra como reducao de custo.",
        "- A aba Auditoria serve somente para controle de coerencia.",
        "- Os resumos sao calculados a partir dos Excel originais de lancamentos.",
        "- Operacoes especiais aparecem em aba separada para controle e nao foram excluidas automaticamente dos resumos.",
        "- Antes de qualquer DRE ou livro definitivo, validar o tratamento das operacoes especiais com o contador.",
        "",
        "Arquivos Excel de origem:",
    ]
    lines.extend(f"- {path}" for path in excel_paths)
    lines.extend(["", "Pastas XML auditadas para operacoes especiais:"])
    if pastas_operacoes:
        lines.extend(f"- {path}" for path in pastas_operacoes)
    else:
        lines.append("- Nenhuma pasta XML informada.")
    lines.extend(
        [
            "",
            "Uso recomendado:",
            "- Usar este arquivo para conferencia operacional e revisao gerencial.",
            "- Nao usar como livro diario definitivo.",
            "- Revisar todos os itens em Lancamentos_REVISAR e Operacoes_Especiais antes de qualquer DRE definitiva.",
        ]
    )

    for line in lines:
        sheet.append([line])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 120
    for row in sheet.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _style_table(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, Decimal):
                cell.number_format = '#,##0.00'


def _set_widths_for_headers(sheet, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = _suggested_width(header)


def _suggested_width(header: str) -> int:
    widths = {
        SOURCE_HEADER: 42,
        "data_documento": 16,
        "competencia": 14,
        "cnpj_empresa": 20,
        "nome_empresa": 34,
        "fornecedor_ou_favorecido": 42,
        "cnpj_cpf_favorecido": 22,
        "documento": 16,
        "conta_id": 16,
        "conta_nome": 36,
        "grupo_conta": 36,
        "descricao_itens": 52,
        "valor": 16,
        "arquivo_origem": 34,
        "chave_origem": 46,
        "hash_registro": 24,
        "status_revisao": 18,
        "alertas": 40,
        "observacoes": 42,
        "duplicatas_json": 40,
        "pasta_xml_origem": 42,
        "tipo_operacao_especial": 26,
        "destino_gerencial_sugerido": 26,
        "status_revisao_sugerido": 24,
        "natureza_operacao": 34,
        "produtos_envolvidos": 60,
        "motivos": 42,
        "Indicador": 34,
        "Valor": 22,
        "Status": 20,
        "Observacao": 76,
    }
    return widths.get(header, max(14, min(len(header) + 4, 40)))


def _validate_file(path: Path, label: str) -> None:
    if not path.exists():
        raise PreLivroDiarioError(f"{label} nao encontrado: {path}")
    if not path.is_file():
        raise PreLivroDiarioError(f"{label} nao e um arquivo: {path}")


def _decimal_value(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0.00")
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _money_text(value: Decimal) -> str:
    return f"{value.quantize(Decimal('0.01')):.2f}".replace(".", ",")


def _audit_value_text(value: object) -> str:
    if isinstance(value, Decimal):
        return _money_text(value)
    return _text(value)


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _has_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""
