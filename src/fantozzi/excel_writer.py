from __future__ import annotations

from dataclasses import asdict
from datetime import date
from decimal import Decimal
import hashlib
import json
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from fantozzi.models import Duplicata, Lancamento


LANCAMENTOS_SHEET = "Lancamentos"
REVISAO_SHEET = "Revisao"

HEADERS = [
    "data_documento",
    "competencia",
    "cnpj_empresa",
    "nome_empresa",
    "fornecedor_ou_favorecido",
    "cnpj_cpf_favorecido",
    "documento",
    "tipo_documento",
    "conta_id",
    "conta_nome",
    "grupo_conta",
    "descricao_itens",
    "valor",
    "origem",
    "arquivo_origem",
    "chave_origem",
    "hash_registro",
    "status_revisao",
    "alertas",
    "observacoes",
    "centro_custo",
    "cfop",
    "natureza_operacao",
    "tipo_operacao_especial",
    "confianca_operacao_especial",
    "tratamento_operacao_especial",
    "tipo_lancamento",
    "dre_linha_id",
    "dre_linha_nome",
    "data_pagamento",
    "duplicatas_json",
]


def write_lancamentos_to_excel(
    excel_path: Path | str, lancamentos: Iterable[Lancamento]
) -> None:
    """Append new Lancamento rows to an .xlsx file without duplicating them."""
    excel_path = Path(excel_path)
    lancamentos = list(lancamentos)

    workbook = _load_or_create_workbook(excel_path)
    lancamentos_sheet = _ensure_sheet(workbook, LANCAMENTOS_SHEET)
    revisao_sheet = _ensure_sheet(workbook, REVISAO_SHEET)

    _ensure_headers(lancamentos_sheet)
    _ensure_headers(revisao_sheet)

    existing_lancamentos = _existing_dedupe_keys(lancamentos_sheet)
    existing_revisao = _existing_dedupe_keys(revisao_sheet)

    for lancamento in lancamentos:
        key = _dedupe_key(lancamento)
        if key not in existing_lancamentos:
            lancamentos_sheet.append(_row_values(lancamento))
            existing_lancamentos.add(key)

        if lancamento.status_revisao != "OK" and key not in existing_revisao:
            revisao_sheet.append(_row_values(lancamento))
            existing_revisao.add(key)

    excel_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(excel_path)


def _load_or_create_workbook(excel_path: Path) -> Workbook:
    if excel_path.exists():
        return load_workbook(excel_path)

    workbook = Workbook()
    workbook.active.title = LANCAMENTOS_SHEET
    return workbook


def _ensure_sheet(workbook: Workbook, sheet_name: str) -> Worksheet:
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return workbook.create_sheet(sheet_name)


def _ensure_headers(sheet: Worksheet) -> None:
    if sheet.max_row == 1 and all(cell.value is None for cell in sheet[1]):
        for column, header in enumerate(HEADERS, start=1):
            sheet.cell(row=1, column=column, value=header)
        return

    existing_headers = [cell.value for cell in sheet[1]]
    if existing_headers[: len(HEADERS)] == HEADERS:
        return

    if existing_headers and existing_headers[0] == HEADERS[0]:
        missing_headers = [header for header in HEADERS if header not in existing_headers]
        next_column = len(existing_headers) + 1
        for offset, header in enumerate(missing_headers):
            sheet.cell(row=1, column=next_column + offset, value=header)
        return

    if sheet.max_row == 0:
        sheet.append(HEADERS)


def _existing_dedupe_keys(sheet: Worksheet) -> set[str]:
    headers = [cell.value for cell in sheet[1]]
    if not headers or headers[0] != HEADERS[0]:
        return set()

    indexes = {header: index for index, header in enumerate(headers)}
    keys: set[str] = set()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        origem = _cell_value(row, indexes, "origem")
        if origem == "MANUAL":
            continue

        hash_registro = _cell_value(row, indexes, "hash_registro")
        if hash_registro:
            keys.add(str(hash_registro))
            continue

        fallback = _fallback_key(
            origem=_cell_value(row, indexes, "origem"),
            chave_origem=_cell_value(row, indexes, "chave_origem"),
            conta_id=_cell_value(row, indexes, "conta_id"),
            valor=_cell_value(row, indexes, "valor"),
        )
        if fallback:
            keys.add(fallback)

    return keys


def _cell_value(row: tuple[object, ...], indexes: dict[str, int], header: str) -> object:
    index = indexes.get(header)
    if index is None or index >= len(row):
        return None
    return row[index]


def _dedupe_key(lancamento: Lancamento) -> str:
    if lancamento.hash_registro:
        return lancamento.hash_registro
    return _fallback_key(
        origem=lancamento.origem,
        chave_origem=lancamento.chave_origem,
        conta_id=lancamento.conta_id,
        valor=lancamento.valor,
    )


def _fallback_key(
    *, origem: object, chave_origem: object, conta_id: object, valor: object
) -> str:
    payload = "|".join(
        [
            str(origem or ""),
            str(chave_origem or ""),
            str(conta_id or ""),
            _normalize_decimal_text(valor),
        ]
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _row_values(lancamento: Lancamento) -> list[object]:
    values = asdict(lancamento)
    values["duplicatas_json"] = _duplicatas_json(lancamento.duplicatas)
    return [_excel_value(values.get(header)) for header in HEADERS]


def _duplicatas_json(duplicatas: list[Duplicata]) -> str:
    payload = [
        {
            "numero": duplicata.numero,
            "vencimento": duplicata.vencimento.isoformat()
            if duplicata.vencimento
            else None,
            "valor": _normalize_decimal_text(duplicata.valor),
        }
        for duplicata in duplicatas
    ]
    return json.dumps(payload, ensure_ascii=False, sort_keys=True)


def _excel_value(value: object) -> object:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, date):
        return value
    return value


def _normalize_decimal_text(value: object) -> str:
    if isinstance(value, Decimal):
        return f"{value.quantize(Decimal('0.01')):.2f}"
    if isinstance(value, float):
        return f"{Decimal(str(value)).quantize(Decimal('0.01')):.2f}"
    if isinstance(value, int):
        return f"{Decimal(value).quantize(Decimal('0.01')):.2f}"
    if value is None:
        return ""

    text = str(value)
    try:
        return f"{Decimal(text).quantize(Decimal('0.01')):.2f}"
    except Exception:
        return text
