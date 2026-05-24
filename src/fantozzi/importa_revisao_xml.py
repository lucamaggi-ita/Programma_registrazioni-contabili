from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fantozzi.csv_utils import sniff_delimiter


EXCEL_NAME = "Analisi_regole_registrazioni_xml_prodotti.xlsx"
MANIFEST_NAME = "manifest_importa_revisao_xml.json"

DECISOES_SHEET = "decisoes_xml"
TECNICO_SHEET = "tecnico_xml"
PLANO_CONTAS_SHEET = "plano_contas"
LISTAS_SHEET = "listas_validacao"
INSTRUCOES_SHEET = "instrucoes"

DECISOES_HEADERS = [
    "CHAVE_ESTAVEL",
    "STATUS_DECISAO",
    "ACAO_REGRA",
    "DESCRICAO_EXEMPLO",
    "DESCRICAO_NORMALIZADA",
    "FORNECEDORES",
    "TIPO_OPERACAO_ESPECIAL",
    "CONTA_ESCOLHIDA",
    "CONTA_ID_APROVADA",
    "CONTA_NOME_APROVADA",
    "PALAVRAS_CHAVE_SUGERIDAS",
    "NIVEL_RISCO_REGRA",
    "QTD_OCORRENCIAS",
    "VALOR_TOTAL",
    "OBSERVACAO_OPERADOR",
]

TECNICO_HEADERS = [
    "CHAVE_ESTAVEL",
    "RUN_ID",
    "CNPJ",
    "ORIGEM_CSV",
    "TIPO_CANDIDATO",
    "DOCUMENTOS",
    "CONTA_ATUAL_ID",
    "CONTA_ATUAL_NOME",
    "STATUS_REVISAO",
    "ALERTAS",
    "CFOP",
    "NATUREZA_OPERACAO",
    "TIPO_OPERACAO_ESPECIAL",
    "CONFIANCA_OPERACAO_ESPECIAL",
    "TRATAMENTO_OPERACAO_ESPECIAL",
    "RISCO_FORNECEDOR_BROAD",
    "RISCO_ALIAS_GENERICO",
    "MOTIVO_FILTRO",
    "QTD_OCORRENCIAS",
    "VALOR_TOTAL",
]

MANUAL_COLUMNS = {
    "STATUS_DECISAO",
    "ACAO_REGRA",
    "CONTA_ESCOLHIDA",
    "PALAVRAS_CHAVE_SUGERIDAS",
    "OBSERVACAO_OPERADOR",
}

SPECIAL_TYPES = {"COMODATO", "BONIFICACAO_RECEBIDA"}
GENERIC_SUPPLIERS = {
    "",
    "SEM FORNECEDOR",
    "FORNECEDOR",
    "FORNECEDORES",
    "DIVERSOS",
    "VARIOS",
    "NAO INFORMADO",
    "N/A",
}
GENERIC_SUPPLIERS_NORMALIZED = frozenset(
    {
        "",
        "SEM FORNECEDOR",
        "FORNECEDOR",
        "FORNECEDORES",
        "DIVERSOS",
        "VARIOS",
        "NAO INFORMADO",
        "N A",
    }
)


class ImportaRevisaoXmlError(ValueError):
    pass


@dataclass(frozen=True)
class ImportaRevisaoXmlResult:
    run_id_origem: str
    run_reports_dir: Path
    cnpjs_processados: tuple[str, ...]
    csvs_lidos: tuple[str, ...]
    excel_path: Path
    manifest_path: Path
    hash_yaml_run: str
    hash_yaml_atual: str
    hash_yaml_run_assente: bool
    stale_run: bool
    allow_stale_run: bool
    linhas_operacionais_geradas: int
    decisoes_preservadas: int
    linhas_novas: int
    comodato_forte_esclusi: int
    bonificacao_forte_esclusa: int
    candidatos_bloqueados_operacao_especial: int
    avisos: tuple[str, ...]
    bloqueios: tuple[str, ...]


@dataclass
class _Group:
    chave_estavel: str
    rows: list[dict[str, str]]
    cnpjs: set[str]
    csv_paths: set[str]


def resolve_run_reports_dir(
    *,
    run_reports_dir: Path | None = None,
    run_id: str | None = None,
    reports_runs_root: Path,
) -> Path:
    if run_reports_dir is not None:
        if run_id:
            raise ImportaRevisaoXmlError(
                "Use --run-reports-dir ou --run-id, nao ambos."
            )
        return Path(run_reports_dir)
    if not run_id:
        raise ImportaRevisaoXmlError("Informe --run-reports-dir ou --run-id.")
    return Path(reports_runs_root) / str(run_id).strip()


def importar_revisao_xml(
    *,
    run_reports_dir: Path,
    output_dir: Path,
    regras_path: Path,
    plano_contas_path: Path,
    allow_stale_run: bool = False,
) -> ImportaRevisaoXmlResult:
    run_reports_dir = Path(run_reports_dir)
    output_dir = Path(output_dir)
    regras_path = Path(regras_path)
    plano_contas_path = Path(plano_contas_path)

    if not run_reports_dir.exists() or not run_reports_dir.is_dir():
        raise ImportaRevisaoXmlError(f"Run reports dir nao encontrada: {run_reports_dir}")
    manifest_run_path = run_reports_dir / "manifest_run.json"
    if not manifest_run_path.exists():
        raise ImportaRevisaoXmlError(f"manifest_run.json nao encontrado: {manifest_run_path}")
    if not regras_path.exists():
        raise ImportaRevisaoXmlError(f"YAML de regras nao encontrado: {regras_path}")
    if not plano_contas_path.exists():
        raise ImportaRevisaoXmlError(
            f"Arquivo de plano de contas nao encontrado: {plano_contas_path}"
        )

    run_manifest = json.loads(manifest_run_path.read_text(encoding="utf-8"))
    run_id = str(run_manifest.get("run_id") or run_reports_dir.name)
    hash_yaml_run = _run_rules_hash(run_manifest)
    hash_yaml_atual = _hash_file(regras_path)
    hash_yaml_run_assente = not bool(hash_yaml_run)
    stale_run = hash_yaml_run_assente or bool(hash_yaml_run != hash_yaml_atual)
    bloqueios: list[str] = []
    avisos: list[str] = []
    if hash_yaml_run_assente:
        message = _missing_hash_message(run_id)
        if not allow_stale_run:
            raise ImportaRevisaoXmlError(message)
        avisos.append(
            "Run sem hash YAML aceita com --allow-stale-run; nao usar para operacao."
        )
        bloqueios.append(message)
    elif stale_run:
        message = _stale_hash_message(run_id, hash_yaml_run, hash_yaml_atual)
        if not allow_stale_run:
            raise ImportaRevisaoXmlError(message)
        avisos.append(
            "Run historica aceita com --allow-stale-run; nao usar para operacao."
        )
        bloqueios.append(message)

    account_options = _load_account_options(plano_contas_path)
    output_dir.mkdir(parents=True, exist_ok=True)
    excel_path = output_dir / EXCEL_NAME
    manifest_path = output_dir / MANIFEST_NAME
    existing = _load_existing_decisions(excel_path)

    groups: dict[str, _Group] = {}
    csvs_lidos: list[str] = []
    cnpjs: set[str] = set()
    comodato_forte = 0
    bonificacao_forte = 0
    bloqueados_especiais = 0
    tecnico_rows: list[dict[str, object]] = []

    for csv_path in _csv_paths(run_reports_dir):
        csvs_lidos.append(str(csv_path))
        cnpj = csv_path.parent.name
        cnpjs.add(cnpj)
        for source_row in _read_csv(csv_path):
            row = _normalized_csv_row(source_row)
            status = _token(row.get("status_revisao"))
            special_type = _token(row.get("tipo_operacao_especial"))
            confidence = _token(row.get("confianca_operacao_especial"))
            key = _stable_key(row)

            if status == "OK":
                tecnico_rows.append(
                    _technical_row(
                        row,
                        key,
                        run_id=run_id,
                        cnpj=cnpj,
                        csv_path=csv_path,
                        motivo_filtro="STATUS_OK",
                    )
                )
                continue

            if special_type == "COMODATO" and confidence == "FORTE":
                comodato_forte += 1
                tecnico_rows.append(
                    _technical_row(
                        row,
                        key,
                        run_id=run_id,
                        cnpj=cnpj,
                        csv_path=csv_path,
                        motivo_filtro="COMODATO_FORTE_EXCLUIDO",
                    )
                )
                continue

            if special_type == "BONIFICACAO_RECEBIDA" and confidence == "FORTE":
                bonificacao_forte += 1
                tecnico_rows.append(
                    _technical_row(
                        row,
                        key,
                        run_id=run_id,
                        cnpj=cnpj,
                        csv_path=csv_path,
                        motivo_filtro="BONIFICACAO_FORTE_EXCLUIDA",
                    )
                )
                continue

            if special_type in SPECIAL_TYPES:
                bloqueados_especiais += 1

            group = groups.setdefault(
                key,
                _Group(
                    chave_estavel=key,
                    rows=[],
                    cnpjs=set(),
                    csv_paths=set(),
                ),
            )
            group.rows.append(row)
            group.cnpjs.add(cnpj)
            group.csv_paths.add(str(csv_path))
            tecnico_rows.append(
                _technical_row(
                    row,
                    key,
                    run_id=run_id,
                    cnpj=cnpj,
                    csv_path=csv_path,
                    motivo_filtro="OPERACIONAL",
                )
            )

    decision_rows: list[dict[str, object]] = []
    preserved_count = 0
    new_count = 0
    for key in sorted(groups):
        group = groups[key]
        decision = _decision_row(group)
        preserved = existing.get(key)
        if preserved:
            _preserve_manual(decision, preserved)
            preserved_count += 1
        else:
            new_count += 1
        decision_rows.append(decision)

    _write_workbook(
        excel_path,
        decision_rows=decision_rows,
        technical_rows=tecnico_rows,
        account_options=account_options,
    )

    result = ImportaRevisaoXmlResult(
        run_id_origem=run_id,
        run_reports_dir=run_reports_dir,
        cnpjs_processados=tuple(sorted(cnpjs)),
        csvs_lidos=tuple(csvs_lidos),
        excel_path=excel_path,
        manifest_path=manifest_path,
        hash_yaml_run=hash_yaml_run,
        hash_yaml_atual=hash_yaml_atual,
        hash_yaml_run_assente=hash_yaml_run_assente,
        stale_run=stale_run,
        allow_stale_run=allow_stale_run,
        linhas_operacionais_geradas=len(decision_rows),
        decisoes_preservadas=preserved_count,
        linhas_novas=new_count,
        comodato_forte_esclusi=comodato_forte,
        bonificacao_forte_esclusa=bonificacao_forte,
        candidatos_bloqueados_operacao_especial=bloqueados_especiais,
        avisos=tuple(avisos),
        bloqueios=tuple(bloqueios),
    )
    _write_manifest(result)
    return result


def console_lines(result: ImportaRevisaoXmlResult) -> list[str]:
    return [
        "[OK] Excel decisional XML prodotti gerado/atualizado.",
        f"Run origem: {result.run_id_origem}",
        f"CNPJs: {', '.join(result.cnpjs_processados)}",
        f"CSVs lidos: {len(result.csvs_lidos)}",
        f"Linhas operacionais: {result.linhas_operacionais_geradas}",
        f"Decisoes preservadas: {result.decisoes_preservadas}",
        f"Linhas novas: {result.linhas_novas}",
        f"COMODATO FORTE excluidos: {result.comodato_forte_esclusi}",
        f"BONIFICACAO FORTE excluidas: {result.bonificacao_forte_esclusa}",
        f"Operacoes especiais bloqueadas: {result.candidatos_bloqueados_operacao_especial}",
        f"Excel: {result.excel_path}",
        f"Manifest: {result.manifest_path}",
    ]


def _csv_paths(run_reports_dir: Path) -> list[Path]:
    paths = sorted(run_reports_dir.glob("*/revisao_classificacao_real.csv"))
    if not paths:
        raise ImportaRevisaoXmlError(
            f"Nenhum revisao_classificacao_real.csv encontrado em: {run_reports_dir}"
        )
    return paths


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        sample = file.read(4096)
        file.seek(0)
        reader = csv.DictReader(file, delimiter=sniff_delimiter(sample))
        return [
            {str(key or "").strip(): str(value or "").strip() for key, value in row.items()}
            for row in reader
        ]


def _normalized_csv_row(row: dict[str, str]) -> dict[str, str]:
    return {_normalize_header(key): value.strip() for key, value in row.items()}


def _normalize_header(value: str) -> str:
    return value.strip().lower()


def _decision_row(group: _Group) -> dict[str, object]:
    rows = group.rows
    first = rows[0]
    special_type = _joined_unique(
        source_row.get("tipo_operacao_especial") for source_row in rows
    )
    blocked_special = any(
        _token(source_row.get("tipo_operacao_especial")) in SPECIAL_TYPES
        for source_row in rows
    )
    decision = {
        "CHAVE_ESTAVEL": group.chave_estavel,
        "STATUS_DECISAO": "BLOQUEADO" if blocked_special else "PENDENTE",
        "ACAO_REGRA": "",
        "DESCRICAO_EXEMPLO": _joined_unique(
            source_row.get("descricao_exemplo") for source_row in rows
        ),
        "DESCRICAO_NORMALIZADA": first.get("descricao_normalizada", ""),
        "FORNECEDORES": _joined_unique(
            source_row.get("fornecedores") for source_row in rows
        ),
        "TIPO_OPERACAO_ESPECIAL": special_type,
        "CONTA_ESCOLHIDA": "",
        "CONTA_ID_APROVADA": "",
        "CONTA_NOME_APROVADA": "",
        "PALAVRAS_CHAVE_SUGERIDAS": _keywords(first.get("descricao_normalizada", "")),
        "NIVEL_RISCO_REGRA": "ALTO" if blocked_special else _risk_level(rows),
        "QTD_OCORRENCIAS": sum(
            _parse_int(source_row.get("ocorrencias")) for source_row in rows
        ),
        "VALOR_TOTAL": sum(
            (_parse_decimal(source_row.get("valor_total")) for source_row in rows),
            Decimal("0"),
        ),
        "OBSERVACAO_OPERADOR": "",
    }
    if blocked_special:
        decision["OBSERVACAO_OPERADOR"] = (
            "Operacao especial nao FORTE: revisar manualmente; nao aprovar "
            "CRIAR_ALIAS/CRIAR_REGRA sem validacao contabil."
        )
    return decision


def _technical_row(
    row: dict[str, str],
    key: str,
    *,
    run_id: str,
    cnpj: str,
    csv_path: Path,
    motivo_filtro: str,
) -> dict[str, object]:
    return {
        "CHAVE_ESTAVEL": key,
        "RUN_ID": run_id,
        "CNPJ": cnpj,
        "ORIGEM_CSV": str(csv_path),
        "TIPO_CANDIDATO": row.get("tipo_candidato", ""),
        "DOCUMENTOS": row.get("documentos", ""),
        "CONTA_ATUAL_ID": row.get("conta_atual_id", ""),
        "CONTA_ATUAL_NOME": row.get("conta_atual_nome", ""),
        "STATUS_REVISAO": row.get("status_revisao", ""),
        "ALERTAS": row.get("alertas", ""),
        "CFOP": row.get("cfop", ""),
        "NATUREZA_OPERACAO": row.get("natureza_operacao", ""),
        "TIPO_OPERACAO_ESPECIAL": row.get("tipo_operacao_especial", ""),
        "CONFIANCA_OPERACAO_ESPECIAL": row.get("confianca_operacao_especial", ""),
        "TRATAMENTO_OPERACAO_ESPECIAL": row.get("tratamento_operacao_especial", ""),
        "RISCO_FORNECEDOR_BROAD": _supplier_broad_risk(row),
        "RISCO_ALIAS_GENERICO": _alias_generic_risk(row),
        "MOTIVO_FILTRO": motivo_filtro,
        "QTD_OCORRENCIAS": _parse_int(row.get("ocorrencias")),
        "VALOR_TOTAL": _parse_decimal(row.get("valor_total")),
    }


def _stable_key(row: dict[str, str]) -> str:
    supplier = _normalize_key(row.get("fornecedores", "")) or "SEM_FORNECEDOR"
    description = _normalize_key(row.get("descricao_normalizada", "")) or "SEM_DESCRICAO"
    return f"XML_PRODUTO|FORNECEDOR={supplier}|DESCRICAO={description}"


def _supplier_broad_risk(row: dict[str, str]) -> str:
    supplier = _normalize_key(row.get("fornecedores", ""))
    if _token(row.get("status_revisao")) == "VALOR_INCONSISTENTE":
        return "SI"
    if not supplier or supplier in GENERIC_SUPPLIERS_NORMALIZED:
        return "SI"
    return "NAO"


def _alias_generic_risk(row: dict[str, str]) -> str:
    description = _normalize_key(row.get("descricao_normalizada", ""))
    tokens = [token for token in re.split(r"\s+", description) if token]
    return "SI" if len(tokens) <= 3 else "NAO"


def _risk_level(rows: Iterable[dict[str, str]]) -> str:
    for row in rows:
        if _supplier_broad_risk(row) == "SI" or _alias_generic_risk(row) == "SI":
            return "MEDIO"
    return "BAIXO"


def _keywords(description: str) -> str:
    words = []
    for word in re.split(r"\s+", _normalize_key(description)):
        if len(word) >= 4 and word not in words:
            words.append(word)
        if len(words) == 5:
            break
    return "; ".join(words)


def _load_existing_decisions(path: Path) -> dict[str, dict[str, object]]:
    if not path.exists():
        return {}
    workbook = load_workbook(path, data_only=False)
    if DECISOES_SHEET not in workbook.sheetnames:
        return {}
    sheet = workbook[DECISOES_SHEET]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {}
    headers = [str(value or "").strip() for value in rows[0]]
    result: dict[str, dict[str, object]] = {}
    for values in rows[1:]:
        row = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
        }
        key = str(row.get("CHAVE_ESTAVEL") or "").strip()
        if key:
            result[key] = row
    return result


def _preserve_manual(target: dict[str, object], existing: dict[str, object]) -> None:
    for column in MANUAL_COLUMNS:
        value = existing.get(column)
        if column in target and value not in (None, ""):
            target[column] = value
    conta_id, conta_nome = _account_from_choice(existing.get("CONTA_ESCOLHIDA"))
    if conta_id:
        target["CONTA_ID_APROVADA"] = conta_id
    if conta_nome:
        target["CONTA_NOME_APROVADA"] = conta_nome


def _account_from_choice(value: object) -> tuple[str, str]:
    choice = str(value or "").strip()
    if not choice or choice.startswith("=") or " - " not in choice:
        return "", ""
    conta_id, conta_nome = choice.split(" - ", 1)
    return conta_id.strip(), conta_nome.strip()


def _write_workbook(
    path: Path,
    *,
    decision_rows: list[dict[str, object]],
    technical_rows: list[dict[str, object]],
    account_options: list[tuple[str, str, str]],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DECISOES_SHEET
    _append_rows(sheet, DECISOES_HEADERS, decision_rows)
    _style_sheet(sheet, DECISOES_HEADERS)
    _add_account_validation(sheet, DECISOES_HEADERS, len(decision_rows), len(account_options))
    _add_account_formulas(sheet, DECISOES_HEADERS, len(decision_rows))

    tech = workbook.create_sheet(TECNICO_SHEET)
    _append_rows(tech, TECNICO_HEADERS, technical_rows)
    _style_sheet(tech, TECNICO_HEADERS)

    plano = workbook.create_sheet(PLANO_CONTAS_SHEET)
    plano.append(["CONTA_ESCOLHIDA", "CONTA_ID", "CONTA_NOME"])
    for option, conta_id, conta_nome in account_options:
        plano.append([option, conta_id, conta_nome])
    _style_sheet(plano, ["CONTA_ESCOLHIDA", "CONTA_ID", "CONTA_NOME"])

    listas = workbook.create_sheet(LISTAS_SHEET)
    _append_rows(
        listas,
        ["STATUS_DECISAO", "ACAO_REGRA", "NIVEL_RISCO_REGRA"],
        [
            {
                "STATUS_DECISAO": "PENDENTE",
                "ACAO_REGRA": "CRIAR_REGRA",
                "NIVEL_RISCO_REGRA": "BAIXO",
            },
            {
                "STATUS_DECISAO": "APROVADO",
                "ACAO_REGRA": "CRIAR_ALIAS",
                "NIVEL_RISCO_REGRA": "MEDIO",
            },
            {
                "STATUS_DECISAO": "BLOQUEADO",
                "ACAO_REGRA": "IGNORAR",
                "NIVEL_RISCO_REGRA": "ALTO",
            },
        ],
    )
    _style_sheet(listas, ["STATUS_DECISAO", "ACAO_REGRA", "NIVEL_RISCO_REGRA"])

    instrucoes = workbook.create_sheet(INSTRUCOES_SHEET)
    for line in _instructions():
        instrucoes.append([line])
    instrucoes.column_dimensions["A"].width = 110
    for row in instrucoes.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    instrucoes["A1"].font = Font(bold=True, size=14)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _append_rows(sheet, headers: list[str], rows: list[dict[str, object]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header, "")) for header in headers])


def _style_sheet(sheet, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = _suggested_width(header)
        for column_cells in sheet.iter_cols(
            min_col=index,
            max_col=index,
            min_row=2,
            max_row=max(sheet.max_row, 2),
        ):
            for cell in column_cells:
                cell.alignment = Alignment(wrap_text=_wrap(header), vertical="top")
                if header == "VALOR_TOTAL":
                    cell.number_format = '#,##0.00'
                if header == "QTD_OCORRENCIAS":
                    cell.number_format = "0"


def _add_account_validation(
    sheet,
    headers: list[str],
    row_count: int,
    account_count: int,
) -> None:
    if row_count < 1 or "CONTA_ESCOLHIDA" not in headers or account_count < 1:
        return
    column = headers.index("CONTA_ESCOLHIDA") + 1
    letter = get_column_letter(column)
    validation = DataValidation(
        type="list",
        formula1=f"={PLANO_CONTAS_SHEET}!$A$2:$A${account_count + 1}",
        allow_blank=True,
    )
    validation.error = "Escolha uma conta da lista."
    validation.errorTitle = "Conta invalida"
    sheet.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{row_count + 1}")


def _add_account_formulas(sheet, headers: list[str], row_count: int) -> None:
    if row_count < 1 or "CONTA_ESCOLHIDA" not in headers:
        return
    escolha_col = get_column_letter(headers.index("CONTA_ESCOLHIDA") + 1)
    conta_id_col = headers.index("CONTA_ID_APROVADA") + 1
    conta_nome_col = headers.index("CONTA_NOME_APROVADA") + 1
    for row_index in range(2, row_count + 2):
        escolha_ref = f"{escolha_col}{row_index}"
        if not sheet.cell(row=row_index, column=conta_id_col).value:
            sheet.cell(row=row_index, column=conta_id_col).value = (
                f'=IFERROR(LEFT({escolha_ref},FIND(" - ",{escolha_ref})-1),"")'
            )
        if not sheet.cell(row=row_index, column=conta_nome_col).value:
            sheet.cell(row=row_index, column=conta_nome_col).value = (
                f'=IFERROR(MID({escolha_ref},FIND(" - ",{escolha_ref})+3,999),"")'
            )


def _load_account_options(path: Path) -> list[tuple[str, str, str]]:
    with Path(path).open(encoding="utf-8") as file:
        data = yaml.safe_load(file) or {}
    options = []
    for item in data.get("contas", []):
        conta_id = str(item.get("id") or "").strip()
        conta_nome = str(item.get("nome") or "").strip()
        if conta_id and conta_nome:
            options.append((f"{conta_id} - {conta_nome}", conta_id, conta_nome))
    return options


def _write_manifest(result: ImportaRevisaoXmlResult) -> None:
    payload = {
        "run_id_origem": result.run_id_origem,
        "run_reports_dir": str(result.run_reports_dir),
        "cnpjs_processados": list(result.cnpjs_processados),
        "csvs_lidos": list(result.csvs_lidos),
        "hash_yaml_run": result.hash_yaml_run,
        "hash_yaml_atual": result.hash_yaml_atual,
        "hash_yaml_run_assente": result.hash_yaml_run_assente,
        "run_verificavel": not result.hash_yaml_run_assente,
        "stale_run": result.stale_run,
        "allow_stale_run": result.allow_stale_run,
        "linhas_operacionais_geradas": result.linhas_operacionais_geradas,
        "decisoes_preservadas": result.decisoes_preservadas,
        "linhas_novas": result.linhas_novas,
        "comodato_forte_esclusi": result.comodato_forte_esclusi,
        "bonificacao_forte_esclusa": result.bonificacao_forte_esclusa,
        "candidatos_bloqueados_operacao_especial": (
            result.candidatos_bloqueados_operacao_especial
        ),
        "excel": str(result.excel_path),
        "avisos": list(result.avisos),
        "bloqueios": list(result.bloqueios),
        "yaml_modificado": False,
        "reports_importacao_real_modificados": False,
    }
    result.manifest_path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def _instructions() -> list[str]:
    return [
        "Revisione XML prodotti: compilare solo le colonne decisionali.",
        "Le righe BLOQUEADO indicano operazioni speciali non FORTE: non approvare automaticamente CRIAR_ALIAS o CRIAR_REGRA.",
        "CONTA_ID_APROVADA e CONTA_NOME_APROVADA sono derivate da CONTA_ESCOLHIDA.",
        "Il YAML non viene modificato da questo comando.",
    ]


def _run_rules_hash(manifest: dict[str, object]) -> str:
    config = manifest.get("config_usado")
    if not isinstance(config, dict):
        return ""
    regras = config.get("regras_classificacao_produtos")
    if not isinstance(regras, dict):
        return ""
    return str(regras.get("hash_sha256") or "").strip()


def _stale_hash_message(run_id: str, hash_run: str, hash_atual: str) -> str:
    return (
        f"BLOCCO: la run {run_id} è stata generata con un YAML diverso da quello attuale.\n"
        f"Hash run:     {hash_run}\n"
        f"Hash attuale: {hash_atual}\n"
        "Eseguire prima reprocessar-importacao-real con il YAML corrente, poi rilanciare la skill.\n"
        "Per audit/debug su dati storici usare --allow-stale-run (non per uso operativo)."
    )


def _missing_hash_message(run_id: str) -> str:
    return (
        f"BLOCCO: la run {run_id} non contiene l'hash del YAML "
        "regras_classificacao_produtos.yaml.\n"
        "Non e possibile garantire che la run sia coerente con le regole attuali.\n"
        "Eseguire prima reprocessar-importacao-real con il YAML corrente, poi rilanciare la skill.\n"
        "Per audit/debug su dati storici usare --allow-stale-run."
    )


def _hash_file(path: Path) -> str:
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()


def _parse_decimal(value: object) -> Decimal:
    text = str(value or "").strip()
    if not text:
        return Decimal("0")
    text = text.replace("R$", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0")


def _parse_int(value: object) -> int:
    try:
        return int(str(value or "0").strip() or "0")
    except ValueError:
        return 0


def _joined_unique(values: Iterable[object], *, limit: int = 8) -> str:
    result: list[str] = []
    for value in values:
        text = str(value or "").strip()
        if text and text not in result:
            result.append(text)
        if len(result) >= limit:
            break
    return "; ".join(result)


def _token(value: object) -> str:
    return _normalize_key(value).replace(" ", "_")


def _normalize_key(value: object) -> str:
    text = str(value or "").strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^A-Za-z0-9]+", " ", text).strip().upper()
    return re.sub(r"\s+", " ", text)


def _cell_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    return value


def _suggested_width(header: str) -> int:
    widths = {
        "CHAVE_ESTAVEL": 58,
        "DESCRICAO_EXEMPLO": 44,
        "DESCRICAO_NORMALIZADA": 38,
        "FORNECEDORES": 34,
        "CONTA_ESCOLHIDA": 42,
        "CONTA_ID_APROVADA": 18,
        "CONTA_NOME_APROVADA": 34,
        "PALAVRAS_CHAVE_SUGERIDAS": 34,
        "OBSERVACAO_OPERADOR": 42,
        "ORIGEM_CSV": 60,
        "ALERTAS": 42,
    }
    return widths.get(header, min(max(len(header) + 2, 12), 28))


def _wrap(header: str) -> bool:
    return header not in {"QTD_OCORRENCIAS", "VALOR_TOTAL"}
