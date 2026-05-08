from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from fantozzi.classification_review_export import CSV_HEADERS
from fantozzi.classification_review_validation import ALLOWED_ACTIONS
from fantozzi.classification_review_workbook import (
    ClassificationReviewWorkbookError,
    DEFAULT_PLANO_CONTAS,
    GENERIC_KEYWORDS,
    REVISAO_SHEET,
    carregar_plano_contas,
)


ACOES_COM_CONTA_OBRIGATORIA = {"CLASSIFICAR", "CRIAR_ALIAS", "CRIAR_REGRA"}
ACAO_COM_PALAVRAS_OBRIGATORIAS = "CRIAR_REGRA"
OPTIONAL_WORKBOOK_HEADERS = {"conta_escolhida"}


class ClassificationReviewWorkbookImportError(Exception):
    def __init__(
        self,
        message: str,
        result: ResultadoImportacaoPlanilhaRevisao | None = None,
    ) -> None:
        super().__init__(message)
        self.result = result


@dataclass(frozen=True)
class ContasValidas:
    por_id: dict[str, str]
    por_opcao: dict[str, tuple[str, str]]


@dataclass(frozen=True)
class ResultadoImportacaoPlanilhaRevisao:
    xlsx_path: Path
    csv_path: Path | None
    relatorio_path: Path | None
    linhas_lidas: int
    linhas_exportadas: int
    erros: tuple[str, ...]
    avisos: tuple[str, ...]
    contagem_acoes: dict[str, int]

    @property
    def is_ok(self) -> bool:
        return not self.erros


def importar_planilha_revisao_classificacao(
    xlsx_path: Path | str,
    output_csv_path: Path | str,
    plano_contas_path: Path | str | None = None,
    relatorio_path: Path | str | None = None,
) -> ResultadoImportacaoPlanilhaRevisao:
    xlsx_path = Path(xlsx_path)
    output_csv_path = Path(output_csv_path)
    plano_contas_path = Path(plano_contas_path or DEFAULT_PLANO_CONTAS)
    relatorio_path = Path(relatorio_path) if relatorio_path else None

    linhas = ler_aba_revisao(xlsx_path)
    contas_validas = carregar_contas_validas(plano_contas_path)
    rows, erros, avisos, contagem_acoes = validar_linhas_revisao(
        linhas, contas_validas
    )

    result = ResultadoImportacaoPlanilhaRevisao(
        xlsx_path=xlsx_path,
        csv_path=None if erros else output_csv_path,
        relatorio_path=relatorio_path,
        linhas_lidas=len(linhas),
        linhas_exportadas=0 if erros else len(rows),
        erros=tuple(erros),
        avisos=tuple(avisos),
        contagem_acoes=contagem_acoes,
    )

    if relatorio_path:
        escrever_relatorio_validacao(result, relatorio_path)

    if erros:
        raise ClassificationReviewWorkbookImportError(erros[0], result)

    escrever_csv_revisado(rows, output_csv_path)
    return result


def ler_aba_revisao(xlsx_path: Path | str) -> list[tuple[int, dict[str, str]]]:
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        raise ClassificationReviewWorkbookImportError(
            f"Planilha assistida nao encontrada: {xlsx_path}"
        )
    if not xlsx_path.is_file():
        raise ClassificationReviewWorkbookImportError(
            f"Planilha assistida nao e um arquivo: {xlsx_path}"
        )

    try:
        workbook = load_workbook(xlsx_path, data_only=True)
    except Exception as exc:
        raise ClassificationReviewWorkbookImportError(
            f"Planilha assistida nao pode ser lida: {xlsx_path}"
        ) from exc

    if REVISAO_SHEET not in workbook.sheetnames:
        raise ClassificationReviewWorkbookImportError(
            f"Aba obrigatoria nao encontrada: {REVISAO_SHEET}"
        )

    sheet = workbook[REVISAO_SHEET]
    raw_headers = [_text(cell.value) for cell in sheet[1]]
    headers = [header.strip() for header in raw_headers]
    _validar_headers_revisao(headers)

    readable_headers = [*CSV_HEADERS, *OPTIONAL_WORKBOOK_HEADERS]
    rows: list[tuple[int, dict[str, str]]] = []
    for row_index, values in enumerate(
        sheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        row_by_header = {
            header: _cell_to_text(values[headers.index(header)])
            for header in readable_headers
            if header in headers
        }
        if _linha_totalmente_vazia(row_by_header):
            continue
        rows.append((row_index, row_by_header))
    return rows


def validar_linhas_revisao(
    linhas: list[tuple[int, dict[str, str]]],
    contas_validas: ContasValidas,
) -> tuple[list[dict[str, str]], list[str], list[str], dict[str, int]]:
    rows: list[dict[str, str]] = []
    erros: list[str] = []
    avisos: list[str] = []
    contagem_acoes = {acao: 0 for acao in sorted(ALLOWED_ACTIONS)}

    for linha_xlsx, row in linhas:
        normalized_row = {header: _text(row.get(header)) for header in CSV_HEADERS}
        acao_original = normalized_row["acao_manual"]
        acao = acao_original.strip().upper()
        normalized_row["acao_manual"] = acao

        if not acao:
            erros.append(f"Linha {linha_xlsx}: acao_manual vazia.")
        elif acao not in ALLOWED_ACTIONS:
            erros.append(f"Linha {linha_xlsx}: acao_manual invalida: {acao_original}")
        else:
            contagem_acoes[acao] += 1

        conta_escolhida = _text(row.get("conta_escolhida"))
        if conta_escolhida:
            conta_extraida = _extract_account_choice(conta_escolhida, contas_validas)
            if conta_extraida is None:
                erros.append(
                    f"Linha {linha_xlsx}: conta_escolhida nao existe "
                    f"no plano de contas: {conta_escolhida}"
                )
                normalized_row["conta_id_sugerida"] = ""
                normalized_row["conta_nome_sugerida"] = ""
            else:
                conta_id, conta_nome = conta_extraida
                normalized_row["conta_id_sugerida"] = conta_id
                normalized_row["conta_nome_sugerida"] = conta_nome
        else:
            _validate_legacy_account_fields(linha_xlsx, normalized_row, contas_validas, erros)

        conta_id = _normalize_account_id(normalized_row["conta_id_sugerida"])
        normalized_row["conta_id_sugerida"] = conta_id

        if acao in ACOES_COM_CONTA_OBRIGATORIA and not conta_id:
            erros.append(
                f"Linha {linha_xlsx}: conta_escolhida obrigatoria "
                f"para acao_manual {acao}."
            )

        if conta_id and conta_id not in contas_validas.por_id:
            erros.append(
                f"Linha {linha_xlsx}: conta_id_sugerida nao existe "
                f"no plano de contas: {conta_id}"
            )

        conta_atual_id = _normalize_account_id(normalized_row["conta_atual_id"])
        normalized_row["conta_atual_id"] = conta_atual_id

        _validate_keywords(linha_xlsx, normalized_row, erros, avisos)
        rows.append(normalized_row)

    return rows, erros, avisos, contagem_acoes


def escrever_csv_revisado(rows: list[dict[str, str]], output_path: Path | str) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=CSV_HEADERS, delimiter=";")
        writer.writeheader()
        writer.writerows(
            [{header: row.get(header, "") for header in CSV_HEADERS} for row in rows]
        )


def carregar_contas_validas(plano_contas_path: Path | str) -> ContasValidas:
    try:
        contas = carregar_plano_contas(plano_contas_path)
    except ClassificationReviewWorkbookError as exc:
        raise ClassificationReviewWorkbookImportError(str(exc)) from exc

    por_id = {
        _normalize_account_id(conta["conta_id"]): conta["conta_nome"]
        for conta in contas
        if conta.get("conta_id")
    }
    por_opcao: dict[str, tuple[str, str]] = {}
    for conta_id, conta_nome in por_id.items():
        for option in _account_options(conta_id, conta_nome):
            por_opcao[_normalize_choice(option)] = (conta_id, conta_nome)
    return ContasValidas(por_id=por_id, por_opcao=por_opcao)


def escrever_relatorio_validacao(
    result: ResultadoImportacaoPlanilhaRevisao, relatorio_path: Path | str
) -> None:
    relatorio_path = Path(relatorio_path)
    relatorio_path.parent.mkdir(parents=True, exist_ok=True)
    relatorio_path.write_text(relatorio_validacao_text(result), encoding="utf-8")


def relatorio_validacao_text(result: ResultadoImportacaoPlanilhaRevisao) -> str:
    csv_path = result.csv_path if result.csv_path else "-"
    lines = [
        "VALIDACAO DA PLANILHA ASSISTIDA DE REVISAO DE CLASSIFICACAO",
        "============================================================",
        "",
        f"Arquivo XLSX: {result.xlsx_path}",
        f"Arquivo CSV gerado: {csv_path}",
        "",
        "Resumo geral:",
        f"- linhas lidas: {result.linhas_lidas}",
        f"- linhas exportadas: {result.linhas_exportadas}",
        f"- erros: {len(result.erros)}",
        f"- avisos: {len(result.avisos)}",
        "",
        "Acoes por tipo:",
    ]
    for action in sorted(ALLOWED_ACTIONS):
        lines.append(f"- {action}: {result.contagem_acoes.get(action, 0)}")

    lines.extend(["", "Erros:", *[f"- {erro}" for erro in result.erros]])
    if not result.erros:
        lines.append("-")
    lines.extend(["", "Avisos:", *[f"- {aviso}" for aviso in result.avisos]])
    if not result.avisos:
        lines.append("-")
    lines.extend(["", f"Resultado: {'OK' if result.is_ok else 'ERRO'}"])
    return "\n".join(lines) + "\n"


def console_lines(result: ResultadoImportacaoPlanilhaRevisao) -> list[str]:
    if result.is_ok:
        return [
            "[OK] Planilha importada com sucesso.",
            f"[OK] CSV revisado gerado: {result.csv_path}",
            f"Linhas exportadas: {result.linhas_exportadas}",
            f"Erros: {len(result.erros)}",
            f"Avisos: {len(result.avisos)}",
        ]
    return [
        "[ERRO] Planilha assistida contem erros bloqueantes.",
        f"Linhas lidas: {result.linhas_lidas}",
        f"Erros: {len(result.erros)}",
        f"Avisos: {len(result.avisos)}",
        *[f"[ERRO] {erro}" for erro in result.erros],
    ]


def _validar_headers_revisao(headers: list[str]) -> None:
    for required_header in CSV_HEADERS:
        if required_header not in headers:
            raise ClassificationReviewWorkbookImportError(
                f"Coluna obrigatoria ausente na aba Revisao: {required_header}"
            )


def _validate_legacy_account_fields(
    linha_xlsx: int,
    row: dict[str, str],
    contas_validas: ContasValidas,
    erros: list[str],
) -> None:
    conta_id = _normalize_account_id(row["conta_id_sugerida"])
    row["conta_id_sugerida"] = conta_id
    if not conta_id:
        return
    conta_nome_esperado = contas_validas.por_id.get(conta_id)
    if conta_nome_esperado is None:
        return
    conta_nome = _text(row["conta_nome_sugerida"])
    if conta_nome and _normalize_choice(conta_nome) != _normalize_choice(conta_nome_esperado):
        erros.append(
            f"Linha {linha_xlsx}: conta_nome_sugerida incoerente "
            f"com conta_id_sugerida {conta_id}."
        )
    row["conta_nome_sugerida"] = conta_nome_esperado


def _validate_keywords(
    linha_xlsx: int,
    row: dict[str, str],
    erros: list[str],
    avisos: list[str],
) -> None:
    acao = row["acao_manual"]
    palavras = _split_keywords(row["palavras_chave_sugeridas"])
    if acao != ACAO_COM_PALAVRAS_OBRIGATORIAS:
        return

    if not palavras:
        erros.append(
            f"Linha {linha_xlsx}: palavras_chave_sugeridas obrigatorias "
            "para acao_manual CRIAR_REGRA."
        )
        return

    generic_terms = [keyword for keyword in palavras if _is_generic_keyword(keyword)]
    useful_terms = [keyword for keyword in palavras if keyword not in generic_terms]

    if generic_terms and not useful_terms:
        erros.append(
            f"Linha {linha_xlsx}: CRIAR_REGRA possui apenas palavras-chave "
            f"genericas: {'; '.join(generic_terms)}."
        )
        return

    if generic_terms:
        avisos.append(
            f"Linha {linha_xlsx}: palavras-chave genericas detectadas: "
            f"{'; '.join(generic_terms)}."
        )
    if len(useful_terms) < 2:
        avisos.append(
            f"Linha {linha_xlsx}: CRIAR_REGRA possui poucas palavras-chave uteis."
        )


def _extract_account_choice(
    conta_escolhida: str, contas_validas: ContasValidas
) -> tuple[str, str] | None:
    normalized_choice = _normalize_choice(conta_escolhida)
    if normalized_choice in contas_validas.por_opcao:
        return contas_validas.por_opcao[normalized_choice]

    conta_id = _extract_account_id(conta_escolhida)
    if conta_id and conta_id in contas_validas.por_id:
        return conta_id, contas_validas.por_id[conta_id]
    return None


def _account_options(conta_id: str, conta_nome: str) -> list[str]:
    return [
        f"{conta_id} - {conta_nome}",
        f"{conta_id} – {conta_nome}",
        f"{conta_id} — {conta_nome}",
        conta_id,
    ]


def _extract_account_id(value: str) -> str:
    text = _text(value)
    for separator in (" - ", " – ", " — "):
        if separator in text:
            text = text.split(separator, 1)[0]
            break
    return _normalize_account_id(text)


def _split_keywords(value: str) -> list[str]:
    keywords = []
    for chunk in _text(value).replace(",", ";").split(";"):
        keyword = " ".join(chunk.strip().upper().split())
        if keyword:
            keywords.append(keyword)
    return keywords


def _is_generic_keyword(keyword: str) -> bool:
    clean_keyword = _normalize_choice(keyword)
    return (
        clean_keyword in GENERIC_KEYWORDS
        or len(clean_keyword) < 3
        or clean_keyword.isnumeric()
    )


def _linha_totalmente_vazia(row: dict[str, str]) -> bool:
    return all(not _text(value) for value in row.values())


def _normalize_account_id(value: str) -> str:
    return "".join(_text(value).split())


def _normalize_choice(value: str) -> str:
    return " ".join(_text(value).upper().split())


def _cell_to_text(value: object) -> str:
    return _text(value)


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()
