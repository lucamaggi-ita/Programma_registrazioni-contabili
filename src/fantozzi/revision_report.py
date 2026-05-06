from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from fantozzi.excel_writer import LANCAMENTOS_SHEET, REVISAO_SHEET


REQUIRED_HEADERS = ("status_revisao", "valor")
IDENTIFICATION_HEADERS = ("descricao", "descricao_itens", "documento", "chave_origem")


class RevisionReportError(Exception):
    pass


@dataclass(frozen=True)
class RevisionRow:
    excel_row: int
    data_documento: str
    documento: str
    fornecedor: str
    conta_id: str
    conta_nome: str
    descricao: str
    valor: Decimal
    origem: str
    status_revisao: str
    alertas: str

    @property
    def is_manual(self) -> bool:
        return self.origem.strip().upper() == "MANUAL"


@dataclass(frozen=True)
class RevisionReportResult:
    excel_path: Path
    lancamentos_totais: int
    lancamentos_em_revisao: int
    fornecedores_com_revisao: int
    documentos_com_revisao: int
    valor_total_em_revisao: Decimal
    status: str
    revision_rows: tuple[RevisionRow, ...]
    warnings: tuple[str, ...]


def generate_revision_report(
    excel_path: Path | str,
    *,
    generated_at: datetime | None = None,
) -> tuple[RevisionReportResult, str]:
    result = analyze_revision_report(excel_path)
    return result, report_text(result, generated_at=generated_at)


def analyze_revision_report(excel_path: Path | str) -> RevisionReportResult:
    excel_path = Path(excel_path)
    _validate_file(excel_path)

    try:
        workbook = load_workbook(excel_path, data_only=True, read_only=True)
    except Exception as exc:
        raise RevisionReportError(f"Arquivo Excel nao pode ser lido: {excel_path}") from exc

    if LANCAMENTOS_SHEET not in workbook.sheetnames:
        raise RevisionReportError(f"Aba obrigatoria ausente: {LANCAMENTOS_SHEET}")

    lancamentos_sheet = workbook[LANCAMENTOS_SHEET]
    headers = _header_indexes(lancamentos_sheet)
    _validate_headers(headers)

    rows = list(_data_rows(lancamentos_sheet))
    revision_rows = tuple(
        _revision_row(excel_row, row, headers)
        for excel_row, row in rows
        if _is_revision(_cell(row, headers, "status_revisao"))
    )

    warnings = _revision_sheet_warnings(workbook, revision_rows)
    fornecedores = {
        row.fornecedor for row in revision_rows if row.fornecedor.strip() != ""
    }
    documentos = {row.documento for row in revision_rows if row.documento.strip() != ""}
    valor_total = sum((row.valor for row in revision_rows), Decimal("0.00"))

    return RevisionReportResult(
        excel_path=excel_path,
        lancamentos_totais=len(rows),
        lancamentos_em_revisao=len(revision_rows),
        fornecedores_com_revisao=len(fornecedores),
        documentos_com_revisao=len(documentos),
        valor_total_em_revisao=valor_total.quantize(Decimal("0.01")),
        status="REVISAR" if revision_rows else "OK",
        revision_rows=revision_rows,
        warnings=warnings,
    )


def report_text(
    result: RevisionReportResult,
    *,
    generated_at: datetime | None = None,
) -> str:
    generated_at = generated_at or datetime.now()
    lines = [
        "FANTOZZI - RELATORIO DE REVISAO",
        "========================================",
        f"Arquivo analisado: {result.excel_path}",
        f"Data/hora relatorio: {generated_at:%Y-%m-%d %H:%M:%S}",
        "",
        "Resumo",
        "----------------------------------------",
        f"Lancamentos totais: {result.lancamentos_totais}",
        f"Lancamentos em revisao: {result.lancamentos_em_revisao}",
        f"Fornecedores com revisao: {result.fornecedores_com_revisao}",
        f"Documentos com revisao: {result.documentos_com_revisao}",
        f"Valor total em revisao: {_format_currency_br(result.valor_total_em_revisao)}",
        "",
        f"Status: {result.status}",
    ]

    if not result.revision_rows:
        lines.extend(
            [
                "",
                "Nenhum lancamento em revisao encontrado.",
            ]
        )
    else:
        lines.extend(
            [
                "",
                "Lancamentos em revisao",
                "----------------------------------------",
            ]
        )
        for index, row in enumerate(result.revision_rows, start=1):
            lines.extend(_revision_row_lines(index, row))

        lines.extend(
            [
                "",
                "Agrupamento por fornecedor",
                "----------------------------------------",
            ]
        )
        lines.extend(_group_lines_by_supplier(result.revision_rows))

        lines.extend(
            [
                "",
                "Agrupamento por conta",
                "----------------------------------------",
            ]
        )
        lines.extend(_group_lines_by_account(result.revision_rows))

    if result.warnings:
        lines.extend(["", "Avisos", "----------------------------------------"])
        lines.extend(f"- {warning}" for warning in result.warnings)

    lines.extend(
        [
            "",
            "Observacoes",
            "----------------------------------------",
            "- Este relatorio e somente leitura.",
            "- Nenhuma correcao automatica foi aplicada.",
            "- Revise as classificacoes antes de alterar regras ou aliases.",
        ]
    )
    return "\n".join(lines) + "\n"


def console_lines(result: RevisionReportResult) -> list[str]:
    return [
        "[OK] Relatorio de revisao concluido",
        f"Lancamentos totais: {result.lancamentos_totais}",
        f"Lancamentos em revisao: {result.lancamentos_em_revisao}",
        f"Valor total em revisao: {_format_currency_br(result.valor_total_em_revisao)}",
        f"Status: {result.status}",
    ]


def save_report(report_path: Path | str, text: str) -> None:
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(text, encoding="utf-8", newline="\n")


def _validate_file(path: Path) -> None:
    if not path.exists():
        raise RevisionReportError(f"Arquivo Excel nao encontrado: {path}")
    if not path.is_file():
        raise RevisionReportError(f"Arquivo Excel nao e um arquivo: {path}")


def _validate_headers(headers: dict[str, int]) -> None:
    missing_required = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing_required:
        raise RevisionReportError(
            "Colunas obrigatorias ausentes em Lancamentos: "
            + ", ".join(missing_required)
        )

    if not any(header in headers for header in IDENTIFICATION_HEADERS):
        raise RevisionReportError(
            "Colunas de identificacao ausentes em Lancamentos: "
            + ", ".join(IDENTIFICATION_HEADERS)
        )


def _header_indexes(sheet: Worksheet) -> dict[str, int]:
    return {
        str(cell.value): index
        for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)))
        if cell.value is not None
    }


def _data_rows(sheet: Worksheet) -> Iterable[tuple[int, tuple[object, ...]]]:
    for excel_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if any(_has_value(value) for value in row):
            yield excel_row, row


def _revision_row(
    excel_row: int, row: tuple[object, ...], headers: dict[str, int]
) -> RevisionRow:
    return RevisionRow(
        excel_row=excel_row,
        data_documento=_text(
            _first_cell(row, headers, ("data_emissao", "data_documento"))
        ),
        documento=_text(_cell(row, headers, "documento")),
        fornecedor=_text(
            _first_cell(row, headers, ("fornecedor_nome", "fornecedor_ou_favorecido"))
        ),
        conta_id=_text(_cell(row, headers, "conta_id")),
        conta_nome=_text(_cell(row, headers, "conta_nome")),
        descricao=_text(_first_cell(row, headers, ("descricao", "descricao_itens"))),
        valor=_decimal_value(_cell(row, headers, "valor")),
        origem=_text(_cell(row, headers, "origem")),
        status_revisao=_text(_cell(row, headers, "status_revisao")) or "(vazio)",
        alertas=_text(_cell(row, headers, "alertas")),
    )


def _revision_row_lines(index: int, row: RevisionRow) -> list[str]:
    manual_note = " (linha manual)" if row.is_manual else ""
    conta = _join_code_name(row.conta_id, row.conta_nome)
    return [
        f"{index}) Linha Excel: {row.excel_row}{manual_note}",
        f"   Documento: {_or_dash(row.documento)}",
        f"   Fornecedor: {_or_dash(row.fornecedor)}",
        f"   Data emissao: {_or_dash(row.data_documento)}",
        f"   Conta: {_or_dash(conta)}",
        f"   Valor: {_format_currency_br(row.valor)}",
        f"   Status revisao: {row.status_revisao}",
        f"   Descricao: {_or_dash(row.descricao)}",
        f"   Motivo/alerta: {_or_dash(row.alertas)}",
        "",
    ]


def _group_lines_by_supplier(rows: tuple[RevisionRow, ...]) -> list[str]:
    values: defaultdict[str, Decimal] = defaultdict(lambda: Decimal("0.00"))
    counts: Counter[str] = Counter()
    for row in rows:
        supplier = row.fornecedor or "(sem fornecedor)"
        counts[supplier] += 1
        values[supplier] += row.valor

    return [
        f"{supplier}: {counts[supplier]} lancamento(s), {_format_currency_br(values[supplier])}"
        for supplier in sorted(counts)
    ]


def _group_lines_by_account(rows: tuple[RevisionRow, ...]) -> list[str]:
    values: defaultdict[str, Decimal] = defaultdict(lambda: Decimal("0.00"))
    counts: Counter[str] = Counter()
    for row in rows:
        account = _join_code_name(row.conta_id, row.conta_nome) or "(sem conta)"
        counts[account] += 1
        values[account] += row.valor

    return [
        f"{account}: {counts[account]} lancamento(s), {_format_currency_br(values[account])}"
        for account in sorted(counts)
    ]


def _revision_sheet_warnings(
    workbook: object, revision_rows: tuple[RevisionRow, ...]
) -> tuple[str, ...]:
    if REVISAO_SHEET not in workbook.sheetnames:
        if revision_rows:
            return (
                f"Aba {REVISAO_SHEET} ausente; a fonte primaria foi a aba {LANCAMENTOS_SHEET}.",
            )
        return ()

    revisao_sheet = workbook[REVISAO_SHEET]
    revisao_rows_count = sum(1 for _ in _data_rows(revisao_sheet))
    if revision_rows and revisao_rows_count == 0:
        return (
            f"Aba {REVISAO_SHEET} esta vazia, mas {LANCAMENTOS_SHEET} contem lancamentos em revisao.",
        )
    if revision_rows and revisao_rows_count != len(revision_rows):
        return (
            f"Aba {REVISAO_SHEET} contem {revisao_rows_count} linha(s), "
            f"mas {LANCAMENTOS_SHEET} contem {len(revision_rows)} lancamento(s) em revisao.",
        )
    return ()


def _cell(row: tuple[object, ...], indexes: dict[str, int], header: str) -> object:
    index = indexes.get(header)
    if index is None or index >= len(row):
        return None
    return row[index]


def _first_cell(
    row: tuple[object, ...], indexes: dict[str, int], headers: tuple[str, ...]
) -> object:
    for header in headers:
        value = _cell(row, indexes, header)
        if _has_value(value):
            return value
    return None


def _is_revision(status_revisao: object) -> bool:
    return _text(status_revisao).upper() != "OK"


def _has_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _decimal_value(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0.00")
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _format_currency_br(value: Decimal) -> str:
    text = f"{value.quantize(Decimal('0.01')):.2f}"
    return "R$ " + text.replace(".", ",")


def _join_code_name(code: str, name: str) -> str:
    if code and name:
        return f"{code} - {name}"
    return code or name


def _or_dash(value: str) -> str:
    return value if value else "-"
