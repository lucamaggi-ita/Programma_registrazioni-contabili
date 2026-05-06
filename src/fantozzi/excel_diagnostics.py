from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from fantozzi.excel_writer import HEADERS, LANCAMENTOS_SHEET, REVISAO_SHEET


REQUIRED_LANCAMENTO_HEADERS = [
    header
    for header in [
        "origem",
        "chave_origem",
        "documento",
        "conta_id",
        "conta_nome",
        "valor",
        "status_revisao",
        "hash_registro",
        "duplicatas_json",
    ]
    if header in HEADERS
]


class ExcelDiagnosticError(Exception):
    pass


@dataclass(frozen=True)
class ExcelDiagnosticResult:
    excel_path: Path
    lancamentos_totais: int
    lancamentos_ok: int
    lancamentos_em_revisao: int
    linhas_manuais: int
    valor_total: Decimal
    duplicatas_json_invalidas: int
    hashes_duplicados: int
    status: str
    errors: tuple[str, ...]

    @property
    def is_ok(self) -> bool:
        return not self.errors


@dataclass(frozen=True)
class ExcelDiagnosticFailure:
    excel_path: Path
    error: str


def diagnose_excel(excel_path: Path | str) -> ExcelDiagnosticResult:
    excel_path = Path(excel_path)
    _validate_file(excel_path)

    try:
        workbook = load_workbook(excel_path, data_only=True)
    except Exception as exc:
        raise ExcelDiagnosticError(f"Arquivo Excel nao pode ser lido: {excel_path}") from exc

    if LANCAMENTOS_SHEET not in workbook.sheetnames:
        raise ExcelDiagnosticError(f"Aba obrigatoria ausente: {LANCAMENTOS_SHEET}")
    if REVISAO_SHEET not in workbook.sheetnames:
        raise ExcelDiagnosticError(f"Aba obrigatoria ausente: {REVISAO_SHEET}")

    lancamentos_sheet = workbook[LANCAMENTOS_SHEET]
    revisao_sheet = workbook[REVISAO_SHEET]
    lancamentos_headers = _header_indexes(lancamentos_sheet)
    revisao_headers = _header_indexes(revisao_sheet)

    missing_headers = [
        header
        for header in REQUIRED_LANCAMENTO_HEADERS
        if header not in lancamentos_headers
    ]
    if missing_headers:
        raise ExcelDiagnosticError(
            "Colunas obrigatorias ausentes em Lancamentos: "
            + ", ".join(missing_headers)
        )

    lancamento_rows = list(_data_rows(lancamentos_sheet))
    revisao_rows = list(_data_rows(revisao_sheet))
    revisao_identities = {
        identity
        for row in revisao_rows
        if (identity := _row_identity(row, revisao_headers)) is not None
    }

    errors: list[str] = []
    hashes: list[str] = []
    lancamentos_ok = 0
    lancamentos_em_revisao = 0
    linhas_manuais = 0
    valor_total = Decimal("0.00")
    duplicatas_json_invalidas = 0

    for row in lancamento_rows:
        origem = _text(_cell(row, lancamentos_headers, "origem"))
        status_revisao = _text(_cell(row, lancamentos_headers, "status_revisao"))
        hash_registro = _text(_cell(row, lancamentos_headers, "hash_registro"))

        if origem.upper() == "MANUAL":
            linhas_manuais += 1
        elif hash_registro:
            hashes.append(hash_registro)

        if status_revisao.upper() == "OK":
            lancamentos_ok += 1
        else:
            lancamentos_em_revisao += 1
            identity = _row_identity(row, lancamentos_headers)
            if identity is not None and identity not in revisao_identities:
                errors.append(
                    "Lancamento em revisao ausente na aba Revisao: "
                    + _identity_label(identity)
                )

        valor_total += _decimal_value(_cell(row, lancamentos_headers, "valor"))

        duplicatas_json = _cell(row, lancamentos_headers, "duplicatas_json")
        if _has_value(duplicatas_json) and not _is_valid_json(duplicatas_json):
            duplicatas_json_invalidas += 1

    duplicated_hashes = [
        hash_registro for hash_registro, count in Counter(hashes).items() if count > 1
    ]
    if duplicatas_json_invalidas:
        errors.append(f"Duplicatas JSON invalidas: {duplicatas_json_invalidas}")
    if duplicated_hashes:
        errors.append("Hashes duplicados: " + ", ".join(sorted(duplicated_hashes)))

    return ExcelDiagnosticResult(
        excel_path=excel_path,
        lancamentos_totais=len(lancamento_rows),
        lancamentos_ok=lancamentos_ok,
        lancamentos_em_revisao=lancamentos_em_revisao,
        linhas_manuais=linhas_manuais,
        valor_total=valor_total.quantize(Decimal("0.01")),
        duplicatas_json_invalidas=duplicatas_json_invalidas,
        hashes_duplicados=len(duplicated_hashes),
        status="REVISAR" if lancamentos_em_revisao else "OK",
        errors=tuple(errors),
    )


def format_currency_br(value: Decimal) -> str:
    text = f"{value.quantize(Decimal('0.01')):.2f}"
    return "R$ " + text.replace(".", ",")


def console_lines(result: ExcelDiagnosticResult) -> list[str]:
    lines = []
    if result.is_ok:
        lines.append("[OK] Diagnostico Excel concluido")
    else:
        lines.append("[ERRO] Diagnostico Excel falhou")

    lines.extend(
        [
            f"Arquivo analisado: {result.excel_path}",
            f"Lancamentos totais: {result.lancamentos_totais}",
            f"Lancamentos OK: {result.lancamentos_ok}",
            f"Lancamentos em revisao: {result.lancamentos_em_revisao}",
            f"Linhas manuais: {result.linhas_manuais}",
            f"Valor total Lancamentos: {format_currency_br(result.valor_total)}",
            f"Duplicatas JSON invalidas: {result.duplicatas_json_invalidas}",
            f"Hashes duplicados: {result.hashes_duplicados}",
            f"Status: {result.status}",
        ]
    )

    lines.extend(f"[ERRO] {diagnostic_error}" for diagnostic_error in result.errors)
    return lines


def report_text(
    result: ExcelDiagnosticResult | ExcelDiagnosticFailure,
    *,
    generated_at: datetime | None = None,
) -> str:
    generated_at = generated_at or datetime.now()
    lines = [
        "FANTOZZI - DIAGNOSTICO EXCEL",
        "========================================",
        f"Arquivo analisado: {result.excel_path}",
        f"Data/hora diagnostico: {generated_at:%Y-%m-%d %H:%M:%S}",
        "",
    ]

    if isinstance(result, ExcelDiagnosticFailure):
        lines.extend(
            [
                "Status: ERRO",
                "Resultado estrutural: FALHOU",
                "",
                "Erros:",
                f"- {result.error}",
            ]
        )
        return "\n".join(lines) + "\n"

    lines.extend(
        [
            f"Lancamentos totais: {result.lancamentos_totais}",
            f"Lancamentos OK: {result.lancamentos_ok}",
            f"Lancamentos em revisao: {result.lancamentos_em_revisao}",
            f"Linhas manuais: {result.linhas_manuais}",
            f"Valor total Lancamentos: {format_currency_br(result.valor_total)}",
            f"Duplicatas JSON invalidas: {result.duplicatas_json_invalidas}",
            f"Hashes duplicados: {result.hashes_duplicados}",
            "",
            f"Status: {result.status}",
            f"Resultado estrutural: {'OK' if result.is_ok else 'FALHOU'}",
        ]
    )

    if result.errors:
        lines.extend(["", "Erros:"])
        lines.extend(f"- {diagnostic_error}" for diagnostic_error in result.errors)

    lines.extend(
        [
            "",
            "Observacoes:",
            "- Diagnostico read-only.",
            "- Nenhuma correcao automatica foi aplicada.",
        ]
    )
    return "\n".join(lines) + "\n"


def save_report(report_path: Path | str, text: str) -> None:
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(text, encoding="utf-8", newline="\n")


def _validate_file(path: Path) -> None:
    if not path.exists():
        raise ExcelDiagnosticError(f"Arquivo Excel nao encontrado: {path}")
    if not path.is_file():
        raise ExcelDiagnosticError(f"Arquivo Excel nao e um arquivo: {path}")


def _header_indexes(sheet: Worksheet) -> dict[str, int]:
    return {
        str(cell.value): index
        for index, cell in enumerate(sheet[1])
        if cell.value is not None
    }


def _data_rows(sheet: Worksheet) -> Iterable[tuple[object, ...]]:
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(_has_value(value) for value in row):
            yield row


def _cell(row: tuple[object, ...], indexes: dict[str, int], header: str) -> object:
    index = indexes.get(header)
    if index is None or index >= len(row):
        return None
    return row[index]


def _row_identity(
    row: tuple[object, ...], indexes: dict[str, int]
) -> tuple[str, str] | None:
    hash_registro = _text(_cell(row, indexes, "hash_registro"))
    if hash_registro:
        return ("hash_registro", hash_registro)

    fallback_parts = [
        _text(_cell(row, indexes, "origem")),
        _text(_cell(row, indexes, "chave_origem")),
        _text(_cell(row, indexes, "documento")),
        _text(_cell(row, indexes, "conta_id")),
        str(_decimal_value(_cell(row, indexes, "valor"))),
    ]
    if not any(fallback_parts):
        return None
    return ("fallback", "|".join(fallback_parts))


def _identity_label(identity: tuple[str, str]) -> str:
    key, value = identity
    if key == "hash_registro":
        return value
    return f"sem hash ({value})"


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _has_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def _decimal_value(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0.00")
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _is_valid_json(value: object) -> bool:
    try:
        json.loads(str(value))
    except (TypeError, json.JSONDecodeError):
        return False
    return True
