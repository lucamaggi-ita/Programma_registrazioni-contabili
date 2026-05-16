from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fantozzi.excel_writer import HEADERS, LANCAMENTOS_SHEET
from fantozzi.operation_classifier import OperationAuditError, auditar_operacoes_especiais


BONIFICACAO_RECEBIDA = "BONIFICACAO_RECEBIDA"
BONIFICACAO_RECEBIDA_ACCOUNT_ID = "2.05.003"
TRATAMENTO_REDUCAO_CUSTO = "REDUCAO_CUSTO"

RESUMO_GERAL_SHEET = "Resumo_Geral"
RESUMO_MENSAL_SHEET = "Resumo_Mensal"
RESUMO_CONTA_SHEET = "Resumo_por_Conta"
RESUMO_FORNECEDOR_SHEET = "Resumo_por_Fornecedor"
RESUMO_CNPJ_SHEET = "Resumo_por_CNPJ"
LANCAMENTOS_OK_BASE_SHEET = "Lancamentos_OK_Base"
EXCLUIDOS_REVISAR_SHEET = "Excluidos_REVISAR"
EXCLUIDOS_OPERACOES_SHEET = "Excluidos_Operacoes_Especiais"
INSTRUCOES_SHEET = "Instrucoes"

SOURCE_HEADER = "arquivo_excel_origem"
IDENTIFIER_HEADER = "identificador_operacional"
OUTPUT_LANCAMENTO_HEADERS = [IDENTIFIER_HEADER, SOURCE_HEADER, *HEADERS]

RESUMO_GERAL_HEADERS = ["indicador", "valor", "observacao"]
RESUMO_MENSAL_HEADERS = [
    "identificador_operacional",
    "ano_mes",
    "quantidade_lancamentos_ok",
    "valor_total_ok",
]
RESUMO_CONTA_HEADERS = [
    "conta_id",
    "conta_nome",
    "quantidade_lancamentos",
    "valor_total",
    "percentual_sobre_total_ok",
]
RESUMO_FORNECEDOR_HEADERS = [
    "fornecedor",
    "quantidade_lancamentos",
    "valor_total",
    "principais_contas",
]
RESUMO_CNPJ_HEADERS = [
    "identificador_operacional",
    "quantidade_lancamentos_ok",
    "valor_total_ok",
]
OPERACOES_HEADERS = [
    "pasta_xml_origem",
    "tipo_operacao_especial",
    "confianca",
    "destino_gerencial_sugerido",
    "status_revisao_sugerido",
    "identificador_operacional",
    "fornecedor",
    "cnpj_fornecedor",
    "documento",
    "serie",
    "chave_nfe",
    "data_emissao",
    "natureza_operacao",
    "cfops_encontrados",
    "valor_total_nfe",
    "motivos",
    "alertas",
    "lancamentos_ok_excluidos",
]


class RelatoriosGerenciaisOkError(Exception):
    pass


@dataclass(frozen=True)
class RelatoriosGerenciaisOkResult:
    entradas: tuple[Path, ...]
    excel_paths: tuple[Path, ...]
    pastas_operacoes_especiais: tuple[Path, ...]
    xlsx_path: Path
    report_path: Path | None
    total_lancamentos_origem: int
    lancamentos_ok_incluidos: int
    lancamentos_revisar_excluidos: int
    operacoes_especiais_excluidas: int
    lancamentos_ok_excluidos_operacoes: int
    valor_total_incluido: Decimal
    valor_total_excluido_revisar: Decimal
    valor_total_excluido_operacoes: Decimal


def gerar_relatorios_gerenciais_ok_provisorios(
    entradas: Iterable[Path | str],
    xlsx_path: Path | str,
    report_path: Path | str | None = None,
    pastas_operacoes_especiais: Iterable[Path | str] | None = None,
) -> RelatoriosGerenciaisOkResult:
    entradas = tuple(Path(path) for path in entradas)
    xlsx_path = Path(xlsx_path)
    report_path = Path(report_path) if report_path else None
    pastas_operacoes = tuple(Path(path) for path in (pastas_operacoes_especiais or ()))
    excel_paths = _resolve_excel_paths(entradas)

    all_lancamentos: list[dict[str, object]] = []
    for excel_path in excel_paths:
        identifier = _derive_identifier(excel_path)
        for row in _read_lancamentos(excel_path):
            all_lancamentos.append(
                {IDENTIFIER_HEADER: identifier, SOURCE_HEADER: str(excel_path), **row}
            )

    operacoes = _read_operacoes_especiais(pastas_operacoes)
    special_keys = _special_keys(operacoes)
    included_ok: list[dict[str, object]] = []
    excluded_revisar: list[dict[str, object]] = []
    excluded_special_ok: list[dict[str, object]] = []

    for row in all_lancamentos:
        status = _text(row.get("status_revisao")).upper()
        if _is_bonificacao_recebida_decidida(row):
            included_ok.append(row)
            continue
        if status == "OK" and _lancamento_matches_special(row, special_keys):
            excluded_special_ok.append(row)
            continue
        if status != "OK":
            excluded_revisar.append(row)
            continue
        included_ok.append(row)

    operacoes_rows = _operation_rows(operacoes, excluded_special_ok)
    workbook = Workbook()
    _write_resumo_geral(
        workbook.active,
        all_lancamentos,
        included_ok,
        excluded_revisar,
        operacoes_rows,
        excluded_special_ok,
    )
    _write_table_sheet(
        workbook.create_sheet(RESUMO_MENSAL_SHEET),
        RESUMO_MENSAL_HEADERS,
        _resumo_mensal_rows(included_ok),
    )
    _write_table_sheet(
        workbook.create_sheet(RESUMO_CONTA_SHEET),
        RESUMO_CONTA_HEADERS,
        _resumo_conta_rows(included_ok),
    )
    _write_table_sheet(
        workbook.create_sheet(RESUMO_FORNECEDOR_SHEET),
        RESUMO_FORNECEDOR_HEADERS,
        _resumo_fornecedor_rows(included_ok),
    )
    _write_table_sheet(
        workbook.create_sheet(RESUMO_CNPJ_SHEET),
        RESUMO_CNPJ_HEADERS,
        _resumo_cnpj_rows(included_ok),
    )
    _write_table_sheet(
        workbook.create_sheet(LANCAMENTOS_OK_BASE_SHEET),
        OUTPUT_LANCAMENTO_HEADERS,
        included_ok,
    )
    _write_table_sheet(
        workbook.create_sheet(EXCLUIDOS_REVISAR_SHEET),
        OUTPUT_LANCAMENTO_HEADERS,
        excluded_revisar,
    )
    _write_table_sheet(
        workbook.create_sheet(EXCLUIDOS_OPERACOES_SHEET),
        OPERACOES_HEADERS,
        operacoes_rows,
    )
    _write_instrucoes_sheet(workbook.create_sheet(INSTRUCOES_SHEET), excel_paths)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)

    result = RelatoriosGerenciaisOkResult(
        entradas=entradas,
        excel_paths=excel_paths,
        pastas_operacoes_especiais=pastas_operacoes,
        xlsx_path=xlsx_path,
        report_path=report_path,
        total_lancamentos_origem=len(all_lancamentos),
        lancamentos_ok_incluidos=len(included_ok),
        lancamentos_revisar_excluidos=len(excluded_revisar),
        operacoes_especiais_excluidas=len(operacoes_rows),
        lancamentos_ok_excluidos_operacoes=len(excluded_special_ok),
        valor_total_incluido=_sum_values(included_ok, "valor"),
        valor_total_excluido_revisar=_sum_values(excluded_revisar, "valor"),
        valor_total_excluido_operacoes=_sum_values(operacoes_rows, "valor_total_nfe"),
    )
    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(
            relatorio_texto(result, included_ok),
            encoding="utf-8",
            newline="\n",
        )
    return result


def console_lines(result: RelatoriosGerenciaisOkResult) -> list[str]:
    lines = [
        "[OK] Relatorios gerenciais OK provisorios gerados.",
        f"     Excel de origem: {len(result.excel_paths)}",
        f"     Total lancamentos origem: {result.total_lancamentos_origem}",
        f"     Incluidos OK: {result.lancamentos_ok_incluidos}",
        f"     Excluidos REVISAR: {result.lancamentos_revisar_excluidos}",
        f"     Operacoes especiais excluidas: {result.operacoes_especiais_excluidas}",
        f"     Valor total incluido: {_format_currency(result.valor_total_incluido)}",
        f"     XLSX gerado: {result.xlsx_path}",
    ]
    if result.report_path:
        lines.append(f"     Relatorio: {result.report_path}")
    return lines


def relatorio_texto(
    result: RelatoriosGerenciaisOkResult,
    included_ok: list[dict[str, object]],
) -> str:
    lines = [
        "RELATORIOS GERENCIAIS OK PROVISORIOS",
        "====================================",
        "",
        "Status do relatorio: PROVISORIO",
        "",
        "Criterio:",
        "- Inclui status_revisao = OK",
        "- Inclui BONIFICACAO_RECEBIDA decidida em 2.05.003 como reducao de custo",
        "- Exclui REVISAR",
        "- Exclui COMODATO e operacoes especiais sem decisao contabil segura",
        "",
        "Resumo:",
        f"- Total lancamentos origem: {result.total_lancamentos_origem}",
        f"- Incluidos OK: {result.lancamentos_ok_incluidos}",
        f"- Excluidos REVISAR: {result.lancamentos_revisar_excluidos}",
        f"- Excluidos operacoes especiais: {result.operacoes_especiais_excluidas}",
        f"- Valor total incluido: {_format_currency(result.valor_total_incluido)}",
        f"- Valor total excluido REVISAR: {_format_currency(result.valor_total_excluido_revisar)}",
        f"- Valor total excluido operacoes especiais: {_format_currency(result.valor_total_excluido_operacoes)}",
        "",
        "Resumo por mes:",
        *_summary_lines(_resumo_mensal_rows(included_ok), ["identificador_operacional", "ano_mes", "quantidade_lancamentos_ok", "valor_total_ok"]),
        "",
        "Resumo por conta:",
        *_summary_lines(_resumo_conta_rows(included_ok)[:20], ["conta_id", "conta_nome", "quantidade_lancamentos", "valor_total"]),
        "",
        "Resumo por fornecedor:",
        *_summary_lines(_resumo_fornecedor_rows(included_ok)[:20], ["fornecedor", "quantidade_lancamentos", "valor_total"]),
        "",
        "Alertas:",
        "- Revisao contabil real pendente.",
        "- Proposta YAML sem regras aprovadas.",
        "- Nao usar como livro diario definitivo.",
        "- Nao usar como DRE definitiva.",
        "",
        "Resultado:",
        "OK",
    ]
    return "\n".join(lines) + "\n"


def _resolve_excel_paths(entradas: tuple[Path, ...]) -> tuple[Path, ...]:
    if not entradas:
        raise RelatoriosGerenciaisOkError("Informe ao menos uma entrada.")
    from fantozzi.pipeline_paths import (  # noqa: PLC0415
        PipelinePathsError,
        resolver_arquivos_lancamentos_xlsx,
    )
    excel_paths: list[Path] = []
    for entrada in entradas:
        try:
            excel_paths.extend(resolver_arquivos_lancamentos_xlsx(entrada))
        except PipelinePathsError as exc:
            raise RelatoriosGerenciaisOkError(str(exc)) from exc
    unique = tuple(dict.fromkeys(excel_paths))
    if not unique:
        raise RelatoriosGerenciaisOkError("Nenhum Excel .xlsx encontrado nas entradas.")
    return unique


def _read_lancamentos(excel_path: Path) -> list[dict[str, object]]:
    try:
        workbook = load_workbook(excel_path, data_only=True)
    except Exception as exc:
        raise RelatoriosGerenciaisOkError(
            f"Excel de lancamentos nao pode ser lido: {excel_path}"
        ) from exc
    if LANCAMENTOS_SHEET not in workbook.sheetnames:
        raise RelatoriosGerenciaisOkError(f"Aba obrigatoria ausente: {LANCAMENTOS_SHEET}")
    sheet = workbook[LANCAMENTOS_SHEET]
    headers = [cell.value for cell in sheet[1]]
    missing = [header for header in HEADERS if header not in headers]
    if missing:
        raise RelatoriosGerenciaisOkError(
            "Colunas obrigatorias ausentes em Lancamentos: " + ", ".join(missing)
        )
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(_has_value(value) for value in values):
            continue
        raw = dict(zip(headers, values))
        rows.append({header: raw.get(header) for header in HEADERS})
    return rows


def _read_operacoes_especiais(pastas: tuple[Path, ...]) -> list[dict[str, object]]:
    rows = []
    for pasta in pastas:
        try:
            result = auditar_operacoes_especiais(pasta)
        except OperationAuditError as exc:
            raise RelatoriosGerenciaisOkError(str(exc)) from exc
        for item in result.itens:
            nfe = item.nfe
            classificacao = item.classificacao
            cfops = ", ".join(
                sorted({produto.cfop for produto in nfe.produtos if produto.cfop})
            )
            rows.append(
                {
                    "pasta_xml_origem": str(pasta),
                    "tipo_operacao_especial": classificacao.tipo,
                    "confianca": classificacao.confianca,
                    "destino_gerencial_sugerido": classificacao.destino_gerencial,
                    "status_revisao_sugerido": classificacao.status_revisao,
                    "identificador_operacional": _derive_identifier(pasta),
                    "fornecedor": nfe.nome_emitente,
                    "cnpj_fornecedor": nfe.cnpj_emitente,
                    "documento": nfe.numero,
                    "serie": nfe.serie,
                    "chave_nfe": nfe.chave,
                    "data_emissao": nfe.data_emissao,
                    "natureza_operacao": nfe.natureza_operacao,
                    "cfops_encontrados": cfops,
                    "valor_total_nfe": nfe.valor_total.quantize(Decimal("0.01")),
                    "motivos": "; ".join(classificacao.motivos),
                    "alertas": "; ".join(classificacao.alertas),
                }
            )
    return rows


def _operation_rows(
    operacoes: list[dict[str, object]],
    excluded_special_ok: list[dict[str, object]],
) -> list[dict[str, object]]:
    counts = Counter(_text(row.get("chave_origem")) for row in excluded_special_ok)
    output = []
    for row in operacoes:
        chave = _text(row.get("chave_nfe"))
        output.append({**row, "lancamentos_ok_excluidos": counts.get(chave, 0)})
    return output


def _special_keys(operacoes: list[dict[str, object]]) -> set[tuple[str, str]]:
    keys = set()
    for row in operacoes:
        chave = _text(row.get("chave_nfe"))
        if chave:
            keys.add(("chave", chave))
        doc = _text(row.get("documento"))
        identifier = _text(row.get("identificador_operacional"))
        if doc and identifier:
            keys.add(("documento_identificador", f"{identifier}|{doc}"))
    return keys


def _lancamento_matches_special(
    row: dict[str, object],
    special_keys: set[tuple[str, str]],
) -> bool:
    chave = _text(row.get("chave_origem"))
    if chave and ("chave", chave) in special_keys:
        return True
    doc = _text(row.get("documento"))
    identifier = _text(row.get(IDENTIFIER_HEADER))
    return bool(
        doc
        and identifier
        and ("documento_identificador", f"{identifier}|{doc}") in special_keys
    )


def _is_bonificacao_recebida_decidida(row: dict[str, object]) -> bool:
    return (
        _text(row.get("tipo_operacao_especial")).upper() == BONIFICACAO_RECEBIDA
        and _text(row.get("conta_id")) == BONIFICACAO_RECEBIDA_ACCOUNT_ID
        and _text(row.get("tratamento_operacao_especial")).upper()
        == TRATAMENTO_REDUCAO_CUSTO
    )


def _write_resumo_geral(
    sheet,
    all_lancamentos: list[dict[str, object]],
    included_ok: list[dict[str, object]],
    excluded_revisar: list[dict[str, object]],
    operacoes_rows: list[dict[str, object]],
    excluded_special_ok: list[dict[str, object]],
) -> None:
    sheet.title = RESUMO_GERAL_SHEET
    rows = [
        {
            "indicador": "Total lancamentos origem",
            "valor": len(all_lancamentos),
            "observacao": "Quantidade total lida nos Excel de origem.",
        },
        {
            "indicador": "Lancamentos incluidos OK",
            "valor": len(included_ok),
            "observacao": "Status OK mais BONIFICACAO_RECEBIDA decidida em 2.05.003.",
        },
        {
            "indicador": "Lancamentos excluidos REVISAR",
            "valor": len(excluded_revisar),
            "observacao": "Transparencia dos itens que ainda dependem de revisao.",
        },
        {
            "indicador": "Operacoes especiais excluidas",
            "valor": len(operacoes_rows),
            "observacao": "Operacoes especiais auditadas; BONIFICACAO_RECEBIDA decidida entra na base gerencial.",
        },
        {
            "indicador": "Lancamentos OK excluidos por operacoes especiais",
            "valor": len(excluded_special_ok),
            "observacao": "Lancamentos OK vinculados a NF-e especiais identificadas.",
        },
        {
            "indicador": "Valor total incluido",
            "valor": _sum_values(included_ok, "valor"),
            "observacao": "Base gerencial provisoria OK.",
        },
        {
            "indicador": "Valor total excluido por revisao",
            "valor": _sum_values(excluded_revisar, "valor"),
            "observacao": "Nao entra neste relatorio provisorio.",
        },
        {
            "indicador": "Valor total excluido por operacoes especiais",
            "valor": _sum_values(operacoes_rows, "valor_total_nfe"),
            "observacao": "Controlado separadamente ate validacao contabil.",
        },
        {
            "indicador": "Status do relatorio",
            "valor": "PROVISORIO",
            "observacao": "Nao e DRE definitiva nem livro diario definitivo.",
        },
    ]
    _write_table_sheet(sheet, RESUMO_GERAL_HEADERS, rows)


def _write_table_sheet(sheet, headers: list[str], rows: list[dict[str, object]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header)) for header in headers])
    _style_table(sheet)
    _set_widths_for_headers(sheet, headers)


def _write_instrucoes_sheet(sheet, excel_paths: tuple[Path, ...]) -> None:
    lines = [
        "RELATORIOS GERENCIAIS OK PROVISORIOS",
        "",
        "Este relatorio e provisorio.",
        "Inclui lancamentos com status OK.",
        "Inclui BONIFICACAO_RECEBIDA decidida em 2.05.003 como reducao de custo.",
        "Exclui lancamentos em revisao.",
        "Exclui COMODATO e operacoes especiais sem decisao contabil segura.",
        "Nao substitui revisao contabil.",
        "Nao e DRE definitiva.",
        "Nao e livro diario definitivo.",
        "Serve para analise gerencial parcial e prudente.",
        "",
        "Arquivos Excel de origem:",
    ]
    lines.extend(f"- {path}" for path in excel_paths)
    for line in lines:
        sheet.append([line])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 120
    for row in sheet.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _resumo_mensal_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    grouped: dict[tuple[str, str], dict[str, object]] = {}
    for row in rows:
        key = (
            _text(row.get(IDENTIFIER_HEADER)),
            _text(row.get("competencia")) or _ano_mes(row.get("data_documento")),
        )
        bucket = grouped.setdefault(
            key,
            {
                "identificador_operacional": key[0],
                "ano_mes": key[1],
                "quantidade_lancamentos_ok": 0,
                "valor_total_ok": Decimal("0.00"),
            },
        )
        bucket["quantidade_lancamentos_ok"] = int(bucket["quantidade_lancamentos_ok"]) + 1
        bucket["valor_total_ok"] = bucket["valor_total_ok"] + _decimal_value(row.get("valor"))
    return sorted(grouped.values(), key=lambda row: (_text(row["identificador_operacional"]), _text(row["ano_mes"])))


def _resumo_conta_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    grouped: dict[tuple[str, str], dict[str, object]] = {}
    total = _sum_values(rows, "valor")
    for row in rows:
        key = (_text(row.get("conta_id")), _text(row.get("conta_nome")))
        bucket = grouped.setdefault(
            key,
            {
                "conta_id": key[0],
                "conta_nome": key[1],
                "quantidade_lancamentos": 0,
                "valor_total": Decimal("0.00"),
                "percentual_sobre_total_ok": Decimal("0.00"),
            },
        )
        bucket["quantidade_lancamentos"] = int(bucket["quantidade_lancamentos"]) + 1
        bucket["valor_total"] = bucket["valor_total"] + _decimal_value(row.get("valor"))
    for bucket in grouped.values():
        if total:
            bucket["percentual_sobre_total_ok"] = (
                _decimal_value(bucket["valor_total"]) / total * Decimal("100")
            ).quantize(Decimal("0.01"))
    return sorted(grouped.values(), key=lambda row: -_decimal_value(row["valor_total"]))


def _resumo_fornecedor_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    grouped: dict[str, dict[str, object]] = {}
    account_counts: dict[str, Counter[str]] = defaultdict(Counter)
    for row in rows:
        fornecedor = _text(row.get("fornecedor_ou_favorecido")) or "-"
        bucket = grouped.setdefault(
            fornecedor,
            {
                "fornecedor": fornecedor,
                "quantidade_lancamentos": 0,
                "valor_total": Decimal("0.00"),
                "principais_contas": "",
            },
        )
        bucket["quantidade_lancamentos"] = int(bucket["quantidade_lancamentos"]) + 1
        bucket["valor_total"] = bucket["valor_total"] + _decimal_value(row.get("valor"))
        account = _text(row.get("conta_nome")) or _text(row.get("conta_id"))
        if account:
            account_counts[fornecedor][account] += 1
    for fornecedor, bucket in grouped.items():
        bucket["principais_contas"] = "; ".join(
            account for account, _count in account_counts[fornecedor].most_common(3)
        )
    return sorted(grouped.values(), key=lambda row: -_decimal_value(row["valor_total"]))


def _resumo_cnpj_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    grouped: dict[str, dict[str, object]] = {}
    for row in rows:
        identifier = _text(row.get(IDENTIFIER_HEADER)) or "-"
        bucket = grouped.setdefault(
            identifier,
            {
                "identificador_operacional": identifier,
                "quantidade_lancamentos_ok": 0,
                "valor_total_ok": Decimal("0.00"),
            },
        )
        bucket["quantidade_lancamentos_ok"] = int(bucket["quantidade_lancamentos_ok"]) + 1
        bucket["valor_total_ok"] = bucket["valor_total_ok"] + _decimal_value(row.get("valor"))
    return sorted(grouped.values(), key=lambda row: _text(row["identificador_operacional"]))


def _summary_lines(rows: list[dict[str, object]], headers: list[str]) -> list[str]:
    if not rows:
        return ["-"]
    lines = []
    for row in rows:
        parts = []
        for header in headers:
            value = row.get(header)
            parts.append(_format_currency(value) if isinstance(value, Decimal) else _text(value))
        lines.append("- " + " | ".join(parts))
    return lines


def _style_table(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, (int, float)) and _looks_like_money_header(sheet, cell.column):
                cell.number_format = '#,##0.00'


def _looks_like_money_header(sheet, column: int) -> bool:
    header = _text(sheet.cell(row=1, column=column).value)
    return "valor" in header or "percentual" in header


def _set_widths_for_headers(sheet, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = _suggested_width(header)


def _suggested_width(header: str) -> int:
    widths = {
        IDENTIFIER_HEADER: 24,
        SOURCE_HEADER: 42,
        "data_documento": 16,
        "competencia": 14,
        "fornecedor_ou_favorecido": 42,
        "documento": 16,
        "conta_id": 16,
        "conta_nome": 36,
        "grupo_conta": 36,
        "descricao_itens": 52,
        "valor": 16,
        "arquivo_origem": 34,
        "chave_origem": 46,
        "status_revisao": 18,
        "alertas": 40,
        "observacoes": 42,
        "duplicatas_json": 40,
        "indicador": 40,
        "observacao": 74,
        "fornecedor": 42,
        "principais_contas": 56,
        "tipo_operacao_especial": 28,
        "natureza_operacao": 34,
        "motivos": 42,
    }
    return widths.get(header, max(14, min(len(header) + 4, 40)))


def _derive_identifier(path: Path) -> str:
    parts = [part for part in path.parts if part]
    for part in reversed(parts):
        if part.isdigit() and len(part) >= 8:
            return part
    parent = path.parent.name.strip()
    return parent or path.stem


def _ano_mes(value: object) -> str:
    text = _text(value)
    if len(text) >= 7:
        return text[:7]
    return "SEM_COMPETENCIA"


def _sum_values(rows: Iterable[dict[str, object]], field: str) -> Decimal:
    total = Decimal("0.00")
    for row in rows:
        total += _decimal_value(row.get(field))
    return total.quantize(Decimal("0.01"))


def _decimal_value(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0.00")
    if isinstance(value, Decimal):
        return value.quantize(Decimal("0.01"))
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _cell_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    return value


def _format_currency(value: object) -> str:
    decimal = _decimal_value(value)
    text = f"{decimal:.2f}"
    integer_part, decimal_part = text.split(".")
    groups = []
    while integer_part:
        groups.append(integer_part[-3:])
        integer_part = integer_part[:-3]
    return "R$ " + ".".join(reversed(groups)) + "," + decimal_part


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _has_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""
