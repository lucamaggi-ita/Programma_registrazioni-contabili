from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


_FONT_REGISTERED = False
_FONT_NORMAL = "Helvetica"
_FONT_BOLD = "Helvetica-Bold"

_COLOR_HEADER = colors.HexColor("#1F4E78")
_COLOR_GROUP = colors.HexColor("#D9E1F2")
_COLOR_SUBTOTAL = colors.HexColor("#EEF2F9")
_COLOR_WHITE = colors.white
_COLOR_GRID = colors.HexColor("#CCCCCC")
_COLOR_TEXT_DARK = colors.HexColor("#222222")
_COLOR_TEXT_META = colors.HexColor("#444444")
_COLOR_TEXT_FOOTER = colors.HexColor("#666666")

DRE_GERAL_SHEET = "DRE_Geral"
DRE_MENSAL_SHEET = "DRE_Mensal"

_DRE_XLSX_CANDIDATES = [
    Path("pacote_contabil_gerencial_provisorio") / "01_DRE_GERENCIAL_CONSOLIDADO.xlsx",
    Path("dre_gerencial_ok_provisoria.xlsx"),
]


class DrePdfGerencialError(Exception):
    pass


@dataclass(frozen=True)
class DrePdfResult:
    pdf_path: Path
    dre_xlsx_source: Path
    run_id: str | None
    periodos: tuple[str, ...]
    total_linhas: int
    status: str


def gerar_pdf_dre(
    run_dir: Path | str,
    output_dir: Path | str | None = None,
) -> DrePdfResult:
    run_dir = Path(run_dir)
    if not run_dir.exists() or not run_dir.is_dir():
        raise DrePdfGerencialError(f"run-dir nao encontrado: {run_dir}")

    dre_xlsx = _find_dre_xlsx(run_dir)
    dre_rows = _load_dre_geral(dre_xlsx)
    if not dre_rows:
        raise DrePdfGerencialError(f"Nenhuma linha DRE encontrada em: {dre_xlsx}")

    periodos = _load_periodos(dre_xlsx)
    manifest = _load_manifest(run_dir)
    run_id = manifest.get("run_id") or None

    if output_dir is None:
        output_dir = run_dir / "output_visual"
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    pdf_path = output_dir / "01_DRE_GERENCIAL.pdf"
    _render_pdf(pdf_path, dre_rows, periodos, run_id)

    return DrePdfResult(
        pdf_path=pdf_path,
        dre_xlsx_source=dre_xlsx,
        run_id=run_id,
        periodos=tuple(periodos),
        total_linhas=len(dre_rows),
        status="TESTE_TECNICO_NAO_OFICIAL",
    )


def console_lines(result: DrePdfResult) -> list[str]:
    lines = [
        "[OK] PDF DRE gerencial gerado.",
        f"     Status: {result.status}",
        f"     PDF: {result.pdf_path}",
        f"     Fonte DRE: {result.dre_xlsx_source}",
        f"     Linhas DRE: {result.total_linhas}",
    ]
    if result.run_id:
        lines.append(f"     Run: {result.run_id}")
    if result.periodos:
        lines.append(f"     Periodos: {' | '.join(result.periodos)}")
    return lines


def _find_dre_xlsx(run_dir: Path) -> Path:
    for candidate in _DRE_XLSX_CANDIDATES:
        p = run_dir / candidate
        if p.exists():
            return p
    raise DrePdfGerencialError(
        f"Nenhum arquivo DRE encontrado em: {run_dir}. "
        f"Candidatos: {[str(c) for c in _DRE_XLSX_CANDIDATES]}"
    )


def _load_dre_geral(xlsx_path: Path) -> list[dict]:
    try:
        wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
    except Exception as exc:
        raise DrePdfGerencialError(f"Nao foi possivel abrir DRE: {xlsx_path}") from exc
    if DRE_GERAL_SHEET not in wb.sheetnames:
        wb.close()
        raise DrePdfGerencialError(
            f"Aba '{DRE_GERAL_SHEET}' nao encontrada em: {xlsx_path}"
        )
    ws = wb[DRE_GERAL_SHEET]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not all_rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in all_rows[0]]
    result = []
    for row in all_rows[1:]:
        d = {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
        if any(v is not None for v in d.values()):
            result.append(d)
    return result


def _load_periodos(xlsx_path: Path) -> list[str]:
    try:
        wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
        if DRE_MENSAL_SHEET not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[DRE_MENSAL_SHEET]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return []
    if len(rows) < 2:
        return []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    if "ano_mes" not in headers:
        return []
    idx = headers.index("ano_mes")
    seen: list[str] = []
    for row in rows[1:]:
        v = str(row[idx]).strip() if idx < len(row) and row[idx] is not None else ""
        if v and v not in seen:
            seen.append(v)
    return sorted(seen)


def _load_manifest(run_dir: Path) -> dict:
    for name in (
        "manifest_run.json",
        str(
            Path("pacote_contabil_gerencial_provisorio")
            / "manifest_pacote_contabil_gerencial.json"
        ),
    ):
        p = run_dir / name
        if p.exists():
            try:
                return json.loads(p.read_text(encoding="utf-8"))
            except Exception:
                pass
    return {}


def _to_decimal(value: object) -> Decimal:
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value or "0"))
    except InvalidOperation:
        return Decimal("0")


def _fmt_brl(value: object) -> str:
    try:
        v = float(_to_decimal(value))
    except (TypeError, ValueError):
        return "R$ -"
    neg = v < 0
    abs_v = abs(v)
    s = f"{abs_v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ -{s}" if neg else f"R$ {s}"


def _ensure_fonts() -> None:
    global _FONT_REGISTERED, _FONT_NORMAL, _FONT_BOLD
    if _FONT_REGISTERED:
        return
    arial = Path(r"C:\Windows\Fonts\arial.ttf")
    arial_bold = Path(r"C:\Windows\Fonts\arialbd.ttf")
    if arial.exists() and arial_bold.exists():
        try:
            pdfmetrics.registerFont(TTFont("FantozziNormal", str(arial)))
            pdfmetrics.registerFont(TTFont("FantozziNormal-Bold", str(arial_bold)))
            _FONT_NORMAL = "FantozziNormal"
            _FONT_BOLD = "FantozziNormal-Bold"
        except Exception:
            pass
    _FONT_REGISTERED = True


def _render_pdf(
    pdf_path: Path,
    dre_rows: list[dict],
    periodos: list[str],
    run_id: str | None,
) -> None:
    _ensure_fonts()
    fn = _FONT_NORMAL
    fb = _FONT_BOLD

    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DreTitle",
        fontName=fb,
        fontSize=16,
        spaceAfter=6,
        textColor=_COLOR_HEADER,
    )
    meta_style = ParagraphStyle(
        "DreMeta",
        fontName=fn,
        fontSize=8,
        spaceAfter=2,
        textColor=_COLOR_TEXT_META,
    )
    footer_style = ParagraphStyle(
        "DreFooter",
        fontName=fn,
        fontSize=7,
        spaceAfter=2,
        textColor=_COLOR_TEXT_FOOTER,
    )

    story = []
    story.append(Paragraph("DRE Gerencial Provisória", title_style))

    periodo_str = " · ".join(periodos) if periodos else "—"
    story.append(Paragraph(f"Período: {periodo_str}", meta_style))
    story.append(
        Paragraph(
            f"Gerado em: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            meta_style,
        )
    )
    if run_id:
        story.append(Paragraph(f"Run: {run_id}", meta_style))
    story.append(
        Paragraph(
            "Status: PROVISÓRIO · GERENCIAL · TESTE TÉCNICO / NÃO OFICIAL",
            meta_style,
        )
    )
    story.append(Spacer(1, 0.4 * cm))

    col_w = [11.5 * cm, 4 * cm]
    table_data: list[list] = [["Grupo / Linha DRE", "Valor (R$)"]]

    # Base styles applied to all rows — overrides added per-row below
    row_styles: list[tuple] = [
        ("BACKGROUND", (0, 0), (-1, 0), _COLOR_HEADER),
        ("TEXTCOLOR", (0, 0), (-1, 0), _COLOR_WHITE),
        ("FONTNAME", (0, 0), (-1, -1), fn),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("FONTNAME", (0, 0), (-1, 0), fb),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.25, _COLOR_GRID),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]

    # Group rows preserving insertion order
    grupos: dict[str, list[dict]] = {}
    for row in dre_rows:
        grupo = str(row.get("grupo_dre") or "—")
        if grupo not in grupos:
            grupos[grupo] = []
        grupos[grupo].append(row)

    ri = 1
    grand_total = Decimal("0.00")

    for grupo, rows in grupos.items():
        table_data.append([grupo, ""])
        row_styles.extend([
            ("BACKGROUND", (0, ri), (-1, ri), _COLOR_GROUP),
            ("FONTNAME", (0, ri), (-1, ri), fb),
        ])
        ri += 1

        subtotal = Decimal("0.00")
        for row in rows:
            linha = str(row.get("linha_dre") or "—")
            valor = _to_decimal(row.get("valor"))
            subtotal += valor
            grand_total += valor
            table_data.append([f"  {linha}", _fmt_brl(valor)])
            ri += 1

        table_data.append([f"  Sub-total {grupo}", _fmt_brl(subtotal)])
        row_styles.extend([
            ("BACKGROUND", (0, ri), (-1, ri), _COLOR_SUBTOTAL),
            ("FONTNAME", (0, ri), (-1, ri), fb),
            ("LINEABOVE", (0, ri), (-1, ri), 0.5, colors.HexColor("#8EA9C1")),
        ])
        ri += 1

    table_data.append(["TOTAL GERAL", _fmt_brl(grand_total)])
    row_styles.extend([
        ("BACKGROUND", (0, ri), (-1, ri), _COLOR_HEADER),
        ("TEXTCOLOR", (0, ri), (-1, ri), _COLOR_WHITE),
        ("FONTNAME", (0, ri), (-1, ri), fb),
        ("FONTSIZE", (0, ri), (-1, ri), 10),
        ("LINEABOVE", (0, ri), (-1, ri), 1.0, _COLOR_HEADER),
    ])

    t = Table(table_data, colWidths=col_w)
    t.setStyle(TableStyle(row_styles))
    story.append(t)
    story.append(Spacer(1, 0.8 * cm))

    story.append(
        Paragraph(
            "Relatório gerencial provisório gerado automaticamente a partir da full run "
            "Fantozzi. Não substitui escrituração contábil oficial sem "
            "revisão do contador. TESTE TÉCNICO / NÃO OFICIAL",
            footer_style,
        )
    )

    doc.build(story)
