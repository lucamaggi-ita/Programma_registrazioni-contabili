from __future__ import annotations

import csv
import hashlib
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Mapping

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fantozzi.chart_of_accounts import ChartOfAccounts
from fantozzi.lancamentos_manuais_gerenciais import (
    CSV_OPTIONAL_HEADERS,
    CSV_REQUIRED_HEADERS,
    LancamentosManuaisGerenciaisError,
    carregar_lancamentos_manuais_gerenciais,
)


OPERATIVE_SHEET = "lancamentos_manuais"
CATEGORIES_SHEET = "categorias"
ACCOUNTS_SHEET = "plano_contas_operativo"
INSTRUCTIONS_SHEET = "instrucoes"

OPERATIVE_HEADERS = (
    "DATA",
    "CONTA_ESCOLHIDA",
    "CONTA_ID_APROVADA",
    "CONTA_NOME_APROVADA",
    "VALOR",
    "DESCRICAO_DA_DESPESA",
    "OBSERVACAO",
)
TECHNICAL_HEADERS = (
    "CONTA_CREDITO_ID",
    "CONTA_CREDITO_NOME",
    "STATUS_VALIDACAO",
    "MENSAGEM_VALIDACAO",
    "CHAVE_PREVIEW",
)
ALL_TEMPLATE_HEADERS = (*OPERATIVE_HEADERS, *TECHNICAL_HEADERS)

CATEGORY_HEADERS = (
    "CONTA_DISPLAY",
    "CONTA_ID",
    "CONTA_NOME",
)
ACCOUNT_HEADERS = (
    "CONTA_ID",
    "CONTA_NOME",
    "TIPO",
    "GRUPO",
    "ENTRA_DRE",
    "USO_OPERADOR",
)

DEFAULT_TEMPLATE_ROWS = 500

NON_SELECTABLE_TIPOS = frozenset({"tecnico", "revisao"})

TECHNICAL_COUNTERPART_CANDIDATES = (
    ("99.99.998", "Contrapartida Tecnica Lancamentos Manuais"),
    ("99.99.999", "Contrapartida Tecnica DRE Manual"),
    ("14.01.001", "Contrapartida Tecnica Lancamentos Manuais"),
)

GENERIC_TERMS = {
    "pagamento",
    "transferencia",
    "pix",
    "boleto",
    "fornecedor",
    "servico",
    "compra",
}


class LancamentosManuaisTemplateError(Exception):
    pass


@dataclass(frozen=True)
class AccountInfo:
    id: str
    nome: str
    grupo: str
    tipo: str
    entra_dre: bool


@dataclass(frozen=True)
class CategoryMapping:
    conta_debito_id: str
    conta_debito_nome: str
    ativa: bool = True


@dataclass(frozen=True)
class TechnicalCounterpart:
    conta_id: str
    conta_nome: str


@dataclass(frozen=True)
class TemplateResult:
    output_xlsx: Path
    categorias: tuple[CategoryMapping, ...]
    conta_credito_tecnica: TechnicalCounterpart | None
    blockers: tuple[str, ...]


@dataclass(frozen=True)
class ValidationResult:
    entrada: Path
    saida_csv: Path
    relatorio: Path | None
    linhas_lidas: int
    linhas_validas: int
    avisos: tuple[str, ...]
    erros: tuple[str, ...]


def gerar_template_lancamentos_manuais_gerenciais(
    output_xlsx_path: Path | str,
    plano_contas_path: Path | str,
    mapeamento_dre_path: Path | str,
) -> TemplateResult:
    output_xlsx_path = Path(output_xlsx_path)
    accounts = _load_accounts(Path(plano_contas_path), Path(mapeamento_dre_path))
    categories = _default_categories(accounts)
    technical_counterpart = _find_technical_counterpart(accounts)
    blockers = _template_blockers(categories, technical_counterpart)

    workbook = Workbook()
    default_sheet = workbook.active
    default_sheet.title = OPERATIVE_SHEET
    _build_operative_sheet(default_sheet, categories, technical_counterpart, blockers)
    _build_categories_sheet(workbook.create_sheet(CATEGORIES_SHEET), categories)
    _build_accounts_sheet(workbook.create_sheet(ACCOUNTS_SHEET), accounts)
    _build_instructions_sheet(workbook.create_sheet(INSTRUCTIONS_SHEET), blockers)

    output_xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx_path)
    return TemplateResult(
        output_xlsx=output_xlsx_path,
        categorias=categories,
        conta_credito_tecnica=technical_counterpart,
        blockers=blockers,
    )


def validar_lancamentos_manuais_gerenciais_operador(
    entrada_path: Path | str,
    saida_csv_path: Path | str,
    relatorio_path: Path | str | None,
    plano_contas_path: Path | str,
    mapeamento_dre_path: Path | str,
) -> ValidationResult:
    entrada_path = Path(entrada_path)
    saida_csv_path = Path(saida_csv_path)
    relatorio = Path(relatorio_path) if relatorio_path else None
    accounts = _load_accounts(Path(plano_contas_path), Path(mapeamento_dre_path))
    technical_counterpart = _find_technical_counterpart(accounts)

    if not entrada_path.exists():
        raise LancamentosManuaisTemplateError(f"Entrada nao encontrada: {entrada_path}")
    if entrada_path.suffix.lower() == ".xlsx":
        rows, categories = _read_workbook_input(entrada_path)
    elif entrada_path.suffix.lower() == ".csv":
        rows = _read_operator_csv_input(entrada_path)
        categories = _default_categories(accounts)
    else:
        raise LancamentosManuaisTemplateError("Entrada deve ser .xlsx ou .csv operacional.")

    technical_rows, warnings, errors = _validate_operator_rows(
        rows,
        categories,
        accounts,
        technical_counterpart,
    )

    if errors:
        result = ValidationResult(
            entrada=entrada_path,
            saida_csv=saida_csv_path,
            relatorio=relatorio,
            linhas_lidas=len(rows),
            linhas_validas=0,
            avisos=tuple(warnings),
            erros=tuple(errors),
        )
        if relatorio:
            _write_validation_report(result)
        raise LancamentosManuaisTemplateError(
            f"Validacao encontrou {len(errors)} erro(s) bloqueante(s)."
        )

    saida_csv_path.parent.mkdir(parents=True, exist_ok=True)
    _write_and_validate_technical_csv(
        saida_csv_path,
        technical_rows,
        Path(plano_contas_path),
        Path(mapeamento_dre_path),
    )
    result = ValidationResult(
        entrada=entrada_path,
        saida_csv=saida_csv_path,
        relatorio=relatorio,
        linhas_lidas=len(rows),
        linhas_validas=len(technical_rows),
        avisos=tuple(warnings),
        erros=(),
    )
    if relatorio:
        _write_validation_report(result)
    return result


def validation_report_text(result: ValidationResult) -> str:
    lines = [
        "FANTOZZI - VALIDACAO LANCAMENTOS MANUAIS GERENCIAIS",
        "===================================================",
        "",
        f"entrada: {result.entrada}",
        f"saida_csv: {result.saida_csv}",
        f"linhas_lidas: {result.linhas_lidas}",
        f"linhas_validas: {result.linhas_validas}",
        f"erros: {len(result.erros)}",
        f"avisos: {len(result.avisos)}",
        "",
        "ATENCAO: o sistema nao garante que esta despesa ja nao exista em XML.",
        "Verifique antes de incluir no DRE.",
        "",
        "Contrapartida tecnica:",
        "- conta neutra fora DRE, quando existente no plano.",
        "- nao representa caixa, banco, cartao ou conta patrimonial real.",
        "",
        "Erros bloqueantes:",
        *(f"- {error}" for error in result.erros),
        *(["- nenhum"] if not result.erros else []),
        "",
        "Avisos:",
        *(f"- {warning}" for warning in result.avisos),
        *(["- nenhum"] if not result.avisos else []),
        "",
        "resultado: " + ("OK" if not result.erros else "BLOQUEADO"),
    ]
    return "\n".join(lines) + "\n"


def _load_accounts(plano_contas_path: Path, mapeamento_dre_path: Path) -> tuple[AccountInfo, ...]:
    chart = ChartOfAccounts.load(plano_contas_path)
    errors = chart.validate()
    if errors:
        raise LancamentosManuaisTemplateError(
            "Plano de contas invalido: " + "; ".join(errors)
        )
    entra_dre_by_id = _load_entra_dre_map(mapeamento_dre_path)
    return tuple(
        AccountInfo(
            id=account.id,
            nome=account.nome,
            grupo=account.grupo,
            tipo=account.tipo,
            entra_dre=entra_dre_by_id.get(account.id, False),
        )
        for account in chart._ordered
    )


def _load_entra_dre_map(mapeamento_dre_path: Path) -> dict[str, bool]:
    data = yaml.safe_load(mapeamento_dre_path.read_text(encoding="utf-8")) or {}
    mapping = data.get("mapeamento", {}) or {}
    output: dict[str, bool] = {}
    if isinstance(mapping, dict):
        for account_id, row in mapping.items():
            if isinstance(row, dict):
                output[_text(account_id)] = bool(row.get("entra_dre", True))
    return output


def _default_categories(accounts: tuple[AccountInfo, ...]) -> tuple[CategoryMapping, ...]:
    return tuple(
        CategoryMapping(conta_debito_id=account.id, conta_debito_nome=account.nome)
        for account in accounts
        if account.tipo not in NON_SELECTABLE_TIPOS
    )


def _find_technical_counterpart(
    accounts: tuple[AccountInfo, ...],
) -> TechnicalCounterpart | None:
    by_id = {account.id: account for account in accounts}
    by_name = {_normalize_text(account.nome): account for account in accounts}
    for account_id, account_name in TECHNICAL_COUNTERPART_CANDIDATES:
        account = by_id.get(account_id) or by_name.get(_normalize_text(account_name))
        if account and _is_safe_technical_counterpart(account):
            return TechnicalCounterpart(account.id, account.nome)
    return None


def _is_safe_technical_counterpart(account: AccountInfo) -> bool:
    text = _normalize_text(f"{account.id} {account.nome} {account.grupo} {account.tipo}")
    forbidden = ("caixa", "banco", "cartao", "itau", "patrimonial")
    required_terms = ("contrapartida", "tecnica")
    return (
        not account.entra_dre
        and all(term in text for term in required_terms)
        and not any(term in text for term in forbidden)
    )


def _template_blockers(
    categories: tuple[CategoryMapping, ...],
    technical_counterpart: TechnicalCounterpart | None,
) -> tuple[str, ...]:
    blockers: list[str] = []
    if not categories:
        blockers.append("Nenhuma categoria operacional segura encontrada no plano de contas.")
    if technical_counterpart is None:
        blockers.append(
            "Nenhuma conta tecnica neutra fora DRE foi encontrada no plano de contas. "
            "Mini-fase necessaria: criar explicitamente uma conta de contrapartida "
            "tecnica auditavel para lancamentos manuais, sem simular caixa/banco/cartao."
        )
    return tuple(blockers)


def _build_operative_sheet(
    sheet,
    categories: tuple[CategoryMapping, ...],
    technical_counterpart: TechnicalCounterpart | None,
    blockers: tuple[str, ...],
) -> None:
    sheet.append(list(ALL_TEMPLATE_HEADERS))
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    # A-G visible operative; H-L hidden technical
    widths = {
        "A": 14,
        "B": 38,
        "C": 16,
        "D": 40,
        "E": 14,
        "F": 44,
        "G": 44,
        "H": 18,
        "I": 36,
        "J": 18,
        "K": 44,
        "L": 24,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for column in range(8, 13):
        sheet.column_dimensions[get_column_letter(column)].hidden = True
    sheet.freeze_panes = "A2"

    for row in range(2, DEFAULT_TEMPLATE_ROWS + 2):
        # C: CONTA_ID_APROVADA — derived from CONTA_ESCOLHIDA via VLOOKUP
        sheet[f"C{row}"] = f'=IFERROR(VLOOKUP(B{row},{CATEGORIES_SHEET}!$A:$B,2,FALSE),"")'
        # D: CONTA_NOME_APROVADA — derived from CONTA_ESCOLHIDA via VLOOKUP
        sheet[f"D{row}"] = f'=IFERROR(VLOOKUP(B{row},{CATEGORIES_SHEET}!$A:$C,3,FALSE),"")'
        # H-I: technical counterpart (hidden)
        sheet[f"H{row}"] = technical_counterpart.conta_id if technical_counterpart else ""
        sheet[f"I{row}"] = technical_counterpart.conta_nome if technical_counterpart else ""
        # J: status (hidden)
        sheet[f"J{row}"] = '=IF(COUNTA(A{0}:G{0})=0,"","PENDENTE")'.format(row)
        # K: mensagem (hidden)
        if blockers:
            sheet[f"K{row}"] = "BLOCKER: " + " | ".join(blockers)
        else:
            sheet[f"K{row}"] = '=IF(J{0}="","",IF(C{0}<>"","OK","REVISAR"))'.format(row)
        # L: chave preview — data|conta_escolhida|descricao|valor (hidden)
        sheet[f"L{row}"] = (
            f'=IF(COUNTA(A{row}:G{row})=0,"",'
            f'TEXT(A{row},"yyyy-mm-dd")&"|"&B{row}&"|"&F{row}&"|"&E{row})'
        )

    if categories:
        dv_categories = DataValidation(
            type="list",
            formula1=f"={CATEGORIES_SHEET}!$A$2:$A${len(categories) + 1}",
            allow_blank=True,
        )
        sheet.add_data_validation(dv_categories)
        dv_categories.add(f"B2:B{DEFAULT_TEMPLATE_ROWS + 1}")


def _build_categories_sheet(sheet, categories: tuple[CategoryMapping, ...]) -> None:
    sheet.append(list(CATEGORY_HEADERS))
    for category in categories:
        display = f"{category.conta_debito_id} - {category.conta_debito_nome}"
        sheet.append([display, category.conta_debito_id, category.conta_debito_nome])
    _format_lookup_sheet(sheet)


def _build_accounts_sheet(sheet, accounts: tuple[AccountInfo, ...]) -> None:
    sheet.append(list(ACCOUNT_HEADERS))
    for account in accounts:
        usage = "categoria" if account.entra_dre else "apoio nao totalizavel"
        sheet.append(
            [
                account.id,
                account.nome,
                account.tipo,
                account.grupo,
                "SIM" if account.entra_dre else "NAO",
                usage,
            ]
        )
    _format_lookup_sheet(sheet)


def _build_instructions_sheet(sheet, blockers: tuple[str, ...]) -> None:
    lines = [
        "Template lancamentos manuais gerenciais",
        "",
        "1. Informar o que foi comprado/contratado, nao como foi pago.",
        "2. Nao inserir despesas ja presentes nos XML.",
        "3. Nao inserir comprovantes que sao apenas pagamento de fatura XML ja registrada.",
        "4. Usar valor positivo.",
        "5. Nao deixar linhas vazias no meio.",
        "6. Campo OBSERVACAO pode indicar comprovante, arquivo_origem ou payment_id.",
        "7. A contrapartida tecnica deve existir no plano, ficar fora DRE e nao representar caixa, banco ou cartao.",
    ]
    if blockers:
        lines.extend(["", "BLOCKERS:"])
        lines.extend(f"- {blocker}" for blocker in blockers)
    for line in lines:
        sheet.append([line])
    sheet.column_dimensions["A"].width = 110
    sheet["A1"].font = Font(bold=True)


def _format_lookup_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="5B9BD5")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"
    for index, width in enumerate((30, 18, 40, 56, 12, 18), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _read_workbook_input(
    workbook_path: Path,
) -> tuple[tuple[dict[str, str], ...], tuple[CategoryMapping, ...]]:
    workbook = load_workbook(workbook_path, data_only=True)
    if OPERATIVE_SHEET not in workbook.sheetnames:
        raise LancamentosManuaisTemplateError(f"Aba obrigatoria ausente: {OPERATIVE_SHEET}")
    rows = _read_sheet_rows(workbook[OPERATIVE_SHEET], OPERATIVE_HEADERS)
    categories = _read_categories_sheet(workbook)
    workbook.close()
    return rows, categories


def _read_sheet_rows(sheet, headers: tuple[str, ...]) -> tuple[dict[str, str], ...]:
    header_values = [_text(cell.value) for cell in sheet[1]]
    missing = [header for header in headers if header not in header_values]
    if missing:
        raise LancamentosManuaisTemplateError(
            "Colunas obrigatorias ausentes no template: " + ", ".join(missing)
        )
    index_by_header = {header: header_values.index(header) + 1 for header in headers}
    last_non_empty = 1
    for row_index in range(2, sheet.max_row + 1):
        if any(_text(sheet.cell(row_index, index_by_header[h]).value) for h in headers):
            last_non_empty = row_index
    rows = []
    for row_index in range(2, last_non_empty + 1):
        row = {
            header: _cell_text(sheet.cell(row_index, index_by_header[header]).value)
            for header in headers
        }
        row["_linha_excel"] = str(row_index)
        rows.append(row)
    return tuple(rows)


def _read_categories_sheet(workbook) -> tuple[CategoryMapping, ...]:
    if CATEGORIES_SHEET not in workbook.sheetnames:
        raise LancamentosManuaisTemplateError(f"Aba obrigatoria ausente: {CATEGORIES_SHEET}")
    rows = _read_lookup_rows(workbook[CATEGORIES_SHEET], CATEGORY_HEADERS)
    return tuple(
        CategoryMapping(
            conta_debito_id=row["CONTA_ID"],
            conta_debito_nome=row["CONTA_NOME"],
        )
        for row in rows
        if row["CONTA_DISPLAY"]
    )


def _read_lookup_rows(sheet, headers: tuple[str, ...]) -> tuple[dict[str, str], ...]:
    header_values = [_text(cell.value) for cell in sheet[1]]
    missing = [header for header in headers if header not in header_values]
    if missing:
        raise LancamentosManuaisTemplateError(
            f"Aba {sheet.title}: colunas ausentes: " + ", ".join(missing)
        )
    index_by_header = {header: header_values.index(header) + 1 for header in headers}
    rows = []
    for row_index in range(2, sheet.max_row + 1):
        row = {
            header: _cell_text(sheet.cell(row_index, index_by_header[header]).value)
            for header in headers
        }
        if any(row.values()):
            rows.append(row)
    return tuple(rows)


def _read_operator_csv_input(csv_path: Path) -> tuple[dict[str, str], ...]:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file, delimiter=";")
        headers = tuple(reader.fieldnames or ())
        missing = [header for header in OPERATIVE_HEADERS if header not in headers]
        if missing:
            raise LancamentosManuaisTemplateError(
                "CSV operacional sem colunas obrigatorias: " + ", ".join(missing)
            )
        rows = []
        for line_number, row in enumerate(reader, start=2):
            values = {header: _text(row.get(header)) for header in OPERATIVE_HEADERS}
            values["_linha_excel"] = str(line_number)
            rows.append(values)
        return tuple(rows)


def _validate_operator_rows(
    rows: tuple[dict[str, str], ...],
    categories: tuple[CategoryMapping, ...],
    accounts: tuple[AccountInfo, ...],
    technical_counterpart: TechnicalCounterpart | None,
) -> tuple[tuple[dict[str, str], ...], list[str], list[str]]:
    active_categories = {
        f"{item.conta_debito_id} - {item.conta_debito_nome}": item
        for item in categories
        if item.ativa
    }
    account_ids = {account.id for account in accounts}
    technical_rows: list[dict[str, str]] = []
    warnings: list[str] = []
    errors: list[str] = []
    seen_keys: dict[str, str] = {}
    seen_non_empty = False
    seen_empty_after_data = False
    if technical_counterpart is None:
        errors.append(
            "Nenhuma conta tecnica neutra fora DRE foi encontrada no plano de contas. "
            "Execute mini-fase separada para cria-la explicitamente."
        )
        return (), warnings, errors

    for index, row in enumerate(rows, start=2):
        line = row.get("_linha_excel", str(index))
        empty = _is_operator_row_empty(row)
        if empty:
            if seen_non_empty:
                seen_empty_after_data = True
            continue
        if seen_empty_after_data:
            errors.append(f"Linha {line}: linha vazia no meio dos dados.")
        seen_non_empty = True

        row_errors, row_warnings = _validate_required_operator_fields(row, line)
        errors.extend(row_errors)
        warnings.extend(row_warnings)
        if row_errors:
            continue

        category = active_categories.get(row["CONTA_ESCOLHIDA"])
        if not category:
            errors.append(
                f"Linha {line}: CONTA_ESCOLHIDA inexistente ou inativa: "
                f"{row['CONTA_ESCOLHIDA']}"
            )
            continue
        if category.conta_debito_id not in account_ids:
            errors.append(
                f"Linha {line}: conta debito inexistente: {category.conta_debito_id}"
            )
            continue
        conta_id_row = row.get("CONTA_ID_APROVADA", "")
        conta_nome_row = row.get("CONTA_NOME_APROVADA", "")
        if conta_id_row and conta_id_row != category.conta_debito_id:
            errors.append(
                f"Linha {line}: CONTA_ID_APROVADA '{conta_id_row}' "
                f"inconsistente com CONTA_ESCOLHIDA (esperado: {category.conta_debito_id})."
            )
            continue
        if conta_nome_row and _normalize_text(conta_nome_row) != _normalize_text(
            category.conta_debito_nome
        ):
            errors.append(
                f"Linha {line}: CONTA_NOME_APROVADA '{conta_nome_row}' "
                f"inconsistente com CONTA_ESCOLHIDA (esperado: {category.conta_debito_nome})."
            )
            continue

        value = _parse_positive_decimal(row["VALOR"], line, errors)
        parsed_date = _parse_iso_date(row["DATA"], line, errors)
        if value is None or parsed_date is None:
            continue

        technical = {
            "data": parsed_date.isoformat(),
            "descricao_da_despesa": row["DESCRICAO_DA_DESPESA"],
            "conta_debito_id": category.conta_debito_id,
            "conta_debito_nome": category.conta_debito_nome,
            "conta_credito_id": technical_counterpart.conta_id,
            "conta_credito_nome": technical_counterpart.conta_nome,
            "valor": _decimal_text(value),
            "observacao": row.get("OBSERVACAO", ""),
        }
        key = _technical_preview_key(technical)
        previous = seen_keys.get(key)
        if previous:
            errors.append(f"Linha {line}: duplicata interna com a linha {previous}.")
            continue
        seen_keys[key] = line
        technical_rows.append(technical)

    if not technical_rows and not errors:
        errors.append("Nenhuma linha operacional preenchida.")
    return tuple(technical_rows), warnings, errors


def _validate_required_operator_fields(
    row: Mapping[str, str],
    line: str,
) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    required = ("DATA", "CONTA_ESCOLHIDA", "VALOR", "DESCRICAO_DA_DESPESA")
    for header in required:
        if not row.get(header):
            errors.append(f"Linha {line}: campo obrigatorio vazio: {header}.")
    if _normalize_text(row.get("DESCRICAO_DA_DESPESA")) in GENERIC_TERMS:
        warnings.append(f"Linha {line}: DESCRICAO_DA_DESPESA muito generica.")
    return errors, warnings


def _write_and_validate_technical_csv(
    saida_csv_path: Path,
    rows: tuple[dict[str, str], ...],
    plano_contas_path: Path,
    mapeamento_dre_path: Path,
) -> None:
    headers = [*CSV_REQUIRED_HEADERS, *CSV_OPTIONAL_HEADERS]
    with NamedTemporaryFile(
        "w",
        encoding="utf-8-sig",
        newline="",
        suffix=".csv",
        dir=str(saida_csv_path.parent),
        delete=False,
    ) as temp_file:
        temp_path = Path(temp_file.name)
        writer = csv.DictWriter(temp_file, fieldnames=headers, delimiter=";")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
    try:
        carregar_lancamentos_manuais_gerenciais(
            temp_path,
            plano_contas_path,
            mapeamento_dre_path,
        )
        saida_csv_path.write_text(
            temp_path.read_text(encoding="utf-8-sig"),
            encoding="utf-8-sig",
            newline="",
        )
    except LancamentosManuaisGerenciaisError as exc:
        raise LancamentosManuaisTemplateError(str(exc)) from exc
    finally:
        temp_path.unlink(missing_ok=True)


def _write_validation_report(result: ValidationResult) -> None:
    if not result.relatorio:
        return
    result.relatorio.parent.mkdir(parents=True, exist_ok=True)
    result.relatorio.write_text(validation_report_text(result), encoding="utf-8")


def _parse_iso_date(value: str, line: str, errors: list[str]) -> date | None:
    try:
        return date.fromisoformat(value)
    except ValueError:
        errors.append(f"Linha {line}: DATA invalida, use YYYY-MM-DD.")
        return None


def _parse_positive_decimal(value: str, line: str, errors: list[str]) -> Decimal | None:
    text = value.replace("R$", "").replace(" ", "").strip()
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        decimal = Decimal(text).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        errors.append(f"Linha {line}: VALOR impossivel de parsear.")
        return None
    if decimal <= 0:
        errors.append(f"Linha {line}: VALOR deve ser maior que zero.")
        return None
    return decimal


def _technical_preview_key(row: Mapping[str, str]) -> str:
    payload = "|".join(
        row.get(key, "")
        for key in ("data", "descricao_da_despesa", "conta_debito_id", "valor")
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _is_operator_row_empty(row: Mapping[str, str]) -> bool:
    return not any(_text(row.get(header)) for header in OPERATIVE_HEADERS)


def _is_active(value: str) -> bool:
    return _normalize_text(value or "SIM") in {"sim", "s", "true", "1", "ativo"}


def _cell_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _decimal_text(value: Decimal) -> str:
    return f"{value.quantize(Decimal('0.01')):.2f}"


def _normalize_text(value: object) -> str:
    return " ".join(_text(value).casefold().split())


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()
