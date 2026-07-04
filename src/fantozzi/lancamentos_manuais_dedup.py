from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook

from fantozzi.excel_writer import HEADERS, LANCAMENTOS_SHEET


AUDITORIA_POSSIVEIS_DUPLICATAS_MANUAIS_XML_FILE = (
    "07_AUDITORIA_POSSIVEIS_DUPLICATAS_MANUAIS_XML.txt"
)
MANUAL_GERENCIAL = "MANUAL_GERENCIAL"
XML_NFE = "XML_NFE"


@dataclass(frozen=True)
class LancamentoDedupRow:
    path: Path
    linha_excel: int
    data: date
    valor_abs: Decimal
    origem: str
    tipo_documento: str
    fornecedor_ou_favorecido: str
    descricao_itens: str
    documento: str
    arquivo_origem: str
    chave_origem: str
    hash_registro: str
    tipo_lancamento: str


@dataclass(frozen=True)
class PossivelDuplicataManualXml:
    manual: LancamentoDedupRow
    xml_candidatos: tuple[LancamentoDedupRow, ...]


@dataclass(frozen=True)
class AuditoriaPossiveisDuplicatasManuaisXmlResult:
    output_path: Path
    possiveis_duplicatas: int
    manuais_lidos: int
    xml_lidos: int
    arquivos_xml_lidos: int
    avisos: tuple[str, ...]
    matches: tuple[PossivelDuplicataManualXml, ...]


def auditar_possiveis_duplicatas_manuais_xml(
    manual_xlsx_path: Path | str,
    xml_xlsx_paths: Iterable[Path | str],
    output_path: Path | str,
) -> AuditoriaPossiveisDuplicatasManuaisXmlResult:
    manual_xlsx_path = Path(manual_xlsx_path)
    xml_paths = tuple(Path(path) for path in xml_xlsx_paths)
    output_path = Path(output_path)

    manual_rows, manual_warnings = _read_lancamento_rows(
        manual_xlsx_path,
        allowed_origens={MANUAL_GERENCIAL},
    )
    xml_rows: list[LancamentoDedupRow] = []
    xml_warnings: list[str] = []
    for xml_path in xml_paths:
        rows, warnings = _read_lancamento_rows(xml_path, allowed_origens={XML_NFE})
        xml_rows.extend(rows)
        xml_warnings.extend(warnings)

    manual_rows = _dedupe_manual_rows(manual_rows)
    matches = _find_matches(manual_rows, xml_rows)
    avisos = (*manual_warnings, *xml_warnings)
    _write_audit(output_path, manual_xlsx_path, xml_paths, manual_rows, xml_rows, matches, avisos)
    return AuditoriaPossiveisDuplicatasManuaisXmlResult(
        output_path=output_path,
        possiveis_duplicatas=len(matches),
        manuais_lidos=len(manual_rows),
        xml_lidos=len(xml_rows),
        arquivos_xml_lidos=len(xml_paths),
        avisos=tuple(avisos),
        matches=matches,
    )


def _read_lancamento_rows(
    xlsx_path: Path,
    *,
    allowed_origens: frozenset[str] | set[str],
) -> tuple[tuple[LancamentoDedupRow, ...], tuple[str, ...]]:
    if not xlsx_path.exists() or not xlsx_path.is_file():
        return (), (f"arquivo_nao_encontrado: {xlsx_path}",)

    try:
        workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception as exc:
        return (), (f"arquivo_nao_lido: {xlsx_path} ({exc})",)

    try:
        if LANCAMENTOS_SHEET not in workbook.sheetnames:
            return (), (f"aba_lancamentos_ausente: {xlsx_path}",)
        sheet = workbook[LANCAMENTOS_SHEET]
        headers = [_text(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        missing = [header for header in ("data_documento", "valor", "origem") if header not in headers]
        if missing:
            return (), (f"colunas_ausentes: {xlsx_path} ({', '.join(missing)})",)
        rows: list[LancamentoDedupRow] = []
        indexes = {header: index for index, header in enumerate(headers)}
        for linha_excel, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            raw = _row_mapping(values, indexes)
            origem = _text(raw.get("origem")).upper()
            tipo_documento = _text(raw.get("tipo_documento")).upper()
            if origem not in allowed_origens and tipo_documento not in allowed_origens:
                continue
            data = _parse_date(raw.get("data_documento"))
            valor = _parse_money(raw.get("valor"))
            if data is None or valor is None:
                continue
            rows.append(
                LancamentoDedupRow(
                    path=xlsx_path,
                    linha_excel=linha_excel,
                    data=data,
                    valor_abs=abs(valor).quantize(Decimal("0.01")),
                    origem=origem,
                    tipo_documento=tipo_documento,
                    fornecedor_ou_favorecido=_text(raw.get("fornecedor_ou_favorecido")),
                    descricao_itens=_text(raw.get("descricao_itens")),
                    documento=_text(raw.get("documento")),
                    arquivo_origem=_text(raw.get("arquivo_origem")),
                    chave_origem=_text(raw.get("chave_origem")),
                    hash_registro=_text(raw.get("hash_registro")),
                    tipo_lancamento=_text(raw.get("tipo_lancamento")),
                )
            )
        return tuple(rows), ()
    finally:
        workbook.close()


def _row_mapping(values: tuple[object, ...], indexes: Mapping[str, int]) -> dict[str, object]:
    return {
        header: values[index] if index < len(values) else None
        for header, index in indexes.items()
        if header in HEADERS
    }


def _dedupe_manual_rows(
    rows: tuple[LancamentoDedupRow, ...],
) -> tuple[LancamentoDedupRow, ...]:
    selected: dict[str, LancamentoDedupRow] = {}
    for row in rows:
        key = _manual_key(row)
        current = selected.get(key)
        if current is None or _manual_priority(row) < _manual_priority(current):
            selected[key] = row
    return tuple(selected[key] for key in sorted(selected))


def _manual_key(row: LancamentoDedupRow) -> str:
    chave = row.chave_origem
    for suffix in (":D", ":C"):
        if chave.endswith(suffix):
            return chave[: -len(suffix)]
    if chave:
        return chave
    return "|".join(
        [
            row.data.isoformat(),
            f"{row.valor_abs:.2f}",
            row.fornecedor_ou_favorecido.casefold(),
            row.arquivo_origem,
        ]
    )


def _manual_priority(row: LancamentoDedupRow) -> int:
    if row.tipo_lancamento.endswith("_TOTALIZAVEL"):
        return 0
    if row.chave_origem.endswith(":D"):
        return 1
    return 2


def _find_matches(
    manual_rows: tuple[LancamentoDedupRow, ...],
    xml_rows: list[LancamentoDedupRow],
) -> tuple[PossivelDuplicataManualXml, ...]:
    xml_by_key: dict[tuple[date, Decimal], list[LancamentoDedupRow]] = {}
    for xml_row in xml_rows:
        xml_by_key.setdefault((xml_row.data, xml_row.valor_abs), []).append(xml_row)

    matches = []
    for manual_row in manual_rows:
        candidatos = tuple(
            sorted(
                xml_by_key.get((manual_row.data, manual_row.valor_abs), ()),
                key=lambda row: (str(row.path).lower(), row.linha_excel),
            )
        )
        if candidatos:
            matches.append(
                PossivelDuplicataManualXml(
                    manual=manual_row,
                    xml_candidatos=candidatos,
                )
            )
    return tuple(matches)


def _write_audit(
    output_path: Path,
    manual_xlsx_path: Path,
    xml_paths: tuple[Path, ...],
    manual_rows: tuple[LancamentoDedupRow, ...],
    xml_rows: list[LancamentoDedupRow],
    matches: tuple[PossivelDuplicataManualXml, ...],
    avisos: tuple[str, ...],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "FANTOZZI - AUDITORIA POSSIVEIS DUPLICATAS MANUAIS XML",
        "====================================================",
        "",
        "status: ATENCAO" if matches else "status: OK",
        f"possiveis_duplicatas_manuais_xml: {len(matches)}",
        f"manual_xlsx: {manual_xlsx_path}",
        f"arquivos_xml_lidos: {len(xml_paths)}",
        f"lancamentos_manuais_lidos: {len(manual_rows)}",
        f"lancamentos_xml_lidos: {len(xml_rows)}",
        "",
        "Contrato:",
        "- Controle informativo, prudencial e nao blocante.",
        "- Luca considera o caso Manuali XML raro na pratica.",
        "- Criterio 002B: data exata + valor absoluto identico a 2 decimais.",
        "- Nao exclui linhas, nao altera giornale/DRE, nao usa fuzzy, nao usa cross-run.",
        "- Nao compara banca/comprovantes e nao trata CFOP 1949.",
        "",
        "Arquivos XML considerados:",
    ]
    lines.extend(f"- {path}" for path in xml_paths)
    lines.extend(["", "Avisos de leitura:"])
    lines.extend((f"- {aviso}" for aviso in avisos) if avisos else ["- nenhum"])
    lines.extend(["", "Possiveis duplicatas:"])
    if not matches:
        lines.append("- nenhuma possivel duplicata encontrada")
    for index, match in enumerate(matches, start=1):
        manual = match.manual
        lines.extend(
            [
                f"{index}. POSSIVEL_DUPLICADO_MANUAL_XML",
                f"   data: {manual.data.isoformat()}",
                f"   valor: {manual.valor_abs:.2f}",
                f"   manual_linha_excel: {manual.linha_excel}",
                f"   manual_descricao: {_clip(manual.fornecedor_ou_favorecido or manual.descricao_itens)}",
                f"   manual_observacao: {_clip(manual.descricao_itens)}",
                f"   manual_chave_origem: {_clip(manual.chave_origem, 96)}",
                f"   xml_candidatos: {len(match.xml_candidatos)}",
            ]
        )
        for xml_row in match.xml_candidatos:
            lines.extend(
                [
                    f"   - xml_linha_excel: {xml_row.linha_excel}",
                    f"     xml_fornecedor: {_clip(xml_row.fornecedor_ou_favorecido)}",
                    f"     xml_documento: {_clip(xml_row.documento, 80)}",
                    f"     xml_descricao: {_clip(xml_row.descricao_itens)}",
                    f"     xml_arquivo: {_clip(str(xml_row.path), 160)}",
                    f"     xml_chave_origem: {_clip(xml_row.chave_origem, 96)}",
                ]
            )
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


def _parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _parse_money(value: object) -> Decimal | None:
    if isinstance(value, Decimal):
        return value.quantize(Decimal("0.01"))
    if isinstance(value, int | float):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = _text(value).replace("R$", "").replace(" ", "")
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def _clip(value: object, limit: int = 120) -> str:
    text = _text(value)
    if len(text) <= limit:
        return text or "-"
    return text[: limit - 3] + "..."


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()
