from __future__ import annotations

import csv
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fantozzi.classification_review_export import CSV_HEADERS


RESUMO_SHEET = "Resumo"
PRIORIDADES_SHEET = "Prioridades"
TOP_VALOR_SHEET = "Top_Valor"
TOP_OCORRENCIAS_SHEET = "Top_Ocorrencias"
POR_FORNECEDOR_SHEET = "Por_Fornecedor"
POSSIVEIS_ESPECIAIS_SHEET = "Possiveis_Especiais"
PARA_CONTADOR_SHEET = "Para_Contador"
BASE_COMPLETA_SHEET = "Base_Completa"
INSTRUCOES_SHEET = "Instrucoes"

CALCULATED_HEADERS = [
    "identificador_operacional",
    "prioridade_revisao",
    "motivo_prioridade",
    "faixa_valor",
    "valor_total_num",
    "ocorrencias_num",
    "requer_contador",
    "possivel_operacao_especial",
    "sugestao_tratamento",
]

PACKAGE_HEADERS = CALCULATED_HEADERS + CSV_HEADERS
SUMMARY_HEADERS = ["indicador", "categoria", "quantidade", "valor_total"]
FORNECEDOR_HEADERS = [
    "fornecedores",
    "quantidade_candidatos",
    "valor_total",
    "ocorrencias_total",
    "prioridade_maxima",
]

PRIORITY_ORDER = {"ALTA": 0, "MEDIA": 1, "BAIXA": 2}
SPECIAL_ORDER = {
    "COMODATO": 0,
    "BONIFICACAO_RECEBIDA": 1,
    "OUTRA_OPERACAO_ESPECIAL_REVISAR": 2,
    "NAO_IDENTIFICADA": 3,
}


class ClassificationReviewPackageError(Exception):
    pass


@dataclass(frozen=True)
class ClassificationReviewPackageResult:
    csv_paths: tuple[Path, ...]
    xlsx_path: Path
    report_path: Path | None
    total_candidatos: int
    total_valor: Decimal
    por_identificador: dict[str, int]
    por_prioridade: dict[str, int]
    por_acao_manual: dict[str, int]
    por_status_revisao: dict[str, int]
    linhas_para_contador: int
    possiveis_operacoes_especiais: int


def gerar_pacote_revisao_contabil(
    csv_paths: list[Path] | tuple[Path, ...],
    xlsx_path: Path,
    report_path: Path | None = None,
    *,
    limite_top: int = 50,
    separador: str = ";",
) -> ClassificationReviewPackageResult:
    if not csv_paths:
        raise ClassificationReviewPackageError("Informe ao menos um CSV de revisao.")
    if limite_top < 1:
        raise ClassificationReviewPackageError("--limite-top deve ser maior que zero.")
    if len(separador) != 1:
        raise ClassificationReviewPackageError("--separador deve ter exatamente 1 caractere.")

    rows: list[dict[str, object]] = []
    for csv_path in csv_paths:
        rows.extend(_read_csv(Path(csv_path), separador))

    rows = [_enrich_row(row) for row in rows]
    workbook = Workbook()
    _create_resumo_sheet(workbook, rows)
    _create_table_sheet(workbook, PRIORIDADES_SHEET, _sort_by_priority(rows), PACKAGE_HEADERS)
    _create_table_sheet(
        workbook,
        TOP_VALOR_SHEET,
        sorted(rows, key=lambda row: row["valor_total_num"], reverse=True)[:limite_top],
        PACKAGE_HEADERS,
    )
    _create_table_sheet(
        workbook,
        TOP_OCORRENCIAS_SHEET,
        sorted(rows, key=lambda row: row["ocorrencias_num"], reverse=True)[:limite_top],
        PACKAGE_HEADERS,
    )
    _create_table_sheet(
        workbook,
        POR_FORNECEDOR_SHEET,
        _fornecedor_rows(rows),
        FORNECEDOR_HEADERS,
    )
    _create_table_sheet(
        workbook,
        POSSIVEIS_ESPECIAIS_SHEET,
        [
            row
            for row in rows
            if row["possivel_operacao_especial"] != "NAO_IDENTIFICADA"
            or _text(row.get("alertas"))
        ],
        PACKAGE_HEADERS,
    )
    _create_table_sheet(
        workbook,
        PARA_CONTADOR_SHEET,
        [row for row in rows if row["requer_contador"] == "SIM"],
        PACKAGE_HEADERS,
    )
    _create_table_sheet(workbook, BASE_COMPLETA_SHEET, rows, PACKAGE_HEADERS)
    _create_instrucoes_sheet(workbook)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)

    result = _build_result(tuple(Path(path) for path in csv_paths), xlsx_path, report_path, rows)
    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(relatorio_pacote(result, rows), encoding="utf-8")
    return result


def console_lines(result: ClassificationReviewPackageResult) -> list[str]:
    lines = [
        "[OK] Pacote assistido para revisao contabil gerado.",
        f"     Candidatos: {result.total_candidatos}",
        f"     Valor total: {_format_decimal_br(result.total_valor)}",
        f"     Para contador: {result.linhas_para_contador}",
        f"     Possiveis operacoes especiais: {result.possiveis_operacoes_especiais}",
        f"     XLSX gerado: {result.xlsx_path}",
    ]
    if result.report_path:
        lines.append(f"     Relatorio: {result.report_path}")
    return lines


def relatorio_pacote(
    result: ClassificationReviewPackageResult, rows: list[dict[str, object]]
) -> str:
    top_valor = sorted(rows, key=lambda row: row["valor_total_num"], reverse=True)[:10]
    top_ocorrencias = sorted(
        rows, key=lambda row: row["ocorrencias_num"], reverse=True
    )[:10]
    para_contador = [row for row in rows if row["requer_contador"] == "SIM"]
    possiveis = [
        row
        for row in rows
        if row["possivel_operacao_especial"] != "NAO_IDENTIFICADA"
    ]
    special_counts = Counter(_text(row["possivel_operacao_especial"]) for row in possiveis)
    contador_motivos = Counter(_text(row["possivel_operacao_especial"]) for row in para_contador)
    fiscal_counts = Counter(
        _text(row.get("tipo_operacao_especial")) or "NAO_INFORMADO"
        for row in rows
    )

    lines = [
        "PACOTE ASSISTIDO PARA REVISAO CONTABIL HUMANA",
        "================================================",
        "",
        "Arquivos CSV lidos:",
        *[f"- {path}" for path in result.csv_paths],
        "",
        f"Total de candidatos: {result.total_candidatos}",
        "",
        "Total por identificador:",
        *_counter_lines(result.por_identificador),
        "",
        "Total por prioridade:",
        *_counter_lines(result.por_prioridade),
        "",
        "Total por acao_manual:",
        *_counter_lines(result.por_acao_manual),
        "",
        "Total por status_revisao:",
        *_counter_lines(result.por_status_revisao),
        "",
        f"Valor total geral: {_format_decimal_br(result.total_valor)}",
        "",
        "Valor total por identificador:",
        *_decimal_total_lines(rows, "identificador_operacional"),
        "",
        "Valor total por prioridade:",
        *_decimal_total_lines(rows, "prioridade_revisao"),
        "",
        "Top 10 por valor:",
        *_top_lines(top_valor, "valor_total"),
        "",
        "Top 10 por ocorrencias:",
        *_top_lines(top_ocorrencias, "ocorrencias"),
        "",
        "Para contador:",
        f"- total de linhas: {len(para_contador)}",
        f"- valor total: {_format_decimal_br(sum((_as_decimal(row['valor_total_num']) for row in para_contador), Decimal('0')))}",
        "- principais motivos:",
        *_counter_lines(dict(contador_motivos)),
        "",
        "Possiveis operacoes especiais:",
        f"- total: {len(possiveis)}",
        "- por tipo:",
        *_counter_lines(dict(special_counts)),
        "",
        "Metadados fiscais nos CSV:",
        "- tipo_operacao_especial:",
        *_counter_lines(dict(fiscal_counts)),
        "",
        "Resultado: OK",
    ]
    return "\n".join(lines) + "\n"


def _read_csv(csv_path: Path, separador: str) -> list[dict[str, object]]:
    if not csv_path.exists():
        raise ClassificationReviewPackageError(f"CSV de revisao nao encontrado: {csv_path}")
    if not csv_path.is_file():
        raise ClassificationReviewPackageError(f"CSV de revisao nao e um arquivo: {csv_path}")
    if csv_path.stat().st_size == 0:
        raise ClassificationReviewPackageError(f"CSV de revisao esta vazio: {csv_path}")

    try:
        with csv_path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.DictReader(file, delimiter=separador)
            if not reader.fieldnames:
                raise ClassificationReviewPackageError(
                    f"CSV de revisao nao possui cabecalho valido: {csv_path}"
                )
            headers = [header.strip() for header in reader.fieldnames if header]
            missing = [header for header in CSV_HEADERS if header not in headers]
            if missing:
                raise ClassificationReviewPackageError(
                    "Colunas obrigatorias ausentes em "
                    f"{csv_path}: {', '.join(missing)}"
                )
            identificador = _derive_identifier(csv_path)
            rows = []
            for row in reader:
                item = {header: _text(row.get(header)) for header in CSV_HEADERS}
                item["identificador_operacional"] = identificador
                rows.append(item)
            return rows
    except UnicodeDecodeError as exc:
        raise ClassificationReviewPackageError(
            f"CSV de revisao nao pode ser lido em UTF-8: {csv_path}"
        ) from exc


def _enrich_row(row: dict[str, object]) -> dict[str, object]:
    valor = _parse_decimal(_text(row.get("valor_total")))
    ocorrencias = _parse_int(_text(row.get("ocorrencias")))
    special = _detect_special_operation(row)
    requer_contador = _requires_accountant(row, special)
    prioridade, motivos = _priority(row, valor, ocorrencias, special)

    enriched = dict(row)
    enriched.update(
        {
            "prioridade_revisao": prioridade,
            "motivo_prioridade": "; ".join(motivos),
            "faixa_valor": _value_range(valor),
            "valor_total_num": valor,
            "ocorrencias_num": ocorrencias,
            "requer_contador": requer_contador,
            "possivel_operacao_especial": special,
            "sugestao_tratamento": _treatment_suggestion(
                prioridade, requer_contador, special, _text(row.get("acao_manual"))
            ),
        }
    )
    return enriched


def _derive_identifier(csv_path: Path) -> str:
    parent = csv_path.parent.name.strip()
    if parent and parent.lower() != "importacao_real":
        return parent
    return csv_path.stem


def _priority(
    row: dict[str, object],
    valor: Decimal,
    ocorrencias: int,
    special: str,
) -> tuple[str, list[str]]:
    motivos: list[str] = []
    if valor >= Decimal("5000"):
        motivos.append("Valor alto")
    if ocorrencias >= 10:
        motivos.append("Muitas ocorrencias")
    if _has_priority_alert(row):
        motivos.append("Alerta relevante")
    if special != "NAO_IDENTIFICADA":
        motivos.append("Possivel operacao especial")

    if motivos:
        return "ALTA", motivos

    motivos = []
    if valor >= Decimal("1000"):
        motivos.append("Valor medio")
    if ocorrencias >= 3:
        motivos.append("Ocorrencias medias")
    if _text(row.get("alertas")):
        motivos.append("Possui alertas")
    if _is_review_account(row):
        motivos.append("Conta atual de revisao")
    if motivos:
        return "MEDIA", motivos

    return "BAIXA", ["Baixo impacto"]


def _is_review_account(row: dict[str, object]) -> bool:
    account_id = _text(row.get("conta_atual_id"))
    account_name = _normalize(_text(row.get("conta_atual_nome")))
    if not account_id:
        return True
    return account_id.startswith("13.") or "REVISAR" in account_name


def _has_priority_alert(row: dict[str, object]) -> bool:
    text = _normalize(_text(row.get("alertas")))
    if not text:
        return False
    return any(
        token in text
        for token in (
            "VALOR INCONSISTENTE",
            "CLASSIFICACAO AMBIGUA",
            "COMODATO",
            "BONIFICACAO",
            "OPERACAO ESPECIAL",
        )
    )


def _value_range(valor: Decimal) -> str:
    if valor >= Decimal("10000"):
        return ">= 10000"
    if valor >= Decimal("5000"):
        return "5000-9999"
    if valor >= Decimal("1000"):
        return "1000-4999"
    if valor >= Decimal("100"):
        return "100-999"
    return "< 100"


def _detect_special_operation(row: dict[str, object]) -> str:
    text = _combined_normalized_text(row)
    if any(token in text for token in ("COMODATO", "POSSIVEL COMODATO", "BENS TERCEIROS")):
        return "COMODATO"
    if any(
        token in text
        for token in (
            "BONIFICACAO",
            "BONIFICACAO RECEBIDA",
            "DOACAO",
            "BRINDE",
        )
    ):
        return "BONIFICACAO_RECEBIDA"
    if "OPERACAO ESPECIAL" in text or "OPERACOES ESPECIAIS" in text:
        return "OUTRA_OPERACAO_ESPECIAL_REVISAR"
    return "NAO_IDENTIFICADA"


def _requires_accountant(row: dict[str, object], special: str) -> str:
    text = _combined_normalized_text(row)
    action = _normalize(_text(row.get("acao_manual")))
    if action == "ANALISAR_COM_CONTADOR":
        return "SIM"
    if special != "NAO_IDENTIFICADA":
        return "SIM"
    if any(token in text for token in ("DUVIDA", "CONTADOR", "FISCAL", "AMBIGU")):
        return "SIM"
    return "NAO"


def _treatment_suggestion(
    prioridade: str,
    requer_contador: str,
    special: str,
    acao_manual: str,
) -> str:
    if requer_contador == "SIM" or special != "NAO_IDENTIFICADA":
        return "Revisar com contador"
    if acao_manual == "MANTER_REVISAO":
        return "Manter em revisao"
    if prioridade == "BAIXA":
        return "Baixo impacto: revisar depois"
    return "Priorizar classificacao"


def _combined_normalized_text(row: dict[str, object]) -> str:
    fields = [
        "descricao_normalizada",
        "descricao_exemplo",
        "fornecedores",
        "documentos",
        "conta_atual_nome",
        "status_revisao",
        "alertas",
        "cfop",
        "natureza_operacao",
        "tipo_operacao_especial",
        "confianca_operacao_especial",
        "tratamento_operacao_especial",
        "acao_manual",
        "observacao_manual",
    ]
    return _normalize(" ".join(_text(row.get(field)) for field in fields))


def _create_resumo_sheet(workbook: Workbook, rows: list[dict[str, object]]) -> None:
    sheet = workbook.active
    sheet.title = RESUMO_SHEET
    sheet.append(SUMMARY_HEADERS)
    for row in _summary_rows(rows):
        sheet.append(row)
    _style_table_sheet(sheet, SUMMARY_HEADERS)
    _set_widths(sheet, {"A": 38, "B": 36, "C": 16, "D": 18})


def _summary_rows(rows: list[dict[str, object]]) -> list[list[object]]:
    summary: list[list[object]] = [
        ["Total de candidatos", "Geral", len(rows), _sum_values(rows)],
        [
            "Quantidade requer contador",
            "SIM",
            sum(1 for row in rows if row["requer_contador"] == "SIM"),
            _sum_values([row for row in rows if row["requer_contador"] == "SIM"]),
        ],
        [
            "Possiveis operacoes especiais",
            "Total",
            sum(1 for row in rows if row["possivel_operacao_especial"] != "NAO_IDENTIFICADA"),
            _sum_values(
                [
                    row
                    for row in rows
                    if row["possivel_operacao_especial"] != "NAO_IDENTIFICADA"
                ]
            ),
        ],
    ]
    for field, label in (
        ("identificador_operacional", "Total por identificador"),
        ("prioridade_revisao", "Total por prioridade"),
        ("acao_manual", "Total por acao_manual"),
        ("status_revisao", "Total por status_revisao"),
    ):
        for value in sorted({_text(row.get(field)) or "-" for row in rows}):
            selected = [row for row in rows if (_text(row.get(field)) or "-") == value]
            summary.append([label, value, len(selected), _sum_values(selected)])
    return summary


def _create_table_sheet(
    workbook: Workbook,
    title: str,
    rows: list[dict[str, object]],
    headers: list[str],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header)) for header in headers])
    _style_table_sheet(sheet, headers)


def _create_instrucoes_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(INSTRUCOES_SHEET)
    instructions = [
        "Pacote assistido para revisao contabil humana",
        "",
        "Este arquivo e um apoio a revisao. Ele nao aplica classificacao automaticamente.",
        "Use a coluna prioridade_revisao para escolher por onde comecar.",
        "ALTA indica maior impacto, maior recorrencia, alertas ou possivel operacao especial.",
        "MEDIA indica impacto intermediario. BAIXA indica menor impacto operacional.",
        "",
        "Depois da analise, preencha as planilhas assistidas originais.",
        "Campos principais: acao_manual, conta_id_sugerida, conta_nome_sugerida, palavras_chave_sugeridas e observacao_manual.",
        "",
        "CLASSIFICAR: propor classificacao direta em conta existente.",
        "CRIAR_ALIAS: propor alias para uma conta existente.",
        "CRIAR_REGRA: propor regra baseada em palavras-chave.",
        "ANALISAR_COM_CONTADOR: manter para validacao contabil/fiscal.",
        "IGNORAR: nao gerar regra nem classificacao.",
        "MANTER_REVISAO: deixar pendente para revisao futura.",
        "",
        "COMODATO e BONIFICACAO_RECEBIDA exigem cuidado e validacao com contador.",
        "Este pacote nao modifica CSV, XML, Excel de lancamentos ou arquivos YAML de config.",
    ]
    for line in instructions:
        sheet.append([line])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 120
    for row in sheet.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _style_table_sheet(sheet, headers: list[str]) -> None:
    _style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = _suggested_width(header)
        for cell in sheet.iter_cols(
            min_col=index,
            max_col=index,
            min_row=2,
            max_row=max(sheet.max_row, 2),
        ):
            for item in cell:
                item.alignment = Alignment(wrap_text=_should_wrap(header), vertical="top")
                if header in {"valor_total_num", "valor_total", "valor_total"}:
                    item.number_format = '#,##0.00'
                if header in {"ocorrencias_num", "ocorrencias_total", "quantidade_candidatos"}:
                    item.number_format = '0'


def _style_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _set_widths(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _suggested_width(header: str) -> int:
    widths = {
        "identificador_operacional": 24,
        "prioridade_revisao": 18,
        "motivo_prioridade": 38,
        "faixa_valor": 14,
        "valor_total_num": 16,
        "ocorrencias_num": 16,
        "requer_contador": 18,
        "possivel_operacao_especial": 30,
        "sugestao_tratamento": 28,
        "descricao_normalizada": 34,
        "descricao_exemplo": 42,
        "fornecedores": 38,
        "documentos": 26,
        "valor_total": 16,
        "conta_atual_nome": 34,
        "alertas": 38,
        "cfop": 16,
        "natureza_operacao": 34,
        "tipo_operacao_especial": 30,
        "confianca_operacao_especial": 28,
        "tratamento_operacao_especial": 30,
        "acao_manual": 28,
        "observacao_manual": 44,
    }
    return widths.get(header, max(14, min(len(header) + 4, 40)))


def _should_wrap(header: str) -> bool:
    return header in {
        "motivo_prioridade",
        "sugestao_tratamento",
        "descricao_normalizada",
        "descricao_exemplo",
        "fornecedores",
        "documentos",
        "conta_atual_nome",
        "alertas",
        "natureza_operacao",
        "tipo_operacao_especial",
        "confianca_operacao_especial",
        "tratamento_operacao_especial",
        "observacao_manual",
    }


def _fornecedor_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        grouped[_text(row.get("fornecedores")) or "-"].append(row)
    output = []
    for fornecedor, items in grouped.items():
        output.append(
            {
                "fornecedores": fornecedor,
                "quantidade_candidatos": len(items),
                "valor_total": _sum_values(items),
                "ocorrencias_total": sum(int(item["ocorrencias_num"]) for item in items),
                "prioridade_maxima": min(
                    (_text(item["prioridade_revisao"]) for item in items),
                    key=lambda value: PRIORITY_ORDER.get(value, 99),
                ),
            }
        )
    return sorted(
        output,
        key=lambda row: (
            PRIORITY_ORDER.get(_text(row["prioridade_maxima"]), 99),
            -_as_decimal(row["valor_total"]),
            _text(row["fornecedores"]),
        ),
    )


def _sort_by_priority(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    return sorted(
        rows,
        key=lambda row: (
            PRIORITY_ORDER.get(_text(row["prioridade_revisao"]), 99),
            -_as_decimal(row["valor_total_num"]),
            -int(row["ocorrencias_num"]),
            _text(row["descricao_normalizada"]),
        ),
    )


def _build_result(
    csv_paths: tuple[Path, ...],
    xlsx_path: Path,
    report_path: Path | None,
    rows: list[dict[str, object]],
) -> ClassificationReviewPackageResult:
    return ClassificationReviewPackageResult(
        csv_paths=csv_paths,
        xlsx_path=xlsx_path,
        report_path=report_path,
        total_candidatos=len(rows),
        total_valor=_sum_values(rows),
        por_identificador=dict(Counter(_text(row["identificador_operacional"]) for row in rows)),
        por_prioridade=dict(Counter(_text(row["prioridade_revisao"]) for row in rows)),
        por_acao_manual=dict(Counter(_text(row.get("acao_manual")) or "-" for row in rows)),
        por_status_revisao=dict(Counter(_text(row.get("status_revisao")) or "-" for row in rows)),
        linhas_para_contador=sum(1 for row in rows if row["requer_contador"] == "SIM"),
        possiveis_operacoes_especiais=sum(
            1 for row in rows if row["possivel_operacao_especial"] != "NAO_IDENTIFICADA"
        ),
    )


def _top_lines(rows: list[dict[str, object]], mode: str) -> list[str]:
    lines = []
    for row in rows:
        value = (
            _format_decimal_br(_as_decimal(row["valor_total_num"]))
            if mode == "valor_total"
            else str(row["ocorrencias_num"])
        )
        lines.append(
            "- "
            f"{_text(row.get('descricao_normalizada'))} | "
            f"{_text(row.get('fornecedores'))} | "
            f"{value} | "
            f"{_text(row.get('prioridade_revisao'))}"
        )
    return lines or ["-"]


def _counter_lines(counter: dict[str, int]) -> list[str]:
    if not counter:
        return ["-"]
    return [f"- {key}: {value}" for key, value in sorted(counter.items())]


def _decimal_total_lines(rows: list[dict[str, object]], field: str) -> list[str]:
    totals: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for row in rows:
        totals[_text(row.get(field)) or "-"] += _as_decimal(row["valor_total_num"])
    return [f"- {key}: {_format_decimal_br(value)}" for key, value in sorted(totals.items())]


def _sum_values(rows: list[dict[str, object]]) -> Decimal:
    return sum((_as_decimal(row["valor_total_num"]) for row in rows), Decimal("0"))


def _cell_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    return value


def _as_decimal(value: object) -> Decimal:
    if isinstance(value, Decimal):
        return value
    return _parse_decimal(_text(value))


def _parse_decimal(value: str) -> Decimal:
    value = value.strip()
    if not value:
        return Decimal("0")
    cleaned = value.replace("R$", "").replace(" ", "")
    if "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return Decimal("0")


def _parse_int(value: str) -> int:
    try:
        return int(value.strip())
    except ValueError:
        return 0


def _format_decimal_br(value: Decimal) -> str:
    text = f"{value.quantize(Decimal('0.01')):.2f}"
    integer_part, decimal_part = text.split(".")
    groups = []
    while integer_part:
        groups.append(integer_part[-3:])
        integer_part = integer_part[:-3]
    return "R$ " + ".".join(reversed(groups)) + "," + decimal_part


def _normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    ascii_text = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join(ascii_text.upper().replace("_", " ").split())


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()
