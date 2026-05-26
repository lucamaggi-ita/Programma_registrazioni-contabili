from __future__ import annotations

import csv
import hashlib
import json
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from fantozzi.banking.payment_rules import normalize_match_text
from fantozzi.banking.revisao_banca import (
    COMPROVANTES_SHEET,
    CONTA_CORRENTE_SHEET,
    TECNICO_CONTA_CORRENTE_SHEET,
)
from fantozzi.excel_writer import write_lancamentos_to_excel
from fantozzi.models import Lancamento


REJEITADOS_CSV_NAME = "lancamentos_banca_rejeitados.csv"
MANIFEST_NAME = "manifest_lancamentos_banca.json"


class LancamentosBancaError(ValueError):
    pass


@dataclass(frozen=True)
class _RejectedRow:
    fonte: str
    sheet: str
    linha: int
    motivo: str
    status_decisao: str
    acao_regra: str
    risco_duplicidade_xml: str
    data: str
    valor: str
    beneficiario: str
    conta_id: str
    conta_nome: str
    fonte_prioritaria: str = ""
    chave_dedup: str = ""
    incluido: str = "false"


@dataclass(frozen=True)
class _DedupKey:
    data: str
    valor_abs: str
    beneficiario_normalizado: str

    def as_text(self) -> str:
        return "|".join(
            [
                f"data={self.data}",
                f"valor_abs={self.valor_abs}",
                f"beneficiario={self.beneficiario_normalizado}",
            ]
        )


def gerar_lancamentos_banca_aprovados(
    comprovantes_xlsx: Path,
    conta_corrente_xlsx: Path,
    output_xlsx: Path,
    periodo_ini: date,
    periodo_fim: date,
    *,
    comprovantes_csv: Path | None = None,
    conta_corrente_csv: Path | None = None,
) -> Path:
    if periodo_ini > periodo_fim:
        raise LancamentosBancaError("periodo_ini nao pode ser maior que periodo_fim.")

    comprovantes_xlsx = Path(comprovantes_xlsx)
    conta_corrente_xlsx = Path(conta_corrente_xlsx)
    output_xlsx = Path(output_xlsx)
    _validate_input_workbook(comprovantes_xlsx, "comprovantes_xlsx")
    _validate_input_workbook(conta_corrente_xlsx, "conta_corrente_xlsx")

    lancamentos: list[Lancamento] = []
    rejeitados: list[_RejectedRow] = []
    manifest = _base_manifest(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        output_xlsx,
        periodo_ini,
        periodo_fim,
    )

    if comprovantes_csv is not None:
        lancamentos.extend(
            _read_comprovantes_csv(
                Path(comprovantes_csv),
                periodo_ini=periodo_ini,
                periodo_fim=periodo_fim,
                manifest=manifest,
                rejeitados=rejeitados,
            )
        )
    if conta_corrente_csv is not None:
        lancamentos.extend(
            _read_conta_corrente_csv(
                Path(conta_corrente_csv),
                periodo_ini=periodo_ini,
                periodo_fim=periodo_fim,
                manifest=manifest,
                rejeitados=rejeitados,
            )
        )
    lancamentos.extend(
        _read_comprovantes(
            comprovantes_xlsx,
            periodo_ini=periodo_ini,
            periodo_fim=periodo_fim,
            manifest=manifest,
            rejeitados=rejeitados,
        )
    )
    lancamentos.extend(
        _read_conta_corrente(
            conta_corrente_xlsx,
            periodo_ini=periodo_ini,
            periodo_fim=periodo_fim,
            manifest=manifest,
            rejeitados=rejeitados,
        )
    )
    lancamentos = _deduplicate_lancamentos(lancamentos, rejeitados, manifest)

    if output_xlsx.exists():
        output_xlsx.unlink()
    write_lancamentos_to_excel(output_xlsx, lancamentos)

    output_rejeitados = output_xlsx.with_name(REJEITADOS_CSV_NAME)
    manifest_path = output_xlsx.with_name(MANIFEST_NAME)
    _write_rejeitados(output_rejeitados, rejeitados)
    manifest["linhas_aprovadas_incluidas"] = len(lancamentos)
    manifest["output_rejeitados_csv"] = str(output_rejeitados)
    manifest["manifest_path"] = str(manifest_path)
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return output_xlsx


def _read_comprovantes_csv(
    path: Path,
    *,
    periodo_ini: date,
    periodo_fim: date,
    manifest: dict[str, object],
    rejeitados: list[_RejectedRow],
) -> list[Lancamento]:
    lancamentos: list[Lancamento] = []
    rows_lidas = 0

    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row_number, row in enumerate(reader, start=2):
            rows_lidas += 1

            if row.get("status_classificacao") != "CLASSIFICADO_POR_REGRA":
                manifest["linhas_rejeitadas_status"] = int(manifest["linhas_rejeitadas_status"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="COMPROVANTES_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="STATUS_NAO_APROVADO",
                    date_header="data_pagamento",
                    beneficiary_header="beneficiario_normalizado",
                )
                continue
            if str(row.get("entra_dre", "")).lower() != "true":
                manifest["linhas_excluidas_escludere"] = int(manifest["linhas_excluidas_escludere"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="COMPROVANTES_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="ENTRA_DRE_FALSE",
                    date_header="data_pagamento",
                    beneficiary_header="beneficiario_normalizado",
                )
                continue

            try:
                data_documento = _parse_date(row.get("data_pagamento"))
            except LancamentosBancaError:
                manifest["linhas_rejeitadas_dados_obrigatorios"] = int(manifest["linhas_rejeitadas_dados_obrigatorios"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="COMPROVANTES_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="SEM_DATA",
                    date_header="data_pagamento",
                    beneficiary_header="beneficiario_normalizado",
                )
                continue
            try:
                valor_assinado = _parse_decimal(row.get("valor"))
                valor = abs(valor_assinado)
            except LancamentosBancaError:
                manifest["linhas_rejeitadas_dados_obrigatorios"] = int(manifest["linhas_rejeitadas_dados_obrigatorios"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="COMPROVANTES_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="SEM_VALOR",
                    date_header="data_pagamento",
                    beneficiary_header="beneficiario_normalizado",
                )
                continue

            if not _text(row.get("beneficiario_normalizado")):
                manifest["linhas_rejeitadas_dados_obrigatorios"] = int(manifest["linhas_rejeitadas_dados_obrigatorios"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="COMPROVANTES_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="DADOS_OBRIGATORIOS_AUSENTES",
                    date_header="data_pagamento",
                    beneficiary_header="beneficiario_normalizado",
                )
                continue
            if not _text(row.get("conta_id")) or not _text(row.get("conta_nome")):
                manifest["linhas_rejeitadas_dados_obrigatorios"] = int(manifest["linhas_rejeitadas_dados_obrigatorios"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="COMPROVANTES_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="SEM_CONTA_ID",
                    date_header="data_pagamento",
                    beneficiary_header="beneficiario_normalizado",
                )
                continue
            if data_documento < periodo_ini or data_documento > periodo_fim:
                manifest["linhas_fora_periodo"] = int(manifest["linhas_fora_periodo"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="COMPROVANTES_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="FORA_PERIODO",
                    date_header="data_pagamento",
                    beneficiary_header="beneficiario_normalizado",
                )
                continue

            raw_hash = _text(row.get("raw_hash") or row.get("payment_id"))
            chave_origem = _chave_origem(
                "COMPROVANTES_CSV",
                row_number,
                raw_hash,
                _text(row.get("regra_id")),
                data_documento,
                valor,
            )
            lancamentos.append(
                Lancamento(
                    data_documento=data_documento,
                    competencia=f"{data_documento.year:04d}-{data_documento.month:02d}",
                    cnpj_empresa=_text(row.get("cnpj")),
                    nome_empresa="",
                    fornecedor_ou_favorecido=_text(row.get("beneficiario_normalizado")),
                    cnpj_cpf_favorecido=_text(row.get("beneficiary_document")),
                    documento=_text(row.get("referencia")),
                    tipo_documento="BANCA_COMPROVANTE_CSV",
                    conta_id=_text(row.get("conta_id")),
                    conta_nome=_text(row.get("conta_nome")),
                    grupo_conta="",
                    descricao_itens=_descricao(
                        row.get("beneficiario_normalizado"),
                        row.get("referencia"),
                        row.get("payment_type"),
                    ),
                    valor=valor,
                    origem="COMPROVANTES_CSV",
                    arquivo_origem=str(path),
                    chave_origem=chave_origem,
                    hash_registro=_hash_registro(chave_origem, row.get("conta_id"), valor),
                    observacoes="",
                    data_pagamento=data_documento,
                )
            )

    manifest["linhas_lidas_comprovantes_csv"] = rows_lidas
    manifest["linhas_aprovadas_comprovantes_csv"] = len(lancamentos)
    manifest["fontes"]["comprovantes_csv"]["linhas_lidas"] = rows_lidas
    return lancamentos


def _read_conta_corrente_csv(
    path: Path,
    *,
    periodo_ini: date,
    periodo_fim: date,
    manifest: dict[str, object],
    rejeitados: list[_RejectedRow],
) -> list[Lancamento]:
    lancamentos: list[Lancamento] = []
    rows_lidas = 0

    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row_number, row in enumerate(reader, start=2):
            rows_lidas += 1

            try:
                data_documento = _parse_date(row.get("data_movimento"))
            except LancamentosBancaError:
                manifest["linhas_rejeitadas_dados_obrigatorios"] = int(manifest["linhas_rejeitadas_dados_obrigatorios"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="CONTA_CORRENTE_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="SEM_DATA",
                    date_header="data_movimento",
                    beneficiary_header="razao_social",
                )
                continue
            try:
                valor_assinado = _parse_decimal(row.get("valor"))
                valor = abs(valor_assinado)
            except LancamentosBancaError:
                manifest["linhas_rejeitadas_dados_obrigatorios"] = int(manifest["linhas_rejeitadas_dados_obrigatorios"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="CONTA_CORRENTE_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="SEM_VALOR",
                    date_header="data_movimento",
                    beneficiary_header="razao_social",
                )
                continue
            # O bridge DRE recebe somente saidas economicas; entradas PIX no CSV
            # de conta corrente ficam fora para evitar lancamento invertido.
            if valor_assinado > 0:
                manifest["linhas_excluidas_escludere"] = int(manifest["linhas_excluidas_escludere"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="CONTA_CORRENTE_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="ACAO_ESCLUIDA",
                    date_header="data_movimento",
                    beneficiary_header="razao_social",
                )
                continue

            fornecedor = _text(row.get("razao_social")) or _text(row.get("lancamento"))
            if not fornecedor:
                manifest["linhas_rejeitadas_dados_obrigatorios"] = int(manifest["linhas_rejeitadas_dados_obrigatorios"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="CONTA_CORRENTE_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="SEM_BENEFICIARIO",
                    date_header="data_movimento",
                    beneficiary_header="razao_social",
                )
                continue
            if not _text(row.get("conta_id")) or not _text(row.get("conta_nome")):
                manifest["linhas_rejeitadas_dados_obrigatorios"] = int(manifest["linhas_rejeitadas_dados_obrigatorios"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="CONTA_CORRENTE_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="SEM_CONTA_ID",
                    date_header="data_movimento",
                    beneficiary_header="razao_social",
                )
                continue
            if data_documento < periodo_ini or data_documento > periodo_fim:
                manifest["linhas_fora_periodo"] = int(manifest["linhas_fora_periodo"]) + 1
                _reject_csv(
                    rejeitados,
                    manifest,
                    fonte="CONTA_CORRENTE_CSV",
                    row_number=row_number,
                    row=row,
                    motivo="FORA_PERIODO",
                    date_header="data_movimento",
                    beneficiary_header="razao_social",
                )
                continue

            chave_origem = _chave_origem(
                "CONTA_CORRENTE_CSV",
                row_number,
                _text(row.get("hash_origem")),
                _text(row.get("chave_estavel")),
                data_documento,
                valor,
            )
            lancamentos.append(
                Lancamento(
                    data_documento=data_documento,
                    competencia=f"{data_documento.year:04d}-{data_documento.month:02d}",
                    cnpj_empresa=_text(row.get("cnpj")),
                    nome_empresa="",
                    fornecedor_ou_favorecido=fornecedor,
                    cnpj_cpf_favorecido="",
                    documento=_text(row.get("chave_estavel")),
                    tipo_documento="BANCA_CONTA_CORRENTE_CSV",
                    conta_id=_text(row.get("conta_id")),
                    conta_nome=_text(row.get("conta_nome")),
                    grupo_conta="",
                    descricao_itens=_descricao(
                        row.get("lancamento"),
                        row.get("razao_social"),
                    ),
                    valor=valor,
                    origem="CONTA_CORRENTE_CSV",
                    arquivo_origem=str(path),
                    chave_origem=chave_origem,
                    hash_registro=_hash_registro(chave_origem, row.get("conta_id"), valor),
                    observacoes="",
                    data_pagamento=data_documento,
                )
            )

    manifest["linhas_lidas_conta_corrente_csv"] = rows_lidas
    manifest["linhas_aprovadas_conta_corrente_csv"] = len(lancamentos)
    manifest["fontes"]["conta_corrente_csv"]["linhas_lidas"] = rows_lidas
    return lancamentos


def _read_comprovantes(
    path: Path,
    *,
    periodo_ini: date,
    periodo_fim: date,
    manifest: dict[str, object],
    rejeitados: list[_RejectedRow],
) -> list[Lancamento]:
    rows = _load_rows(path, COMPROVANTES_SHEET)
    manifest["linhas_lidas_comprovantes"] = len(rows)
    manifest["fontes"]["decisoes_comprovantes_aprovadas"]["linhas_lidas"] = len(rows)
    lancamentos: list[Lancamento] = []

    for row_number, row in rows:
        decision = _decision(
            row,
            date_header="DATA_PAGAMENTO",
            value_header="VALOR",
            beneficiary_header="BENEFICIARIO_FAVORECIDO",
            periodo_ini=periodo_ini,
            periodo_fim=periodo_fim,
        )
        if decision:
            _reject(
                rejeitados,
                manifest,
                fonte="COMPROVANTES_BANCA",
                sheet=COMPROVANTES_SHEET,
                row_number=row_number,
                row=row,
                motivo=decision,
                date_header="DATA_PAGAMENTO",
                beneficiary_header="BENEFICIARIO_FAVORECIDO",
            )
            continue

        data_documento = _parse_date(row.get("DATA_PAGAMENTO"))
        valor = abs(_parse_decimal(row.get("VALOR")))
        chave_origem = _chave_origem(
            "COMPROVANTES_BANCA",
            row_number,
            row.get("CHAVE_REGRA_BENEFICIARIO"),
            row.get("CHAVE_REGRA_REFERENCIA"),
            data_documento,
            valor,
        )
        lancamentos.append(
            Lancamento(
                data_documento=data_documento,
                competencia=f"{data_documento.year:04d}-{data_documento.month:02d}",
                cnpj_empresa=_text(row.get("CNPJ")),
                nome_empresa="",
                fornecedor_ou_favorecido=_text(row.get("BENEFICIARIO_FAVORECIDO")),
                cnpj_cpf_favorecido="",
                documento=_text(row.get("CHAVE_REGRA_REFERENCIA")),
                tipo_documento="BANCA_COMPROVANTE",
                conta_id=_text(row.get("CONTA_ID_APROVADA")),
                conta_nome=_text(row.get("CONTA_NOME_APROVADA")),
                grupo_conta="",
                descricao_itens=_descricao(
                    row.get("BENEFICIARIO_FAVORECIDO"),
                    row.get("REFERENCIA_EMPRESA"),
                    row.get("TIPO_PAGAMENTO"),
                ),
                valor=valor,
                origem="COMPROVANTES_BANCA",
                arquivo_origem=str(path),
                chave_origem=chave_origem,
                hash_registro=_hash_registro(chave_origem, row.get("CONTA_ID_APROVADA"), valor),
                observacoes=_text(row.get("OBSERVACAO_OPERADOR")),
                data_pagamento=data_documento,
            )
        )
    return lancamentos


def _read_conta_corrente(
    path: Path,
    *,
    periodo_ini: date,
    periodo_fim: date,
    manifest: dict[str, object],
    rejeitados: list[_RejectedRow],
) -> list[Lancamento]:
    rows = _load_rows(path, CONTA_CORRENTE_SHEET)
    tecnico_by_key = _load_tecnico_conta_corrente(path)
    manifest["linhas_lidas_conta_corrente"] = len(rows)
    manifest["fontes"]["decisoes_conta_corrente_aprovadas"]["linhas_lidas"] = len(rows)
    lancamentos: list[Lancamento] = []

    for row_number, row in rows:
        decision = _decision(
            row,
            date_header="DATA",
            value_header="VALOR",
            beneficiary_header="RAZAO_SOCIAL",
            periodo_ini=periodo_ini,
            periodo_fim=periodo_fim,
            rejeitar_valor_positivo=True,
        )
        if decision:
            _reject(
                rejeitados,
                manifest,
                fonte="CONTA_CORRENTE",
                sheet=CONTA_CORRENTE_SHEET,
                row_number=row_number,
                row=row,
                motivo=decision,
                date_header="DATA",
                beneficiary_header="RAZAO_SOCIAL",
            )
            continue

        data_documento = _parse_date(row.get("DATA"))
        valor = abs(_parse_decimal(row.get("VALOR")))
        key = _stable_key_conta_corrente(row)
        tecnico = tecnico_by_key.get(key, {})
        chave_origem = _chave_origem(
            "CONTA_CORRENTE",
            row_number,
            row.get("CHAVE_REGRA_LANCAMENTO"),
            row.get("CHAVE_REGRA_RAZAO_SOCIAL"),
            data_documento,
            valor,
        )
        lancamentos.append(
            Lancamento(
                data_documento=data_documento,
                competencia=f"{data_documento.year:04d}-{data_documento.month:02d}",
                cnpj_empresa=_text(tecnico.get("CNPJ")),
                nome_empresa="",
                fornecedor_ou_favorecido=_text(row.get("RAZAO_SOCIAL")),
                cnpj_cpf_favorecido="",
                documento=_text(row.get("CHAVE_REGRA_LANCAMENTO")),
                tipo_documento="BANCA_CONTA_CORRENTE",
                conta_id=_text(row.get("CONTA_ID_APROVADA")),
                conta_nome=_text(row.get("CONTA_NOME_APROVADA")),
                grupo_conta="",
                descricao_itens=_descricao(
                    row.get("LANCAMENTO"),
                    row.get("RAZAO_SOCIAL"),
                    row.get("EXEMPLOS"),
                ),
                valor=valor,
                origem="CONTA_CORRENTE",
                arquivo_origem=str(path),
                chave_origem=chave_origem,
                hash_registro=_hash_registro(chave_origem, row.get("CONTA_ID_APROVADA"), valor),
                observacoes=_text(row.get("OBSERVACAO_OPERADOR")),
                data_pagamento=data_documento,
            )
        )
    return lancamentos


def _deduplicate_lancamentos(
    lancamentos: list[Lancamento],
    rejeitados: list[_RejectedRow],
    manifest: dict[str, object],
) -> list[Lancamento]:
    accepted: list[Lancamento] = []
    by_key: dict[_DedupKey, Lancamento] = {}
    ordered = sorted(enumerate(lancamentos), key=_dedup_sort_key)

    for _, lancamento in ordered:
        key = _dedup_key(lancamento)
        winner = by_key.get(key)
        if winner is not None:
            same_source = winner.origem == lancamento.origem
            motivo = "DUPLICADO_INTERNO_FONTE" if same_source else "DUPLICADO_FONTE_PRIORITARIA"
            _reject_lancamento(
                rejeitados,
                manifest,
                lancamento,
                motivo=motivo,
                fonte_prioritaria=winner.origem,
                chave_dedup=key.as_text(),
                increment_source=True,
            )
            manifest["deduplicacao"]["duplicados_excluidos"] += 1
            manifest["deduplicacao"]["valor_duplicado_excluido"] = round(
                float(manifest["deduplicacao"]["valor_duplicado_excluido"])
                + float(abs(lancamento.valor)),
                2,
            )
            continue

        for accepted_lancamento in accepted:
            possible_key = _dedup_key(accepted_lancamento)
            if (
                possible_key.valor_abs == key.valor_abs
                and possible_key.beneficiario_normalizado == key.beneficiario_normalizado
                and abs((accepted_lancamento.data_documento - lancamento.data_documento).days) == 1
            ):
                _reject_lancamento(
                    rejeitados,
                    manifest,
                    lancamento,
                    motivo="POSSIVEL_DUPLICADO_DATA_DIVERGENTE",
                    fonte_prioritaria=accepted_lancamento.origem,
                    chave_dedup=key.as_text(),
                    increment_source=False,
                    incluido=True,
                )
                manifest["deduplicacao"]["possiveis_duplicados_data_divergente"] += 1
                break

        by_key[key] = lancamento
        accepted.append(lancamento)

    _refresh_source_included_counts(manifest, accepted)
    return accepted


def _dedup_key(lancamento: Lancamento) -> _DedupKey:
    return _DedupKey(
        data=lancamento.data_documento.isoformat(),
        valor_abs=f"{abs(lancamento.valor).quantize(Decimal('0.01')):.2f}",
        beneficiario_normalizado=normalize_match_text(lancamento.fornecedor_ou_favorecido),
    )


SOURCE_PRIORITY: dict[str, int] = {
    "COMPROVANTES_CSV": 10,
    "COMPROVANTES_BANCA": 20,
    "CONTA_CORRENTE": 30,
    "CONTA_CORRENTE_CSV": 40,
}


def _dedup_sort_key(item: tuple[int, Lancamento]) -> tuple[int, int]:
    original_index, lancamento = item
    return (SOURCE_PRIORITY.get(lancamento.origem, 99), original_index)


def _source_priority(origem: str) -> int:
    return SOURCE_PRIORITY.get(origem, 99)


def _source_manifest_key(fonte: str) -> str:
    return {
        "COMPROVANTES_CSV": "comprovantes_csv",
        "CONTA_CORRENTE_CSV": "conta_corrente_csv",
        "COMPROVANTES_BANCA": "decisoes_comprovantes_aprovadas",
        "CONTA_CORRENTE": "decisoes_conta_corrente_aprovadas",
    }.get(fonte, fonte.lower())


def _refresh_source_included_counts(
    manifest: dict[str, object],
    lancamentos: list[Lancamento],
) -> None:
    included_by_source = {key: 0 for key in manifest["fontes"]}
    for lancamento in lancamentos:
        source_key = _source_manifest_key(lancamento.origem)
        if source_key in included_by_source:
            included_by_source[source_key] += 1

    for source_key, total in included_by_source.items():
        manifest["fontes"][source_key]["linhas_incluidas"] = total

    manifest["linhas_aprovadas_comprovantes_csv"] = included_by_source["comprovantes_csv"]
    manifest["linhas_aprovadas_conta_corrente_csv"] = included_by_source["conta_corrente_csv"]


def _decision(
    row: dict[str, object],
    *,
    date_header: str,
    value_header: str,
    beneficiary_header: str,
    periodo_ini: date,
    periodo_fim: date,
    rejeitar_valor_positivo: bool = False,
) -> str:
    status = _upper(row.get("STATUS_DECISAO"))
    acao = _upper(row.get("ACAO_REGRA"))

    if status != "APROVADO":
        return f"STATUS_{status or 'VAZIO'}"
    if not acao:
        return "ACAO_REGRA_VAZIA"
    if acao == "XML":
        return "ACAO_XML"
    if acao == "ESCLUDERE":
        return "ACAO_ESCLUDERE"
    if acao == "SISPAG":
        return "ACAO_SISPAG"
    if acao in {"BLOQUEADO", "REVISAR", "NAO_CRIAR_REGRA"}:
        return f"ACAO_{acao}"
    if acao and acao != "CLASSIFICAR":
        return f"ACAO_{acao}"
    if _duplication_risk(row.get("RISCO_DUPLICIDADE_XML")):
        return "RISCO_DUPLICIDADE_XML"

    if not _text(row.get(date_header)):
        return "SEM_DATA"
    try:
        data_documento = _parse_date(row.get(date_header))
    except LancamentosBancaError:
        return "SEM_DATA"
    if not _text(row.get(value_header)):
        return "SEM_VALOR"
    try:
        valor = _parse_decimal(row.get(value_header))
    except LancamentosBancaError:
        return "SEM_VALOR"
    if rejeitar_valor_positivo and valor > 0:
        return "ACAO_ESCLUIDA"

    if not _text(row.get(beneficiary_header)):
        return "SEM_BENEFICIARIO"
    if not _text(row.get("CONTA_ID_APROVADA")) or not _text(row.get("CONTA_NOME_APROVADA")):
        return "SEM_CONTA_ID"
    if data_documento < periodo_ini or data_documento > periodo_fim:
        return "FORA_PERIODO"
    return ""


def _load_rows(path: Path, sheet_name: str) -> list[tuple[int, dict[str, object]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise LancamentosBancaError(f"Aba ausente em {path}: {sheet_name}")
    sheet = workbook[sheet_name]
    headers = [str(value or "").strip() for value in next(sheet.iter_rows(values_only=True))]
    rows: list[tuple[int, dict[str, object]]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
        }
        if any(value not in (None, "") for value in row.values()):
            rows.append((row_number, row))
    workbook.close()
    return rows


def _load_tecnico_conta_corrente(path: Path) -> dict[str, dict[str, object]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if TECNICO_CONTA_CORRENTE_SHEET not in workbook.sheetnames:
        workbook.close()
        return {}
    sheet = workbook[TECNICO_CONTA_CORRENTE_SHEET]
    headers = [str(value or "").strip() for value in next(sheet.iter_rows(values_only=True))]
    result: dict[str, dict[str, object]] = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
        }
        key = _text(row.get("CHAVE_ESTAVEL"))
        if key:
            result[key] = row
    workbook.close()
    return result


def _stable_key_conta_corrente(row: dict[str, object]) -> str:
    history = normalize_match_text(row.get("LANCAMENTO"))
    razao = normalize_match_text(row.get("RAZAO_SOCIAL")) or "SEM_RAZAO"
    tipo = _tipo_movimento_text(history, _parse_decimal(row.get("VALOR")))
    if not history:
        return ""
    return (
        f"CONTA_CORRENTE|HISTORICO={history}|RAZAO={razao}|"
        f"TIPO={tipo}|DIRECAO=SAIDA"
    )


def _tipo_movimento_text(description: object, amount: Decimal | None) -> str:
    key = normalize_match_text(description)
    if amount is not None and amount > 0:
        return "ENTRADA_FORA_ESCOPO"
    if "SISPAG" in key:
        return "SISPAG_AGREGADO"
    if "TARIFA" in key or "TAXA" in key:
        return "TARIFA_BANCARIA"
    if "PIX" in key:
        return "PIX_TRANSFERENCIA_SAIDA"
    if any(token in key for token in ("BOLETO", "TED", "TRANSFERENCIA", "TRANSF")):
        return "BOLETO_TED_TRANSFERENCIA_SAIDA"
    return "SAIDA_PAGAMENTO"


def _reject(
    rejeitados: list[_RejectedRow],
    manifest: dict[str, object],
    *,
    fonte: str,
    sheet: str,
    row_number: int,
    row: dict[str, object],
    motivo: str,
    date_header: str,
    beneficiary_header: str,
) -> None:
    counter = {
        "ACAO_XML": "linhas_excluidas_xml",
        "ACAO_ESCLUDERE": "linhas_excluidas_escludere",
        "ACAO_SISPAG": "linhas_excluidas_sispag",
        "ACAO_ESCLUIDA": "linhas_excluidas_escludere",
        "RISCO_DUPLICIDADE_XML": "linhas_rejeitadas_risco_duplicidade",
        "FORA_PERIODO": "linhas_fora_periodo",
        "DADOS_OBRIGATORIOS_AUSENTES": "linhas_rejeitadas_dados_obrigatorios",
        "SEM_DATA": "linhas_rejeitadas_dados_obrigatorios",
        "SEM_VALOR": "linhas_rejeitadas_dados_obrigatorios",
        "SEM_BENEFICIARIO": "linhas_rejeitadas_dados_obrigatorios",
        "SEM_CONTA_ID": "linhas_rejeitadas_dados_obrigatorios",
    }.get(motivo, "linhas_rejeitadas_status")
    manifest[counter] = int(manifest[counter]) + 1
    _bump_source_rejected(manifest, fonte)
    rejeitados.append(
        _RejectedRow(
            fonte=fonte,
            sheet=sheet,
            linha=row_number,
            motivo=_canonical_rejection_motivo(motivo),
            status_decisao=_text(row.get("STATUS_DECISAO")),
            acao_regra=_text(row.get("ACAO_REGRA")),
            risco_duplicidade_xml=_text(row.get("RISCO_DUPLICIDADE_XML")),
            data=_text(row.get(date_header)),
            valor=_text(row.get("VALOR")),
            beneficiario=_text(row.get(beneficiary_header)),
            conta_id=_text(row.get("CONTA_ID_APROVADA")),
            conta_nome=_text(row.get("CONTA_NOME_APROVADA")),
        )
    )


def _reject_csv(
    rejeitados: list[_RejectedRow],
    manifest: dict[str, object],
    *,
    fonte: str,
    row_number: int,
    row: dict[str, object],
    motivo: str,
    date_header: str,
    beneficiary_header: str,
) -> None:
    _bump_source_rejected(manifest, fonte)
    rejeitados.append(
        _RejectedRow(
            fonte=fonte,
            sheet="CSV",
            linha=row_number,
            motivo=motivo,
            status_decisao=_text(row.get("status_classificacao")),
            acao_regra=_text(row.get("acao_regra")),
            risco_duplicidade_xml="",
            data=_text(row.get(date_header)),
            valor=_text(row.get("valor")),
            beneficiario=_text(row.get(beneficiary_header)) or _text(row.get("lancamento")),
            conta_id=_text(row.get("conta_id")),
            conta_nome=_text(row.get("conta_nome")),
        )
    )


def _reject_lancamento(
    rejeitados: list[_RejectedRow],
    manifest: dict[str, object],
    lancamento: Lancamento,
    *,
    motivo: str,
    fonte_prioritaria: str,
    chave_dedup: str,
    increment_source: bool,
    incluido: bool = False,
) -> None:
    if increment_source:
        _bump_source_rejected(manifest, lancamento.origem)
    rejeitados.append(
        _RejectedRow(
            fonte=lancamento.origem,
            sheet="DEDUP",
            linha=_linha_from_chave_origem(lancamento.chave_origem),
            motivo=motivo,
            status_decisao="",
            acao_regra="",
            risco_duplicidade_xml="",
            data=lancamento.data_documento.isoformat(),
            valor=f"{lancamento.valor.quantize(Decimal('0.01')):.2f}",
            beneficiario=lancamento.fornecedor_ou_favorecido,
            conta_id=lancamento.conta_id,
            conta_nome=lancamento.conta_nome,
            fonte_prioritaria=fonte_prioritaria,
            chave_dedup=chave_dedup,
            incluido="true" if incluido else "false",
        )
    )


def _bump_source_rejected(manifest: dict[str, object], fonte: str) -> None:
    source_key = _source_manifest_key(fonte)
    fontes = manifest.get("fontes")
    if isinstance(fontes, dict) and source_key in fontes:
        fontes[source_key]["linhas_rejeitadas"] += 1


def _canonical_rejection_motivo(motivo: str) -> str:
    if motivo.startswith("STATUS_"):
        return "STATUS_NAO_APROVADO"
    if motivo == "ACAO_REGRA_VAZIA":
        return motivo
    if motivo.startswith("ACAO_") or motivo == "RISCO_DUPLICIDADE_XML":
        return "ACAO_ESCLUIDA"
    if motivo == "DADOS_OBRIGATORIOS_AUSENTES":
        return "SEM_CONTA_ID"
    return motivo


def _linha_from_chave_origem(chave_origem: str) -> int:
    for part in chave_origem.split("|"):
        if part.startswith("linha="):
            try:
                return int(part.removeprefix("linha="))
            except ValueError:
                return 0
    return 0


def _write_rejeitados(path: Path, rejeitados: Iterable[_RejectedRow]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "fonte",
        "sheet",
        "linha",
        "motivo",
        "status_decisao",
        "acao_regra",
        "risco_duplicidade_xml",
        "data",
        "valor",
        "beneficiario",
        "conta_id",
        "conta_nome",
        "fonte_prioritaria",
        "chave_dedup",
        "incluido",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        for row in rejeitados:
            writer.writerow(row.__dict__)


def _base_manifest(
    comprovantes_xlsx: Path,
    conta_corrente_xlsx: Path,
    output_xlsx: Path,
    periodo_ini: date,
    periodo_fim: date,
) -> dict[str, object]:
    return {
        "periodo_ini": periodo_ini.isoformat(),
        "periodo_fim": periodo_fim.isoformat(),
        "comprovantes_xlsx": str(comprovantes_xlsx),
        "conta_corrente_xlsx": str(conta_corrente_xlsx),
        "linhas_lidas_comprovantes_csv": 0,
        "linhas_aprovadas_comprovantes_csv": 0,
        "linhas_lidas_conta_corrente_csv": 0,
        "linhas_aprovadas_conta_corrente_csv": 0,
        "linhas_lidas_comprovantes": 0,
        "linhas_lidas_conta_corrente": 0,
        "linhas_aprovadas_incluidas": 0,
        "linhas_excluidas_xml": 0,
        "linhas_excluidas_escludere": 0,
        "linhas_excluidas_sispag": 0,
        "linhas_rejeitadas_risco_duplicidade": 0,
        "linhas_fora_periodo": 0,
        "linhas_rejeitadas_status": 0,
        "linhas_rejeitadas_dados_obrigatorios": 0,
        "fontes": {
            "comprovantes_csv": {
                "linhas_lidas": 0,
                "linhas_incluidas": 0,
                "linhas_rejeitadas": 0,
            },
            "conta_corrente_csv": {
                "linhas_lidas": 0,
                "linhas_incluidas": 0,
                "linhas_rejeitadas": 0,
            },
            "decisoes_comprovantes_aprovadas": {
                "linhas_lidas": 0,
                "linhas_incluidas": 0,
                "linhas_rejeitadas": 0,
            },
            "decisoes_conta_corrente_aprovadas": {
                "linhas_lidas": 0,
                "linhas_incluidas": 0,
                "linhas_rejeitadas": 0,
            },
        },
        "deduplicacao": {
            "duplicados_excluidos": 0,
            "valor_duplicado_excluido": 0.0,
            "possiveis_duplicados_data_divergente": 0,
        },
        "output_xlsx": str(output_xlsx),
    }


def _validate_input_workbook(path: Path, label: str) -> None:
    if not path.exists():
        raise LancamentosBancaError(f"{label} nao encontrado: {path}")
    if not path.is_file():
        raise LancamentosBancaError(f"{label} nao e arquivo: {path}")


def _parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        raise LancamentosBancaError("Data obrigatoria ausente.")
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(text[:10])
    except ValueError as exc:
        raise LancamentosBancaError(f"Data invalida: {text}") from exc


def _parse_decimal(value: object) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    text = _text(value)
    if not text:
        raise LancamentosBancaError("Valor obrigatorio ausente.")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise LancamentosBancaError(f"Valor invalido: {value}") from exc


def _duplication_risk(value: object) -> bool:
    text = _upper(value)
    return text not in {"", "OK", "NAO", "NÃO", "FALSE", "0", "SEM_RISCO"}


def _chave_origem(
    fonte: str,
    row_number: int,
    key_a: object,
    key_b: object,
    data_documento: date,
    valor: Decimal,
) -> str:
    parts = [
        fonte,
        f"linha={row_number}",
        _text(key_a),
        _text(key_b),
        data_documento.isoformat(),
        f"{valor.quantize(Decimal('0.01')):.2f}",
    ]
    return "|".join(parts)


def _hash_registro(chave_origem: str, conta_id: object, valor: Decimal) -> str:
    payload = "|".join(
        [
            chave_origem,
            _text(conta_id),
            f"{valor.quantize(Decimal('0.01')):.2f}",
        ]
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _descricao(*values: object) -> str:
    return " | ".join(_text(value) for value in values if _text(value))


def _upper(value: object) -> str:
    return _text(value).upper()


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()
