from __future__ import annotations

import csv
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fantozzi.csv_utils import sniff_delimiter


REVISAO_SHEET = "REVISAO_COMPROVANTES"
RESUMO_BENEFICIARIO_SHEET = "RESUMO_POR_BENEFICIARIO"
RESUMO_MES_CNPJ_SHEET = "RESUMO_POR_MES_CNPJ"
SISPAG_SHEET = "SISPAG_FORA_ESCOPO"
INSTRUCOES_SHEET = "INSTRUCOES_OPERADOR"

OUTPUT_XLSX = "pacote_revisao_comprovantes_13T.xlsx"
OUTPUT_TXT = "pacote_revisao_comprovantes_13T.txt"
OUTPUT_SISPAG_CSV = "sispag_fora_escopo.csv"

FORA_ESCOPO_SISPAG = "FORA_ESCOPO_SISPAG"

REVISAO_HEADERS = [
    "cnpj",
    "mes",
    "data_pagamento",
    "beneficiario_normalizado",
    "referencia",
    "valor",
    "arquivo_origem",
    "status_classificacao",
    "motivo_revisao",
    "regras_candidatas",
    "contas_candidatas",
    "DECISAO_OPERADOR",
    "CONTA_ID_APROVADA",
    "CONTA_NOME_APROVADA",
    "CRIAR_REGRA_FUTURA",
    "TIPO_REGRA_SUGERIDA",
    "ALIAS_O_KEYWORD_APROVADO",
    "OBSERVACOES",
]

SISPAG_HEADERS = [
    "cnpj",
    "mes",
    "data_pagamento",
    "beneficiario_normalizado",
    "referencia",
    "valor",
    "arquivo_origem",
    "status_classificacao",
    "motivo_revisao",
    "regras_candidatas",
    "contas_candidatas",
    "DECISAO_OPERADOR",
    "CONTA_ID_APROVADA",
    "CONTA_NOME_APROVADA",
    "CRIAR_REGRA_FUTURA",
    "TIPO_REGRA_SUGERIDA",
    "ALIAS_O_KEYWORD_APROVADO",
    "OBSERVACOES",
    "payment_id",
    "beneficiary_name",
    "payment_type",
    "status_pagamento",
    "source_file",
    "raw_hash",
]

RESUMO_BENEFICIARIO_HEADERS = [
    "beneficiario_normalizado",
    "quantidade",
    "valor_total",
    "cnpjs_envolvidos",
    "meses_envolvidos",
    "exemplos_referencia",
]

RESUMO_MES_CNPJ_HEADERS = [
    "cnpj",
    "mes",
    "quantidade_a_revisar",
    "valor_total_a_revisar",
]

DECISOES_VALIDAS = [
    "APROVAR_CONTA",
    "MANTER_REVISAO",
    "IGNORAR_FORA_DRE",
    "NAO_CRIAR_REGRA",
]
CRIAR_REGRA_VALIDOS = ["sim", "nao"]
TIPOS_REGRA_VALIDOS = ["beneficiario", "keyword", "referencia"]


class PaymentReviewPackageError(ValueError):
    pass


@dataclass(frozen=True)
class PaymentReviewPackageResult:
    run_dir: Path
    xlsx_path: Path
    report_path: Path
    sispag_path: Path
    total_input_rows: int
    total_revisao_rows: int
    total_sispag_rows: int
    matching_automatico_executado: bool
    yaml_modificado: bool


def gerar_pacote_revisao_comprovantes(
    run_dir: Path,
    *,
    csv_revisao: Path | None = None,
    xlsx_path: Path | None = None,
    report_path: Path | None = None,
    sispag_csv_path: Path | None = None,
) -> PaymentReviewPackageResult:
    run_dir = Path(run_dir)
    csv_revisao = Path(csv_revisao or run_dir / "comprovantes_a_revisar.csv")
    xlsx_path = Path(xlsx_path or run_dir / OUTPUT_XLSX)
    report_path = Path(report_path or run_dir / OUTPUT_TXT)
    sispag_csv_path = Path(sispag_csv_path or run_dir / OUTPUT_SISPAG_CSV)

    headers, rows = _read_csv(csv_revisao)
    if not rows:
        raise PaymentReviewPackageError(f"CSV de revisao sem linhas: {csv_revisao}")

    review_rows: list[dict[str, object]] = []
    sispag_rows: list[dict[str, object]] = []
    for row in rows:
        output_row = _review_output_row(row)
        if _is_sispag(row):
            output_row.update(
                {
                    "status_classificacao": FORA_ESCOPO_SISPAG,
                    "motivo_revisao": "SISPAG_FORA_ESCOPO_OPERACIONAL",
                    "regras_candidatas": "",
                    "contas_candidatas": "",
                    "DECISAO_OPERADOR": "",
                    "CONTA_ID_APROVADA": "",
                    "CONTA_NOME_APROVADA": "",
                    "CRIAR_REGRA_FUTURA": "nao",
                    "TIPO_REGRA_SUGERIDA": "",
                    "ALIAS_O_KEYWORD_APROVADO": "",
                    "OBSERVACOES": "Fora da revisao contabil 13T; controle financeiro separado.",
                }
            )
            sispag_rows.append(output_row)
        else:
            review_rows.append(output_row)

    workbook = Workbook()
    _create_revisao_sheet(workbook, review_rows)
    _create_resumo_beneficiario_sheet(workbook, review_rows)
    _create_resumo_mes_cnpj_sheet(workbook, review_rows)
    _create_sispag_sheet(workbook, sispag_rows)
    _create_instrucoes_sheet(workbook)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)
    _write_csv(sispag_csv_path, SISPAG_HEADERS, sispag_rows)

    result = PaymentReviewPackageResult(
        run_dir=run_dir,
        xlsx_path=xlsx_path,
        report_path=report_path,
        sispag_path=sispag_csv_path,
        total_input_rows=len(rows),
        total_revisao_rows=len(review_rows),
        total_sispag_rows=len(sispag_rows),
        matching_automatico_executado=False,
        yaml_modificado=False,
    )
    report_path.write_text(_report_text(result, review_rows, sispag_rows), encoding="utf-8")
    return result


def console_lines(result: PaymentReviewPackageResult) -> list[str]:
    return [
        "[OK] Pacote 13T de revisao de comprovantes gerado.",
        f"     Run origem: {result.run_dir.name}",
        f"     Linhas entrada: {result.total_input_rows}",
        f"     Revisao operacional: {result.total_revisao_rows}",
        f"     SISPAG fora de escopo: {result.total_sispag_rows}",
        f"     XLSX: {result.xlsx_path}",
        f"     Relatorio: {result.report_path}",
        f"     SISPAG CSV: {result.sispag_path}",
        "     matching_automatico_executado=false",
        "     YAML de regras nao modificado.",
    ]


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.exists():
        raise PaymentReviewPackageError(f"CSV de revisao nao encontrado: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        sample = file.read(4096)
        file.seek(0)
        reader = csv.DictReader(file, delimiter=sniff_delimiter(sample))
        headers = list(reader.fieldnames or [])
        if not headers:
            raise PaymentReviewPackageError(f"CSV de revisao sem cabecalho: {path}")
        rows = [
            {header: (row.get(header) or "").strip() for header in headers}
            for row in reader
        ]
    return headers, rows


def _write_csv(path: Path, headers: list[str], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers, delimiter=";")
        writer.writeheader()
        for row in rows:
            writer.writerow({header: _csv_value(row.get(header, "")) for header in headers})


def _review_output_row(row: dict[str, str]) -> dict[str, object]:
    return {
        "cnpj": row.get("cnpj", ""),
        "mes": row.get("mes", ""),
        "data_pagamento": row.get("data_pagamento", ""),
        "beneficiario_normalizado": row.get("beneficiario_normalizado", ""),
        "referencia": row.get("referencia", ""),
        "valor": row.get("valor", ""),
        "arquivo_origem": row.get("arquivo_origem", ""),
        "status_classificacao": row.get("status_classificacao", ""),
        "motivo_revisao": row.get("motivo_revisao", ""),
        "regras_candidatas": row.get("regras_candidatas", ""),
        "contas_candidatas": row.get("contas_candidatas", ""),
        "DECISAO_OPERADOR": "",
        "CONTA_ID_APROVADA": "",
        "CONTA_NOME_APROVADA": "",
        "CRIAR_REGRA_FUTURA": "",
        "TIPO_REGRA_SUGERIDA": "",
        "ALIAS_O_KEYWORD_APROVADO": "",
        "OBSERVACOES": "",
        "payment_id": row.get("payment_id", ""),
        "beneficiary_name": row.get("beneficiary_name", ""),
        "payment_type": row.get("payment_type", ""),
        "status_pagamento": row.get("status_pagamento", ""),
        "source_file": row.get("source_file", ""),
        "raw_hash": row.get("raw_hash", ""),
    }


def _is_sispag(row: dict[str, str]) -> bool:
    text = " ".join(
        row.get(field, "")
        for field in (
            "beneficiario_normalizado",
            "referencia",
            "arquivo_origem",
            "motivo_revisao",
            "payment_id",
            "beneficiary_name",
            "payment_type",
            "status_pagamento",
            "source_file",
        )
    )
    return "SISPAG" in _normalize(text)


def _create_revisao_sheet(workbook: Workbook, rows: list[dict[str, object]]) -> None:
    sheet = workbook.active
    sheet.title = REVISAO_SHEET
    _append_table(sheet, REVISAO_HEADERS, rows)
    _apply_validations(sheet, len(rows))
    _style_table_sheet(sheet, REVISAO_HEADERS)


def _create_resumo_beneficiario_sheet(
    workbook: Workbook,
    rows: list[dict[str, object]],
) -> None:
    sheet = workbook.create_sheet(RESUMO_BENEFICIARIO_SHEET)
    summary = _summary_by_beneficiary(rows)
    _append_table(sheet, RESUMO_BENEFICIARIO_HEADERS, summary)
    _style_table_sheet(sheet, RESUMO_BENEFICIARIO_HEADERS)


def _create_resumo_mes_cnpj_sheet(
    workbook: Workbook,
    rows: list[dict[str, object]],
) -> None:
    sheet = workbook.create_sheet(RESUMO_MES_CNPJ_SHEET)
    summary = _summary_by_month_cnpj(rows)
    _append_table(sheet, RESUMO_MES_CNPJ_HEADERS, summary)
    _style_table_sheet(sheet, RESUMO_MES_CNPJ_HEADERS)


def _create_sispag_sheet(workbook: Workbook, rows: list[dict[str, object]]) -> None:
    sheet = workbook.create_sheet(SISPAG_SHEET)
    _append_table(sheet, SISPAG_HEADERS, rows)
    _style_table_sheet(sheet, SISPAG_HEADERS)


def _create_instrucoes_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(INSTRUCOES_SHEET)
    rows = [
        ["Fase 13T - revisao manual de comprovantes"],
        [""],
        ["Nenhuma regra YAML e criada automaticamente por esta planilha."],
        ["matching automatico bancario nao executado."],
        ["SISPAG fica fora da aba REVISAO_COMPROVANTES e aparece em SISPAG_FORA_ESCOPO."],
        [""],
        ["Quando usar CRIAR_REGRA_FUTURA = sim"],
        ["Use somente quando o operador validou que o pagamento sempre cai no mesmo conto."],
        ["Preencha CONTA_ID_APROVADA, CONTA_NOME_APROVADA, TIPO_REGRA_SUGERIDA e alias/keyword."],
        [""],
        ["Alias vs keyword"],
        ["Alias/beneficiario: usar quando o beneficiario recorrente identifica claramente o custo."],
        ["Exemplo: COELBA -> energia eletrica; EMBASA -> agua e esgoto."],
        ["Keyword: usar somente se a palavra for especifica e conferida em varios exemplos."],
        ["Exemplo: FGTS, DARF, DAS, EMBASA."],
        [""],
        ["Nao usar keyword generica"],
        ["pagamento, transferencia, boleto, pix, fornecedor, taxa, servico, compra."],
        [""],
        ["DECISAO_OPERADOR"],
        ["APROVAR_CONTA: operador aprovou uma conta para futura regra declarada."],
        ["MANTER_REVISAO: ainda nao ha decisao segura."],
        ["IGNORAR_FORA_DRE: item nao deve entrar no DRE."],
        ["NAO_CRIAR_REGRA: revisado, mas nao deve gerar regra futura."],
    ]
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 120
    for row in sheet.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _append_table(sheet, headers: list[str], rows: list[dict[str, object]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header, "")) for header in headers])


def _apply_validations(sheet, row_count: int) -> None:
    if row_count < 1:
        return
    _add_list_validation(sheet, "DECISAO_OPERADOR", DECISOES_VALIDAS, row_count)
    _add_list_validation(sheet, "CRIAR_REGRA_FUTURA", CRIAR_REGRA_VALIDOS, row_count)
    _add_list_validation(sheet, "TIPO_REGRA_SUGERIDA", TIPOS_REGRA_VALIDOS, row_count)


def _add_list_validation(
    sheet,
    header: str,
    values: list[str],
    row_count: int,
) -> None:
    column = REVISAO_HEADERS.index(header) + 1
    letter = get_column_letter(column)
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(values) + '"',
        allow_blank=True,
    )
    validation.error = f"Escolha um valor valido para {header}."
    validation.errorTitle = "Valor invalido"
    sheet.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{row_count + 1}")


def _style_table_sheet(sheet, headers: list[str]) -> None:
    _style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = _suggested_width(header)
        if sheet.max_row < 2:
            continue
        for column in sheet.iter_cols(
            min_col=index,
            max_col=index,
            min_row=2,
            max_row=sheet.max_row,
        ):
            for cell in column:
                cell.alignment = Alignment(wrap_text=_should_wrap(header), vertical="top")
                if header in {"valor", "valor_total", "valor_total_a_revisar"}:
                    cell.number_format = '#,##0.00'
                if header in {"quantidade", "quantidade_a_revisar"}:
                    cell.number_format = '0'


def _style_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _summary_by_beneficiary(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("beneficiario_normalizado") or "-")].append(row)

    output = []
    for beneficiary, items in grouped.items():
        output.append(
            {
                "beneficiario_normalizado": beneficiary,
                "quantidade": len(items),
                "valor_total": _sum_amounts(items),
                "cnpjs_envolvidos": _joined_unique(items, "cnpj"),
                "meses_envolvidos": _joined_unique(items, "mes"),
                "exemplos_referencia": _joined_unique(items, "referencia", limit=3),
            }
        )
    return sorted(output, key=lambda row: (-int(row["quantidade"]), str(row["beneficiario_normalizado"])))


def _summary_by_month_cnpj(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    grouped: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        grouped[(str(row.get("cnpj", "")), str(row.get("mes", "")))].append(row)

    output = []
    for (cnpj, mes), items in grouped.items():
        output.append(
            {
                "cnpj": cnpj,
                "mes": mes,
                "quantidade_a_revisar": len(items),
                "valor_total_a_revisar": _sum_amounts(items),
            }
        )
    return sorted(output, key=lambda row: (str(row["mes"]), str(row["cnpj"])))


def _report_text(
    result: PaymentReviewPackageResult,
    review_rows: list[dict[str, object]],
    sispag_rows: list[dict[str, object]],
) -> str:
    lines = [
        "PACOTE 13T - REVISAO MANUAL DE COMPROVANTES",
        "================================================",
        "",
        f"run_id origem: {result.run_dir.name}",
        f"total linhas a revisar na 13S: {result.total_input_rows}",
        f"linhas na REVISAO_COMPROVANTES: {result.total_revisao_rows}",
        f"linhas SISPAG_FORA_ESCOPO: {result.total_sispag_rows}",
        "",
        "Total por CNPJ:",
        *_counter_lines(Counter(str(row.get("cnpj", "-")) for row in review_rows)),
        "",
        "Total por mes:",
        *_counter_lines(Counter(str(row.get("mes", "-")) for row in review_rows)),
        "",
        "Top beneficiarios por quantidade:",
        *_top_beneficiary_lines(review_rows, by_value=False),
        "",
        "Top beneficiarios por valor:",
        *_top_beneficiary_lines(review_rows, by_value=True),
        "",
        "Avisos:",
        "- nenhuma regra foi criada automaticamente",
        "- matching automatico bancario nao executado",
        "- SISPAG foi marcado como FORA_ESCOPO_SISPAG e separado da revisao operacional",
        "- YAML de regras de comprovantes nao foi modificado",
        "",
        "Instrucoes sinteticas:",
        "- use CRIAR_REGRA_FUTURA=sim apenas para regra declarada e validada",
        "- prefira alias/beneficiario quando o favorecido identifica claramente o custo",
        "- use keyword somente quando a palavra for especifica",
        "- nao use palavras genericas como pagamento, transferencia, boleto, pix ou servico",
    ]
    return "\n".join(lines) + "\n"


def _top_beneficiary_lines(
    rows: list[dict[str, object]],
    *,
    by_value: bool,
) -> list[str]:
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("beneficiario_normalizado") or "-")].append(row)
    summary = [
        (beneficiary, len(items), _sum_amounts(items))
        for beneficiary, items in grouped.items()
    ]
    if by_value:
        summary.sort(key=lambda item: (-item[2], item[0]))
    else:
        summary.sort(key=lambda item: (-item[1], item[0]))
    return [
        f"- {beneficiary}: {qty} linhas | {_format_decimal(value)}"
        for beneficiary, qty, value in summary[:10]
    ] or ["-"]


def _counter_lines(counter: Counter[str]) -> list[str]:
    if not counter:
        return ["-"]
    return [f"- {key}: {value}" for key, value in sorted(counter.items())]


def _sum_amounts(rows: list[dict[str, object]]) -> Decimal:
    return sum((_parse_decimal(row.get("valor", "")) for row in rows), Decimal("0"))


def _joined_unique(
    rows: list[dict[str, object]],
    field: str,
    *,
    limit: int | None = None,
) -> str:
    values = []
    for row in rows:
        value = str(row.get(field, "")).strip()
        if value and value not in values:
            values.append(value)
    if limit is not None:
        values = values[:limit]
    return " | ".join(values)


def _parse_decimal(value: object) -> Decimal:
    text = str(value or "").strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text or "0")
    except InvalidOperation:
        return Decimal("0")


def _format_decimal(value: Decimal) -> str:
    return f"R$ {value:.2f}"


def _normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    ascii_text = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join(ascii_text.upper().split())


def _suggested_width(header: str) -> int:
    widths = {
        "beneficiario_normalizado": 34,
        "referencia": 42,
        "arquivo_origem": 58,
        "motivo_revisao": 30,
        "DECISAO_OPERADOR": 24,
        "CONTA_NOME_APROVADA": 34,
        "CRIAR_REGRA_FUTURA": 20,
        "TIPO_REGRA_SUGERIDA": 22,
        "ALIAS_O_KEYWORD_APROVADO": 38,
        "OBSERVACOES": 52,
        "exemplos_referencia": 62,
    }
    return widths.get(header, max(14, min(len(header) + 4, 32)))


def _should_wrap(header: str) -> bool:
    return header in {
        "beneficiario_normalizado",
        "referencia",
        "arquivo_origem",
        "motivo_revisao",
        "ALIAS_O_KEYWORD_APROVADO",
        "OBSERVACOES",
        "exemplos_referencia",
    }


def _cell_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str) and value.replace(".", "", 1).isdigit():
        return value
    return value


def _csv_value(value: object) -> str:
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if value is None:
        return ""
    return str(value)
