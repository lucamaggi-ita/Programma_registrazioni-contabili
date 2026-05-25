from __future__ import annotations

import json
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Iterable

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fantozzi.banking.common import only_digits
from fantozzi.banking.models import BankLaunchEnrichment, PaymentReceiptRaw
from fantozzi.banking.payment_receipt_parser import parse_payment_receipts
from fantozzi.banking.payment_rules import (
    STATUS_CLASSIFICADO,
    STATUS_EXCLUIDO,
    STATUS_SEM_REGRA,
    PaymentClassificationResult,
    PaymentRulesError,
    classify_payment_receipt,
    load_payment_exclusion_rules,
    load_payment_rules,
    normalize_match_text,
    validate_payment_rules,
)
from fantozzi.banking.xlsx_bank_parser import parse_bank_launches


COMPROVANTES_XLSX = "Analisi_regole_registrazioni_banca_comprovanti.xlsx"
CONTA_CORRENTE_XLSX = "Analisi_regole_registrazioni_banca_conto_corrente.xlsx"
MANIFEST_NAME = "manifest_importa_revisao_banca.json"
MANIFEST_HISTORICO_TEMPLATE = "manifest_importa_revisao_banca_{periodo}.json"

COMPROVANTES_SHEET = "decisoes_comprovantes"
CONTA_CORRENTE_SHEET = "decisoes_conta_corrente"
TECNICO_COMPROVANTES_SHEET = "tecnico_comprovantes"
TECNICO_CONTA_CORRENTE_SHEET = "tecnico_conta_corrente"
PLANO_CONTAS_SHEET = "plano_contas"
LISTAS_SHEET = "listas_validacao"
INSTRUCOES_SHEET = "instrucoes"

COMPROVANTES_HEADERS = [
    "STATUS_DECISAO",
    "ACAO_REGRA",
    "PERIODO",
    "CNPJ",
    "DATA_PAGAMENTO",
    "VALOR",
    "BENEFICIARIO_FAVORECIDO",
    "REFERENCIA_EMPRESA",
    "TIPO_PAGAMENTO",
    "CONTA_ESCOLHIDA",
    "CONTA_ID_APROVADA",
    "CONTA_NOME_APROVADA",
    "CHAVE_REGRA_BENEFICIARIO",
    "CHAVE_REGRA_REFERENCIA",
    "RISCO_DUPLICIDADE_XML",
    "ESCOPO_REGRA",
    "QTD_OCORRENCIAS",
    "VALOR_TOTAL",
    "EXEMPLOS",
    "OBSERVACAO_OPERADOR",
]

CONTA_CORRENTE_HEADERS = [
    "STATUS_DECISAO",
    "ACAO_REGRA",
    "DATA",
    "VALOR",
    "LANCAMENTO",
    "RAZAO_SOCIAL",
    "CONTA_ESCOLHIDA",
    "CONTA_ID_APROVADA",
    "CONTA_NOME_APROVADA",
    "CHAVE_REGRA_LANCAMENTO",
    "CHAVE_REGRA_RAZAO_SOCIAL",
    "POSSIVEL_COMPROVANTE_RELACIONADO",
    "COMPROVANTE_RELACIONADO_BENEFICIARIO",
    "COMPROVANTE_RELACIONADO_DOCUMENTO",
    "COMPROVANTE_RELACIONADO_DATA",
    "COMPROVANTE_RELACIONADO_VALOR",
    "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE",
    "USAR_SOMENTE_SEM_COMPROVANTE",
    "RISCO_DUPLICIDADE_XML",
    "ESCOPO_REGRA",
    "QTD_OCORRENCIAS",
    "VALOR_TOTAL",
    "EXEMPLOS",
    "OBSERVACAO_OPERADOR",
]

TECNICO_COMPROVANTES_HEADERS = [
    "CHAVE_ESTAVEL",
    "ORIGEM_ARQUIVO",
    "TIPO_FONTE",
    "DATAS_EXEMPLO",
    "VALORES_EXEMPLO",
    "REGRA_EXISTENTE",
    "MOTIVO_REVISAO",
    "HASH_ORIGEM",
    "CNPJS_ENVOLVIDOS",
]

TECNICO_CONTA_CORRENTE_HEADERS = [
    "CHAVE_ESTAVEL",
    "CNPJ",
    "MES",
    "ORIGEM_ARQUIVO",
    "DATA_MOVIMENTO",
    "VALORES_EXEMPLO",
    "DIRECAO",
    "TIPO_MOVIMENTO",
    "REGRA_EXISTENTE",
    "MOTIVO_REVISAO",
    "HASH_ORIGEM",
    "COMPROVANTE_RELACIONADO_BENEFICIARIO",
    "COMPROVANTE_RELACIONADO_DOCUMENTO",
    "COMPROVANTE_RELACIONADO_DATA",
    "COMPROVANTE_RELACIONADO_VALOR",
    "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE",
    "COMPETENCIA_REFERENCIA",
    "OBSERVACAO_AUDIT",
]

MANUAL_COLUMNS = {
    "STATUS_DECISAO",
    "ACAO_REGRA",
    "CONTA_ESCOLHIDA",
    "CONTA_ID_APROVADA",
    "CONTA_NOME_APROVADA",
    "OBSERVACAO_OPERADOR",
    "CHAVE_REGRA_BENEFICIARIO",
    "CHAVE_REGRA_REFERENCIA",
    "CHAVE_REGRA_LANCAMENTO",
    "CHAVE_REGRA_RAZAO_SOCIAL",
    "POSSIVEL_COMPROVANTE_RELACIONADO",
    "COMPROVANTE_RELACIONADO_BENEFICIARIO",
    "COMPROVANTE_RELACIONADO_DOCUMENTO",
    "COMPROVANTE_RELACIONADO_DATA",
    "COMPROVANTE_RELACIONADO_VALOR",
    "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE",
    "USAR_SOMENTE_SEM_COMPROVANTE",
    "RISCO_DUPLICIDADE_XML",
    "ESCOPO_REGRA",
}
BANK_STATUS_DECISAO_VALUES = (
    "PENDENTE",
    "APROVADO",
    "REVISAR",
    "NAO_CRIAR_REGRA",
    "BLOQUEADO",
)
BANK_ACAO_REGRA_VALUES = (
    "CLASSIFICAR",
    "ESCLUDERE",
    "XML",
    "NAO_CRIAR_REGRA",
    "REVISAR",
    "BLOQUEADO",
)
RISCO_DUPLICIDADE_XML_VALUES = ("OK", "DUVIDA")
ESCOPO_REGRA_VALUES = ("COMPROVANTE", "CONTA_CORRENTE")
USAR_SOMENTE_SEM_COMPROVANTE_VALUES = ("SIM", "NAO")
BOLETO_PAGO_FALLBACK_REGRA_ID = "EXCLUSAO_XML_CONTA_CORRENTE_BOLETO_PAGO"


class ImportaRevisaoBancaError(ValueError):
    pass


@dataclass(frozen=True)
class ImportaRevisaoBancaResult:
    bank_raw_dir: Path
    periodo: str
    cnpjs_processados: tuple[str, ...]
    arquivos_lidos: tuple[str, ...]
    comprovantes_lidos: int
    lancamentos_lidos: int
    comprovantes_candidatos: int
    conta_corrente_candidatos: int
    excel_comprovantes: Path
    excel_conta_corrente: Path
    manifest_path: Path
    manifest_historico_path: Path
    ofx_encontrado: bool
    ofx_usado: bool
    decisoes_preservadas: int
    novas_chaves_adicionadas: int
    conta_corrente_regras_yaml_aplicadas: int
    conta_corrente_boleto_pago_fallback: int
    conta_corrente_revisao_sem_regra: int
    conta_corrente_pix_enviado: int
    conta_corrente_pix_comprovante_exato: int
    conta_corrente_pix_yaml_via_comprovante: int
    conta_corrente_pix_revisao_sem_regra_comprovante: int
    conta_corrente_pix_revisao_match_ambiguo: int
    conta_corrente_pix_revisao_sem_comprovante: int
    conta_corrente_pix_fallback_aplicado: int
    possiveis_comprovantes_relacionados: int
    alertas: tuple[str, ...]


@dataclass
class _Group:
    chave_estavel: str
    rows: list[dict[str, object]]
    tecnico_rows: list[dict[str, object]]
    preserved: dict[str, object] | None = None


@dataclass(frozen=True)
class _ReceiptBridge:
    status: str
    receipt: PaymentReceiptRaw | None
    result: PaymentClassificationResult
    candidate_count: int = 0


def resolve_bank_raw_dir(
    *,
    bank_raw_dir: Path | None = None,
    ano: str | None = None,
    mes: str | None = None,
    project_root: Path | None = None,
) -> Path:
    if bank_raw_dir is not None:
        if ano or mes:
            raise ImportaRevisaoBancaError(
                "Use --bank-raw-dir ou --ano/--mes, nao ambos."
            )
        return Path(bank_raw_dir)

    if not ano or not mes:
        raise ImportaRevisaoBancaError(
            "Informe --bank-raw-dir ou ambos --ano e --mes."
        )
    year = str(ano).strip()
    month = _normalize_month_arg(mes, year)
    root = Path(project_root or Path.cwd())
    return root / "bank_raw" / year / month


def importar_revisao_banca(
    *,
    bank_raw_dir: Path,
    output_dir: Path,
    regras_path: Path,
    plano_contas_path: Path,
    mapeamento_dre_path: Path,
) -> ImportaRevisaoBancaResult:
    bank_raw_dir = Path(bank_raw_dir)
    output_dir = Path(output_dir)
    _validate_month_dir(bank_raw_dir)
    periodo = _periodo_from_dir(bank_raw_dir)

    rules = load_payment_rules(regras_path)
    exclusion_rules = load_payment_exclusion_rules(regras_path)
    validate_payment_rules(
        rules,
        plano_contas_path=plano_contas_path,
        mapeamento_dre_path=mapeamento_dre_path,
    )

    account_options = _load_account_options(plano_contas_path)
    output_dir.mkdir(parents=True, exist_ok=True)
    comprovantes_path = output_dir / COMPROVANTES_XLSX
    conta_corrente_path = output_dir / CONTA_CORRENTE_XLSX
    manifest_path = output_dir / MANIFEST_NAME
    manifest_historico_path = output_dir / MANIFEST_HISTORICO_TEMPLATE.format(
        periodo=periodo
    )

    existing_comprovantes = _load_existing_decisions(
        comprovantes_path,
        COMPROVANTES_SHEET,
        kind="comprovantes",
    )
    existing_conta = _load_existing_decisions(
        conta_corrente_path,
        CONTA_CORRENTE_SHEET,
        kind="conta_corrente",
    )

    cnpjs: list[str] = []
    arquivos_lidos: list[str] = []
    ofx_found = False
    comprovantes_lidos = 0
    lancamentos_lidos = 0
    conta_corrente_regras_yaml_aplicadas = 0
    conta_corrente_boleto_pago_fallback = 0
    conta_corrente_revisao_sem_regra = 0
    conta_corrente_pix_enviado = 0
    conta_corrente_pix_comprovante_exato = 0
    conta_corrente_pix_yaml_via_comprovante = 0
    conta_corrente_pix_revisao_sem_regra_comprovante = 0
    conta_corrente_pix_revisao_match_ambiguo = 0
    conta_corrente_pix_revisao_sem_comprovante = 0
    conta_corrente_pix_fallback_aplicado = 0
    comprovante_groups: dict[str, _Group] = {}
    conta_groups: dict[str, _Group] = {}
    conta_tecnico_regra_existente: list[dict[str, object]] = []
    receipts_by_exact_context: dict[tuple[str, str, str, str], list[PaymentReceiptRaw]] = {}
    pix_receipt_hashes_usados: set[str] = set()

    for cnpj_dir in _cnpj_dirs(bank_raw_dir):
        cnpj = only_digits(cnpj_dir.name)
        cnpjs.append(cnpj)
        consulta_files = sorted(cnpj_dir.glob("ConsultaPagamentos*.xls"))
        lancamento_files = sorted(cnpj_dir.glob("Extrato_Lan*amentos*.xlsx"))
        ofx_found = ofx_found or any(cnpj_dir.glob("Extrato*.ofx"))
        if not consulta_files and not lancamento_files:
            raise ImportaRevisaoBancaError(
                f"{cnpj_dir} nao contem ConsultaPagamentos*.xls nem "
                "Extrato_Lancamentos*.xlsx."
            )

        for consulta_path in consulta_files:
            arquivos_lidos.append(str(consulta_path))
            receipts = parse_payment_receipts(consulta_path)
            comprovantes_lidos += len(receipts)
            for receipt in receipts:
                if _is_non_operational_receipt(receipt):
                    continue
                _register_receipt_exact_context(
                    receipts_by_exact_context,
                    cnpj=cnpj,
                    periodo=periodo,
                    receipt=receipt,
                )
                row = _receipt_dict(receipt, cnpj)
                result = classify_payment_receipt(row, rules, exclusion_rules)
                if result.status_classificacao in {STATUS_CLASSIFICADO, STATUS_EXCLUIDO}:
                    continue
                key = _stable_key_comprovante(receipt)
                group = comprovante_groups.setdefault(
                    key,
                    _Group(chave_estavel=key, rows=[], tecnico_rows=[]),
                )
                group.rows.append(
                    _comprovante_source_row(receipt, cnpj, periodo, consulta_path, result)
                )
                group.tecnico_rows.append(
                    _tecnico_comprovante_row(receipt, cnpj, consulta_path, key, result)
                )

        for lancamentos_path in lancamento_files:
            arquivos_lidos.append(str(lancamentos_path))
            launches, _audits = parse_bank_launches(lancamentos_path)
            lancamentos_lidos += len(launches)
            for launch in launches:
                if launch.amount is None or launch.amount >= 0:
                    continue
                tipo_movimento = _tipo_movimento(launch)
                bridge = _empty_receipt_bridge()
                if tipo_movimento == "PIX_TRANSFERENCIA_SAIDA":
                    conta_corrente_pix_enviado += 1
                    bridge = _find_pix_receipt_bridge(
                        launch,
                        cnpj,
                        periodo,
                        receipts_by_exact_context,
                        pix_receipt_hashes_usados,
                        rules,
                        exclusion_rules,
                    )
                    if bridge.receipt is not None:
                        pix_receipt_hashes_usados.add(bridge.receipt.raw_hash)
                    if bridge.status != "SEM_COMPROVANTE":
                        conta_corrente_pix_comprovante_exato += 1
                    if bridge.status == "MATCH_AMBIGUO":
                        conta_corrente_pix_revisao_match_ambiguo += 1
                        result = bridge.result
                    elif bridge.status == "COMPROVANTE_SEM_REGRA":
                        conta_corrente_pix_revisao_sem_regra_comprovante += 1
                        result = bridge.result
                    elif bridge.status == "CLASSIFICADO":
                        conta_corrente_pix_yaml_via_comprovante += 1
                        result = bridge.result
                    else:
                        result = _classify_conta_corrente(
                            _conta_rule_input(launch),
                            rules,
                            exclusion_rules,
                        )
                else:
                    result = _classify_conta_corrente(
                        _conta_rule_input(launch),
                        rules,
                        exclusion_rules,
                    )
                key = _stable_key_conta_corrente(launch)
                tecnico_row = _tecnico_conta_row(
                    launch,
                    cnpj,
                    periodo,
                    lancamentos_path,
                    key,
                    result,
                    bridge,
                )
                if result.status_classificacao in {STATUS_CLASSIFICADO, STATUS_EXCLUIDO}:
                    if result.regra_id == BOLETO_PAGO_FALLBACK_REGRA_ID:
                        conta_corrente_boleto_pago_fallback += 1
                        if tipo_movimento == "PIX_TRANSFERENCIA_SAIDA":
                            conta_corrente_pix_fallback_aplicado += 1
                    else:
                        conta_corrente_regras_yaml_aplicadas += 1
                    conta_tecnico_regra_existente.append(tecnico_row)
                    continue
                if result.status_classificacao == STATUS_SEM_REGRA:
                    if (
                        tipo_movimento == "PIX_TRANSFERENCIA_SAIDA"
                        and bridge.status == "SEM_COMPROVANTE"
                    ):
                        conta_corrente_pix_revisao_sem_comprovante += 1
                    conta_corrente_revisao_sem_regra += 1
                group = conta_groups.setdefault(
                    key,
                    _Group(chave_estavel=key, rows=[], tecnico_rows=[]),
                )
                group.rows.append(
                    _conta_source_row(
                        launch,
                        cnpj,
                        periodo,
                        lancamentos_path,
                        result,
                        bridge,
                    )
                )
                group.tecnico_rows.append(tecnico_row)

    comprovante_rows, comprovante_tecnico, preserved_comprovantes, new_comprovantes = (
        _build_comprovante_output_rows(comprovante_groups, existing_comprovantes)
    )
    conta_rows, conta_tecnico, preserved_conta, new_conta = _build_conta_output_rows(
        conta_groups,
        existing_conta,
        conta_tecnico_regra_existente,
    )
    possiveis_comprovantes_relacionados = sum(
        1
        for row in conta_rows
        if row.get("POSSIVEL_COMPROVANTE_RELACIONADO") in {"SIM", "MATCH_AMBIGUO"}
    )
    alertas = _build_alertas(possiveis_comprovantes_relacionados)

    _write_decision_workbook(
        comprovantes_path,
        sheet_name=COMPROVANTES_SHEET,
        headers=COMPROVANTES_HEADERS,
        rows=comprovante_rows,
        technical_sheet=TECNICO_COMPROVANTES_SHEET,
        technical_headers=TECNICO_COMPROVANTES_HEADERS,
        technical_rows=comprovante_tecnico,
        account_options=account_options,
    )
    _write_decision_workbook(
        conta_corrente_path,
        sheet_name=CONTA_CORRENTE_SHEET,
        headers=CONTA_CORRENTE_HEADERS,
        rows=conta_rows,
        technical_sheet=TECNICO_CONTA_CORRENTE_SHEET,
        technical_headers=TECNICO_CONTA_CORRENTE_HEADERS,
        technical_rows=conta_tecnico,
        account_options=account_options,
    )

    result = ImportaRevisaoBancaResult(
        bank_raw_dir=bank_raw_dir,
        periodo=periodo,
        cnpjs_processados=tuple(sorted(set(cnpjs))),
        arquivos_lidos=tuple(arquivos_lidos),
        comprovantes_lidos=comprovantes_lidos,
        lancamentos_lidos=lancamentos_lidos,
        comprovantes_candidatos=len(comprovante_rows),
        conta_corrente_candidatos=len(conta_rows),
        excel_comprovantes=comprovantes_path,
        excel_conta_corrente=conta_corrente_path,
        manifest_path=manifest_path,
        manifest_historico_path=manifest_historico_path,
        ofx_encontrado=ofx_found,
        ofx_usado=False,
        decisoes_preservadas=preserved_comprovantes + preserved_conta,
        novas_chaves_adicionadas=new_comprovantes + new_conta,
        conta_corrente_regras_yaml_aplicadas=conta_corrente_regras_yaml_aplicadas,
        conta_corrente_boleto_pago_fallback=conta_corrente_boleto_pago_fallback,
        conta_corrente_revisao_sem_regra=conta_corrente_revisao_sem_regra,
        conta_corrente_pix_enviado=conta_corrente_pix_enviado,
        conta_corrente_pix_comprovante_exato=conta_corrente_pix_comprovante_exato,
        conta_corrente_pix_yaml_via_comprovante=conta_corrente_pix_yaml_via_comprovante,
        conta_corrente_pix_revisao_sem_regra_comprovante=conta_corrente_pix_revisao_sem_regra_comprovante,
        conta_corrente_pix_revisao_match_ambiguo=conta_corrente_pix_revisao_match_ambiguo,
        conta_corrente_pix_revisao_sem_comprovante=conta_corrente_pix_revisao_sem_comprovante,
        conta_corrente_pix_fallback_aplicado=conta_corrente_pix_fallback_aplicado,
        possiveis_comprovantes_relacionados=possiveis_comprovantes_relacionados,
        alertas=alertas,
    )
    _write_manifest(result)
    return result


def console_lines(result: ImportaRevisaoBancaResult) -> list[str]:
    return [
        "[OK] Importa revisao banca concluido.",
        f"Periodo: {result.periodo}",
        f"CNPJs: {', '.join(result.cnpjs_processados)}",
        f"Comprovantes lidos: {result.comprovantes_lidos}",
        f"Lancamentos lidos: {result.lancamentos_lidos}",
        f"Candidatos comprovantes: {result.comprovantes_candidatos}",
        f"Candidatos conta corrente: {result.conta_corrente_candidatos}",
        f"Conta corrente com regra YAML aplicada: {result.conta_corrente_regras_yaml_aplicadas}",
        f"Conta corrente excluida por fallback BOLETO PAGO: {result.conta_corrente_boleto_pago_fallback}",
        f"Conta corrente sem regra para revisao: {result.conta_corrente_revisao_sem_regra}",
        f"PIX ENVIADO no conta corrente: {result.conta_corrente_pix_enviado}",
        f"PIX ENVIADO com comprovante exato: {result.conta_corrente_pix_comprovante_exato}",
        f"PIX ENVIADO com regra YAML via comprovante: {result.conta_corrente_pix_yaml_via_comprovante}",
        f"PIX ENVIADO em revisao por comprovante sem regra: {result.conta_corrente_pix_revisao_sem_regra_comprovante}",
        f"PIX ENVIADO em revisao por match ambiguo: {result.conta_corrente_pix_revisao_match_ambiguo}",
        f"PIX ENVIADO em revisao por ausencia de comprovante: {result.conta_corrente_pix_revisao_sem_comprovante}",
        f"PIX ENVIADO classificados/excluidos por fallback: {result.conta_corrente_pix_fallback_aplicado}",
        f"Decisoes preservadas: {result.decisoes_preservadas}",
        f"Novas chaves adicionadas: {result.novas_chaves_adicionadas}",
        f"Excel comprovantes: {result.excel_comprovantes}",
        f"Excel conta corrente: {result.excel_conta_corrente}",
        f"Manifest: {result.manifest_path}",
        f"Manifest historico: {result.manifest_historico_path}",
        f"Possiveis comprovantes relacionados no conta corrente: {result.possiveis_comprovantes_relacionados}",
        *[f"ALERTA: {alerta}" for alerta in result.alertas],
        "OFX usado: false",
        "YAML de regras nao modificado.",
    ]


def _normalize_month_arg(value: str, year: str) -> str:
    raw = str(value).strip()
    match = re.fullmatch(r"(\d{4})-(\d{2})", raw)
    if match:
        if match.group(1) != year:
            raise ImportaRevisaoBancaError("--ano diverge de --mes.")
        return raw
    if re.fullmatch(r"\d{1,2}", raw):
        month = int(raw)
        if 1 <= month <= 12:
            return f"{year}-{month:02d}"
    raise ImportaRevisaoBancaError("--mes deve ser MM ou YYYY-MM.")


def _validate_month_dir(path: Path) -> None:
    if not path.exists() or not path.is_dir():
        raise ImportaRevisaoBancaError(f"Pasta mes nao encontrada: {path}")
    if not _cnpj_dirs(path):
        raise ImportaRevisaoBancaError(f"Pasta mes sem subpastas CNPJ: {path}")


def _cnpj_dirs(path: Path) -> list[Path]:
    return [
        entry
        for entry in sorted(path.iterdir(), key=lambda item: item.name)
        if entry.is_dir() and only_digits(entry.name)
    ]


def _periodo_from_dir(path: Path) -> str:
    if re.fullmatch(r"\d{4}-\d{2}", path.name):
        return path.name
    raise ImportaRevisaoBancaError(
        f"Pasta mes deve terminar com YYYY-MM: {path}"
    )


def _receipt_dict(receipt: PaymentReceiptRaw, cnpj: str) -> dict[str, str]:
    return {
        "cnpj": cnpj,
        "payment_date": _date_text(receipt.payment_date),
        "amount": _decimal_text(receipt.amount),
        "beneficiary_name": receipt.beneficiary_name,
        "beneficiary_document": receipt.beneficiary_document,
        "payment_type": receipt.payment_type,
        "reference_id": receipt.reference_id,
        "status": receipt.status,
        "source_file": receipt.source_file,
        "raw_hash": receipt.raw_hash,
    }


def _conta_rule_input(launch: BankLaunchEnrichment) -> dict[str, str]:
    return {
        "beneficiary_name": launch.counterparty_name,
        "reference_id": launch.description,
        "payment_type": _tipo_movimento(launch),
        "description_receipt": launch.description,
        "source_file": launch.source_file,
        "raw_hash": launch.raw_hash,
    }


def _classify_conta_corrente(
    row: dict[str, str],
    rules,
    exclusion_rules,
):
    specific_exclusions, fallback_exclusions = _split_boleto_fallback_exclusions(
        exclusion_rules
    )
    specific_result = classify_payment_receipt(row, rules, specific_exclusions)
    if specific_result.status_classificacao != STATUS_SEM_REGRA:
        return specific_result
    return classify_payment_receipt(row, [], fallback_exclusions)


def _split_boleto_fallback_exclusions(exclusion_rules):
    specific_exclusions = []
    fallback_exclusions = []
    for rule in exclusion_rules:
        if getattr(rule, "regra_id", "") == BOLETO_PAGO_FALLBACK_REGRA_ID:
            fallback_exclusions.append(rule)
        else:
            specific_exclusions.append(rule)
    return specific_exclusions, fallback_exclusions


def _register_receipt_exact_context(
    receipts_by_exact_context: dict[tuple[str, str, str, str], list[PaymentReceiptRaw]],
    *,
    cnpj: str,
    periodo: str,
    receipt: PaymentReceiptRaw,
) -> None:
    if not receipt.payment_date:
        return
    key = (
        cnpj,
        periodo,
        _date_text(receipt.payment_date),
        _decimal_text(receipt.amount),
    )
    receipts_by_exact_context.setdefault(key, []).append(receipt)


def _find_pix_receipt_bridge(
    launch: BankLaunchEnrichment,
    cnpj: str,
    periodo: str,
    receipts_by_exact_context: dict[tuple[str, str, str, str], list[PaymentReceiptRaw]],
    used_receipt_hashes: set[str],
    rules,
    exclusion_rules,
) -> _ReceiptBridge:
    if not launch.transaction_date or launch.amount is None:
        return _empty_receipt_bridge()
    key = (
        cnpj,
        periodo,
        _date_text(launch.transaction_date),
        _decimal_text(abs(launch.amount)),
    )
    exact_candidates = [
        receipt
        for receipt in receipts_by_exact_context.get(key, [])
        if receipt.raw_hash not in used_receipt_hashes
    ]
    candidates = [
        receipt
        for receipt in exact_candidates
        if _receipt_matches_launch_identity(launch, receipt)
    ]
    if not candidates:
        return _empty_receipt_bridge()
    if len(candidates) > 1:
        return _ReceiptBridge(
            status="MATCH_AMBIGUO",
            receipt=None,
            result=_review_result("MATCH_AMBIGUO_COMPROVANTE"),
            candidate_count=len(candidates),
        )

    receipt = candidates[0]
    result = classify_payment_receipt(_receipt_dict(receipt, cnpj), rules, exclusion_rules)
    if result.status_classificacao in {STATUS_CLASSIFICADO, STATUS_EXCLUIDO}:
        return _ReceiptBridge(
            status="CLASSIFICADO",
            receipt=receipt,
            result=result,
            candidate_count=1,
        )
    return _ReceiptBridge(
        status="COMPROVANTE_SEM_REGRA",
        receipt=receipt,
        result=_review_result(
            "COMPROVANTE_RELACIONADO_SEM_REGRA",
            result.regras_candidatas,
        ),
        candidate_count=1,
    )


def _empty_receipt_bridge() -> _ReceiptBridge:
    return _ReceiptBridge(
        status="SEM_COMPROVANTE",
        receipt=None,
        result=_review_result("SEM_COMPROVANTE_EXATO"),
    )


def _review_result(
    motivo: str,
    regras_candidatas: tuple[str, ...] = (),
) -> PaymentClassificationResult:
    return PaymentClassificationResult(
        status_classificacao=STATUS_SEM_REGRA,
        conta_id="",
        conta_nome="",
        dre_linha_id="",
        entra_dre=None,
        regra_id="",
        motivo=motivo,
        regras_candidatas=regras_candidatas,
    )


def _receipt_matches_launch_identity(
    launch: BankLaunchEnrichment,
    receipt: PaymentReceiptRaw,
) -> bool:
    launch_document = only_digits(launch.counterparty_document)
    receipt_document = only_digits(receipt.beneficiary_document)
    if launch_document and receipt_document:
        # Both sides have documents: strict match required.
        return launch_document == receipt_document
    receipt_name = normalize_match_text(receipt.beneficiary_name)
    if not receipt_name:
        return False
    launch_name = normalize_match_text(launch.counterparty_name)
    launch_description = normalize_match_text(launch.description)
    if launch_name and _names_compatible(launch_name, receipt_name):
        return True
    if launch_description and receipt_name in launch_description:
        return True
    # One or both sides lack document; date+value+CNPJ (context key) already match.
    # Accept the bridge: multiple candidates produce MATCH_AMBIGUO, not arbitrary picks.
    return True


def _names_compatible(left: str, right: str) -> bool:
    if left == right:
        return True
    return len(left) >= 8 and len(right) >= 8 and (left in right or right in left)


def _is_non_operational_receipt(receipt: PaymentReceiptRaw) -> bool:
    return normalize_match_text(receipt.beneficiary_name) == "TOTAL"


def _stable_key_comprovante(receipt: PaymentReceiptRaw) -> str:
    beneficiary = normalize_match_text(receipt.beneficiary_name)
    reference = normalize_match_text(receipt.reference_id) or "SEM_REFERENCIA"
    payment_type = normalize_match_text(receipt.payment_type) or "SEM_TIPO"
    return (
        f"COMPROVANTE|BENEFICIARIO={beneficiary}|REFERENCIA={reference}|"
        f"TIPO={payment_type}"
    )


def _stable_key_conta_corrente(launch: BankLaunchEnrichment) -> str:
    history = normalize_match_text(launch.description)
    razao = normalize_match_text(launch.counterparty_name) or "SEM_RAZAO"
    tipo = _tipo_movimento(launch)
    return (
        f"CONTA_CORRENTE|HISTORICO={history}|RAZAO={razao}|"
        f"TIPO={tipo}|DIRECAO=SAIDA"
    )


def _comprovante_source_row(
    receipt: PaymentReceiptRaw,
    cnpj: str,
    periodo: str,
    source_path: Path,
    result: object,
) -> dict[str, object]:
    return {
        "PERIODO": periodo,
        "CNPJ": cnpj,
        "DATA_PAGAMENTO": _date_text(receipt.payment_date),
        "VALOR": receipt.amount,
        "BENEFICIARIO_FAVORECIDO": normalize_match_text(receipt.beneficiary_name),
        "REFERENCIA_EMPRESA": receipt.reference_id,
        "TIPO_PAGAMENTO": receipt.payment_type,
        "CHAVE_REGRA_BENEFICIARIO": normalize_match_text(receipt.beneficiary_name),
        "CHAVE_REGRA_REFERENCIA": normalize_match_text(receipt.reference_id),
        "MOTIVO_REVISAO": getattr(result, "motivo", "SEM_REGRA"),
        "REGRA_EXISTENTE": "|".join(getattr(result, "regras_candidatas", ()) or ()),
        "ORIGEM_ARQUIVO": str(source_path),
        "HASH_ORIGEM": receipt.raw_hash,
    }


def _conta_source_row(
    launch: BankLaunchEnrichment,
    cnpj: str,
    periodo: str,
    source_path: Path,
    result: object,
    bridge: _ReceiptBridge,
) -> dict[str, object]:
    return {
        "CNPJ": cnpj,
        "MES": periodo,
        "DATA": _date_text(launch.transaction_date),
        "VALOR": launch.amount,
        "LANCAMENTO": launch.description,
        "RAZAO_SOCIAL": launch.counterparty_name,
        "CHAVE_REGRA_LANCAMENTO": normalize_match_text(launch.description),
        "CHAVE_REGRA_RAZAO_SOCIAL": normalize_match_text(launch.counterparty_name),
        "POSSIVEL_COMPROVANTE_RELACIONADO": _bridge_related_status(bridge),
        "COMPROVANTE_RELACIONADO_BENEFICIARIO": _bridge_receipt_value(
            bridge,
            "beneficiary_name",
        ),
        "COMPROVANTE_RELACIONADO_DOCUMENTO": _bridge_receipt_value(
            bridge,
            "beneficiary_document",
        ),
        "COMPROVANTE_RELACIONADO_DATA": _bridge_receipt_value(bridge, "payment_date"),
        "COMPROVANTE_RELACIONADO_VALOR": _bridge_receipt_value(bridge, "amount"),
        "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE": _bridge_regra_existente(bridge),
        "USAR_SOMENTE_SEM_COMPROVANTE": "SIM",
        "MOTIVO_REVISAO": getattr(result, "motivo", "SEM_REGRA"),
        "REGRA_EXISTENTE": "|".join(getattr(result, "regras_candidatas", ()) or ()),
        "ORIGEM_ARQUIVO": str(source_path),
        "HASH_ORIGEM": launch.raw_hash,
        "TIPO_MOVIMENTO": _tipo_movimento(launch),
    }


def _bridge_related_status(bridge: _ReceiptBridge) -> str:
    if bridge.status == "MATCH_AMBIGUO":
        return "MATCH_AMBIGUO"
    if bridge.status in {"CLASSIFICADO", "COMPROVANTE_SEM_REGRA"}:
        return "SIM"
    return "NAO"


def _bridge_receipt_value(bridge: _ReceiptBridge, field_name: str) -> str:
    if bridge.receipt is None:
        return ""
    value = getattr(bridge.receipt, field_name)
    if field_name == "payment_date":
        return _date_text(value)
    if field_name == "amount":
        return _decimal_text(value)
    return str(value or "")


def _bridge_regra_existente(bridge: _ReceiptBridge) -> str:
    if bridge.status == "MATCH_AMBIGUO":
        return f"MATCH_AMBIGUO:{bridge.candidate_count}"
    return "|".join(bridge.result.regras_candidatas)


def _tecnico_comprovante_row(
    receipt: PaymentReceiptRaw,
    cnpj: str,
    source_path: Path,
    key: str,
    result: object,
) -> dict[str, object]:
    return {
        "CHAVE_ESTAVEL": key,
        "ORIGEM_ARQUIVO": str(source_path),
        "TIPO_FONTE": "ConsultaPagamentos.xls",
        "DATAS_EXEMPLO": _date_text(receipt.payment_date),
        "VALORES_EXEMPLO": _decimal_text(receipt.amount),
        "REGRA_EXISTENTE": "|".join(getattr(result, "regras_candidatas", ()) or ()),
        "MOTIVO_REVISAO": getattr(result, "motivo", "SEM_REGRA"),
        "HASH_ORIGEM": receipt.raw_hash,
        "CNPJS_ENVOLVIDOS": cnpj,
    }


def _tecnico_conta_row(
    launch: BankLaunchEnrichment,
    cnpj: str,
    periodo: str,
    source_path: Path,
    key: str,
    result: object,
    bridge: _ReceiptBridge,
) -> dict[str, object]:
    return {
        "CHAVE_ESTAVEL": key,
        "CNPJ": cnpj,
        "MES": periodo,
        "ORIGEM_ARQUIVO": str(source_path),
        "DATA_MOVIMENTO": _date_text(launch.transaction_date),
        "VALORES_EXEMPLO": _decimal_text(launch.amount),
        "DIRECAO": "SAIDA",
        "TIPO_MOVIMENTO": _tipo_movimento(launch),
        "REGRA_EXISTENTE": "|".join(getattr(result, "regras_candidatas", ()) or ()),
        "MOTIVO_REVISAO": getattr(result, "motivo", "SEM_REGRA"),
        "HASH_ORIGEM": launch.raw_hash,
        "COMPROVANTE_RELACIONADO_BENEFICIARIO": _bridge_receipt_value(
            bridge,
            "beneficiary_name",
        ),
        "COMPROVANTE_RELACIONADO_DOCUMENTO": _bridge_receipt_value(
            bridge,
            "beneficiary_document",
        ),
        "COMPROVANTE_RELACIONADO_DATA": _bridge_receipt_value(bridge, "payment_date"),
        "COMPROVANTE_RELACIONADO_VALOR": _bridge_receipt_value(bridge, "amount"),
        "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE": _bridge_regra_existente(bridge),
        "COMPETENCIA_REFERENCIA": _competencia_referencia_conta_corrente(
            launch,
            result,
        ),
        "OBSERVACAO_AUDIT": _observacao_audit_conta_corrente(launch, result),
    }


def _build_comprovante_output_rows(
    groups: dict[str, _Group],
    existing: dict[str, dict[str, object]],
) -> tuple[list[dict[str, object]], list[dict[str, object]], int, int]:
    rows: list[dict[str, object]] = []
    technical_rows: list[dict[str, object]] = []
    preserved = 0
    new = 0
    for key in sorted(groups):
        group = groups[key]
        preserved_row = existing.get(key)
        if preserved_row:
            preserved += 1
        else:
            new += 1
        first = group.rows[0]
        base = {
            "STATUS_DECISAO": "",
            "ACAO_REGRA": "",
            "PERIODO": first.get("PERIODO", ""),
            "CNPJ": _joined_unique(row.get("CNPJ") for row in group.rows),
            "DATA_PAGAMENTO": first.get("DATA_PAGAMENTO", ""),
            "VALOR": first.get("VALOR", ""),
            "BENEFICIARIO_FAVORECIDO": first.get("BENEFICIARIO_FAVORECIDO", ""),
            "REFERENCIA_EMPRESA": first.get("REFERENCIA_EMPRESA", ""),
            "TIPO_PAGAMENTO": first.get("TIPO_PAGAMENTO", ""),
            "CONTA_ESCOLHIDA": "",
            "CONTA_ID_APROVADA": "",
            "CONTA_NOME_APROVADA": "",
            "CHAVE_REGRA_BENEFICIARIO": first.get("CHAVE_REGRA_BENEFICIARIO", ""),
            "CHAVE_REGRA_REFERENCIA": first.get("CHAVE_REGRA_REFERENCIA", ""),
            "RISCO_DUPLICIDADE_XML": "",
            "ESCOPO_REGRA": "",
            "QTD_OCORRENCIAS": len(group.rows),
            "VALOR_TOTAL": _sum_decimals(row.get("VALOR") for row in group.rows),
            "EXEMPLOS": _examples(
                f"{row.get('DATA_PAGAMENTO')} | {row.get('VALOR')} | "
                f"{row.get('BENEFICIARIO_FAVORECIDO')} | {row.get('REFERENCIA_EMPRESA')}"
                for row in group.rows
            ),
            "OBSERVACAO_OPERADOR": "",
        }
        _preserve_manual(base, preserved_row)
        rows.append(base)
        technical_rows.extend(_merge_technical_rows(group.tecnico_rows, key))
    return rows, technical_rows, preserved, new


def _build_conta_output_rows(
    groups: dict[str, _Group],
    existing: dict[str, dict[str, object]],
    applied_technical_rows: list[dict[str, object]] | None = None,
) -> tuple[list[dict[str, object]], list[dict[str, object]], int, int]:
    rows: list[dict[str, object]] = []
    technical_rows: list[dict[str, object]] = list(applied_technical_rows or [])
    preserved = 0
    new = 0
    for key in sorted(groups):
        group = groups[key]
        preserved_row = existing.get(key)
        if preserved_row:
            preserved += 1
        else:
            new += 1
        first = group.rows[0]
        base = {
            "STATUS_DECISAO": "",
            "ACAO_REGRA": "",
            "DATA": first.get("DATA", ""),
            "VALOR": first.get("VALOR", ""),
            "LANCAMENTO": first.get("LANCAMENTO", ""),
            "RAZAO_SOCIAL": first.get("RAZAO_SOCIAL", ""),
            "CONTA_ESCOLHIDA": "",
            "CONTA_ID_APROVADA": "",
            "CONTA_NOME_APROVADA": "",
            "CHAVE_REGRA_LANCAMENTO": first.get("CHAVE_REGRA_LANCAMENTO", ""),
            "CHAVE_REGRA_RAZAO_SOCIAL": first.get("CHAVE_REGRA_RAZAO_SOCIAL", ""),
            "POSSIVEL_COMPROVANTE_RELACIONADO": first.get(
                "POSSIVEL_COMPROVANTE_RELACIONADO", "NAO"
            ),
            "COMPROVANTE_RELACIONADO_BENEFICIARIO": first.get(
                "COMPROVANTE_RELACIONADO_BENEFICIARIO", ""
            ),
            "COMPROVANTE_RELACIONADO_DOCUMENTO": first.get(
                "COMPROVANTE_RELACIONADO_DOCUMENTO", ""
            ),
            "COMPROVANTE_RELACIONADO_DATA": first.get(
                "COMPROVANTE_RELACIONADO_DATA", ""
            ),
            "COMPROVANTE_RELACIONADO_VALOR": first.get(
                "COMPROVANTE_RELACIONADO_VALOR", ""
            ),
            "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE": first.get(
                "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE", ""
            ),
            "USAR_SOMENTE_SEM_COMPROVANTE": "SIM",
            "RISCO_DUPLICIDADE_XML": "",
            "ESCOPO_REGRA": "",
            "QTD_OCORRENCIAS": len(group.rows),
            "VALOR_TOTAL": _sum_decimals(row.get("VALOR") for row in group.rows),
            "EXEMPLOS": _examples(
                f"{row.get('DATA')} | {row.get('VALOR')} | {row.get('LANCAMENTO')} | "
                f"{row.get('RAZAO_SOCIAL')}"
                for row in group.rows
            ),
            "OBSERVACAO_OPERADOR": "",
        }
        _preserve_manual(base, preserved_row)
        rows.append(base)
        technical_rows.extend(_merge_technical_rows(group.tecnico_rows, key))
    return rows, technical_rows, preserved, new


def _merge_technical_rows(rows: list[dict[str, object]], key: str) -> list[dict[str, object]]:
    if not rows:
        return []
    merged = dict(rows[0])
    merged["CHAVE_ESTAVEL"] = key
    for column in (
        "ORIGEM_ARQUIVO",
        "DATAS_EXEMPLO",
        "VALORES_EXEMPLO",
        "HASH_ORIGEM",
        "CNPJS_ENVOLVIDOS",
        "CNPJ",
        "DATA_MOVIMENTO",
        "COMPROVANTE_RELACIONADO_BENEFICIARIO",
        "COMPROVANTE_RELACIONADO_DOCUMENTO",
        "COMPROVANTE_RELACIONADO_DATA",
        "COMPROVANTE_RELACIONADO_VALOR",
        "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE",
        "COMPETENCIA_REFERENCIA",
        "OBSERVACAO_AUDIT",
    ):
        if column in merged:
            merged[column] = _joined_unique((row.get(column) for row in rows), limit=5)
    return [merged]


def _load_existing_decisions(
    path: Path,
    sheet_name: str,
    *,
    kind: str,
) -> dict[str, dict[str, object]]:
    if not path.exists():
        return {}
    workbook = load_workbook(path, data_only=False)
    if sheet_name not in workbook.sheetnames:
        return {}
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {}
    headers = [str(value or "").strip() for value in rows[0]]
    result: dict[str, dict[str, object]] = {}
    for values in rows[1:]:
        row = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        if not any(value not in (None, "") for value in row.values()):
            continue
        key = _existing_key(row, kind)
        if key:
            result[key] = row
    return result


def _existing_key(row: dict[str, object], kind: str) -> str:
    if row.get("CHAVE_ESTAVEL"):
        return str(row["CHAVE_ESTAVEL"])
    if kind == "comprovantes":
        beneficiary = normalize_match_text(
            row.get("BENEFICIARIO_FAVORECIDO")
            or row.get("BENEFICIARIO_NORMALIZADO")
            or row.get("CHAVE_REGRA_BENEFICIARIO")
        )
        reference = normalize_match_text(
            row.get("REFERENCIA_EMPRESA") or row.get("CHAVE_REGRA_REFERENCIA")
        ) or "SEM_REFERENCIA"
        payment_type = normalize_match_text(row.get("TIPO_PAGAMENTO")) or "SEM_TIPO"
        if beneficiary:
            return (
                f"COMPROVANTE|BENEFICIARIO={beneficiary}|REFERENCIA={reference}|"
                f"TIPO={payment_type}"
            )
        return ""
    history = normalize_match_text(row.get("LANCAMENTO") or row.get("HISTORICO_NORMALIZADO"))
    razao = normalize_match_text(row.get("RAZAO_SOCIAL")) or "SEM_RAZAO"
    tipo = _tipo_movimento_text(history, _parse_decimal(row.get("VALOR")))
    if history:
        return (
            f"CONTA_CORRENTE|HISTORICO={history}|RAZAO={razao}|"
            f"TIPO={tipo}|DIRECAO=SAIDA"
        )
    return ""


def _preserve_manual(target: dict[str, object], existing: dict[str, object] | None) -> None:
    if not existing:
        return
    for column in MANUAL_COLUMNS:
        if column in target and existing.get(column) not in (None, ""):
            if column in {"CONTA_ID_APROVADA", "CONTA_NOME_APROVADA"}:
                continue
            target[column] = existing[column]
    conta_id, conta_nome = _account_from_choice(existing.get("CONTA_ESCOLHIDA"))
    if conta_id and "CONTA_ID_APROVADA" in target:
        target["CONTA_ID_APROVADA"] = conta_id
    elif (
        "CONTA_ID_APROVADA" in target
        and existing.get("CONTA_ID_APROVADA") not in (None, "")
        and not _is_formula(existing.get("CONTA_ID_APROVADA"))
    ):
        target["CONTA_ID_APROVADA"] = existing["CONTA_ID_APROVADA"]

    if conta_nome and "CONTA_NOME_APROVADA" in target:
        target["CONTA_NOME_APROVADA"] = conta_nome
    elif (
        "CONTA_NOME_APROVADA" in target
        and existing.get("CONTA_NOME_APROVADA") not in (None, "")
        and not _is_formula(existing.get("CONTA_NOME_APROVADA"))
    ):
        target["CONTA_NOME_APROVADA"] = existing["CONTA_NOME_APROVADA"]


def _account_from_choice(value: object) -> tuple[str, str]:
    choice = str(value or "").strip()
    if " - " not in choice or _is_formula(choice):
        return "", ""
    conta_id, conta_nome = choice.split(" - ", 1)
    return conta_id.strip(), conta_nome.strip()


def _is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _write_decision_workbook(
    path: Path,
    *,
    sheet_name: str,
    headers: list[str],
    rows: list[dict[str, object]],
    technical_sheet: str,
    technical_headers: list[str],
    technical_rows: list[dict[str, object]],
    account_options: list[tuple[str, str, str]],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    _append_rows(sheet, headers, rows)
    _style_sheet(sheet, headers)
    _add_decision_validations(sheet, headers, len(rows))
    _add_account_validation(sheet, headers, len(rows), len(account_options))
    _add_account_formulas(sheet, headers, len(rows))

    tech = workbook.create_sheet(technical_sheet)
    _append_rows(tech, technical_headers, technical_rows)
    _style_sheet(tech, technical_headers)

    plano = workbook.create_sheet(PLANO_CONTAS_SHEET)
    plano.append(["CONTA_ESCOLHIDA", "CONTA_ID", "CONTA_NOME"])
    for option, conta_id, conta_nome in account_options:
        plano.append([option, conta_id, conta_nome])
    _style_sheet(plano, ["CONTA_ESCOLHIDA", "CONTA_ID", "CONTA_NOME"])

    listas = workbook.create_sheet(LISTAS_SHEET)
    _append_rows(
        listas,
        ["STATUS_DECISAO", "ACAO_REGRA", "RISCO_DUPLICIDADE_XML", "ESCOPO_REGRA", "USAR_SOMENTE_SEM_COMPROVANTE"],
        [
            {
                "STATUS_DECISAO": "PENDENTE",
                "ACAO_REGRA": "CLASSIFICAR",
                "RISCO_DUPLICIDADE_XML": "OK",
                "ESCOPO_REGRA": "COMPROVANTE",
                "USAR_SOMENTE_SEM_COMPROVANTE": "SIM",
            },
            {
                "STATUS_DECISAO": "APROVADO",
                "ACAO_REGRA": "ESCLUDERE",
                "RISCO_DUPLICIDADE_XML": "DUVIDA",
                "ESCOPO_REGRA": "CONTA_CORRENTE",
                "USAR_SOMENTE_SEM_COMPROVANTE": "NAO",
            },
            {
                "STATUS_DECISAO": "REVISAR",
                "ACAO_REGRA": "XML",
                "RISCO_DUPLICIDADE_XML": "",
                "ESCOPO_REGRA": "",
                "USAR_SOMENTE_SEM_COMPROVANTE": "",
            },
            {
                "STATUS_DECISAO": "NAO_CRIAR_REGRA",
                "ACAO_REGRA": "NAO_CRIAR_REGRA",
                "RISCO_DUPLICIDADE_XML": "",
                "ESCOPO_REGRA": "",
                "USAR_SOMENTE_SEM_COMPROVANTE": "",
            },
            {
                "STATUS_DECISAO": "BLOQUEADO",
                "ACAO_REGRA": "REVISAR",
                "RISCO_DUPLICIDADE_XML": "",
                "ESCOPO_REGRA": "",
                "USAR_SOMENTE_SEM_COMPROVANTE": "",
            },
            {
                "STATUS_DECISAO": "",
                "ACAO_REGRA": "BLOQUEADO",
                "RISCO_DUPLICIDADE_XML": "",
                "ESCOPO_REGRA": "",
                "USAR_SOMENTE_SEM_COMPROVANTE": "",
            },
        ],
    )
    _style_sheet(listas, ["STATUS_DECISAO", "ACAO_REGRA", "RISCO_DUPLICIDADE_XML", "ESCOPO_REGRA", "USAR_SOMENTE_SEM_COMPROVANTE"])

    instrucoes = workbook.create_sheet(INSTRUCOES_SHEET)
    for line in _instructions(sheet_name):
        instrucoes.append([line])
    instrucoes.column_dimensions["A"].width = 110
    for row in instrucoes.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    instrucoes["A1"].font = Font(bold=True, size=14)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _append_rows(sheet, headers: list[str], rows: list[dict[str, object]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header, "")) for header in headers])


def _style_sheet(sheet, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        width = _suggested_width(header)
        sheet.column_dimensions[get_column_letter(index)].width = width
        for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=2, max_row=max(sheet.max_row, 2)):
            for item in cell:
                item.alignment = Alignment(wrap_text=_wrap(header), vertical="top")
                if header in {"VALOR", "VALOR_TOTAL"}:
                    item.number_format = '#,##0.00'
                if header == "QTD_OCORRENCIAS":
                    item.number_format = "0"


def _add_account_validation(sheet, headers: list[str], row_count: int, account_count: int) -> None:
    if row_count < 1 or "CONTA_ESCOLHIDA" not in headers or account_count < 1:
        return
    column = headers.index("CONTA_ESCOLHIDA") + 1
    letter = get_column_letter(column)
    validation = DataValidation(
        type="list",
        formula1=f"={PLANO_CONTAS_SHEET}!$A$2:$A${account_count + 1}",
        allow_blank=True,
    )
    validation.error = "Escolha uma conta da lista."
    validation.errorTitle = "Conta invalida"
    sheet.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{row_count + 1}")


def _add_decision_validations(sheet, headers: list[str], row_count: int) -> None:
    _add_list_validation(
        sheet,
        headers,
        row_count,
        "STATUS_DECISAO",
        BANK_STATUS_DECISAO_VALUES,
        "Escolha um status valido.",
        "Status invalido",
    )
    _add_list_validation(
        sheet,
        headers,
        row_count,
        "ACAO_REGRA",
        BANK_ACAO_REGRA_VALUES,
        "Escolha uma acao valida.",
        "Acao invalida",
    )
    _add_list_validation(
        sheet,
        headers,
        row_count,
        "RISCO_DUPLICIDADE_XML",
        RISCO_DUPLICIDADE_XML_VALUES,
        "Escolha um risco valido.",
        "Risco invalido",
    )
    _add_list_validation(
        sheet,
        headers,
        row_count,
        "ESCOPO_REGRA",
        ESCOPO_REGRA_VALUES,
        "Escolha um escopo valido.",
        "Escopo invalido",
    )
    _add_list_validation(
        sheet,
        headers,
        row_count,
        "USAR_SOMENTE_SEM_COMPROVANTE",
        USAR_SOMENTE_SEM_COMPROVANTE_VALUES,
        "Escolha SIM ou NAO.",
        "Valor invalido",
    )


def _add_list_validation(
    sheet,
    headers: list[str],
    row_count: int,
    column_name: str,
    values: tuple[str, ...],
    error: str,
    error_title: str,
) -> None:
    if row_count < 1 or column_name not in headers:
        return
    letter = get_column_letter(headers.index(column_name) + 1)
    validation = DataValidation(
        type="list",
        formula1=f'"{",".join(values)}"',
        allow_blank=True,
    )
    validation.error = error
    validation.errorTitle = error_title
    sheet.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{row_count + 1}")


def _add_account_formulas(sheet, headers: list[str], row_count: int) -> None:
    if row_count < 1:
        return
    if "CONTA_ESCOLHIDA" not in headers:
        return
    escolha_col = get_column_letter(headers.index("CONTA_ESCOLHIDA") + 1)
    conta_id_col = headers.index("CONTA_ID_APROVADA") + 1 if "CONTA_ID_APROVADA" in headers else None
    conta_nome_col = headers.index("CONTA_NOME_APROVADA") + 1 if "CONTA_NOME_APROVADA" in headers else None
    for row_index in range(2, row_count + 2):
        escolha_ref = f"{escolha_col}{row_index}"
        if conta_id_col and not sheet.cell(row=row_index, column=conta_id_col).value:
            sheet.cell(row=row_index, column=conta_id_col).value = (
                f'=IFERROR(LEFT({escolha_ref},FIND(" - ",{escolha_ref})-1),"")'
            )
        if conta_nome_col and not sheet.cell(row=row_index, column=conta_nome_col).value:
            sheet.cell(row=row_index, column=conta_nome_col).value = (
                f'=IFERROR(MID({escolha_ref},FIND(" - ",{escolha_ref})+3,999),"")'
            )


def _load_account_options(path: Path) -> list[tuple[str, str, str]]:
    with Path(path).open(encoding="utf-8") as file:
        data = yaml.safe_load(file) or {}
    options = []
    for item in data.get("contas", []):
        conta_id = str(item.get("id") or "").strip()
        conta_nome = str(item.get("nome") or "").strip()
        if conta_id and conta_nome:
            options.append((f"{conta_id} - {conta_nome}", conta_id, conta_nome))
    return options


def _write_manifest(result: ImportaRevisaoBancaResult) -> None:
    payload = {
        "bank_raw_dir": str(result.bank_raw_dir),
        "periodo": result.periodo,
        "cnpjs_processados": list(result.cnpjs_processados),
        "arquivos_lidos": list(result.arquivos_lidos),
        "comprovantes_lidos": result.comprovantes_lidos,
        "lancamentos_lidos": result.lancamentos_lidos,
        "comprovantes_candidatos": result.comprovantes_candidatos,
        "conta_corrente_candidatos": result.conta_corrente_candidatos,
        "excel_comprovantes": str(result.excel_comprovantes),
        "excel_conta_corrente": str(result.excel_conta_corrente),
        "ofx_encontrado": result.ofx_encontrado,
        "ofx_usado": result.ofx_usado,
        "decisoes_preservadas": result.decisoes_preservadas,
        "novas_chaves_adicionadas": result.novas_chaves_adicionadas,
        "conta_corrente_regras_yaml_aplicadas": result.conta_corrente_regras_yaml_aplicadas,
        "conta_corrente_boleto_pago_fallback": result.conta_corrente_boleto_pago_fallback,
        "conta_corrente_revisao_sem_regra": result.conta_corrente_revisao_sem_regra,
        "conta_corrente_pix_enviado": result.conta_corrente_pix_enviado,
        "conta_corrente_pix_comprovante_exato": result.conta_corrente_pix_comprovante_exato,
        "conta_corrente_pix_yaml_via_comprovante": result.conta_corrente_pix_yaml_via_comprovante,
        "conta_corrente_pix_revisao_sem_regra_comprovante": result.conta_corrente_pix_revisao_sem_regra_comprovante,
        "conta_corrente_pix_revisao_match_ambiguo": result.conta_corrente_pix_revisao_match_ambiguo,
        "conta_corrente_pix_revisao_sem_comprovante": result.conta_corrente_pix_revisao_sem_comprovante,
        "conta_corrente_pix_fallback_aplicado": result.conta_corrente_pix_fallback_aplicado,
        "manifest_historico": str(result.manifest_historico_path),
        "possiveis_comprovantes_relacionados": result.possiveis_comprovantes_relacionados,
        "alertas": list(result.alertas),
        "yaml_modificado": False,
    }
    text = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
    result.manifest_path.write_text(text, encoding="utf-8")
    result.manifest_historico_path.write_text(text, encoding="utf-8")


def _build_alertas(possiveis_comprovantes_relacionados: int) -> tuple[str, ...]:
    if possiveis_comprovantes_relacionados <= 0:
        return ()
    return (
        "Existem movimentos de conta corrente possivelmente relacionados a "
        "comprovantes. Verificar USAR_SOMENTE_SEM_COMPROVANTE para evitar "
        "duplicidade no DRE.",
    )


def _competencia_referencia_conta_corrente(
    launch: BankLaunchEnrichment,
    result: object,
) -> str:
    if not _is_simples_nacional_result(launch, result):
        return ""
    movement_date = launch.transaction_date
    if not hasattr(movement_date, "year") or not hasattr(movement_date, "month"):
        return ""
    year = int(movement_date.year)
    month = int(movement_date.month) - 1
    if month == 0:
        year -= 1
        month = 12
    return f"{year:04d}-{month:02d}"


def _observacao_audit_conta_corrente(
    launch: BankLaunchEnrichment,
    result: object,
) -> str:
    if not _is_simples_nacional_result(launch, result):
        return ""
    competencia = _competencia_referencia_conta_corrente(launch, result)
    suffix = f"; competencia_referencia={competencia}" if competencia else ""
    return (
        "SIMPLES NACIONAL referente ao mes anterior ao pagamento "
        f"(competencia M-1){suffix}"
    )


def _is_simples_nacional_result(launch: BankLaunchEnrichment, result: object) -> bool:
    if getattr(result, "conta_id", "") != "5.01.001":
        return False
    text = " ".join(
        (
            normalize_match_text(launch.description),
            normalize_match_text(launch.counterparty_name),
        )
    )
    return "SIMPLES" in text


def _tipo_movimento(launch: BankLaunchEnrichment) -> str:
    return _tipo_movimento_text(launch.description, launch.amount)


def _tipo_movimento_text(description: object, amount: Decimal | None) -> str:
    key = normalize_match_text(description)
    if amount is not None and amount > 0:
        return "ENTRADA_FORA_ESCOPO"
    if "SISPAG" in key:
        return "SISPAG_AGREGADO"
    if "TARIFA" in key or "TAXA" in key:
        return "TARIFA_BANCARIA"
    if "PIX" in key:
        return "PIX_TRANSFERENCIA_SAIDA"
    if any(token in key for token in ("BOLETO", "TED", "TRANSFERENCIA", "TRANSF")):
        return "BOLETO_TED_TRANSFERENCIA_SAIDA"
    return "SAIDA_PAGAMENTO"


def _date_text(value: object) -> str:
    return value.isoformat() if hasattr(value, "isoformat") else str(value or "")


def _decimal_text(value: object) -> str:
    decimal = _parse_decimal(value)
    return f"{decimal:.2f}"


def _parse_decimal(value: object) -> Decimal:
    if isinstance(value, Decimal):
        return value
    text = str(value or "").strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text or "0")
    except Exception:
        return Decimal("0")


def _sum_decimals(values: Iterable[object]) -> Decimal:
    return sum((_parse_decimal(value) for value in values), Decimal("0"))


def _joined_unique(values: Iterable[object], *, limit: int | None = None) -> str:
    output: list[str] = []
    for value in values:
        text = str(value or "").strip()
        if text and text not in output:
            output.append(text)
    if limit is not None:
        output = output[:limit]
    return " | ".join(output)


def _examples(values: Iterable[str]) -> str:
    return _joined_unique(values, limit=5)


def _cell_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    return value


def _suggested_width(header: str) -> int:
    widths = {
        "BENEFICIARIO_FAVORECIDO": 34,
        "REFERENCIA_EMPRESA": 42,
        "LANCAMENTO": 42,
        "RAZAO_SOCIAL": 34,
        "CONTA_ESCOLHIDA": 42,
        "CONTA_NOME_APROVADA": 34,
        "POSSIVEL_COMPROVANTE_RELACIONADO": 28,
        "COMPROVANTE_RELACIONADO_BENEFICIARIO": 38,
        "COMPROVANTE_RELACIONADO_DOCUMENTO": 24,
        "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE": 42,
        "EXEMPLOS": 70,
        "OBSERVACAO_OPERADOR": 45,
        "ORIGEM_ARQUIVO": 58,
        "CHAVE_ESTAVEL": 80,
    }
    return widths.get(header, max(14, min(len(header) + 4, 32)))


def _wrap(header: str) -> bool:
    return header in {
        "BENEFICIARIO_FAVORECIDO",
        "REFERENCIA_EMPRESA",
        "LANCAMENTO",
        "RAZAO_SOCIAL",
        "POSSIVEL_COMPROVANTE_RELACIONADO",
        "COMPROVANTE_RELACIONADO_BENEFICIARIO",
        "COMPROVANTE_RELACIONADO_DOCUMENTO",
        "COMPROVANTE_RELACIONADO_REGRA_EXISTENTE",
        "EXEMPLOS",
        "OBSERVACAO_OPERADOR",
        "ORIGEM_ARQUIVO",
        "CHAVE_ESTAVEL",
    }


def _instructions(sheet_name: str) -> list[str]:
    return [
        f"Importa revisao banca - {sheet_name}",
        "",
        "Este arquivo prepara decisoes para o operador. Ele nao altera YAML.",
        "Preencha STATUS_DECISAO=APROVADO apenas quando a regra estiver segura.",
        "CONTA_ESCOLHIDA e um dropdown operacional; CONTA_ID_APROVADA e o campo tecnico usado pela skill de regras.",
        "OFX nao e obrigatorio nesta fase e nao foi usado para gerar as linhas.",
        "As decisoes manuais existentes sao preservadas por CHAVE_ESTAVEL quando o arquivo e regenerado.",
        "Para PIX ENVIADO, POSSIVEL_COMPROVANTE_RELACIONADO=SIM exige mesmo CNPJ, data exata e valor exato.",
        "MATCH_AMBIGUO sinaliza mais de um comprovante possivel; a linha fica para revisao.",
    ]
