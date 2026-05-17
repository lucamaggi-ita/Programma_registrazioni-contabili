from __future__ import annotations

from pathlib import Path
import re

from openpyxl import load_workbook

from fantozzi.banking.common import (
    is_balance_description,
    normalize_key,
    parse_date,
    parse_decimal,
    stable_hash,
    text,
)
from fantozzi.banking.models import BankBalanceAudit, BankLaunchEnrichment


COLUMN_ALIASES = {
    "data": {"data"},
    "description": {"lancamento", "lançamento"},
    "counterparty_name": {"razaosocial", "razaosocialcontraparte"},
    "counterparty_document": {"cpfcnpj", "cpfcnpjcontraparte", "cpfcnpjbeneficiario"},
    "amount": {"valor", "valorr"},
    "balance": {"saldo", "saldor"},
}


def parse_bank_launches(path: Path) -> tuple[list[BankLaunchEnrichment], list[BankBalanceAudit]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header_index, columns = _detect_header(rows)

    enrichments: list[BankLaunchEnrichment] = []
    balance_rows: list[BankLaunchEnrichment] = []
    for row in rows[header_index + 1 :]:
        record = _record_from_row(row, columns)
        if not any(text(value) for value in record.values()):
            continue
        description = text(record.get("description"))
        tx_date = parse_date(record.get("data"))
        amount = _optional_decimal(record.get("amount"))
        balance = _optional_decimal(record.get("balance"))
        entry = BankLaunchEnrichment(
            transaction_date=tx_date,
            amount=amount,
            description=description,
            counterparty_name=text(record.get("counterparty_name")),
            counterparty_document=text(record.get("counterparty_document")),
            row_balance=balance,
            source_file=path.name,
            raw_hash=stable_hash(path.name, tx_date, amount, balance, description),
        )
        if is_balance_description(description) or amount is None:
            balance_rows.append(entry)
        else:
            enrichments.append(entry)

    return enrichments, [_balance_audit_from_rows(path, balance_rows)]


def detect_header_row(path: Path) -> int:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header_index, _ = _detect_header(rows)
    return header_index


def summarize_bank_launch_rows(path: Path) -> dict[str, int]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header_index, columns = _detect_header(rows)
    summary = {
        "xlsx_lancamentos_balance_rows_filtered": 0,
        "xlsx_lancamentos_empty_rows_discarded": 0,
    }
    for row in rows[header_index + 1 :]:
        record = _record_from_row(row, columns)
        if not any(text(value) for value in record.values()):
            summary["xlsx_lancamentos_empty_rows_discarded"] += 1
            continue
        description = text(record.get("description"))
        amount = _optional_decimal(record.get("amount"))
        if is_balance_description(description) or amount is None:
            summary["xlsx_lancamentos_balance_rows_filtered"] += 1
    return summary


def _detect_header(rows: list[tuple[object, ...]]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows):
        normalized = [normalize_key(cell) for cell in row]
        if "data" in normalized and any(
            value in normalized for value in ("lancamento", "lançamento")
        ) and any(value in normalized for value in ("valor", "valorr")):
            return index, _columns_from_header(row)
    raise ValueError("Header XLSX de lancamentos nao encontrado.")


def _columns_from_header(row: tuple[object, ...]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, header in enumerate(row):
        key = normalize_key(header)
        for canonical, aliases in COLUMN_ALIASES.items():
            if key in aliases:
                columns[canonical] = index
    required = {"data", "description", "amount"}
    missing = sorted(required - set(columns))
    if missing:
        raise ValueError("Colunas obrigatorias ausentes no XLSX: " + ", ".join(missing))
    return columns


def _record_from_row(row: tuple[object, ...], columns: dict[str, int]) -> dict[str, object]:
    return {
        name: row[index] if index < len(row) else None
        for name, index in columns.items()
    }


def _optional_decimal(value: object) -> object:
    if value is None or text(value) == "":
        return None
    return parse_decimal(value)


def _balance_audit_from_rows(
    path: Path,
    rows: list[BankLaunchEnrichment],
) -> BankBalanceAudit:
    opening = None
    closing = None
    for row in rows:
        desc = re.sub(r"\s+", " ", row.description.upper())
        value = row.row_balance if row.row_balance is not None else row.amount
        if "SALDO ANTERIOR" in desc and opening is None:
            opening = value
        if "SALDO TOTAL" in desc or "DISPON" in desc or "SALDO DO DIA" in desc:
            closing = value
    return BankBalanceAudit(
        cnpj="",
        account_id="",
        period="",
        opening_balance=opening,
        closing_balance=closing,
        source_file=path.name,
    )
