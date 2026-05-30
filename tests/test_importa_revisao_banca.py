from __future__ import annotations

import csv
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from fantozzi.banking.models import BankLaunchEnrichment, PaymentReceiptRaw
from fantozzi.banking.revisao_banca import (
    COMPROVANTES_GREY_COLUMNS,
    COMPROVANTES_HEADERS,
    COMPROVANTES_HIDDEN_COLUMNS,
    COMPROVANTES_SHEET,
    COMPROVANTES_XLSX,
    CONTA_CORRENTE_CLASSIFICADOS_CSV,
    CONTA_CORRENTE_GREY_COLUMNS,
    CONTA_CORRENTE_HEADERS,
    CONTA_CORRENTE_HIDDEN_COLUMNS,
    CONTA_CORRENTE_SHEET,
    CONTA_CORRENTE_XLSX,
    ImportaRevisaoBancaError,
    _existing_key,
    _sheet_has_x14_list_validation_prefix,
    importar_revisao_banca,
    resolve_bank_raw_dir,
    validate_excel_template_banca,
)
from fantozzi.cli import build_parser, main


def test_resolve_bank_raw_dir_from_year_and_month(tmp_path):
    path = resolve_bank_raw_dir(
        ano="2026",
        mes="04",
        project_root=tmp_path,
    )

    assert path == tmp_path / "bank_raw" / "2026" / "2026-04"


def test_importa_revisao_banca_preserves_existing_decisions_and_ignores_missing_ofx(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir = tmp_path / "bank_raw" / "2026" / "2026-04"
    cnpj_dir = bank_raw_dir / "08884180000123"
    cnpj_dir.mkdir(parents=True)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    output_dir.mkdir(parents=True)
    _write_existing_comprovantes(output_dir / COMPROVANTES_XLSX)

    monkeypatch.setattr(
        "fantozzi.banking.revisao_banca.parse_payment_receipts",
        lambda path: [
            PaymentReceiptRaw(
                payment_date=date(2026, 4, 5),
                amount=Decimal("20.00"),
                beneficiary_name="Fornecedor Teste",
                beneficiary_document="12.345.678/0001-90",
                payment_type="PIX Transferencias",
                reference_id="-",
                status="Efetuado",
                source_file=path.name,
                raw_hash="receipt-1",
            )
        ],
    )
    monkeypatch.setattr(
        "fantozzi.banking.revisao_banca.parse_bank_launches",
        lambda path: (
            [
                BankLaunchEnrichment(
                    transaction_date=date(2026, 4, 5),
                    amount=Decimal("-20.00"),
                    description="PIX Fornecedor Teste",
                    counterparty_name="Fornecedor Teste",
                    counterparty_document="12.345.678/0001-90",
                    row_balance=None,
                    source_file=path.name,
                    raw_hash="launch-1",
                )
            ],
            [],
        ),
    )

    result = importar_revisao_banca(
        bank_raw_dir=bank_raw_dir,
        output_dir=output_dir,
        regras_path=_write_rules(tmp_path),
        plano_contas_path=_write_plan(tmp_path),
        mapeamento_dre_path=_write_mapping(tmp_path),
    )

    assert result.ofx_encontrado is False
    assert result.ofx_usado is False
    assert result.comprovantes_lidos == 1
    assert result.lancamentos_lidos == 1
    assert result.comprovantes_candidatos == 1
    assert result.conta_corrente_candidatos == 1
    assert result.decisoes_preservadas == 1
    assert result.novas_chaves_adicionadas == 1

    workbook = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    sheet = workbook[COMPROVANTES_SHEET]
    headers = [cell.value for cell in sheet[1]]
    status_col = headers.index("STATUS_DECISAO") + 1
    conta_col = headers.index("CONTA_ID_APROVADA") + 1
    qtd_col = headers.index("QTD_OCORRENCIAS") + 1
    assert sheet.cell(row=2, column=status_col).value == "APROVADO"
    assert sheet.cell(row=2, column=conta_col).value == "4.01.012"
    assert sheet.cell(row=2, column=qtd_col).value == 1

    conta_workbook = load_workbook(output_dir / CONTA_CORRENTE_XLSX, data_only=False)
    assert CONTA_CORRENTE_SHEET in conta_workbook.sheetnames
    manifest = json.loads((output_dir / "manifest_importa_revisao_banca.json").read_text())
    assert manifest["ofx_usado"] is False
    assert manifest["yaml_modificado"] is False


def test_importa_revisao_banca_adds_decision_dropdown_validations(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [_launch("PIX enviado", "Fornecedor Teste", raw_hash="launch-1")],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    comprovantes_workbook = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    comprovantes = comprovantes_workbook[COMPROVANTES_SHEET]
    _assert_dropdown(
        comprovantes,
        "STATUS_DECISAO",
        '"PENDENTE,APROVADO,REVISAR,NAO_CRIAR_REGRA,BLOQUEADO"',
    )
    _assert_dropdown(
        comprovantes,
        "ACAO_REGRA",
        '"CLASSIFICAR,ESCLUDERE,XML,NAO_CRIAR_REGRA,REVISAR,BLOQUEADO"',
    )
    _assert_dropdown(comprovantes, "RISCO_DUPLICIDADE_XML", '"OK,DUVIDA"')
    _assert_dropdown(comprovantes, "ESCOPO_REGRA", '"COMPROVANTE,CONTA_CORRENTE"')
    _assert_dropdown_prefix(comprovantes, "CONTA_ESCOLHIDA", "=plano_contas!")

    conta_workbook = load_workbook(output_dir / CONTA_CORRENTE_XLSX, data_only=False)
    conta = conta_workbook[CONTA_CORRENTE_SHEET]
    conta_headers = [cell.value for cell in conta[1]]
    assert "POSSIVEL_COMPROVANTE_RELACIONADO" in conta_headers
    _assert_dropdown(
        conta,
        "STATUS_DECISAO",
        '"PENDENTE,APROVADO,REVISAR,NAO_CRIAR_REGRA,BLOQUEADO"',
    )
    _assert_dropdown(
        conta,
        "ACAO_REGRA",
        '"CLASSIFICAR,ESCLUDERE,XML,NAO_CRIAR_REGRA,REVISAR,BLOQUEADO"',
    )
    _assert_dropdown(conta, "USAR_SOMENTE_SEM_COMPROVANTE", '"SIM,NAO"')
    _assert_dropdown(conta, "RISCO_DUPLICIDADE_XML", '"OK,DUVIDA"')
    _assert_dropdown(conta, "ESCOPO_REGRA", '"COMPROVANTE,CONTA_CORRENTE"')
    _assert_dropdown_prefix(conta, "CONTA_ESCOLHIDA", "=plano_contas!")


def test_importa_revisao_banca_reuses_yaml_rules_before_boleto_fallback(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [
            _launch(
                "BOLETO PAGO ABRASEL ASS",
                "ABRASEL ASSOCIACAO BRASILEIRA DE BARES E RESTAUR",
                raw_hash="launch-abrasel",
            ),
            _launch(
                "BOLETO PAGO FORNECEDOR SEM REGRA",
                "FORNECEDOR SEM REGRA",
                raw_hash="launch-boleto",
            ),
            _launch(
                "BOLETO PAGO FIBRANET BA",
                "FIBRANET BA TELECOM LTDA",
                raw_hash="launch-fibranet",
            ),
            _launch(
                "PIX FORNECEDOR SEM REGRA",
                "FORNECEDOR SEM REGRA",
                raw_hash="launch-pix",
            ),
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_supplier_and_boleto_fallback(
                tmp_path,
                "ABRASEL ASSOCIACAO BRASILEIRA",
            ),
        }
    )

    assert result.conta_corrente_regras_yaml_aplicadas == 2
    assert result.conta_corrente_boleto_pago_fallback == 1
    assert result.conta_corrente_revisao_sem_regra == 1
    assert result.conta_corrente_candidatos == 1

    decisoes = _data_rows(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET)
    assert [row["LANCAMENTO"] for row in decisoes] == ["PIX FORNECEDOR SEM REGRA"]

    tecnico = _data_rows(
        output_dir / CONTA_CORRENTE_XLSX,
        "tecnico_conta_corrente",
    )
    by_hash = {row["HASH_ORIGEM"]: row for row in tecnico}
    assert by_hash["launch-abrasel"]["REGRA_EXISTENTE"] == "REGRA_ABRASEL"
    assert by_hash["launch-fibranet"]["REGRA_EXISTENTE"] == (
        "REGRA_4_01_014_FIBRANET_BA_TELECOM"
    )
    assert by_hash["launch-boleto"]["REGRA_EXISTENTE"] == (
        "EXCLUSAO_XML_CONTA_CORRENTE_BOLETO_PAGO"
    )
    assert by_hash["launch-pix"]["REGRA_EXISTENTE"] in (None, "")
    assert by_hash["launch-boleto"]["MOTIVO_REVISAO"] == "EXCLUSAO_OPERADOR_REGRA"

    manifest = json.loads((output_dir / "manifest_importa_revisao_banca.json").read_text())
    assert manifest["conta_corrente_regras_yaml_aplicadas"] == 2
    assert manifest["conta_corrente_boleto_pago_fallback"] == 1
    assert manifest["conta_corrente_revisao_sem_regra"] == 1


def test_conta_corrente_specific_exclusion_preempts_boleto_fallback(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [
            _launch(
                "BOLETO PAGO SINTHOTESB",
                "SINTHOTESB SIND TRAB BARES HOT",
                raw_hash="launch-sinthotesb",
            ),
            _launch(
                "BOLETO PAGO FORNECEDOR SEM REGRA",
                "FORNECEDOR SEM REGRA",
                raw_hash="launch-boleto",
            ),
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_sinthotesb_and_boleto_fallback(tmp_path),
        }
    )

    assert result.conta_corrente_regras_yaml_aplicadas == 1
    assert result.conta_corrente_boleto_pago_fallback == 1
    assert result.conta_corrente_revisao_sem_regra == 0
    assert result.conta_corrente_candidatos == 0

    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {row["HASH_ORIGEM"]: row for row in tecnico}
    assert by_hash["launch-sinthotesb"]["REGRA_EXISTENTE"] == (
        "EXCLUSAO_ESCLUDERE_SINTHOTESB_SIND_TRAB_BARES_HOT"
    )
    assert by_hash["launch-sinthotesb"]["MOTIVO_REVISAO"] == (
        "EXCLUSAO_OPERADOR_REGRA"
    )
    assert "EXCLUSAO_XML_CONTA_CORRENTE_BOLETO_PAGO" not in (
        by_hash["launch-sinthotesb"]["REGRA_EXISTENTE"] or ""
    )
    assert by_hash["launch-boleto"]["REGRA_EXISTENTE"] == (
        "EXCLUSAO_XML_CONTA_CORRENTE_BOLETO_PAGO"
    )


def test_importa_revisao_banca_classifies_recurring_conta_corrente_directly(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [
            _launch(
                "DA COELBA 7014914336",
                "",
                raw_hash="launch-coelba",
                amount="-9164.48",
                transaction_date=date(2026, 4, 23),
            ),
            _launch(
                "EMBASA 142698660",
                "",
                raw_hash="launch-embasa",
                amount="-351.23",
                transaction_date=date(2026, 4, 13),
            ),
            _launch(
                "DA VIVO-BA 20304230426",
                "",
                raw_hash="launch-vivo",
                amount="-380.97",
                transaction_date=date(2026, 4, 6),
            ),
            _launch(
                "DEBITO SEGURO 06/57",
                "",
                raw_hash="launch-debito-seguro",
                amount="-138.04",
                transaction_date=date(2026, 4, 20),
            ),
            _launch(
                "PAGAMENTOS TRIB COD BARRAS",
                "SIMPLES NACIONAL",
                raw_hash="launch-simples",
                amount="-1027.55",
                transaction_date=date(2026, 4, 17),
            ),
            _launch(
                "PAGAMENTOS TRIB COD BARRAS",
                "",
                raw_hash="launch-tributo-generico",
                amount="-10.00",
                transaction_date=date(2026, 4, 18),
            ),
            _launch(
                "SISPAG SALARIOS",
                "",
                raw_hash="launch-sispag",
                amount="-29641.68",
                transaction_date=date(2026, 4, 6),
            ),
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_recurring_conta_corrente(tmp_path),
        }
    )

    assert result.conta_corrente_regras_yaml_aplicadas == 6
    assert result.conta_corrente_revisao_sem_regra == 1
    assert result.conta_corrente_candidatos == 1

    rows = _data_rows(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET)
    assert [row["LANCAMENTO"] for row in rows] == ["PAGAMENTOS TRIB COD BARRAS"]

    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {row["HASH_ORIGEM"]: row for row in tecnico}
    assert by_hash["launch-coelba"]["REGRA_EXISTENTE"] == (
        "REGRA_4_01_011_CONTA_CORRENTE_COELBA"
    )
    assert by_hash["launch-embasa"]["REGRA_EXISTENTE"] == (
        "REGRA_4_01_012_CONTA_CORRENTE_EMBASA"
    )
    assert by_hash["launch-vivo"]["REGRA_EXISTENTE"] == (
        "REGRA_4_01_014_CONTA_CORRENTE_VIVO"
    )
    assert by_hash["launch-debito-seguro"]["REGRA_EXISTENTE"] == (
        "REGRA_7_01_001_CONTA_CORRENTE_DEBITO_SEGURO"
    )
    assert by_hash["launch-simples"]["REGRA_EXISTENTE"] == (
        "REGRA_5_01_001_CONTA_CORRENTE_SIMPLES_NACIONAL"
    )
    assert by_hash["launch-simples"]["COMPETENCIA_REFERENCIA"] == "2026-03"
    assert "competencia M-1" in by_hash["launch-simples"]["OBSERVACAO_AUDIT"]
    assert by_hash["launch-sispag"]["REGRA_EXISTENTE"] == (
        "EXCLUSAO_ESCLUDERE_CONTA_CORRENTE_SISPAG_SALARIOS"
    )
    assert by_hash["launch-sispag"]["MOTIVO_REVISAO"] == "EXCLUSAO_OPERADOR_REGRA"
    assert by_hash["launch-tributo-generico"]["REGRA_EXISTENTE"] in (None, "")
    assert by_hash["launch-tributo-generico"]["MOTIVO_REVISAO"] == "SEM_REGRA"


def test_importa_revisao_banca_filters_total_receipt_from_decision_sheets(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir = tmp_path / "bank_raw" / "2026" / "2026-04"
    cnpj_dir = bank_raw_dir / "08884180000123"
    cnpj_dir.mkdir(parents=True)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"

    monkeypatch.setattr(
        "fantozzi.banking.revisao_banca.parse_payment_receipts",
        lambda path: [
            PaymentReceiptRaw(
                payment_date=date(2026, 4, 5),
                amount=Decimal("20.00"),
                beneficiary_name="Fornecedor Teste",
                beneficiary_document="12.345.678/0001-90",
                payment_type="PIX Transferencias",
                reference_id="-",
                status="Efetuado",
                source_file=path.name,
                raw_hash="receipt-1",
            ),
            PaymentReceiptRaw(
                payment_date=None,
                amount=Decimal("1000.00"),
                beneficiary_name="TOTAL",
                beneficiary_document="",
                payment_type="",
                reference_id="",
                status="",
                source_file=path.name,
                raw_hash="receipt-total",
            ),
        ],
    )

    result = importar_revisao_banca(
        bank_raw_dir=bank_raw_dir,
        output_dir=output_dir,
        regras_path=_write_rules(tmp_path),
        plano_contas_path=_write_plan(tmp_path),
        mapeamento_dre_path=_write_mapping(tmp_path),
    )

    assert result.comprovantes_lidos == 2
    assert result.comprovantes_candidatos == 1

    workbook = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    decisoes = workbook[COMPROVANTES_SHEET]
    decisoes_rows = list(decisoes.iter_rows(values_only=True))
    headers = [cell or "" for cell in decisoes_rows[0]]
    beneficiary_col = headers.index("BENEFICIARIO_FAVORECIDO")
    assert len(decisoes_rows) == 2
    assert all(row[beneficiary_col] != "TOTAL" for row in decisoes_rows[1:])

    tecnico = workbook["tecnico_comprovantes"]
    tecnico_rows = list(tecnico.iter_rows(values_only=True))
    tecnico_headers = [cell or "" for cell in tecnico_rows[0]]
    key_col = tecnico_headers.index("CHAVE_ESTAVEL")
    assert len(tecnico_rows) == 2
    assert all("BENEFICIARIO=TOTAL" not in row[key_col] for row in tecnico_rows[1:])


def test_importa_revisao_banca_derives_account_fields_from_choice_not_old_formulas(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir = tmp_path / "bank_raw" / "2026" / "2026-04"
    cnpj_dir = bank_raw_dir / "08884180000123"
    cnpj_dir.mkdir(parents=True)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    output_dir.mkdir(parents=True)
    _write_existing_comprovantes_with_stale_formula(output_dir / COMPROVANTES_XLSX)

    monkeypatch.setattr(
        "fantozzi.banking.revisao_banca.parse_payment_receipts",
        lambda path: [
            PaymentReceiptRaw(
                payment_date=date(2026, 4, 5),
                amount=Decimal("20.00"),
                beneficiary_name="Fornecedor Teste",
                beneficiary_document="12.345.678/0001-90",
                payment_type="PIX Transferencias",
                reference_id="-",
                status="Efetuado",
                source_file=path.name,
                raw_hash="receipt-1",
            )
        ],
    )

    importar_revisao_banca(
        bank_raw_dir=bank_raw_dir,
        output_dir=output_dir,
        regras_path=_write_rules(tmp_path),
        plano_contas_path=_write_plan(tmp_path),
        mapeamento_dre_path=_write_mapping(tmp_path),
    )

    workbook = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    sheet = workbook[COMPROVANTES_SHEET]
    headers = [cell.value for cell in sheet[1]]
    conta_id_col = headers.index("CONTA_ID_APROVADA") + 1
    conta_nome_col = headers.index("CONTA_NOME_APROVADA") + 1

    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=conta_id_col).value == "4.01.012"
    assert sheet.cell(row=2, column=conta_nome_col).value == "Agua e Esgoto"


def test_existing_key_for_comprovantes_without_beneficiary_does_not_fall_through():
    assert (
        _existing_key(
            {
                "BENEFICIARIO_FAVORECIDO": "",
                "LANCAMENTO": "PIX ENVIADO",
                "RAZAO_SOCIAL": "Fornecedor Teste",
                "VALOR": "-20.00",
            },
            "comprovantes",
        )
        == ""
    )


def test_cli_accepts_importar_revisao_banca_modes():
    parser = build_parser()

    args = parser.parse_args(
        [
            "importar-revisao-banca",
            "--bank-raw-dir",
            "bank_raw/2026/2026-04",
        ]
    )

    assert args.command == "importar-revisao-banca"
    assert args.bank_raw_dir == "bank_raw/2026/2026-04"

    args = parser.parse_args(
        [
            "importar-revisao-banca",
            "--ano",
            "2026",
            "--mes",
            "04",
        ]
    )

    assert args.ano == "2026"
    assert args.mes == "04"


def test_cli_importar_revisao_banca_requires_period():
    assert main(["importar-revisao-banca"]) == 1


def test_resolve_bank_raw_dir_rejects_ambiguous_inputs(tmp_path):
    try:
        resolve_bank_raw_dir(
            bank_raw_dir=tmp_path / "bank_raw" / "2026" / "2026-04",
            ano="2026",
            mes="04",
        )
    except ImportaRevisaoBancaError as exc:
        assert "nao ambos" in str(exc)
    else:
        raise AssertionError("expected ImportaRevisaoBancaError")


def test_importar_revisao_banca_idempotente(monkeypatch, tmp_path):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    receipt = _receipt("Fornecedor Teste", raw_hash="receipt-1")
    launch = _launch("PIX Fornecedor Teste", "Fornecedor Teste", raw_hash="launch-1")
    _patch_bank_sources(monkeypatch, [receipt], [launch])

    kwargs = _import_kwargs(tmp_path, bank_raw_dir, output_dir)
    importar_revisao_banca(**kwargs)
    first_comprovantes = _sheet_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    first_conta = _sheet_rows(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET)

    importar_revisao_banca(**kwargs)
    second_comprovantes = _sheet_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    second_conta = _sheet_rows(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET)

    assert second_comprovantes == first_comprovantes
    assert second_conta == first_conta


def test_importar_revisao_banca_preserva_decisoes_apos_rerun(monkeypatch, tmp_path):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    receipt = _receipt("Fornecedor Teste", raw_hash="receipt-1")
    launch = _launch("PIX Fornecedor Teste", "Fornecedor Teste", raw_hash="launch-1")
    _patch_bank_sources(monkeypatch, [receipt], [launch])

    kwargs = _import_kwargs(tmp_path, bank_raw_dir, output_dir)
    importar_revisao_banca(**kwargs)
    _set_cell_by_header(
        output_dir / COMPROVANTES_XLSX,
        COMPROVANTES_SHEET,
        row=2,
        values={
            "STATUS_DECISAO": "APROVADO",
            "ACAO_REGRA": "CLASSIFICAR",
            "CONTA_ESCOLHIDA": "4.01.012 - Agua e Esgoto",
            "OBSERVACAO_OPERADOR": "decisao preservada",
            "RISCO_DUPLICIDADE_XML": "OK",
            "ESCOPO_REGRA": "COMPROVANTE",
        },
    )
    _set_cell_by_header(
        output_dir / CONTA_CORRENTE_XLSX,
        CONTA_CORRENTE_SHEET,
        row=2,
        values={
            "STATUS_DECISAO": "APROVADO",
            "ACAO_REGRA": "CLASSIFICAR",
            "CONTA_ESCOLHIDA": "4.01.012 - Agua e Esgoto",
            "USAR_SOMENTE_SEM_COMPROVANTE": "SIM",
            "OBSERVACAO_OPERADOR": "conta preservada",
        },
    )

    importar_revisao_banca(**kwargs)

    comprovante = _row_by_index(
        output_dir / COMPROVANTES_XLSX,
        COMPROVANTES_SHEET,
        row=2,
    )
    conta = _row_by_index(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET, row=2)
    assert comprovante["STATUS_DECISAO"] == "APROVADO"
    assert comprovante["ACAO_REGRA"] == "CLASSIFICAR"
    assert comprovante["CONTA_ID_APROVADA"] == "4.01.012"
    assert comprovante["CONTA_NOME_APROVADA"] == "Agua e Esgoto"
    assert comprovante["OBSERVACAO_OPERADOR"] == "decisao preservada"
    assert comprovante["RISCO_DUPLICIDADE_XML"] == "OK"
    assert comprovante["ESCOPO_REGRA"] == "COMPROVANTE"
    assert conta["USAR_SOMENTE_SEM_COMPROVANTE"] == "SIM"
    assert conta["OBSERVACAO_OPERADOR"] == "conta preservada"


def test_importar_revisao_banca_aggiunge_nova_chave_sem_perder_decisoes(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    receipts = [_receipt("Fornecedor A", raw_hash="receipt-a")]
    _patch_bank_sources(monkeypatch, receipts, [])

    kwargs = _import_kwargs(tmp_path, bank_raw_dir, output_dir)
    importar_revisao_banca(**kwargs)
    _set_cell_by_header(
        output_dir / COMPROVANTES_XLSX,
        COMPROVANTES_SHEET,
        row=2,
        values={"STATUS_DECISAO": "APROVADO", "OBSERVACAO_OPERADOR": "manter"},
    )
    receipts.append(_receipt("Fornecedor B", raw_hash="receipt-b"))

    importar_revisao_banca(**kwargs)

    rows = _data_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    assert len(rows) == 2
    fornecedor_a = next(row for row in rows if row["BENEFICIARIO_FAVORECIDO"] == "FORNECEDOR A")
    fornecedor_b = next(row for row in rows if row["BENEFICIARIO_FAVORECIDO"] == "FORNECEDOR B")
    assert fornecedor_a["STATUS_DECISAO"] == "APROVADO"
    assert fornecedor_a["OBSERVACAO_OPERADOR"] == "manter"
    assert fornecedor_b["STATUS_DECISAO"] in (None, "")


def test_importar_revisao_banca_remove_chave_desaparecida_do_operacional(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    receipts = [
        _receipt("Fornecedor A", raw_hash="receipt-a"),
        _receipt("Fornecedor B", raw_hash="receipt-b"),
    ]
    _patch_bank_sources(monkeypatch, receipts, [])

    kwargs = _import_kwargs(tmp_path, bank_raw_dir, output_dir)
    importar_revisao_banca(**kwargs)
    receipts.pop()

    importar_revisao_banca(**kwargs)

    rows = _data_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    assert [row["BENEFICIARIO_FAVORECIDO"] for row in rows] == ["FORNECEDOR A"]


def test_importar_revisao_banca_skip_classificado_por_yaml(monkeypatch, tmp_path):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [],
    )

    result = importar_revisao_banca(
        bank_raw_dir=bank_raw_dir,
        output_dir=output_dir,
        regras_path=_write_rules_with_supplier(tmp_path, "FORNECEDOR TESTE"),
        plano_contas_path=_write_plan(tmp_path),
        mapeamento_dre_path=_write_mapping(tmp_path),
    )

    assert result.comprovantes_candidatos == 0
    data_rows = [
        row
        for row in _data_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
        if any(value not in (None, "") for value in row.values())
    ]
    assert data_rows == []


def test_importar_revisao_banca_escreve_manifest_historico_por_mes(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [],
    )

    result = importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    current = output_dir / "manifest_importa_revisao_banca.json"
    historic = output_dir / "manifest_importa_revisao_banca_2026-04.json"
    assert result.manifest_path == current
    assert result.manifest_historico_path == historic
    assert current.exists()
    assert historic.exists()
    assert json.loads(current.read_text(encoding="utf-8")) == json.loads(
        historic.read_text(encoding="utf-8")
    )


def test_importar_revisao_banca_indica_possivel_comprovante_relacionado(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [_launch("PIX enviado", "Fornecedor Teste", raw_hash="launch-1")],
    )

    result = importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    row = _row_by_index(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET, row=2)
    assert row["POSSIVEL_COMPROVANTE_RELACIONADO"] == "SIM"
    assert row["STATUS_DECISAO"] in (None, "")
    assert row["CONTA_ESCOLHIDA"] in (None, "")
    assert result.possiveis_comprovantes_relacionados == 1
    assert result.alertas
    manifest = json.loads((output_dir / "manifest_importa_revisao_banca.json").read_text())
    assert manifest["possiveis_comprovantes_relacionados"] == 1
    assert manifest["alertas"]


def test_pix_enviado_uses_exact_receipt_beneficiary_rule_as_bridge(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "Beneficiario Com Regra",
                raw_hash="receipt-bridge",
                amount="100.00",
                payment_date=date(2026, 4, 10),
                document="12.345.678/0001-90",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "Nome Banco Diferente",
                raw_hash="launch-bridge",
                amount="-100.00",
                transaction_date=date(2026, 4, 10),
                document="12.345.678/0001-90",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_supplier(tmp_path, "BENEFICIARIO COM REGRA"),
        }
    )

    assert result.conta_corrente_pix_enviado == 1
    assert result.conta_corrente_pix_comprovante_exato == 1
    assert result.conta_corrente_pix_yaml_via_comprovante == 1
    assert result.conta_corrente_candidatos == 0

    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {row["HASH_ORIGEM"]: row for row in tecnico}
    assert by_hash["launch-bridge"]["REGRA_EXISTENTE"] == "REGRA_FORNECEDOR_TESTE"
    assert by_hash["launch-bridge"]["COMPROVANTE_RELACIONADO_BENEFICIARIO"] == (
        "Beneficiario Com Regra"
    )
    assert by_hash["launch-bridge"]["COMPROVANTE_RELACIONADO_DATA"] == "2026-04-10"
    assert by_hash["launch-bridge"]["COMPROVANTE_RELACIONADO_VALOR"] == "100.00"


def test_pix_enviado_does_not_bridge_receipt_with_different_date(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "Beneficiario Com Regra",
                raw_hash="receipt-date-off",
                amount="175.00",
                payment_date=date(2026, 4, 21),
                document="12.345.678/0001-90",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "Nome Banco Diferente",
                raw_hash="launch-date-off",
                amount="-175.00",
                transaction_date=date(2026, 4, 23),
                document="12.345.678/0001-90",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_supplier(tmp_path, "BENEFICIARIO COM REGRA"),
        }
    )

    assert result.conta_corrente_pix_comprovante_exato == 0
    assert result.conta_corrente_pix_revisao_sem_comprovante == 1
    assert result.conta_corrente_candidatos == 1
    row = _row_by_index(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET, row=2)
    assert row["POSSIVEL_COMPROVANTE_RELACIONADO"] == "NAO"
    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {item["HASH_ORIGEM"]: item for item in tecnico}
    assert by_hash["launch-date-off"]["REGRA_EXISTENTE"] in (None, "")


def test_pix_enviado_does_not_reuse_exact_receipt_for_duplicate_launch(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "OS MORENOS DO COCO LTDA",
                raw_hash="receipt-os-morenos-exact",
                amount="175.00",
                payment_date=date(2026, 4, 23),
                document="12.345.678/0001-90",
            ),
            _receipt(
                "OS MORENOS DO COCO LTDA",
                raw_hash="receipt-os-morenos-date-off",
                amount="175.00",
                payment_date=date(2026, 4, 21),
                document="12.345.678/0001-90",
            ),
        ],
        [
            _launch(
                "PIX ENVIADO",
                "OS MORENOS DO COCO",
                raw_hash="launch-os-morenos-exact",
                amount="-175.00",
                transaction_date=date(2026, 4, 23),
                document="12.345.678/0001-90",
            ),
            _launch(
                "PIX ENVIADO",
                "OS MORENOS DO COCO",
                raw_hash="launch-os-morenos-review",
                amount="-175.00",
                transaction_date=date(2026, 4, 23),
                document="12.345.678/0001-90",
            ),
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_supplier(tmp_path, "OS MORENOS DO COCO LTDA"),
        }
    )

    assert result.conta_corrente_pix_enviado == 2
    assert result.conta_corrente_pix_comprovante_exato == 1
    assert result.conta_corrente_pix_yaml_via_comprovante == 1
    assert result.conta_corrente_pix_revisao_sem_comprovante == 1
    assert result.conta_corrente_candidatos == 1

    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {item["HASH_ORIGEM"]: item for item in tecnico}
    assert by_hash["launch-os-morenos-exact"]["REGRA_EXISTENTE"] == "REGRA_FORNECEDOR_TESTE"
    assert by_hash["launch-os-morenos-review"]["REGRA_EXISTENTE"] in (None, "")
    assert by_hash["launch-os-morenos-review"]["COMPROVANTE_RELACIONADO_DATA"] in (None, "")


def test_pix_enviado_exact_receipt_without_yaml_rule_stays_review(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "Beneficiario Sem Regra",
                raw_hash="receipt-no-rule",
                amount="100.00",
                payment_date=date(2026, 4, 10),
                document="12.345.678/0001-90",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "Nome Banco Diferente",
                raw_hash="launch-no-rule",
                amount="-100.00",
                transaction_date=date(2026, 4, 10),
                document="12.345.678/0001-90",
            )
        ],
    )

    result = importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    assert result.conta_corrente_pix_comprovante_exato == 1
    assert result.conta_corrente_pix_revisao_sem_regra_comprovante == 1
    assert result.conta_corrente_candidatos == 1
    row = _row_by_index(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET, row=2)
    assert row["POSSIVEL_COMPROVANTE_RELACIONADO"] == "SIM"
    assert row["COMPROVANTE_RELACIONADO_BENEFICIARIO"] == "Beneficiario Sem Regra"
    assert row["COMPROVANTE_RELACIONADO_REGRA_EXISTENTE"] in (None, "")
    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {item["HASH_ORIGEM"]: item for item in tecnico}
    assert by_hash["launch-no-rule"]["REGRA_EXISTENTE"] in (None, "")


def test_pix_enviado_ambiguous_exact_receipts_stay_review(
    monkeypatch,
    tmp_path,
):
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "Beneficiario A",
                raw_hash="receipt-amb-a",
                amount="100.00",
                payment_date=date(2026, 4, 10),
                document="12.345.678/0001-90",
            ),
            _receipt(
                "Beneficiario B",
                raw_hash="receipt-amb-b",
                amount="100.00",
                payment_date=date(2026, 4, 10),
                document="12.345.678/0001-90",
            ),
        ],
        [
            _launch(
                "PIX ENVIADO",
                "Nome Banco Diferente",
                raw_hash="launch-amb",
                amount="-100.00",
                transaction_date=date(2026, 4, 10),
                document="12.345.678/0001-90",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_supplier(tmp_path, "BENEFICIARIO A"),
        }
    )

    assert result.conta_corrente_pix_comprovante_exato == 1
    assert result.conta_corrente_pix_revisao_match_ambiguo == 1
    assert result.conta_corrente_pix_yaml_via_comprovante == 0
    assert result.conta_corrente_candidatos == 1
    row = _row_by_index(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET, row=2)
    assert row["POSSIVEL_COMPROVANTE_RELACIONADO"] == "MATCH_AMBIGUO"
    assert row["COMPROVANTE_RELACIONADO_REGRA_EXISTENTE"] == "MATCH_AMBIGUO:2"
    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {item["HASH_ORIGEM"]: item for item in tecnico}
    assert by_hash["launch-amb"]["REGRA_EXISTENTE"] in (None, "")


def test_pix_enviado_xml_exclusion_bridge_exact_match_closes_movement(
    monkeypatch,
    tmp_path,
):
    """PIX ENVIADO com comprovante exato cujo beneficiario tem regra XML deve fechar o movimento."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "ATTC COMERCIO DE BEBIDAS LTDA",
                raw_hash="receipt-attc",
                amount="350.00",
                payment_date=date(2026, 4, 15),
                document="",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "VINHO VERITAS",
                raw_hash="launch-vinho-veritas",
                amount="-350.00",
                transaction_date=date(2026, 4, 15),
                document="",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_xml_exclusion(tmp_path, "ATTC COMERCIO DE BEBIDAS LTDA"),
        }
    )

    assert result.conta_corrente_pix_enviado == 1
    assert result.conta_corrente_pix_comprovante_exato == 1
    assert result.conta_corrente_pix_yaml_via_comprovante == 1
    assert result.conta_corrente_candidatos == 0

    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {row["HASH_ORIGEM"]: row for row in tecnico}
    assert "EXCLUSAO_XML_TEST_ATTC_COMERCIO_DE_BEBIDAS_LTDA" in (
        by_hash["launch-vinho-veritas"]["REGRA_EXISTENTE"] or ""
    )
    assert by_hash["launch-vinho-veritas"]["COMPROVANTE_RELACIONADO_BENEFICIARIO"] == (
        "ATTC COMERCIO DE BEBIDAS LTDA"
    )


def test_pix_enviado_xml_exclusion_bridge_different_date_stays_review(
    monkeypatch,
    tmp_path,
):
    """PIX ENVIADO com comprovante de data diferente NAO deve fechar pelo bridge XML."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "ATTC COMERCIO DE BEBIDAS LTDA",
                raw_hash="receipt-attc-date-off",
                amount="350.00",
                payment_date=date(2026, 4, 13),
                document="",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "VINHO VERITAS",
                raw_hash="launch-vinho-veritas-date-off",
                amount="-350.00",
                transaction_date=date(2026, 4, 15),
                document="",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_xml_exclusion(tmp_path, "ATTC COMERCIO DE BEBIDAS LTDA"),
        }
    )

    assert result.conta_corrente_pix_comprovante_exato == 0
    assert result.conta_corrente_pix_yaml_via_comprovante == 0
    assert result.conta_corrente_pix_revisao_sem_comprovante == 1
    assert result.conta_corrente_candidatos == 1

    row = _row_by_index(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET, row=2)
    assert row["POSSIVEL_COMPROVANTE_RELACIONADO"] == "NAO"


def test_pix_enviado_embox_xml_exclusion_bridge_exact_match(
    monkeypatch,
    tmp_path,
):
    """BOCA RESTAURANTE / EMBOX BOCA COMERCIO EIRELI: bridge XML com data e valor exatos."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "EMBOX BOCA COMERCIO EIRELI",
                raw_hash="receipt-embox",
                amount="480.00",
                payment_date=date(2026, 4, 18),
                document="",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "BOCA RESTAURANTE E LANCHONETE",
                raw_hash="launch-boca",
                amount="-480.00",
                transaction_date=date(2026, 4, 18),
                document="",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_xml_exclusion(tmp_path, "EMBOX BOCA COMERCIO EIRELI"),
        }
    )

    assert result.conta_corrente_pix_yaml_via_comprovante == 1
    assert result.conta_corrente_candidatos == 0

    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {row["HASH_ORIGEM"]: row for row in tecnico}
    assert "EXCLUSAO_XML_TEST_EMBOX_BOCA_COMERCIO_EIRELI" in (
        by_hash["launch-boca"]["REGRA_EXISTENTE"] or ""
    )


def test_pix_enviado_e_santos_xml_exclusion_bridge_exact_match(
    monkeypatch,
    tmp_path,
):
    """RAINHA DO FATO / E SANTOS DE OLIVEIRA - CARNES: bridge XML normaliza o hifen."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "E SANTOS DE OLIVEIRA - CARNES",
                raw_hash="receipt-e-santos",
                amount="600.00",
                payment_date=date(2026, 4, 20),
                document="",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "RAINHA DO FATO",
                raw_hash="launch-rainha",
                amount="-600.00",
                transaction_date=date(2026, 4, 20),
                document="",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_xml_exclusion(tmp_path, "E SANTOS DE OLIVEIRA CARNES"),
        }
    )

    assert result.conta_corrente_pix_yaml_via_comprovante == 1
    assert result.conta_corrente_candidatos == 0

    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {row["HASH_ORIGEM"]: row for row in tecnico}
    assert by_hash["launch-rainha"]["COMPROVANTE_RELACIONADO_BENEFICIARIO"] == (
        "E SANTOS DE OLIVEIRA - CARNES"
    )


def test_pix_enviado_rainha_do_fato_without_exact_receipt_stays_review(
    monkeypatch,
    tmp_path,
):
    """RAINHA DO FATO sem comprovante exato nao recebe XML generico."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "E SANTOS DE OLIVEIRA - CARNES",
                raw_hash="receipt-e-santos-date-off",
                amount="600.00",
                payment_date=date(2026, 4, 19),
                document="",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "RAINHA DO FATO",
                raw_hash="launch-rainha-date-off",
                amount="-600.00",
                transaction_date=date(2026, 4, 20),
                document="",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_xml_exclusion(
                tmp_path,
                "E SANTOS DE OLIVEIRA CARNES",
            ),
        }
    )

    assert result.conta_corrente_pix_comprovante_exato == 0
    assert result.conta_corrente_pix_yaml_via_comprovante == 0
    assert result.conta_corrente_pix_revisao_sem_comprovante == 1
    assert result.conta_corrente_candidatos == 1

    row = _row_by_index(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET, row=2)
    assert row["LANCAMENTO"] == "PIX ENVIADO"
    assert row["RAZAO_SOCIAL"] == "RAINHA DO FATO"
    assert row["POSSIVEL_COMPROVANTE_RELACIONADO"] == "NAO"


def test_pix_enviado_julio_cesar_positive_rule_bridge_exact_match(
    monkeypatch,
    tmp_path,
):
    """LITORAL REFRIGERACAO / JULIO CESAR RODRIGUES DE SOUZA: classifica em 8.02.003."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "JULIO CESAR RODRIGUES DE SOUZA",
                raw_hash="receipt-julio",
                amount="250.00",
                payment_date=date(2026, 4, 22),
                document="",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "LITORAL REFRIGERACAO E MANUTENCAO",
                raw_hash="launch-litoral",
                amount="-250.00",
                transaction_date=date(2026, 4, 22),
                document="",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_positive_rule_julio_cesar(tmp_path),
        }
    )

    assert result.conta_corrente_pix_yaml_via_comprovante == 1
    assert result.conta_corrente_candidatos == 0

    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {row["HASH_ORIGEM"]: row for row in tecnico}
    assert by_hash["launch-litoral"]["REGRA_EXISTENTE"] == (
        "REGRA_8_02_003_JULIO_CESAR_RODRIGUES_DE_SOUZA"
    )
    assert by_hash["launch-litoral"]["COMPROVANTE_RELACIONADO_BENEFICIARIO"] == (
        "JULIO CESAR RODRIGUES DE SOUZA"
    )


def test_pix_enviado_julio_cesar_different_date_stays_review(
    monkeypatch,
    tmp_path,
):
    """LITORAL / JULIO CESAR com data diferente NAO deve fechar automaticamente."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "JULIO CESAR RODRIGUES DE SOUZA",
                raw_hash="receipt-julio-date-off",
                amount="250.00",
                payment_date=date(2026, 4, 20),
                document="",
            )
        ],
        [
            _launch(
                "PIX ENVIADO",
                "LITORAL REFRIGERACAO E MANUTENCAO",
                raw_hash="launch-litoral-date-off",
                amount="-250.00",
                transaction_date=date(2026, 4, 22),
                document="",
            )
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_positive_rule_julio_cesar(tmp_path),
        }
    )

    assert result.conta_corrente_pix_comprovante_exato == 0
    assert result.conta_corrente_pix_yaml_via_comprovante == 0
    assert result.conta_corrente_pix_revisao_sem_comprovante == 1
    assert result.conta_corrente_candidatos == 1

    row = _row_by_index(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET, row=2)
    assert row["POSSIVEL_COMPROVANTE_RELACIONADO"] == "NAO"


# ---------------------------------------------------------------------------
# Testes: conta_corrente_classificados.csv (FZ-GIORNALE-DRE-BRIDGE-BANCA-002B)
# ---------------------------------------------------------------------------


def test_conta_corrente_classificados_csv_criado_com_classificados_yaml(
    monkeypatch, tmp_path
):
    """Righe YAML-classificate (COELBA, EMBASA, VIVO DA) entrano nel CSV, non nell'Excel."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [
            _launch(
                "DA COELBA 7014914336", "", raw_hash="launch-coelba",
                amount="-9164.48", transaction_date=date(2026, 4, 23),
            ),
            _launch(
                "EMBASA 142698660", "", raw_hash="launch-embasa",
                amount="-351.23", transaction_date=date(2026, 4, 13),
            ),
            _launch(
                "DA VIVO-BA 20304230426", "", raw_hash="launch-vivo",
                amount="-380.97", transaction_date=date(2026, 4, 6),
            ),
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_recurring_conta_corrente(tmp_path),
        }
    )

    csv_path = output_dir / CONTA_CORRENTE_CLASSIFICADOS_CSV
    assert csv_path.exists()
    assert result.csv_conta_corrente_classificados == csv_path
    assert result.conta_corrente_classificados_linhas == 3

    rows = _read_csv(csv_path)
    lancamentos_no_csv = {r["lancamento"] for r in rows}
    assert "DA COELBA 7014914336" in lancamentos_no_csv
    assert "EMBASA 142698660" in lancamentos_no_csv
    assert "DA VIVO-BA 20304230426" in lancamentos_no_csv

    coelba = next(r for r in rows if "COELBA" in r["lancamento"])
    assert coelba["conta_id"] == "4.01.011"
    assert coelba["mes"] == "2026-04"
    assert coelba["cnpj"] == "08884180000123"
    assert coelba["regra_id"] == "REGRA_4_01_011_CONTA_CORRENTE_COELBA"

    assert result.conta_corrente_candidatos == 0
    decisoes = _data_rows(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET)
    assert all(
        "COELBA" not in str(r.get("LANCAMENTO", ""))
        and "EMBASA" not in str(r.get("LANCAMENTO", ""))
        and "VIVO" not in str(r.get("LANCAMENTO", ""))
        for r in decisoes
    )


def test_conta_corrente_classificados_csv_nao_inclui_sem_regra(monkeypatch, tmp_path):
    """Righe senza regola vanno all'Excel decisionale, non al CSV."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [
            _launch(
                "DA COELBA 7014914336", "", raw_hash="launch-coelba",
                amount="-9164.48", transaction_date=date(2026, 4, 23),
            ),
            _launch(
                "PAGAMENTOS TRIB COD BARRAS", "", raw_hash="launch-sem-regra",
                amount="-10.00", transaction_date=date(2026, 4, 18),
            ),
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_recurring_conta_corrente(tmp_path),
        }
    )

    csv_rows = _read_csv(output_dir / CONTA_CORRENTE_CLASSIFICADOS_CSV)
    lancamentos_csv = {r["lancamento"] for r in csv_rows}
    assert "DA COELBA 7014914336" in lancamentos_csv
    assert "PAGAMENTOS TRIB COD BARRAS" not in lancamentos_csv

    assert result.conta_corrente_candidatos == 1
    decisoes = _data_rows(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET)
    assert any("PAGAMENTOS TRIB" in str(r.get("LANCAMENTO", "")) for r in decisoes)


def test_conta_corrente_classificados_csv_nao_inclui_escludere(monkeypatch, tmp_path):
    """Righe ESCLUDERE (SISPAG) non devono entrare nel CSV classificados."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [
            _launch(
                "DA COELBA 7014914336", "", raw_hash="launch-coelba",
                amount="-9164.48", transaction_date=date(2026, 4, 23),
            ),
            _launch(
                "SISPAG SALARIOS", "", raw_hash="launch-sispag",
                amount="-29641.68", transaction_date=date(2026, 4, 6),
            ),
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_recurring_conta_corrente(tmp_path),
        }
    )

    csv_rows = _read_csv(output_dir / CONTA_CORRENTE_CLASSIFICADOS_CSV)
    lancamentos_csv = {r["lancamento"] for r in csv_rows}
    assert "DA COELBA 7014914336" in lancamentos_csv
    assert "SISPAG SALARIOS" not in lancamentos_csv
    assert result.conta_corrente_classificados_linhas == 1


def test_conta_corrente_classificados_csv_nao_inclui_pix_bridge_comprovante(
    monkeypatch, tmp_path
):
    """PIX classificado via bridge comprovante NON deve entrare nel CSV (il comprovante lo copre)."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "Beneficiario Com Regra", raw_hash="receipt-bridge",
                amount="100.00", payment_date=date(2026, 4, 10),
                document="12.345.678/0001-90",
            ),
        ],
        [
            _launch(
                "PIX ENVIADO", "Nome Banco Diferente", raw_hash="launch-bridge",
                amount="-100.00", transaction_date=date(2026, 4, 10),
                document="12.345.678/0001-90",
            ),
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_supplier(tmp_path, "BENEFICIARIO COM REGRA"),
        }
    )

    assert result.conta_corrente_pix_yaml_via_comprovante == 1
    csv_path = output_dir / CONTA_CORRENTE_CLASSIFICADOS_CSV
    assert csv_path.exists()
    csv_rows = _read_csv(csv_path)
    assert csv_rows == []
    assert result.conta_corrente_classificados_linhas == 0


def test_conta_corrente_classificados_csv_manifest_atualizado(monkeypatch, tmp_path):
    """Il manifest deve includere il path CSV e il conteggio delle linhas classificate."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [
            _launch(
                "DA COELBA 7014914336", "", raw_hash="launch-coelba",
                amount="-9164.48", transaction_date=date(2026, 4, 23),
            ),
        ],
    )

    importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_recurring_conta_corrente(tmp_path),
        }
    )

    manifest = json.loads(
        (output_dir / "manifest_importa_revisao_banca.json").read_text(encoding="utf-8")
    )
    expected_csv = str(output_dir / CONTA_CORRENTE_CLASSIFICADOS_CSV)
    assert manifest["csv_conta_corrente_classificados"] == expected_csv
    assert manifest["conta_corrente_classificados_linhas"] == 1


def test_excel_decisional_nao_muda_com_novo_csv(monkeypatch, tmp_path):
    """A geração do CSV não deve alterar o conteúdo dos Excel decisionais."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [
            _launch(
                "DA COELBA 7014914336", "", raw_hash="launch-coelba",
                amount="-9164.48", transaction_date=date(2026, 4, 23),
            ),
            _launch(
                "PAGAMENTOS TRIB COD BARRAS", "", raw_hash="launch-sem-regra",
                amount="-10.00", transaction_date=date(2026, 4, 18),
            ),
        ],
    )

    result = importar_revisao_banca(
        **{
            **_import_kwargs(tmp_path, bank_raw_dir, output_dir),
            "regras_path": _write_rules_with_recurring_conta_corrente(tmp_path),
        }
    )

    # Classificato via YAML → non nell'Excel
    assert result.conta_corrente_regras_yaml_aplicadas == 1
    # Senza regola → nell'Excel
    assert result.conta_corrente_candidatos == 1
    decisoes = _data_rows(output_dir / CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET)
    assert len(decisoes) == 1
    assert decisoes[0]["LANCAMENTO"] == "PAGAMENTOS TRIB COD BARRAS"
    assert "COELBA" not in str(decisoes[0].get("LANCAMENTO", ""))


# ---------------------------------------------------------------------------
# Testes: conformidade template Excel decisionale banca (FZ-SKILLS-BANCA-EXCEL-011A)
# ---------------------------------------------------------------------------

_REQUIRED_COMPROVANTES_SHEETS = {
    "decisoes_comprovantes",
    "tecnico_comprovantes",
    "plano_contas",
    "listas_validacao",
    "instrucoes",
}
_REQUIRED_CONTA_CORRENTE_SHEETS = {
    "decisoes_conta_corrente",
    "tecnico_conta_corrente",
    "plano_contas",
    "listas_validacao",
    "instrucoes",
}


def test_template_excel_banca_fogli_obbligatori(monkeypatch, tmp_path):
    """Entrambi i workbook generati hanno esattamente i fogli obbligatori."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [_launch("PIX enviado", "Fornecedor Teste", raw_hash="launch-1")],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    comp_wb = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    assert set(comp_wb.sheetnames) == _REQUIRED_COMPROVANTES_SHEETS

    conta_wb = load_workbook(output_dir / CONTA_CORRENTE_XLSX, data_only=False)
    assert set(conta_wb.sheetnames) == _REQUIRED_CONTA_CORRENTE_SHEETS


def test_template_excel_banca_colonne_canoniche_comprovantes(monkeypatch, tmp_path):
    """Il foglio decisoes_comprovantes ha esattamente le 20 colonne canoniche nell'ordine corretto."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    workbook = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    headers = [cell.value for cell in workbook[COMPROVANTES_SHEET][1]]
    assert headers == COMPROVANTES_HEADERS


def test_template_excel_banca_colonne_canoniche_conta_corrente(monkeypatch, tmp_path):
    """Il foglio decisoes_conta_corrente ha esattamente le 24 colonne canoniche nell'ordine corretto."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [_launch("PIX enviado", "Fornecedor Teste", raw_hash="launch-1")],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    workbook = load_workbook(output_dir / CONTA_CORRENTE_XLSX, data_only=False)
    headers = [cell.value for cell in workbook[CONTA_CORRENTE_SHEET][1]]
    assert headers == CONTA_CORRENTE_HEADERS


def test_template_excel_banca_todas_colunas_visiveis(monkeypatch, tmp_path):
    """Tutte le colonne (operative, avanzate e tecniche) sono visibili — nessuna nascosta."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [_launch("PIX enviado", "Fornecedor Teste", raw_hash="launch-1")],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    comp_wb = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    comp_sheet = comp_wb[COMPROVANTES_SHEET]
    comp_headers = [cell.value for cell in comp_sheet[1]]
    for col, header in enumerate(comp_headers, start=1):
        dim = comp_sheet.column_dimensions.get(get_column_letter(col))
        is_hidden = dim is not None and bool(dim.hidden)
        assert not is_hidden, f"Comprovantes: colonna '{header}' NON deve essere nascosta"

    conta_wb = load_workbook(output_dir / CONTA_CORRENTE_XLSX, data_only=False)
    conta_sheet = conta_wb[CONTA_CORRENTE_SHEET]
    conta_headers = [cell.value for cell in conta_sheet[1]]
    for col, header in enumerate(conta_headers, start=1):
        dim = conta_sheet.column_dimensions.get(get_column_letter(col))
        is_hidden = dim is not None and bool(dim.hidden)
        assert not is_hidden, f"Conta corrente: colonna '{header}' NON deve essere nascosta"


def test_template_excel_banca_autofilter_cobre_area_completa(monkeypatch, tmp_path):
    """L'AutoFilter copre l'intera area dati incluse tutte le righe, non solo l'intestazione."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt("Fornecedor A", raw_hash="receipt-a"),
            _receipt("Fornecedor B", raw_hash="receipt-b"),
            _receipt("Fornecedor C", raw_hash="receipt-c"),
        ],
        [
            _launch("PIX A", "Fornecedor A", raw_hash="launch-a"),
            _launch("PIX B", "Fornecedor B", raw_hash="launch-b"),
        ],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    comp_wb = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    comp_sheet = comp_wb[COMPROVANTES_SHEET]
    assert comp_sheet.auto_filter.ref == comp_sheet.dimensions

    conta_wb = load_workbook(output_dir / CONTA_CORRENTE_XLSX, data_only=False)
    conta_sheet = conta_wb[CONTA_CORRENTE_SHEET]
    assert conta_sheet.auto_filter.ref == conta_sheet.dimensions


# ---------------------------------------------------------------------------
# Testes: validate_excel_template_banca (FZ-SKILLS-BANCA-EXCEL-011A.1)
# ---------------------------------------------------------------------------


def test_validate_template_banca_arquivo_ausente(tmp_path):
    """File Excel mancanti producono violazioni — nessuna conformanza possibile."""
    output_dir = tmp_path / "decisoes_regras"
    output_dir.mkdir()

    violations = validate_excel_template_banca(output_dir)

    assert any(COMPROVANTES_XLSX in v for v in violations)
    assert any(CONTA_CORRENTE_XLSX in v for v in violations)


def test_validate_template_banca_foglio_ausente(tmp_path):
    """Foglio operativo con nome sbagliato produce violazione."""
    output_dir = tmp_path / "decisoes_regras"
    output_dir.mkdir()

    wb = Workbook()
    wb.active.title = "foglio_sbagliato"
    wb.active.append(COMPROVANTES_HEADERS)
    wb.save(output_dir / COMPROVANTES_XLSX)

    wb2 = Workbook()
    wb2.active.title = CONTA_CORRENTE_SHEET
    wb2.active.append(CONTA_CORRENTE_HEADERS)
    wb2.save(output_dir / CONTA_CORRENTE_XLSX)

    violations = validate_excel_template_banca(output_dir)

    assert any(COMPROVANTES_SHEET in v for v in violations)


def test_validate_template_banca_coluna_extra(tmp_path):
    """Colonna extra nel foglio decisionale produce violazione BLOCCO."""
    output_dir = tmp_path / "decisoes_regras"
    output_dir.mkdir()

    wb = Workbook()
    sheet = wb.active
    sheet.title = COMPROVANTES_SHEET
    sheet.append(COMPROVANTES_HEADERS + ["COLONNA_EXTRA"])
    wb.save(output_dir / COMPROVANTES_XLSX)

    wb2 = Workbook()
    sheet2 = wb2.active
    sheet2.title = CONTA_CORRENTE_SHEET
    sheet2.append(CONTA_CORRENTE_HEADERS)
    wb2.save(output_dir / CONTA_CORRENTE_XLSX)

    violations = validate_excel_template_banca(output_dir)

    assert any("COLONNA_EXTRA" in v for v in violations)


def test_validate_template_banca_coluna_ausente(tmp_path):
    """Colonna mancante nel foglio decisionale produce violazione BLOCCO."""
    output_dir = tmp_path / "decisoes_regras"
    output_dir.mkdir()

    wb = Workbook()
    sheet = wb.active
    sheet.title = COMPROVANTES_SHEET
    sheet.append(COMPROVANTES_HEADERS[:-1])
    wb.save(output_dir / COMPROVANTES_XLSX)

    wb2 = Workbook()
    sheet2 = wb2.active
    sheet2.title = CONTA_CORRENTE_SHEET
    sheet2.append(CONTA_CORRENTE_HEADERS)
    wb2.save(output_dir / CONTA_CORRENTE_XLSX)

    violations = validate_excel_template_banca(output_dir)

    assert any(COMPROVANTES_HEADERS[-1] in v for v in violations)


def test_validate_template_banca_ordine_errato(tmp_path):
    """Colonne in ordine invertito (stesso set, ordine sbagliato) produce violazione BLOCCO."""
    output_dir = tmp_path / "decisoes_regras"
    output_dir.mkdir()

    wb = Workbook()
    sheet = wb.active
    sheet.title = COMPROVANTES_SHEET
    sheet.append(list(reversed(COMPROVANTES_HEADERS)))
    wb.save(output_dir / COMPROVANTES_XLSX)

    wb2 = Workbook()
    sheet2 = wb2.active
    sheet2.title = CONTA_CORRENTE_SHEET
    sheet2.append(CONTA_CORRENTE_HEADERS)
    wb2.save(output_dir / CONTA_CORRENTE_XLSX)

    violations = validate_excel_template_banca(output_dir)

    assert any(COMPROVANTES_SHEET in v for v in violations)


def test_validate_template_banca_conforme_com_dados(monkeypatch, tmp_path):
    """Workbook generato da importar_revisao_banca supera la validazione senza violazioni."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [_launch("PIX enviado", "Fornecedor Teste", raw_hash="launch-1")],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    violations = validate_excel_template_banca(output_dir)

    assert violations == [], f"Violazioni inattese: {violations}"


# ---------------------------------------------------------------------------
# Testes: colonne visibili — template senza colonne nascoste (FZ-SKILLS-BANCA-EXCEL-011A.5)
# ---------------------------------------------------------------------------


def test_validate_template_banca_nao_bloqueia_colonna_tecnica_visivel(tmp_path):
    """validate_excel_template_banca non segnala violazione se le colonne tecniche sono visibili."""
    output_dir = tmp_path / "decisoes_regras"
    output_dir.mkdir()

    for xlsx_name, headers, sheet_name in [
        (COMPROVANTES_XLSX, COMPROVANTES_HEADERS, COMPROVANTES_SHEET),
        (CONTA_CORRENTE_XLSX, CONTA_CORRENTE_HEADERS, CONTA_CORRENTE_SHEET),
    ]:
        wb = Workbook()
        sheet = wb.active
        sheet.title = sheet_name
        sheet.append(headers)
        sheet.append(["val"] * len(headers))
        sheet.auto_filter.ref = sheet.dimensions
        _add_test_dropdowns(sheet, headers, sheet_name == CONTA_CORRENTE_SHEET)
        wb.save(output_dir / xlsx_name)

    violations = validate_excel_template_banca(output_dir)

    assert violations == [], f"Violazioni inattese con colonne tecniche visibili: {violations}"


def test_template_excel_banca_chave_regra_visiveis(monkeypatch, tmp_path):
    """CHAVE_REGRA_* sono visibili nel template generato (non nascoste)."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [_launch("PIX enviado", "Fornecedor Teste", raw_hash="launch-1")],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    comp_wb = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    comp_sheet = comp_wb[COMPROVANTES_SHEET]
    comp_headers = [cell.value for cell in comp_sheet[1]]
    for chave_col in ("CHAVE_REGRA_BENEFICIARIO", "CHAVE_REGRA_REFERENCIA"):
        assert chave_col in comp_headers, f"{chave_col} non trovata nelle colonne comprovantes"
        col_idx = comp_headers.index(chave_col) + 1
        dim = comp_sheet.column_dimensions.get(get_column_letter(col_idx))
        is_hidden = dim is not None and bool(dim.hidden)
        assert not is_hidden, f"Comprovantes: {chave_col} deve essere visibile"

    conta_wb = load_workbook(output_dir / CONTA_CORRENTE_XLSX, data_only=False)
    conta_sheet = conta_wb[CONTA_CORRENTE_SHEET]
    conta_headers = [cell.value for cell in conta_sheet[1]]
    for chave_col in ("CHAVE_REGRA_LANCAMENTO", "CHAVE_REGRA_RAZAO_SOCIAL"):
        assert chave_col in conta_headers, f"{chave_col} non trovata nelle colonne conta corrente"
        col_idx = conta_headers.index(chave_col) + 1
        dim = conta_sheet.column_dimensions.get(get_column_letter(col_idx))
        is_hidden = dim is not None and bool(dim.hidden)
        assert not is_hidden, f"Conta corrente: {chave_col} deve essere visibile"


def test_template_excel_banca_conta_id_aprovada_visivel(monkeypatch, tmp_path):
    """CONTA_ID_APROVADA e CONTA_NOME_APROVADA sono visibili nel template generato."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [_launch("PIX enviado", "Fornecedor Teste", raw_hash="launch-1")],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    for xlsx_name, sheet_name in [
        (COMPROVANTES_XLSX, COMPROVANTES_SHEET),
        (CONTA_CORRENTE_XLSX, CONTA_CORRENTE_SHEET),
    ]:
        wb = load_workbook(output_dir / xlsx_name, data_only=False)
        sheet = wb[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        for col_name in ("CONTA_ID_APROVADA", "CONTA_NOME_APROVADA"):
            assert col_name in headers, f"{col_name} non trovata in {sheet_name}"
            col_idx = headers.index(col_name) + 1
            dim = sheet.column_dimensions.get(get_column_letter(col_idx))
            is_hidden = dim is not None and bool(dim.hidden)
            assert not is_hidden, f"{sheet_name}: {col_name} deve essere visibile"


def test_template_excel_banca_colonne_grigie_stilizzate(monkeypatch, tmp_path):
    """Le colonne tecniche/derivate hanno fill grigio nel template generato."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [_receipt("Fornecedor Teste", raw_hash="receipt-1")],
        [_launch("PIX enviado", "Fornecedor Teste", raw_hash="launch-1")],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    comp_wb = load_workbook(output_dir / COMPROVANTES_XLSX, data_only=False)
    comp_sheet = comp_wb[COMPROVANTES_SHEET]
    comp_headers = [cell.value for cell in comp_sheet[1]]
    for col_idx, header in enumerate(comp_headers, start=1):
        if header in COMPROVANTES_GREY_COLUMNS:
            cell_fill = comp_sheet.cell(row=1, column=col_idx).fill
            assert "D9D9D9" in cell_fill.fgColor.rgb.upper(), (
                f"Comprovantes: colonna '{header}' dovrebbe avere fill grigio (rgb={cell_fill.fgColor.rgb!r})"
            )

    conta_wb = load_workbook(output_dir / CONTA_CORRENTE_XLSX, data_only=False)
    conta_sheet = conta_wb[CONTA_CORRENTE_SHEET]
    conta_headers = [cell.value for cell in conta_sheet[1]]
    for col_idx, header in enumerate(conta_headers, start=1):
        if header in CONTA_CORRENTE_GREY_COLUMNS:
            cell_fill = conta_sheet.cell(row=1, column=col_idx).fill
            assert "D9D9D9" in cell_fill.fgColor.rgb.upper(), (
                f"Conta corrente: colonna '{header}' dovrebbe avere fill grigio (rgb={cell_fill.fgColor.rgb!r})"
            )


def test_validate_template_banca_dropdown_conta_escolhida_assente_blocca(tmp_path):
    """Dropdown CONTA_ESCOLHIDA assente è ancora un blocco (bloccante se assente)."""
    output_dir = tmp_path / "decisoes_regras"
    output_dir.mkdir()

    for xlsx_name, headers, sheet_name in [
        (COMPROVANTES_XLSX, COMPROVANTES_HEADERS, COMPROVANTES_SHEET),
        (CONTA_CORRENTE_XLSX, CONTA_CORRENTE_HEADERS, CONTA_CORRENTE_SHEET),
    ]:
        wb = Workbook()
        sheet = wb.active
        sheet.title = sheet_name
        sheet.append(headers)
        sheet.append(["val"] * len(headers))
        sheet.auto_filter.ref = sheet.dimensions
        wb.save(output_dir / xlsx_name)

    violations = validate_excel_template_banca(output_dir)

    assert any("CONTA_ESCOLHIDA" in v for v in violations), (
        f"Atteso blocco per CONTA_ESCOLHIDA assente, violazioni: {violations}"
    )


def _add_test_dropdowns(sheet, headers: list[str], dropdown_usar_somente: bool) -> None:
    """Aggiunge dropdown minimi per far passare la validazione template."""
    from openpyxl.worksheet.datavalidation import DataValidation as DV
    row_count = max(sheet.max_row - 1, 1)

    for col_name, formula in [
        ("STATUS_DECISAO", '"PENDENTE,APROVADO,REVISAR,NAO_CRIAR_REGRA,BLOQUEADO"'),
        ("ACAO_REGRA", '"CLASSIFICAR,ESCLUDERE,XML,NAO_CRIAR_REGRA,REVISAR,BLOQUEADO"'),
        ("RISCO_DUPLICIDADE_XML", '"OK,DUVIDA"'),
        ("ESCOPO_REGRA", '"COMPROVANTE,CONTA_CORRENTE"'),
    ]:
        if col_name not in headers:
            continue
        col_idx = headers.index(col_name) + 1
        letter = get_column_letter(col_idx)
        val = DV(type="list", formula1=formula, allow_blank=True)
        sheet.add_data_validation(val)
        val.add(f"{letter}2:{letter}{row_count + 1}")

    if dropdown_usar_somente and "USAR_SOMENTE_SEM_COMPROVANTE" in headers:
        col_idx = headers.index("USAR_SOMENTE_SEM_COMPROVANTE") + 1
        letter = get_column_letter(col_idx)
        val = DV(type="list", formula1='"SIM,NAO"', allow_blank=True)
        sheet.add_data_validation(val)
        val.add(f"{letter}2:{letter}{row_count + 1}")

    if "CONTA_ESCOLHIDA" in headers:
        col_idx = headers.index("CONTA_ESCOLHIDA") + 1
        letter = get_column_letter(col_idx)
        val = DV(type="list", formula1="=plano_contas!$A$2:$A$2", allow_blank=True)
        sheet.add_data_validation(val)
        val.add(f"{letter}2:{letter}{row_count + 1}")


# ---------------------------------------------------------------------------
# Testes: revisione banca mensile — coerenza PERIODO (FZ-SKILLS-BANCA-EXCEL-011A.4)
# ---------------------------------------------------------------------------


def test_manifest_declara_periodo_do_mes_importado(monkeypatch, tmp_path):
    """Il manifest dichiara esattamente il mese importato (YYYY-MM), non un intervallo."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(monkeypatch, [_receipt("Fornecedor Teste", raw_hash="receipt-1")], [])

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    manifest = json.loads((output_dir / "manifest_importa_revisao_banca.json").read_text(encoding="utf-8"))
    assert manifest["periodo"] == "2026-04"
    import re
    assert re.fullmatch(r"\d{4}-\d{2}", manifest["periodo"]), (
        f"periodo deve essere YYYY-MM (singolo mese), trovato: {manifest['periodo']!r}"
    )


def test_excel_comprovantes_periodo_coerente_com_banco_raw_dir(monkeypatch, tmp_path):
    """Ogni riga del foglio decisoes_comprovantes ha PERIODO == mese della bank_raw_dir."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt("Fornecedor A", raw_hash="receipt-a"),
            _receipt("Fornecedor B", raw_hash="receipt-b"),
        ],
        [],
    )

    importar_revisao_banca(**_import_kwargs(tmp_path, bank_raw_dir, output_dir))

    rows = _data_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    assert rows, "Nessuna riga dati trovata nel foglio decisoes_comprovantes"
    for row in rows:
        assert row["PERIODO"] == "2026-04", (
            f"PERIODO atteso '2026-04', trovato {row['PERIODO']!r} per beneficiario {row.get('BENEFICIARIO_FAVORECIDO')!r}"
        )


def test_importacao_mes_posterior_substitui_periodo_no_manifest(monkeypatch, tmp_path):
    """Importando marzo e poi aprile, il manifest dichiara solo aprile (ultimo mese importato)."""
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"

    # Prima importazione: marzo 2026
    mar_dir = tmp_path / "bank_raw" / "2026" / "2026-03"
    mar_cnpj = mar_dir / "08884180000123"
    mar_cnpj.mkdir(parents=True)
    (mar_cnpj / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    _patch_bank_sources(monkeypatch, [_receipt("Fornecedor Marzo", raw_hash="receipt-mar")], [])
    importar_revisao_banca(
        bank_raw_dir=mar_dir,
        output_dir=output_dir,
        regras_path=_write_rules(tmp_path),
        plano_contas_path=_write_plan(tmp_path),
        mapeamento_dre_path=_write_mapping(tmp_path),
    )

    manifest_dopo_marzo = json.loads(
        (output_dir / "manifest_importa_revisao_banca.json").read_text(encoding="utf-8")
    )
    assert manifest_dopo_marzo["periodo"] == "2026-03"

    # Seconda importazione: aprile 2026
    apr_dir = tmp_path / "bank_raw" / "2026" / "2026-04"
    apr_cnpj = apr_dir / "08884180000123"
    apr_cnpj.mkdir(parents=True)
    (apr_cnpj / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    _patch_bank_sources(monkeypatch, [_receipt("Fornecedor Abril", raw_hash="receipt-apr")], [])
    importar_revisao_banca(
        bank_raw_dir=apr_dir,
        output_dir=output_dir,
        regras_path=_write_rules(tmp_path),
        plano_contas_path=_write_plan(tmp_path),
        mapeamento_dre_path=_write_mapping(tmp_path),
    )

    manifest_dopo_aprile = json.loads(
        (output_dir / "manifest_importa_revisao_banca.json").read_text(encoding="utf-8")
    )
    # Il manifest corrente deve dichiarare SOLO aprile — non un intervallo o "marzo-aprile"
    assert manifest_dopo_aprile["periodo"] == "2026-04", (
        f"Dopo importazione aprile, manifest deve dichiarare '2026-04', trovato: {manifest_dopo_aprile['periodo']!r}"
    )

    # Le righe Excel devono contenere PERIODO == "2026-04", non "2026-03"
    rows = _data_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    assert rows, "Nessuna riga trovata negli Excel dopo importazione aprile"
    for row in rows:
        assert row["PERIODO"] == "2026-04", (
            f"Dopo importazione aprile, PERIODO atteso '2026-04', trovato {row['PERIODO']!r}"
        )


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=";")
        return list(reader)


def _write_existing_comprovantes(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = COMPROVANTES_SHEET
    headers = [
        "STATUS_DECISAO",
        "ACAO_REGRA",
        "PERIODO",
        "CNPJ",
        "DATA_PAGAMENTO",
        "VALOR",
        "BENEFICIARIO_FAVORECIDO",
        "REFERENCIA_EMPRESA",
        "TIPO_PAGAMENTO",
        "CONTA_ESCOLHIDA",
        "CONTA_ID_APROVADA",
        "CONTA_NOME_APROVADA",
        "CHAVE_REGRA_BENEFICIARIO",
        "CHAVE_REGRA_REFERENCIA",
        "RISCO_DUPLICIDADE_XML",
        "ESCOPO_REGRA",
        "QTD_OCORRENCIAS",
        "VALOR_TOTAL",
        "EXEMPLOS",
        "OBSERVACAO_OPERADOR",
    ]
    sheet.append(headers)
    sheet.append(
        [
            "APROVADO",
            "CLASSIFICAR",
            "2026-04",
            "08884180000123",
            "2026-04-05",
            20,
            "FORNECEDOR TESTE",
            "-",
            "PIX Transferencias",
            "4.01.012 - Agua e Esgoto",
            "4.01.012",
            "Agua e Esgoto",
            "FORNECEDOR TESTE",
            "",
            "OK",
            "COMPROVANTE",
            99,
            999,
            "old",
            "preservar",
        ]
    )
    workbook.save(path)


def _write_existing_comprovantes_with_stale_formula(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = COMPROVANTES_SHEET
    headers = [
        "STATUS_DECISAO",
        "ACAO_REGRA",
        "PERIODO",
        "CNPJ",
        "DATA_PAGAMENTO",
        "VALOR",
        "BENEFICIARIO_FAVORECIDO",
        "REFERENCIA_EMPRESA",
        "TIPO_PAGAMENTO",
        "CONTA_ESCOLHIDA",
        "CONTA_ID_APROVADA",
        "CONTA_NOME_APROVADA",
        "CHAVE_REGRA_BENEFICIARIO",
        "CHAVE_REGRA_REFERENCIA",
        "RISCO_DUPLICIDADE_XML",
        "ESCOPO_REGRA",
        "QTD_OCORRENCIAS",
        "VALOR_TOTAL",
        "EXEMPLOS",
        "OBSERVACAO_OPERADOR",
    ]
    sheet.append(headers)
    sheet.append(
        [
            "APROVADO",
            "CLASSIFICAR",
            "2026-04",
            "08884180000123",
            "2026-04-04",
            10,
            "OUTRO FORNECEDOR",
            "-",
            "PIX Transferencias",
            "7.01.001 - Tarifas Bancarias",
            '=IFERROR(LEFT(J2,FIND(" - ",J2)-1),"")',
            '=IFERROR(MID(J2,FIND(" - ",J2)+3,999),"")',
            "OUTRO FORNECEDOR",
            "",
            "OK",
            "COMPROVANTE",
            1,
            10,
            "old",
            "",
        ]
    )
    sheet.append(
        [
            "APROVADO",
            "CLASSIFICAR",
            "2026-04",
            "08884180000123",
            "2026-04-05",
            20,
            "FORNECEDOR TESTE",
            "-",
            "PIX Transferencias",
            "4.01.012 - Agua e Esgoto",
            '=IFERROR(LEFT(J2,FIND(" - ",J2)-1),"")',
            '=IFERROR(MID(J2,FIND(" - ",J2)+3,999),"")',
            "FORNECEDOR TESTE",
            "",
            "OK",
            "COMPROVANTE",
            1,
            20,
            "old",
            "",
        ]
    )
    workbook.save(path)


def _write_rules(tmp_path: Path) -> Path:
    path = tmp_path / "regras.yaml"
    path.write_text("regras: []\nexclusoes: []\n", encoding="utf-8")
    return path


def _write_rules_with_supplier(tmp_path: Path, supplier: str) -> Path:
    path = tmp_path / "regras.yaml"
    path.write_text(
        f"""regras:
- regra_id: REGRA_FORNECEDOR_TESTE
  conta_id: "4.01.012"
  conta_nome: Agua e Esgoto
  dre_linha_id: "06.02"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
    - {supplier}
    referencia_empresa_contem_algum: []
exclusoes: []
""",
        encoding="utf-8",
    )
    return path


def _write_rules_with_supplier_and_boleto_fallback(tmp_path: Path, supplier: str) -> Path:
    path = tmp_path / "regras.yaml"
    path.write_text(
        f"""regras:
- regra_id: REGRA_ABRASEL
  conta_id: "4.01.012"
  conta_nome: Agua e Esgoto
  dre_linha_id: "06.02"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
    - {supplier}
    referencia_empresa_contem_algum: []
- regra_id: REGRA_4_01_014_FIBRANET_BA_TELECOM
  conta_id: "4.01.014"
  conta_nome: Internet e Telefonia
  dre_linha_id: "06.02"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
    - FIBRANET BA TELECOM LTDA
    referencia_empresa_contem_algum: []
exclusoes:
- regra_id: EXCLUSAO_XML_CONTA_CORRENTE_BOLETO_PAGO
  acao: XML
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum: []
    descricao_comprovante_contem_algum:
    - BOLETO PAGO
""",
        encoding="utf-8",
    )
    return path


def _write_rules_with_sinthotesb_and_boleto_fallback(tmp_path: Path) -> Path:
    path = tmp_path / "regras.yaml"
    path.write_text(
        """regras: []
exclusoes:
- regra_id: EXCLUSAO_ESCLUDERE_SINTHOTESB_SIND_TRAB_BARES_HOT
  acao: ESCLUDERE
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
    - SINTHOTESB SIND TRAB BARES HOT
    referencia_empresa_contem_algum: []
- regra_id: EXCLUSAO_XML_CONTA_CORRENTE_BOLETO_PAGO
  acao: XML
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum: []
    descricao_comprovante_contem_algum:
    - BOLETO PAGO
""",
        encoding="utf-8",
    )
    return path


def _write_rules_with_recurring_conta_corrente(tmp_path: Path) -> Path:
    path = tmp_path / "regras.yaml"
    path.write_text(
        """regras:
- regra_id: REGRA_4_01_011_CONTA_CORRENTE_COELBA
  conta_id: "4.01.011"
  conta_nome: Energia Eletrica
  dre_linha_id: "06.02"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - COELBA
- regra_id: REGRA_4_01_012_CONTA_CORRENTE_EMBASA
  conta_id: "4.01.012"
  conta_nome: Agua e Esgoto
  dre_linha_id: "06.02"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - EMBASA
- regra_id: REGRA_4_01_014_CONTA_CORRENTE_VIVO
  conta_id: "4.01.014"
  conta_nome: Internet e Telefonia
  dre_linha_id: "06.02"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - VIVO
- regra_id: REGRA_7_01_001_CONTA_CORRENTE_DEBITO_SEGURO
  conta_id: "7.01.001"
  conta_nome: Tarifas Bancarias
  dre_linha_id: "09.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - DEBITO SEGURO
- regra_id: REGRA_5_01_001_CONTA_CORRENTE_SIMPLES_NACIONAL
  conta_id: "5.01.001"
  conta_nome: Simples Nacional - DAS Mensal - Anexo I
  dre_linha_id: "02.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
    - SIMPLES NACIONAL
    referencia_empresa_contem_algum: []
exclusoes:
- regra_id: EXCLUSAO_ESCLUDERE_CONTA_CORRENTE_SISPAG_SALARIOS
  acao: ESCLUDERE
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - SISPAG SALARIOS
    tipo_pagamento_contem_algum:
    - SISPAG_AGREGADO
""",
        encoding="utf-8",
    )
    return path


def _write_rules_with_xml_exclusion(tmp_path: Path, beneficiary: str) -> Path:
    path = tmp_path / "regras.yaml"
    path.write_text(
        f"""regras: []
exclusoes:
- regra_id: EXCLUSAO_XML_TEST_{beneficiary.replace(' ', '_')}
  acao: XML
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
    - {beneficiary}
    referencia_empresa_contem_algum: []
""",
        encoding="utf-8",
    )
    return path


def _write_rules_with_positive_rule_julio_cesar(tmp_path: Path) -> Path:
    path = tmp_path / "regras.yaml"
    path.write_text(
        """regras:
- regra_id: REGRA_8_02_003_JULIO_CESAR_RODRIGUES_DE_SOUZA
  conta_id: "8.02.003"
  conta_nome: Manutencao de Geladeiras e Freezers
  dre_linha_id: "10.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
    - JULIO CESAR RODRIGUES DE SOUZA
    referencia_empresa_contem_algum: []
exclusoes: []
""",
        encoding="utf-8",
    )
    return path


def _write_plan(tmp_path: Path) -> Path:
    path = tmp_path / "plano.yaml"
    path.write_text(
        """contas:
- id: "4.01.011"
  nome: Energia Eletrica
- id: "4.01.012"
  nome: Agua e Esgoto
- id: "4.01.014"
  nome: Internet e Telefonia
- id: "5.01.001"
  nome: Simples Nacional - DAS Mensal - Anexo I
- id: "7.01.001"
  nome: Tarifas Bancarias
- id: "8.02.003"
  nome: Manutencao de Geladeiras e Freezers
""",
        encoding="utf-8",
    )
    return path


def _make_bank_raw_month(tmp_path: Path) -> tuple[Path, Path]:
    bank_raw_dir = tmp_path / "bank_raw" / "2026" / "2026-04"
    cnpj_dir = bank_raw_dir / "08884180000123"
    cnpj_dir.mkdir(parents=True)
    return bank_raw_dir, cnpj_dir


def _receipt(
    beneficiary: str,
    *,
    raw_hash: str,
    reference: str = "-",
    amount: str = "20.00",
    payment_date: date = date(2026, 4, 5),
    document: str = "12.345.678/0001-90",
    payment_type: str = "PIX Transferencias",
) -> PaymentReceiptRaw:
    return PaymentReceiptRaw(
        payment_date=payment_date,
        amount=Decimal(amount),
        beneficiary_name=beneficiary,
        beneficiary_document=document,
        payment_type=payment_type,
        reference_id=reference,
        status="Efetuado",
        source_file="ConsultaPagamentos.xls",
        raw_hash=raw_hash,
    )


def _launch(
    description: str,
    counterparty: str,
    *,
    raw_hash: str,
    amount: str = "-20.00",
    transaction_date: date = date(2026, 4, 5),
    document: str = "12.345.678/0001-90",
) -> BankLaunchEnrichment:
    return BankLaunchEnrichment(
        transaction_date=transaction_date,
        amount=Decimal(amount),
        description=description,
        counterparty_name=counterparty,
        counterparty_document=document,
        row_balance=None,
        source_file="Extrato_Lancamentos.xlsx",
        raw_hash=raw_hash,
    )


def _patch_bank_sources(
    monkeypatch,
    receipts: list[PaymentReceiptRaw],
    launches: list[BankLaunchEnrichment],
) -> None:
    monkeypatch.setattr(
        "fantozzi.banking.revisao_banca.parse_payment_receipts",
        lambda path: list(receipts),
    )
    monkeypatch.setattr(
        "fantozzi.banking.revisao_banca.parse_bank_launches",
        lambda path: (list(launches), []),
    )


def _import_kwargs(tmp_path: Path, bank_raw_dir: Path, output_dir: Path) -> dict[str, Path]:
    return {
        "bank_raw_dir": bank_raw_dir,
        "output_dir": output_dir,
        "regras_path": _write_rules(tmp_path),
        "plano_contas_path": _write_plan(tmp_path),
        "mapeamento_dre_path": _write_mapping(tmp_path),
    }


def _sheet_rows(path: Path, sheet_name: str) -> list[tuple[object, ...]]:
    workbook = load_workbook(path, data_only=False)
    return list(workbook[sheet_name].iter_rows(values_only=True))


def _data_rows(path: Path, sheet_name: str) -> list[dict[str, object]]:
    rows = _sheet_rows(path, sheet_name)
    headers = [str(value or "") for value in rows[0]]
    return [
        {headers[index]: value for index, value in enumerate(row)}
        for row in rows[1:]
    ]


def _row_by_index(path: Path, sheet_name: str, *, row: int) -> dict[str, object]:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    return {
        str(headers[index]): sheet.cell(row=row, column=index + 1).value
        for index in range(len(headers))
    }


def _set_cell_by_header(
    path: Path,
    sheet_name: str,
    *,
    row: int,
    values: dict[str, object],
) -> None:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    for header, value in values.items():
        sheet.cell(row=row, column=headers.index(header) + 1).value = value
    workbook.save(path)


def _assert_dropdown(sheet, header: str, formula1: str, *, row: int = 2) -> None:
    assert _has_dropdown(sheet, header, formula1, row=row)


def _assert_dropdown_prefix(sheet, header: str, formula_prefix: str, *, row: int = 2) -> None:
    assert _has_dropdown_prefix(sheet, header, formula_prefix, row=row)


def _has_dropdown(sheet, header: str, formula1: str, *, row: int = 2) -> bool:
    cell = _cell_ref(sheet, header, row)
    return any(
        validation.type == "list"
        and validation.formula1 == formula1
        and cell in str(validation.sqref)
        for validation in sheet.data_validations.dataValidation
    )


def _has_dropdown_prefix(sheet, header: str, formula_prefix: str, *, row: int = 2) -> bool:
    cell = _cell_ref(sheet, header, row)
    return any(
        validation.type == "list"
        and str(validation.formula1).startswith(formula_prefix)
        and cell in str(validation.sqref)
        for validation in sheet.data_validations.dataValidation
    )


def _cell_ref(sheet, header: str, row: int) -> str:
    headers = [cell.value for cell in sheet[1]]
    column = chr(ord("A") + headers.index(header))
    return f"{column}{row}"


# ---------------------------------------------------------------------------
# FZ-SKILLS-BANCA-EXCEL-011A.6 — dropdown x14/extLst riconosciuto dal validatore
# ---------------------------------------------------------------------------


def _inject_x14_conta_escolhida(xlsx_path: Path, sheet_name: str, col_letter: str) -> None:
    """Inject x14:dataValidations extLst for CONTA_ESCOLHIDA via ZIP manipulation (no openpyxl save)."""
    import io as _io
    import re as _re
    import zipfile as _zf

    x14_block = (
        "<extLst>"
        '<ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}"'
        ' xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        '<x14:dataValidations count="1"'
        ' xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
        '<x14:dataValidation type="list" allowBlank="1"'
        ' errorTitle="Conta invalida" error="Escolha uma conta da lista.">'
        "<x14:formula1><xm:f>plano_contas!$A$2:$A$2</xm:f></x14:formula1>"
        f"<xm:sqref>{col_letter}2:{col_letter}2</xm:sqref>"
        "</x14:dataValidation>"
        "</x14:dataValidations>"
        "</ext>"
        "</extLst>"
    )

    with _zf.ZipFile(xlsx_path, "r") as zf:
        entries = {name: zf.read(name) for name in zf.namelist()}

    # Resolve sheet XML path via sequential index in workbook.xml (openpyxl always uses sheetN.xml in order)
    wb_xml = entries["xl/workbook.xml"].decode("utf-8")
    sheet_elems = _re.findall(r'<sheet\b[^>]+/?>', wb_xml)
    sheet_idx = None
    for i, elem in enumerate(sheet_elems, start=1):
        if f'name="{sheet_name}"' in elem or f"name='{sheet_name}'" in elem:
            sheet_idx = i
            break
    assert sheet_idx is not None, f"Sheet {sheet_name!r} not found in workbook.xml"
    target = f"xl/worksheets/sheet{sheet_idx}.xml"

    sheet_xml = entries[target].decode("utf-8")
    assert "</worksheet>" in sheet_xml
    entries[target] = sheet_xml.replace("</worksheet>", x14_block + "</worksheet>").encode("utf-8")

    buf = _io.BytesIO()
    with _zf.ZipFile(buf, "w", _zf.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    xlsx_path.write_bytes(buf.getvalue())


def _make_minimal_xlsx_for_validation(
    path: Path,
    xlsx_name: str,
    headers: list[str],
    sheet_name: str,
    dropdown_usar_somente: bool,
    include_standard_conta_escolhida: bool = True,
) -> None:
    """Create a minimal XLSX that passes validate_excel_template_banca."""
    from openpyxl.worksheet.datavalidation import DataValidation as DV

    wb = Workbook()
    sheet = wb.active
    sheet.title = sheet_name
    sheet.append(headers)
    sheet.append(["val"] * len(headers))
    sheet.auto_filter.ref = sheet.dimensions

    row_count = 1
    for col_name, formula in [
        ("STATUS_DECISAO", '"PENDENTE,APROVADO,REVISAR,NAO_CRIAR_REGRA,BLOQUEADO"'),
        ("ACAO_REGRA", '"CLASSIFICAR,ESCLUDERE,XML,NAO_CRIAR_REGRA,REVISAR,BLOQUEADO"'),
        ("RISCO_DUPLICIDADE_XML", '"OK,DUVIDA"'),
        ("ESCOPO_REGRA", '"COMPROVANTE,CONTA_CORRENTE"'),
    ]:
        if col_name not in headers:
            continue
        col_idx = headers.index(col_name) + 1
        letter = get_column_letter(col_idx)
        val = DV(type="list", formula1=formula, allow_blank=True)
        sheet.add_data_validation(val)
        val.add(f"{letter}2:{letter}{row_count + 1}")

    if dropdown_usar_somente and "USAR_SOMENTE_SEM_COMPROVANTE" in headers:
        col_idx = headers.index("USAR_SOMENTE_SEM_COMPROVANTE") + 1
        letter = get_column_letter(col_idx)
        val = DV(type="list", formula1='"SIM,NAO"', allow_blank=True)
        sheet.add_data_validation(val)
        val.add(f"{letter}2:{letter}{row_count + 1}")

    if include_standard_conta_escolhida and "CONTA_ESCOLHIDA" in headers:
        col_idx = headers.index("CONTA_ESCOLHIDA") + 1
        letter = get_column_letter(col_idx)
        val = DV(type="list", formula1="=plano_contas!$A$2:$A$2", allow_blank=True)
        sheet.add_data_validation(val)
        val.add(f"{letter}2:{letter}{row_count + 1}")

    # plano_contas sheet (required by validator)
    plano = wb.create_sheet("plano_contas")
    plano.append(["CONTA_ESCOLHIDA", "CONTA_ID", "CONTA_NOME"])
    plano.append(["1.01.001 - Cassa", "1.01.001", "Cassa"])

    wb.create_sheet("tecnico_comprovantes" if sheet_name == COMPROVANTES_SHEET else "tecnico_conta_corrente")
    wb.create_sheet("listas_validacao")
    wb.create_sheet("instrucoes")

    wb.save(path / xlsx_name)


def test_validate_template_banca_conta_escolhida_standard_ancora_ok(tmp_path):
    """Standard CONTA_ESCOLHIDA validation (openpyxl-generated) continua a passare."""
    output_dir = tmp_path / "decisoes_regras"
    output_dir.mkdir()

    for xlsx_name, headers, sheet_name, usar_somente in [
        (COMPROVANTES_XLSX, COMPROVANTES_HEADERS, COMPROVANTES_SHEET, False),
        (CONTA_CORRENTE_XLSX, CONTA_CORRENTE_HEADERS, CONTA_CORRENTE_SHEET, True),
    ]:
        _make_minimal_xlsx_for_validation(
            output_dir, xlsx_name, headers, sheet_name, usar_somente,
            include_standard_conta_escolhida=True,
        )

    violations = validate_excel_template_banca(output_dir)
    assert violations == [], f"Atteso nessun blocco, trovato: {violations}"


def test_validate_template_banca_conta_escolhida_x14_riconosciuto(tmp_path):
    """CONTA_ESCOLHIDA in formato x14/extLst (come Excel salva cross-sheet ref) è riconosciuta."""
    output_dir = tmp_path / "decisoes_regras"
    output_dir.mkdir()

    for xlsx_name, headers, sheet_name, usar_somente, col_letter in [
        (COMPROVANTES_XLSX, COMPROVANTES_HEADERS, COMPROVANTES_SHEET, False, "J"),
        (CONTA_CORRENTE_XLSX, CONTA_CORRENTE_HEADERS, CONTA_CORRENTE_SHEET, True, "G"),
    ]:
        _make_minimal_xlsx_for_validation(
            output_dir, xlsx_name, headers, sheet_name, usar_somente,
            include_standard_conta_escolhida=False,  # no standard dropdown
        )
        _inject_x14_conta_escolhida(output_dir / xlsx_name, sheet_name, col_letter)

    violations = validate_excel_template_banca(output_dir)
    assert violations == [], (
        f"x14 CONTA_ESCOLHIDA deve essere riconosciuta dal fallback, violazioni: {violations}"
    )


def test_sheet_has_x14_list_validation_prefix_assente_restituisce_false(tmp_path):
    """_sheet_has_x14_list_validation_prefix restituisce False se x14 non è presente."""
    from openpyxl.worksheet.datavalidation import DataValidation as DV

    # XLSX senza nessun x14 (solo standard openpyxl)
    xlsx_path = tmp_path / "no_x14.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = COMPROVANTES_SHEET
    ws.append(COMPROVANTES_HEADERS)
    ws.append(["val"] * len(COMPROVANTES_HEADERS))
    wb.save(xlsx_path)

    result = _sheet_has_x14_list_validation_prefix(
        xlsx_path, COMPROVANTES_SHEET, "CONTA_ESCOLHIDA", COMPROVANTES_HEADERS, "=plano_contas!"
    )
    assert result is False, "Atteso False: nessun x14 presente"


def _write_mapping(tmp_path: Path) -> Path:
    path = tmp_path / "mapeamento.yaml"
    path.write_text(
        """mapeamento:
  "4.01.011":
    dre_linha_id: "06.02"
    entra_dre: true
  "4.01.012":
    dre_linha_id: "06.02"
    entra_dre: true
  "4.01.014":
    dre_linha_id: "06.02"
    entra_dre: true
  "5.01.001":
    dre_linha_id: "02.01"
    entra_dre: true
  "7.01.001":
    dre_linha_id: "09.01"
    entra_dre: true
  "8.02.003":
    dre_linha_id: "10.01"
    entra_dre: true
""",
        encoding="utf-8",
    )
    return path


# ---------------------------------------------------------------------------
# Helpers: plano + mapping con conto 5.01.003 (parcelamento federal)
# ---------------------------------------------------------------------------


def _write_plan_with_parcelamento(tmp_path: Path) -> Path:
    path = tmp_path / "plano_parcelamento.yaml"
    path.write_text(
        """contas:
- id: "4.01.012"
  nome: Agua e Esgoto
- id: "5.01.001"
  nome: Simples Nacional - DAS Mensal - Anexo I
- id: "5.01.003"
  nome: Parcelamento Federal
""",
        encoding="utf-8",
    )
    return path


def _write_mapping_with_parcelamento(tmp_path: Path) -> Path:
    path = tmp_path / "mapeamento_parcelamento.yaml"
    path.write_text(
        """mapeamento:
  "4.01.012":
    dre_linha_id: "06.02"
    entra_dre: true
  "5.01.001":
    dre_linha_id: "02.01"
    entra_dre: true
  "5.01.003":
    dre_linha_id: "07.01"
    entra_dre: true
""",
        encoding="utf-8",
    )
    return path


def _write_rules_simples_parcelamento_conflict(tmp_path: Path) -> Path:
    """Regras che simulano il conflitto reale PARCSN/PERTSN/SISPAR vs SIMPLES_NACIONAL."""
    path = tmp_path / "regras_conflict.yaml"
    path.write_text(
        """regras:
- regra_id: REGRA_5_01_001_SIMPLES_NACIONAL_GENERICO
  conta_id: "5.01.001"
  conta_nome: Simples Nacional - DAS Mensal - Anexo I
  dre_linha_id: "02.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
    - SIMPLES NACIONAL
    referencia_empresa_contem_algum: []
- regra_id: REGRA_5_01_003_PARCSN
  conta_id: "5.01.003"
  conta_nome: Parcelamento Federal
  dre_linha_id: "07.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - parcsn
- regra_id: REGRA_5_01_003_PERTSN
  conta_id: "5.01.003"
  conta_nome: Parcelamento Federal
  dre_linha_id: "07.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - pertsn
- regra_id: REGRA_5_01_003_SISPAR_PGFN
  conta_id: "5.01.003"
  conta_nome: Parcelamento Federal
  dre_linha_id: "07.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - sispar
    - pgfn
exclusoes: []
""",
        encoding="utf-8",
    )
    return path


# ---------------------------------------------------------------------------
# Testes: FZ-SKILLS-BANCA-REVISIONE-011C — filtro comprovantes ja cobertos YAML
# ---------------------------------------------------------------------------


def test_comprovante_parcsn_coberto_por_yaml_nao_aparece_como_candidato(
    monkeypatch,
    tmp_path,
):
    """PARCSN coberto por regra especifica em referencia_empresa nao deve ser candidato operador."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "DAS SIMPLES NACIONAL COD BARRA",
                raw_hash="receipt-parcsn",
                reference="PARCSN 2 RATA 26/60",
                payment_type="DAS codigo barras",
            )
        ],
        [],
    )

    result = importar_revisao_banca(
        bank_raw_dir=bank_raw_dir,
        output_dir=output_dir,
        regras_path=_write_rules_simples_parcelamento_conflict(tmp_path),
        plano_contas_path=_write_plan_with_parcelamento(tmp_path),
        mapeamento_dre_path=_write_mapping_with_parcelamento(tmp_path),
    )

    assert result.comprovantes_candidatos == 0
    rows = _data_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    assert not any("PARCSN" in str(r.get("REFERENCIA_EMPRESA", "")) for r in rows)


def test_comprovante_pertsn_coberto_por_yaml_nao_aparece_como_candidato(
    monkeypatch,
    tmp_path,
):
    """PERTSN coberto por regra especifica em referencia_empresa nao deve ser candidato operador."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "DAS SIMPLES NACIONAL COD BARRA",
                raw_hash="receipt-pertsn",
                reference="PERTSN 9101 RATA 92",
                payment_type="DAS codigo barras",
            )
        ],
        [],
    )

    result = importar_revisao_banca(
        bank_raw_dir=bank_raw_dir,
        output_dir=output_dir,
        regras_path=_write_rules_simples_parcelamento_conflict(tmp_path),
        plano_contas_path=_write_plan_with_parcelamento(tmp_path),
        mapeamento_dre_path=_write_mapping_with_parcelamento(tmp_path),
    )

    assert result.comprovantes_candidatos == 0
    rows = _data_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    assert not any("PERTSN" in str(r.get("REFERENCIA_EMPRESA", "")) for r in rows)


def test_comprovante_pgfn_sispar_coberto_por_yaml_nao_aparece_como_candidato(
    monkeypatch,
    tmp_path,
):
    """PGFN-SISPAR coberto por regra especifica em referencia_empresa nao deve ser candidato operador."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "DAS SIMPLES NACIONAL COD BARRA",
                raw_hash="receipt-sispar",
                reference="PGFN-SISPAR 8916087",
                payment_type="DAS codigo barras",
            )
        ],
        [],
    )

    result = importar_revisao_banca(
        bank_raw_dir=bank_raw_dir,
        output_dir=output_dir,
        regras_path=_write_rules_simples_parcelamento_conflict(tmp_path),
        plano_contas_path=_write_plan_with_parcelamento(tmp_path),
        mapeamento_dre_path=_write_mapping_with_parcelamento(tmp_path),
    )

    assert result.comprovantes_candidatos == 0
    rows = _data_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    assert not any("SISPAR" in str(r.get("REFERENCIA_EMPRESA", "")) for r in rows)


def test_comprovante_sem_regra_akira_permanece_candidato(
    monkeypatch,
    tmp_path,
):
    """Comprovante sem cobertura YAML (AKIRA) deve permanecer candidato no Excel operativo."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "ConsultaPagamentos.xls").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [
            _receipt(
                "AKIRA DESINSETIZADORA E DESENT",
                raw_hash="receipt-akira",
                reference="-",
                payment_type="PIX Transferencias",
            )
        ],
        [],
    )

    result = importar_revisao_banca(
        bank_raw_dir=bank_raw_dir,
        output_dir=output_dir,
        regras_path=_write_rules_simples_parcelamento_conflict(tmp_path),
        plano_contas_path=_write_plan_with_parcelamento(tmp_path),
        mapeamento_dre_path=_write_mapping_with_parcelamento(tmp_path),
    )

    assert result.comprovantes_candidatos == 1
    rows = _data_rows(output_dir / COMPROVANTES_XLSX, COMPROVANTES_SHEET)
    assert any("AKIRA" in str(r.get("BENEFICIARIO_FAVORECIDO", "")) for r in rows)


def test_conta_corrente_simples_nacional_nao_regride_com_tiebreaker(
    monkeypatch,
    tmp_path,
):
    """Lancamento conta corrente SIMPLES NACIONAL ainda classifica corretamente — sem regressao."""
    bank_raw_dir, cnpj_dir = _make_bank_raw_month(tmp_path)
    (cnpj_dir / "Extrato_Lancamentos.xlsx").write_bytes(b"placeholder")
    output_dir = tmp_path / "bank_raw" / "_decisoes_regras"
    _patch_bank_sources(
        monkeypatch,
        [],
        [
            _launch(
                "PAGAMENTOS TRIB COD BARRAS",
                "SIMPLES NACIONAL",
                raw_hash="launch-simples",
                amount="-1027.55",
                transaction_date=date(2026, 4, 17),
            )
        ],
    )

    result = importar_revisao_banca(
        bank_raw_dir=bank_raw_dir,
        output_dir=output_dir,
        regras_path=_write_rules_simples_parcelamento_conflict(tmp_path),
        plano_contas_path=_write_plan_with_parcelamento(tmp_path),
        mapeamento_dre_path=_write_mapping_with_parcelamento(tmp_path),
    )

    assert result.conta_corrente_regras_yaml_aplicadas == 1
    assert result.conta_corrente_candidatos == 0

    tecnico = _data_rows(output_dir / CONTA_CORRENTE_XLSX, "tecnico_conta_corrente")
    by_hash = {row["HASH_ORIGEM"]: row for row in tecnico}
    assert by_hash["launch-simples"]["REGRA_EXISTENTE"] == "REGRA_5_01_001_SIMPLES_NACIONAL_GENERICO"
