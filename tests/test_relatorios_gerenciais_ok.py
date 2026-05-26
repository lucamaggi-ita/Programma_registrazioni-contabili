from datetime import date
from decimal import Decimal
from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook
import pytest

from fantozzi.excel_writer import write_lancamentos_to_excel
from fantozzi.models import Duplicata, Lancamento
from fantozzi.nfe_importer import parse_nfe
from fantozzi.relatorios_gerenciais_ok import (
    EXCLUIDOS_OPERACOES_SHEET,
    EXCLUIDOS_REVISAR_SHEET,
    INSTRUCOES_SHEET,
    LANCAMENTOS_OK_BASE_SHEET,
    RESUMO_CNPJ_SHEET,
    RESUMO_CONTA_SHEET,
    RESUMO_FORNECEDOR_SHEET,
    RESUMO_GERAL_SHEET,
    RESUMO_MENSAL_SHEET,
    RelatoriosGerenciaisOkError,
    gerar_relatorios_gerenciais_ok_provisorios,
)


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_XML = ROOT / "samples" / "nfe_exemplo" / "nfe_oesa_12566_20260505.xml"


EXPECTED_SHEETS = [
    RESUMO_GERAL_SHEET,
    RESUMO_MENSAL_SHEET,
    RESUMO_CONTA_SHEET,
    RESUMO_FORNECEDOR_SHEET,
    RESUMO_CNPJ_SHEET,
    LANCAMENTOS_OK_BASE_SHEET,
    EXCLUIDOS_REVISAR_SHEET,
    EXCLUIDOS_OPERACOES_SHEET,
    INSTRUCOES_SHEET,
]


def test_gera_xlsx_com_abas_e_filtra_somente_ok(tmp_path):
    excel_path = tmp_path / "55746722000181" / "lancamentos.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento("ok-1", valor=Decimal("10.00")),
            _lancamento("revisar-1", valor=Decimal("7.00"), status_revisao="REVISAR"),
        ],
    )
    output_path = tmp_path / "reports" / "relatorios_ok.xlsx"

    result = gerar_relatorios_gerenciais_ok_provisorios([excel_path], output_path)

    workbook = load_workbook(output_path, data_only=True)
    assert result.total_lancamentos_origem == 2
    assert result.lancamentos_ok_incluidos == 1
    assert result.lancamentos_revisar_excluidos == 1
    assert workbook.sheetnames == EXPECTED_SHEETS
    assert workbook[LANCAMENTOS_OK_BASE_SHEET].max_row == 2
    assert workbook[EXCLUIDOS_REVISAR_SHEET].max_row == 2


def test_exclui_comodato_mesmo_com_status_ok(tmp_path):
    xml_dir = _xml_dir_with_special_operation(tmp_path, natureza="Comodato", cfop="5908")
    nfe = parse_nfe(next(xml_dir.glob("*.xml")))
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento("ok-normal", chave_origem="NFE_NORMAL"),
            _lancamento("ok-comodato", chave_origem=nfe.chave),
        ],
    )
    output_path = tmp_path / "reports" / "relatorios_ok.xlsx"

    result = gerar_relatorios_gerenciais_ok_provisorios(
        [excel_path], output_path, pastas_operacoes_especiais=[xml_dir]
    )

    operacoes = _sheet_rows(load_workbook(output_path, data_only=True), EXCLUIDOS_OPERACOES_SHEET)
    assert result.lancamentos_ok_incluidos == 1
    assert result.operacoes_especiais_excluidas == 1
    assert result.lancamentos_ok_excluidos_operacoes == 1
    assert operacoes[0]["tipo_operacao_especial"] == "COMODATO"
    assert operacoes[0]["lancamentos_ok_excluidos"] == 1


def test_inclui_bonificacao_decidida_mesmo_em_revisao(tmp_path):
    xml_dir = _xml_dir_with_special_operation(
        tmp_path, natureza="Bonificacao recebida", cfop="5910"
    )
    nfe = parse_nfe(next(xml_dir.glob("*.xml")))
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento(
                "bonif-decidida",
                status_revisao="REVISAR",
                conta_id="2.05.003",
                conta_nome="Reducao de Custos - Bonificacoes Recebidas",
                chave_origem=nfe.chave,
                tipo_operacao_especial="BONIFICACAO_RECEBIDA",
                tratamento_operacao_especial="REDUCAO_CUSTO",
            )
        ],
    )
    output_path = tmp_path / "reports" / "relatorios_ok.xlsx"

    result = gerar_relatorios_gerenciais_ok_provisorios(
        [excel_path], output_path, pastas_operacoes_especiais=[xml_dir]
    )

    workbook = load_workbook(output_path, data_only=True)
    ok_rows = _sheet_rows(workbook, LANCAMENTOS_OK_BASE_SHEET)
    revisar_rows = _sheet_rows(workbook, EXCLUIDOS_REVISAR_SHEET)
    operacoes = _sheet_rows(workbook, EXCLUIDOS_OPERACOES_SHEET)
    assert result.lancamentos_ok_incluidos == 1
    assert result.lancamentos_revisar_excluidos == 0
    assert ok_rows[0]["conta_id"] == "2.05.003"
    assert ok_rows[0]["status_revisao"] == "OK"
    assert revisar_rows == []
    assert operacoes[0]["tipo_operacao_especial"] == "BONIFICACAO_RECEBIDA"
    assert operacoes[0]["lancamentos_ok_excluidos"] == 0


def test_gera_txt_com_resumo_e_alertas(tmp_path):
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(excel_path, [_lancamento("ok-1")])
    output_path = tmp_path / "reports" / "relatorios_ok.xlsx"
    report_path = tmp_path / "reports" / "relatorios_ok.txt"

    gerar_relatorios_gerenciais_ok_provisorios([excel_path], output_path, report_path)

    text = report_path.read_text(encoding="utf-8")
    assert "RELATORIOS GERENCIAIS OK PROVISORIOS" in text
    assert "Status do relatorio: PROVISORIO" in text
    assert "Nao usar como DRE definitiva" in text
    assert "Resultado:\nOK" in text


def test_calcula_resumos_por_conta_mes_fornecedor_e_cnpj(tmp_path):
    excel_path = tmp_path / "55746722000181" / "lancamentos.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento("ok-1", valor=Decimal("10.00"), fornecedor="Fornecedor A"),
            _lancamento("ok-2", valor=Decimal("15.00"), fornecedor="Fornecedor A"),
        ],
    )
    output_path = tmp_path / "reports" / "relatorios_ok.xlsx"

    gerar_relatorios_gerenciais_ok_provisorios([excel_path], output_path)

    workbook = load_workbook(output_path, data_only=True)
    conta = _sheet_rows(workbook, RESUMO_CONTA_SHEET)[0]
    mensal = _sheet_rows(workbook, RESUMO_MENSAL_SHEET)[0]
    fornecedor = _sheet_rows(workbook, RESUMO_FORNECEDOR_SHEET)[0]
    cnpj = _sheet_rows(workbook, RESUMO_CNPJ_SHEET)[0]
    assert conta["quantidade_lancamentos"] == 2
    assert conta["valor_total"] == 25
    assert mensal["ano_mes"] == "2026-05"
    assert fornecedor["fornecedor"] == "Fornecedor A"
    assert cnpj["identificador_operacional"] == "55746722000181"


def test_retorna_erro_quando_entrada_nao_existe(tmp_path):
    with pytest.raises(RelatoriosGerenciaisOkError) as exc_info:
        gerar_relatorios_gerenciais_ok_provisorios(
            [tmp_path / "inexistente"], tmp_path / "out.xlsx"
        )

    assert "Entrada nao encontrada" in str(exc_info.value)


def _xml_dir_with_special_operation(tmp_path: Path, *, natureza: str, cfop: str) -> Path:
    xml_dir = tmp_path / f"xml_{cfop}"
    xml_dir.mkdir()
    copyfile(SAMPLE_XML, xml_dir / "especial.xml")
    xml_path = xml_dir / "especial.xml"
    text = xml_path.read_text(encoding="utf-8")
    text = text.replace("<natOp>VENDA DE MERCADORIAS</natOp>", f"<natOp>{natureza}</natOp>")
    text = text.replace("<CFOP>5102</CFOP>", f"<CFOP>{cfop}</CFOP>", 1)
    xml_path.write_text(text, encoding="utf-8")
    return xml_dir


def _sheet_rows(workbook, sheet_name: str) -> list[dict[str, object]]:
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    return [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in values)
    ]


def _lancamento(
    hash_registro: str,
    *,
    valor: Decimal = Decimal("10.00"),
    status_revisao: str = "OK",
    fornecedor: str = "Fornecedor teste",
    conta_id: str = "2.01.004",
    conta_nome: str = "Refrigerantes",
    chave_origem: str = "NFE123",
    tipo_operacao_especial: str = "",
    tratamento_operacao_especial: str = "",
) -> Lancamento:
    return Lancamento(
        data_documento=date(2026, 5, 5),
        competencia="2026-05",
        cnpj_empresa="55746722000181",
        nome_empresa="TRATTORIA DO SOSSEGO RESTAURANTE LTDA",
        fornecedor_ou_favorecido=fornecedor,
        cnpj_cpf_favorecido="12345678000195",
        documento="12566",
        tipo_documento="XML_NFE",
        conta_id=conta_id,
        conta_nome=conta_nome,
        grupo_conta="Custos de Mercadorias e Insumos",
        descricao_itens="COCA-COLA KS 290ML",
        valor=valor,
        duplicatas=[
            Duplicata(numero="001", vencimento=date(2026, 5, 20), valor=valor)
        ],
        origem="XML_NFE",
        arquivo_origem="nfe.xml",
        chave_origem=chave_origem,
        hash_registro=hash_registro,
        status_revisao=status_revisao,
        alertas="" if status_revisao == "OK" else "CLASSIFICACAO_NAO_ENCONTRADA",
        tipo_operacao_especial=tipo_operacao_especial,
        tratamento_operacao_especial=tratamento_operacao_especial,
        tipo_lancamento="custo",
    )
