from decimal import Decimal
import json
from pathlib import Path

from openpyxl import load_workbook

from fantozzi.allocation import allocate_nfe_to_lancamentos
from fantozzi.chart_of_accounts import ChartOfAccounts
from fantozzi.excel_writer import write_lancamentos_to_excel
from fantozzi.nfe_importer import parse_nfe
from fantozzi.product_classifier import ProductClassifier


ROOT = Path(__file__).resolve().parents[1]
CONFIG_DIR = ROOT / "config"
SAMPLE_XML = ROOT / "samples" / "nfe_exemplo" / "nfe_oesa_12566_20260505.xml"

REQUIRED_HEADERS = {
    "data_documento",
    "documento",
    "fornecedor_ou_favorecido",
    "conta_id",
    "conta_nome",
    "valor",
    "origem",
    "chave_origem",
    "hash_registro",
    "status_revisao",
    "duplicatas_json",
}


def test_xml_nfe_to_excel_end_to_end(tmp_path):
    nfe = parse_nfe(SAMPLE_XML)

    assert nfe.numero
    assert nfe.valor_total == Decimal("35.00")
    assert len(nfe.produtos) == 3
    assert len(nfe.duplicatas) == 2

    classifier = _product_classifier()
    lancamentos = allocate_nfe_to_lancamentos(nfe, classifier)

    assert len(lancamentos) >= 1
    assert sum((lancamento.valor for lancamento in lancamentos), Decimal("0")) == (
        nfe.valor_total
    )
    assert all(lancamento.conta_id for lancamento in lancamentos)
    assert all(lancamento.hash_registro for lancamento in lancamentos)
    assert all(isinstance(lancamento.valor, Decimal) for lancamento in lancamentos)
    assert _duplicata_totals_by_vencimento(lancamentos) == _original_duplicata_totals(nfe)

    excel_path = tmp_path / "fantozzi_end_to_end.xlsx"
    write_lancamentos_to_excel(excel_path, lancamentos)

    assert excel_path.exists()
    workbook = load_workbook(excel_path)
    assert "Lancamentos" in workbook.sheetnames

    sheet = workbook["Lancamentos"]
    headers = [cell.value for cell in sheet[1]]
    assert REQUIRED_HEADERS.issubset(set(headers))
    assert sheet.max_row - 1 == len(lancamentos)
    assert _excel_total(sheet, headers) == nfe.valor_total

    first_write_rows = sheet.max_row
    write_lancamentos_to_excel(excel_path, lancamentos)
    workbook = load_workbook(excel_path)
    sheet = workbook["Lancamentos"]
    assert sheet.max_row == first_write_rows

    headers = [cell.value for cell in sheet[1]]
    duplicatas = _first_duplicatas_json(sheet, headers)
    assert isinstance(duplicatas, list)
    if duplicatas:
        assert {"numero", "vencimento", "valor"}.issubset(duplicatas[0])

    if any(lancamento.status_revisao != "OK" for lancamento in lancamentos):
        assert "Revisao" in workbook.sheetnames
        assert workbook["Revisao"].max_row >= 2
    else:
        assert "Revisao" in workbook.sheetnames
        assert workbook["Revisao"].max_row == 1


def _product_classifier() -> ProductClassifier:
    chart = ChartOfAccounts.load(CONFIG_DIR / "plano_contas_gerencial.yaml")
    return ProductClassifier(chart).load(CONFIG_DIR / "regras_classificacao_produtos.yaml")


def _duplicata_totals_by_vencimento(lancamentos):
    totals = {}
    for lancamento in lancamentos:
        for duplicata in lancamento.duplicatas:
            totals.setdefault(duplicata.vencimento, Decimal("0"))
            totals[duplicata.vencimento] += duplicata.valor
    return totals


def _original_duplicata_totals(nfe):
    totals = {}
    for duplicata in nfe.duplicatas:
        totals.setdefault(duplicata.vencimento, Decimal("0"))
        totals[duplicata.vencimento] += duplicata.valor
    return totals


def _excel_total(sheet, headers) -> Decimal:
    valor_index = headers.index("valor")
    total = Decimal("0")
    for row in sheet.iter_rows(min_row=2, values_only=True):
        total += Decimal(str(row[valor_index]))
    return total.quantize(Decimal("0.01"))


def _first_duplicatas_json(sheet, headers):
    duplicatas_index = headers.index("duplicatas_json")
    for row in sheet.iter_rows(min_row=2, values_only=True):
        return json.loads(row[duplicatas_index])
    return []
