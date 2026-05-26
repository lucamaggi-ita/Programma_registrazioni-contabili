from datetime import date
from decimal import Decimal
from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook
import pytest

from fantozzi.excel_writer import write_lancamentos_to_excel
from fantozzi.models import Duplicata, Lancamento
from fantozzi.pre_livro_diario import (
    AUDITORIA_SHEET,
    INSTRUCOES_SHEET,
    LANCAMENTOS_OK_SHEET,
    LANCAMENTOS_REVISAR_SHEET,
    OPERACOES_ESPECIAIS_SHEET,
    RESUMO_CONTA_SHEET,
    RESUMO_MENSAL_SHEET,
    PreLivroDiarioError,
    gerar_pre_livro_diario_teste,
)


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_XML = ROOT / "samples" / "nfe_exemplo" / "nfe_oesa_12566_20260505.xml"


def test_gera_pre_livro_diario_com_abas_e_resumos(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"
    output_path = tmp_path / "reports" / "pre_livro.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento("hash-ok", conta_id="2.01.004", valor=Decimal("10.00")),
            _lancamento(
                "hash-revisar",
                conta_id="13.01.001",
                conta_nome="Revisar - Produto Nao Classificado",
                grupo_conta="Contas de Revisao",
                valor=Decimal("7.50"),
                status_revisao="REVISAR",
                alertas="CLASSIFICACAO_NAO_ENCONTRADA",
            ),
        ],
    )
    original_excel = excel_path.read_bytes()

    result = gerar_pre_livro_diario_teste([excel_path], output_path)

    workbook = load_workbook(output_path, data_only=True)
    assert result.lancamentos_ok == 1
    assert result.lancamentos_revisar == 1
    assert result.operacoes_especiais == 0
    assert workbook.sheetnames == [
        LANCAMENTOS_OK_SHEET,
        LANCAMENTOS_REVISAR_SHEET,
        OPERACOES_ESPECIAIS_SHEET,
        RESUMO_CONTA_SHEET,
        RESUMO_MENSAL_SHEET,
        AUDITORIA_SHEET,
        INSTRUCOES_SHEET,
    ]
    assert workbook[LANCAMENTOS_OK_SHEET].max_row == 2
    assert workbook[LANCAMENTOS_REVISAR_SHEET].max_row == 2
    assert workbook[RESUMO_CONTA_SHEET].max_row == 3
    assert workbook[RESUMO_MENSAL_SHEET].max_row == 2
    auditoria = _audit_rows(workbook)
    assert set(auditoria) >= {
        "total_lancamentos_ok",
        "total_lancamentos_revisar",
        "total_lancamentos_geral",
        "total_operacoes_especiais",
        "total_resumo_por_conta",
        "total_resumo_mensal",
        "valor_total_lancamentos_ok",
        "valor_total_lancamentos_revisar",
        "valor_total_lancamentos_geral",
        "valor_total_operacoes_especiais",
        "status_geral_auditoria",
    }
    assert auditoria["total_lancamentos_geral"]["Valor"] == 2
    assert auditoria["total_lancamentos_geral"]["Status"] == "OK"
    assert auditoria["total_lancamentos_revisar"]["Status"] == "ATENCAO"
    assert auditoria["status_geral_auditoria"]["Valor"] == "ATENCAO"
    assert "nao um livro diario definitivo" in workbook[INSTRUCOES_SHEET]["A3"].value
    assert excel_path.read_bytes() == original_excel


def test_inclui_operacoes_especiais_quando_pasta_xml_e_informada(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"
    output_path = tmp_path / "reports" / "pre_livro.xlsx"
    xml_dir = _xml_dir_with_special_operation(tmp_path)
    write_lancamentos_to_excel(excel_path, [_lancamento("hash-ok")])

    result = gerar_pre_livro_diario_teste([excel_path], output_path, [xml_dir])

    sheet = load_workbook(output_path, data_only=True)[OPERACOES_ESPECIAIS_SHEET]
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(sheet.iter_rows(min_row=2, values_only=True))))
    assert result.operacoes_especiais == 1
    assert row["tipo_operacao_especial"] == "COMODATO"
    assert row["destino_gerencial_sugerido"] == "FORA_DRE"
    assert row["status_revisao_sugerido"] == "ANALISAR_COM_CONTADOR"
    auditoria = _audit_rows(load_workbook(output_path, data_only=True))
    assert auditoria["total_operacoes_especiais"]["Status"] == "ATENCAO_CONTABIL"
    assert auditoria["status_geral_auditoria"]["Valor"] == "ATENCAO_CONTABIL"


def test_bonificacao_decidida_entra_em_ok_e_comodato_permanece_revisar(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"
    output_path = tmp_path / "reports" / "pre_livro.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento("hash-ok", conta_id="2.01.004", valor=Decimal("10.00")),
            _lancamento(
                "hash-bonif",
                conta_id="2.05.003",
                conta_nome="Reducao de Custos - Bonificacoes Recebidas",
                valor=Decimal("12.50"),
                status_revisao="REVISAR",
                alertas="OPERACAO_ESPECIAL_DECIDIDA",
                tipo_operacao_especial="BONIFICACAO_RECEBIDA",
                tratamento_operacao_especial="REDUCAO_CUSTO",
            ),
            _lancamento(
                "hash-comodato",
                conta_id="13.01.001",
                conta_nome="Comodato - Fora DRE",
                valor=Decimal("20.00"),
                status_revisao="REVISAR",
                alertas="OPERACAO_ESPECIAL_COMODATO",
                tipo_operacao_especial="COMODATO",
            ),
        ],
    )

    result = gerar_pre_livro_diario_teste([excel_path], output_path)

    workbook = load_workbook(output_path, data_only=True)
    ok_rows = _sheet_rows(workbook, LANCAMENTOS_OK_SHEET)
    revisar_rows = _sheet_rows(workbook, LANCAMENTOS_REVISAR_SHEET)
    resumo_rows = _sheet_rows(workbook, RESUMO_CONTA_SHEET)
    bonif_resumo = next(row for row in resumo_rows if row["conta_id"] == "2.05.003")
    assert result.lancamentos_ok == 2
    assert result.lancamentos_revisar == 1
    bonif_row = next(row for row in ok_rows if row["conta_id"] == "2.05.003")
    assert bonif_row["status_revisao"] == "OK"
    assert all(row["tipo_operacao_especial"] != "COMODATO" for row in ok_rows)
    assert [row["tipo_operacao_especial"] for row in revisar_rows] == ["COMODATO"]
    assert bonif_resumo["qtd_lancamentos_ok"] == 1
    assert bonif_resumo["qtd_lancamentos_revisar"] == 0


def test_salva_relatorio_txt_de_auditoria_quando_informado(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"
    output_path = tmp_path / "reports" / "pre_livro.xlsx"
    report_path = tmp_path / "reports" / "auditoria.txt"
    write_lancamentos_to_excel(excel_path, [_lancamento("hash-ok")])

    gerar_pre_livro_diario_teste([excel_path], output_path, relatorio_auditoria_path=report_path)

    text = report_path.read_text(encoding="utf-8")
    assert "AUDITORIA DO PRE-LIVRO DIARIO" in text
    assert "total_lancamentos_ok;1;OK" in text


def test_retorna_erro_quando_excel_nao_existe(tmp_path):
    output_path = tmp_path / "reports" / "pre_livro.xlsx"

    with pytest.raises(PreLivroDiarioError) as exc_info:
        gerar_pre_livro_diario_teste([tmp_path / "inexistente.xlsx"], output_path)

    assert "Excel de lancamentos nao encontrado" in str(exc_info.value)
    assert not output_path.exists()


def _xml_dir_with_special_operation(tmp_path: Path) -> Path:
    xml_dir = tmp_path / "nfe_xml_especial"
    xml_dir.mkdir()
    copyfile(SAMPLE_XML, xml_dir / "comodato.xml")
    xml_path = xml_dir / "comodato.xml"
    text = xml_path.read_text(encoding="utf-8")
    text = text.replace("<natOp>VENDA DE MERCADORIAS</natOp>", "<natOp>Comodato</natOp>")
    text = text.replace("<CFOP>5102</CFOP>", "<CFOP>5908</CFOP>", 1)
    xml_path.write_text(text, encoding="utf-8")
    return xml_dir


def _audit_rows(workbook) -> dict[str, dict[str, object]]:
    sheet = workbook[AUDITORIA_SHEET]
    headers = [cell.value for cell in sheet[1]]
    rows = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        rows[row["Indicador"]] = row
    return rows


def _sheet_rows(workbook, sheet_name: str) -> list[dict[str, object]]:
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    return [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in values)
    ]


def _lancamento(
    hash_registro: str,
    *,
    conta_id: str = "2.01.004",
    conta_nome: str = "Refrigerantes",
    grupo_conta: str = "Custos de Mercadorias e Insumos",
    valor: Decimal = Decimal("10.00"),
    status_revisao: str = "OK",
    alertas: str = "",
    tipo_operacao_especial: str = "",
    tratamento_operacao_especial: str = "",
) -> Lancamento:
    return Lancamento(
        data_documento=date(2026, 5, 5),
        competencia="2026-05",
        cnpj_empresa="55746722000181",
        nome_empresa="TRATTORIA DO SOSSEGO RESTAURANTE LTDA",
        fornecedor_ou_favorecido="OESA COMERCIO E REPRESENTACOES S/A",
        cnpj_cpf_favorecido="12345678000195",
        documento="12566",
        tipo_documento="XML_NFE",
        conta_id=conta_id,
        conta_nome=conta_nome,
        grupo_conta=grupo_conta,
        descricao_itens="COCA-COLA KS 290ML",
        valor=valor,
        duplicatas=[
            Duplicata(
                numero="001",
                vencimento=date(2026, 5, 20),
                valor=valor,
            )
        ],
        origem="XML_NFE",
        arquivo_origem="nfe.xml",
        chave_origem="NFE123",
        hash_registro=hash_registro,
        status_revisao=status_revisao,
        alertas=alertas,
        tipo_operacao_especial=tipo_operacao_especial,
        tratamento_operacao_especial=tratamento_operacao_especial,
        tipo_lancamento="custo",
    )
