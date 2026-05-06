from pathlib import Path

from openpyxl import Workbook
import pytest

from fantozzi.excel_writer import HEADERS, LANCAMENTOS_SHEET, REVISAO_SHEET
from fantozzi.revision_report import (
    RevisionReportError,
    analyze_revision_report,
    generate_revision_report,
    save_report,
)


def test_valid_excel_without_revisions_generates_ok_report(tmp_path):
    excel_path = tmp_path / "sem_revisao.xlsx"
    _save_workbook(excel_path, lancamentos=[_row()])

    result, text = generate_revision_report(excel_path)

    assert result.lancamentos_em_revisao == 0
    assert "Status: OK" in text
    assert "Lancamentos em revisao: 0" in text
    assert "Nenhum lancamento em revisao encontrado" in text


def test_excel_with_revision_row_generates_review_report(tmp_path):
    excel_path = tmp_path / "com_revisao.xlsx"
    _save_workbook(excel_path, lancamentos=[_row(status_revisao="REVISAR")])

    result, text = generate_revision_report(excel_path)

    assert result.lancamentos_em_revisao == 1
    assert "Status: REVISAR" in text
    assert "Lancamentos em revisao: 1" in text


def test_report_includes_main_revision_row_fields(tmp_path):
    excel_path = tmp_path / "campos_revisao.xlsx"
    _save_workbook(
        excel_path,
        lancamentos=[
            _row(
                data_documento="2026-05-05",
                documento="NF-e 12566",
                fornecedor_ou_favorecido="OESA COMERCIO",
                conta_id="13.01.001",
                conta_nome="Revisar - Produto Nao Classificado",
                descricao_itens="PRODUTO DESCONHECIDO",
                valor=35,
                status_revisao="REVISAR",
                alertas="CLASSIFICACAO_NAO_ENCONTRADA",
            )
        ],
    )

    _, text = generate_revision_report(excel_path)

    assert "Documento: NF-e 12566" in text
    assert "Fornecedor: OESA COMERCIO" in text
    assert "Conta: 13.01.001 - Revisar - Produto Nao Classificado" in text
    assert "Valor: R$ 35,00" in text
    assert "Descricao: PRODUTO DESCONHECIDO" in text
    assert "Status revisao: REVISAR" in text


def test_empty_status_is_treated_as_revision(tmp_path):
    excel_path = tmp_path / "status_vazio.xlsx"
    _save_workbook(excel_path, lancamentos=[_row(status_revisao="")])

    result = analyze_revision_report(excel_path)

    assert result.lancamentos_em_revisao == 1


def test_manual_ok_row_does_not_enter_revision_detail(tmp_path):
    excel_path = tmp_path / "manual_ok.xlsx"
    _save_workbook(excel_path, lancamentos=[_row(origem="MANUAL", status_revisao="OK")])

    result, text = generate_revision_report(excel_path)

    assert result.lancamentos_totais == 1
    assert result.lancamentos_em_revisao == 0
    assert "Linha Excel: 2" not in text


def test_manual_non_ok_row_enters_revision_detail_as_manual(tmp_path):
    excel_path = tmp_path / "manual_revisar.xlsx"
    _save_workbook(
        excel_path,
        lancamentos=[_row(origem="MANUAL", status_revisao="REVISAR")],
    )

    result, text = generate_revision_report(excel_path)

    assert result.lancamentos_em_revisao == 1
    assert "Linha Excel: 2 (linha manual)" in text


def test_missing_excel_file_raises_error(tmp_path):
    missing_excel = tmp_path / "inexistente.xlsx"

    with pytest.raises(RevisionReportError):
        analyze_revision_report(missing_excel)


def test_missing_lancamentos_sheet_raises_error(tmp_path):
    excel_path = tmp_path / "sem_lancamentos.xlsx"
    workbook = Workbook()
    workbook.active.title = REVISAO_SHEET
    workbook.save(excel_path)

    with pytest.raises(RevisionReportError, match="Aba obrigatoria ausente"):
        analyze_revision_report(excel_path)


def test_missing_required_columns_raise_error(tmp_path):
    excel_path = tmp_path / "sem_colunas.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = LANCAMENTOS_SHEET
    sheet.append(["documento", "valor"])
    sheet.append(["NF-1", 10])
    workbook.save(excel_path)

    with pytest.raises(RevisionReportError, match="Colunas obrigatorias ausentes"):
        analyze_revision_report(excel_path)


def test_missing_identification_columns_raise_error(tmp_path):
    excel_path = tmp_path / "sem_identificacao.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = LANCAMENTOS_SHEET
    sheet.append(["status_revisao", "valor"])
    sheet.append(["OK", 10])
    workbook.save(excel_path)

    with pytest.raises(RevisionReportError, match="Colunas de identificacao ausentes"):
        analyze_revision_report(excel_path)


def test_save_report_creates_txt_file(tmp_path):
    report_path = tmp_path / "revisao.txt"

    save_report(report_path, "conteudo")

    assert report_path.exists()
    assert report_path.read_text(encoding="utf-8") == "conteudo"


def test_save_report_creates_missing_directory(tmp_path):
    report_path = tmp_path / "reports" / "revisao.txt"

    save_report(report_path, "conteudo")

    assert report_path.exists()


def _save_workbook(
    excel_path: Path,
    *,
    lancamentos: list[list[object]],
    revisao: list[list[object]] | None = None,
) -> None:
    workbook = Workbook()
    lancamentos_sheet = workbook.active
    lancamentos_sheet.title = LANCAMENTOS_SHEET
    revisao_sheet = workbook.create_sheet(REVISAO_SHEET)

    lancamentos_sheet.append(HEADERS)
    revisao_sheet.append(HEADERS)
    for lancamento in lancamentos:
        lancamentos_sheet.append(lancamento)
    for lancamento in revisao or []:
        revisao_sheet.append(lancamento)

    workbook.save(excel_path)


def _row(**overrides: object) -> list[object]:
    values = {
        "data_documento": "2026-05-05",
        "fornecedor_ou_favorecido": "Fornecedor teste",
        "documento": "NF-1",
        "conta_id": "2.01.004",
        "conta_nome": "Refrigerantes",
        "descricao_itens": "COCA COLA",
        "valor": 10,
        "origem": "XML_NFE",
        "chave_origem": "chave-1",
        "hash_registro": "hash-1",
        "status_revisao": "OK",
        "duplicatas_json": "[]",
    }
    values.update(overrides)
    return [values.get(header) for header in HEADERS]
