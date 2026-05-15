from pathlib import Path

from openpyxl import Workbook
import pytest

from fantozzi.classification_candidates import (
    ClassificationCandidatesError,
    analyze_classification_candidates,
    generate_classification_candidates_report,
    save_report,
)
from fantozzi.excel_writer import HEADERS, LANCAMENTOS_SHEET, REVISAO_SHEET


def test_report_from_minimal_revisao_sheet(tmp_path):
    excel_path = tmp_path / "somente_revisao.xlsx"
    _save_workbook(excel_path, lancamentos=False, revisao_rows=[_row()])

    result, text = generate_classification_candidates_report(excel_path)

    assert result.source_sheet == REVISAO_SHEET
    assert result.total_candidates == 1
    assert "RELATORIO DE CANDIDATOS PARA CLASSIFICACAO" in text
    assert "PRODUTO DESCONHECIDO" in text
    assert "Texto normalizado:" in text
    assert "produto desconhecido" in text


def test_aggregates_by_normalized_description(tmp_path):
    excel_path = tmp_path / "agregado.xlsx"
    _save_workbook(
        excel_path,
        revisao_rows=[
            _row(descricao_itens="Coca-Cola Lata", valor=10, documento="NF-1"),
            _row(descricao_itens="COCA COLA LATA", valor=15, documento="NF-2"),
        ],
    )

    result = analyze_classification_candidates(excel_path)

    assert result.total_candidates == 1
    candidate = result.candidates[0]
    assert candidate.normalized_description == "coca cola lata"
    assert candidate.occurrences == 2
    assert candidate.total_value == 25
    assert candidate.documents == ("NF-1", "NF-2")


def test_orders_by_total_value_occurrences_and_description(tmp_path):
    excel_path = tmp_path / "ordenado.xlsx"
    _save_workbook(
        excel_path,
        revisao_rows=[
            _row(descricao_itens="Produto Baixo", valor=10),
            _row(descricao_itens="Produto Alto", valor=100),
            _row(descricao_itens="Produto Medio", valor=20),
            _row(descricao_itens="Produto Medio", valor=20, documento="NF-2"),
        ],
    )

    result = analyze_classification_candidates(excel_path)

    assert [candidate.normalized_description for candidate in result.candidates] == [
        "produto alto",
        "produto medio",
        "produto baixo",
    ]


def test_report_includes_manual_yaml_draft_only(tmp_path):
    excel_path = tmp_path / "yaml_rascunho.xlsx"
    _save_workbook(excel_path, revisao_rows=[_row(descricao_itens="Produto X")])

    _, text = generate_classification_candidates_report(excel_path)

    assert "Bloco YAML sugerido somente como rascunho:" in text
    assert "PREENCHER_MANUALMENTE" in text
    assert "Nenhuma regra YAML foi alterada" in text


def test_report_includes_fiscal_metadata(tmp_path):
    excel_path = tmp_path / "fiscal.xlsx"
    _save_workbook(
        excel_path,
        revisao_rows=[
            _row(
                cfop="5910",
                natureza_operacao="BONIFICACAO",
                tipo_operacao_especial="BONIFICACAO_RECEBIDA",
                confianca_operacao_especial="FORTE",
                tratamento_operacao_especial="REDUCAO_CUSTO",
            )
        ],
    )

    _, text = generate_classification_candidates_report(excel_path)

    assert "CFOP: 5910" in text
    assert "Natureza operacao: BONIFICACAO" in text
    assert "Tipo operacao especial: BONIFICACAO_RECEBIDA" in text


def test_missing_excel_file_raises_clear_error(tmp_path):
    missing_excel = tmp_path / "inexistente.xlsx"

    with pytest.raises(ClassificationCandidatesError, match="Arquivo Excel nao encontrado"):
        analyze_classification_candidates(missing_excel)


def test_missing_required_sheets_raise_clear_error(tmp_path):
    excel_path = tmp_path / "sem_abas.xlsx"
    workbook = Workbook()
    workbook.active.title = "OutraAba"
    workbook.save(excel_path)

    with pytest.raises(ClassificationCandidatesError, match="Aba obrigatoria ausente"):
        analyze_classification_candidates(excel_path)


def test_missing_value_column_raises_clear_error(tmp_path):
    excel_path = tmp_path / "sem_valor.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = REVISAO_SHEET
    sheet.append(["descricao_itens", "status_revisao"])
    sheet.append(["Produto", "REVISAR"])
    workbook.save(excel_path)

    with pytest.raises(ClassificationCandidatesError, match="Colunas obrigatorias"):
        analyze_classification_candidates(excel_path)


def test_missing_description_columns_raise_clear_error(tmp_path):
    excel_path = tmp_path / "sem_descricao.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = REVISAO_SHEET
    sheet.append(["valor", "status_revisao"])
    sheet.append([10, "REVISAR"])
    workbook.save(excel_path)

    with pytest.raises(ClassificationCandidatesError, match="Colunas de descricao"):
        analyze_classification_candidates(excel_path)


def test_save_report_creates_txt_file(tmp_path):
    report_path = tmp_path / "candidatos.txt"

    save_report(report_path, "conteudo")

    assert report_path.exists()
    assert report_path.read_text(encoding="utf-8") == "conteudo"


def test_save_report_creates_missing_directory(tmp_path):
    report_path = tmp_path / "reports" / "candidatos.txt"

    save_report(report_path, "conteudo")

    assert report_path.exists()


def _save_workbook(
    excel_path: Path,
    *,
    lancamentos: bool = True,
    lancamento_rows: list[list[object]] | None = None,
    revisao_rows: list[list[object]] | None = None,
) -> None:
    workbook = Workbook()
    if lancamentos:
        lancamentos_sheet = workbook.active
        lancamentos_sheet.title = LANCAMENTOS_SHEET
        lancamentos_sheet.append(HEADERS)
        for row in lancamento_rows or []:
            lancamentos_sheet.append(row)
        revisao_sheet = workbook.create_sheet(REVISAO_SHEET)
    else:
        revisao_sheet = workbook.active
        revisao_sheet.title = REVISAO_SHEET

    revisao_sheet.append(HEADERS)
    for row in revisao_rows or []:
        revisao_sheet.append(row)

    workbook.save(excel_path)


def _row(**overrides: object) -> list[object]:
    values = {
        "data_documento": "2026-05-05",
        "fornecedor_ou_favorecido": "Fornecedor teste",
        "documento": "NF-1",
        "conta_id": "13.01.001",
        "conta_nome": "Revisar - Produto Nao Classificado",
        "descricao_itens": "PRODUTO DESCONHECIDO",
        "valor": 10,
        "origem": "XML_NFE",
        "chave_origem": "chave-1",
        "hash_registro": "hash-1",
        "status_revisao": "REVISAR",
        "alertas": "CLASSIFICACAO_NAO_ENCONTRADA",
        "cfop": "",
        "natureza_operacao": "",
        "tipo_operacao_especial": "",
        "confianca_operacao_especial": "",
        "tratamento_operacao_especial": "",
        "duplicatas_json": "[]",
    }
    values.update(overrides)
    return [values.get(header) for header in HEADERS]
