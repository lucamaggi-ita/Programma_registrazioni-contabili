from pathlib import Path

from openpyxl import Workbook

from fantozzi.cli import main
from fantozzi.excel_writer import HEADERS, LANCAMENTOS_SHEET, REVISAO_SHEET


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_XML = ROOT / "samples" / "nfe_exemplo" / "nfe_oesa_12566_20260505.xml"


def test_diagnosticar_excel_returns_success_for_generated_excel(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"

    import_result = main(
        ["importar-nfe", "--xml", str(SAMPLE_XML), "--excel", str(excel_path)]
    )
    diagnostic_result = main(["diagnosticar-excel", "--excel", str(excel_path)])

    assert import_result == 0
    assert diagnostic_result == 0


def test_diagnosticar_excel_returns_error_for_missing_file(tmp_path):
    missing_excel = tmp_path / "inexistente.xlsx"

    result = main(["diagnosticar-excel", "--excel", str(missing_excel)])

    assert result == 1


def test_diagnosticar_excel_returns_error_without_lancamentos_sheet(tmp_path):
    excel_path = tmp_path / "sem_lancamentos.xlsx"
    workbook = Workbook()
    workbook.active.title = REVISAO_SHEET
    workbook.save(excel_path)

    result = main(["diagnosticar-excel", "--excel", str(excel_path)])

    assert result == 1


def test_diagnosticar_excel_returns_error_for_invalid_duplicatas_json(tmp_path):
    excel_path = tmp_path / "duplicatas_invalidas.xlsx"
    _save_workbook(
        excel_path,
        lancamentos=[
            _row(
                hash_registro="hash-1",
                duplicatas_json="{nao-e-json",
            )
        ],
    )

    result = main(["diagnosticar-excel", "--excel", str(excel_path)])

    assert result == 1


def test_diagnosticar_excel_returns_error_for_non_manual_duplicate_hash(tmp_path):
    excel_path = tmp_path / "hash_duplicado.xlsx"
    _save_workbook(
        excel_path,
        lancamentos=[
            _row(hash_registro="hash-duplicado", valor=10),
            _row(hash_registro="hash-duplicado", valor=20),
        ],
    )

    result = main(["diagnosticar-excel", "--excel", str(excel_path)])

    assert result == 1


def test_diagnosticar_excel_allows_duplicate_hash_when_one_row_is_manual(tmp_path):
    excel_path = tmp_path / "hash_manual.xlsx"
    _save_workbook(
        excel_path,
        lancamentos=[
            _row(hash_registro="hash-igual", valor=10),
            _row(origem="MANUAL", hash_registro="hash-igual", valor=20),
        ],
    )

    result = main(["diagnosticar-excel", "--excel", str(excel_path)])

    assert result == 0


def test_diagnosticar_excel_allows_review_row_when_present_in_revisao(tmp_path):
    excel_path = tmp_path / "revisao_ok.xlsx"
    lancamento = _row(status_revisao="REVISAR", hash_registro="hash-revisao")
    _save_workbook(
        excel_path,
        lancamentos=[lancamento],
        revisao=[lancamento],
    )

    result = main(["diagnosticar-excel", "--excel", str(excel_path)])

    assert result == 0


def test_diagnosticar_excel_returns_error_when_review_row_is_missing(tmp_path):
    excel_path = tmp_path / "revisao_ausente.xlsx"
    _save_workbook(
        excel_path,
        lancamentos=[_row(status_revisao="REVISAR", hash_registro="hash-revisao")],
    )

    result = main(["diagnosticar-excel", "--excel", str(excel_path)])

    assert result == 1


def _save_workbook(
    excel_path: Path,
    *,
    lancamentos: list[list[object]],
    revisao: list[list[object]] | None = None,
) -> None:
    workbook = Workbook()
    lancamentos_sheet = workbook.active
    lancamentos_sheet.title = LANCAMENTOS_SHEET
    revisao_sheet = workbook.create_sheet(REVISAO_SHEET)

    lancamentos_sheet.append(HEADERS)
    revisao_sheet.append(HEADERS)
    for lancamento in lancamentos:
        lancamentos_sheet.append(lancamento)
    for lancamento in revisao or []:
        revisao_sheet.append(lancamento)

    workbook.save(excel_path)


def _row(**overrides: object) -> list[object]:
    values = {
        "documento": "NF-1",
        "conta_id": "1.1",
        "conta_nome": "Conta teste",
        "valor": 10,
        "origem": "NFE",
        "chave_origem": "chave-1",
        "hash_registro": "hash-1",
        "status_revisao": "OK",
        "duplicatas_json": "[]",
    }
    values.update(overrides)
    return [values.get(header) for header in HEADERS]
