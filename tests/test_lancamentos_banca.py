from __future__ import annotations

import csv
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from fantozzi.banking.lancamentos_banca import (
    MANIFEST_NAME,
    REJEITADOS_CSV_NAME,
    gerar_lancamentos_banca_aprovados,
)
from fantozzi.banking.revisao_banca import (
    COMPROVANTES_HEADERS,
    COMPROVANTES_SHEET,
    CONTA_CORRENTE_HEADERS,
    CONTA_CORRENTE_SHEET,
    TECNICO_CONTA_CORRENTE_HEADERS,
    TECNICO_CONTA_CORRENTE_SHEET,
)
from fantozzi.cli import build_parser, main
from fantozzi.excel_writer import HEADERS, LANCAMENTOS_SHEET


def test_gerar_lancamentos_banca_filters_decisions_and_writes_standard_output(
    tmp_path,
):
    comprovantes = _write_comprovantes(
        tmp_path / "Analisi_regole_registrazioni_banca_comprovanti.xlsx",
        [
            _comprovante_row("APROVADO", "CLASSIFICAR", "2026-02-05", 2000, "IZALTINO JOSE ZANI JUNIOR"),
            _comprovante_row("PENDENTE", "CLASSIFICAR", "2026-02-05", 10, "PENDENTE LTDA"),
            _comprovante_row("APROVADO", "XML", "2026-02-06", 11, "XML LTDA"),
            _comprovante_row("APROVADO", "ESCLUDERE", "2026-02-07", 12, "ESCLUDERE LTDA"),
            _comprovante_row("APROVADO", "SISPAG", "2026-02-08", 13, "SISPAG LTDA"),
            _comprovante_row("BLOQUEADO", "CLASSIFICAR", "2026-02-09", 14, "BLOQUEADO LTDA"),
            _comprovante_row("REVISAR", "CLASSIFICAR", "2026-02-10", 15, "REVISAR LTDA"),
            _comprovante_row("NAO_CRIAR_REGRA", "CLASSIFICAR", "2026-02-11", 16, "NAO CRIAR LTDA"),
            _comprovante_row("APROVADO", "CLASSIFICAR", "2026-03-01", 17, "FORA PERIODO LTDA"),
            _comprovante_row(
                "APROVADO",
                "CLASSIFICAR",
                "2026-02-12",
                18,
                "RISCO LTDA",
                risco="DUVIDA",
            ),
        ],
    )
    conta_corrente = _write_conta_corrente(
        tmp_path / "Analisi_regole_registrazioni_banca_conto_corrente.xlsx",
        [
            _conta_row(
                "APROVADO",
                "CLASSIFICAR",
                "2026-02-18",
                -3290,
                "PIX ENVIADO",
                "TOP LEVEL DIGITAL",
                "9.01.001",
                "Publicidade e Propaganda",
            ),
            _conta_row(
                "APROVADO",
                "CLASSIFICAR",
                "2026-02-19",
                "-1.234,56",
                "DA COELBA",
                "COELBA",
                "4.01.011",
                "Energia Eletrica",
                risco="OK",
            ),
        ],
    )
    original_comprovantes = comprovantes.read_bytes()
    original_conta = conta_corrente.read_bytes()

    output = gerar_lancamentos_banca_aprovados(
        comprovantes,
        conta_corrente,
        tmp_path / "lancamentos_banca_aprovados.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
    )

    rows = _lancamentos(output)
    assert [row["fornecedor_ou_favorecido"] for row in rows] == [
        "IZALTINO JOSE ZANI JUNIOR",
        "TOP LEVEL DIGITAL",
        "COELBA",
    ]
    assert rows[0]["origem"] == "COMPROVANTES_BANCA"
    assert rows[0]["conta_id"] == "6.01.002"
    assert rows[0]["valor"] == 2000
    assert rows[1]["origem"] == "CONTA_CORRENTE"
    assert rows[1]["conta_id"] == "9.01.001"
    assert rows[1]["valor"] == 3290
    assert rows[2]["valor"] == 1234.56

    workbook = load_workbook(output, data_only=True)
    headers = [cell.value for cell in workbook[LANCAMENTOS_SHEET][1]]
    assert headers[: len(HEADERS)] == HEADERS

    manifest = json.loads((tmp_path / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert manifest["linhas_lidas_comprovantes"] == 10
    assert manifest["linhas_lidas_conta_corrente"] == 2
    assert manifest["linhas_aprovadas_incluidas"] == 3
    assert manifest["linhas_excluidas_xml"] == 1
    assert manifest["linhas_excluidas_escludere"] == 1
    assert manifest["linhas_excluidas_sispag"] == 1
    assert manifest["linhas_rejeitadas_status"] == 4
    assert manifest["linhas_fora_periodo"] == 1
    assert manifest["linhas_rejeitadas_risco_duplicidade"] == 1
    assert (tmp_path / REJEITADOS_CSV_NAME).exists()
    assert comprovantes.read_bytes() == original_comprovantes
    assert conta_corrente.read_bytes() == original_conta


def test_cli_gerar_lancamentos_banca_aprovados(tmp_path):
    comprovantes = _write_comprovantes(
        tmp_path / "comprovantes.xlsx",
        [_comprovante_row("APROVADO", "CLASSIFICAR", "2026-02-05", 100, "M A BRITO COMERCIO LTDA")],
    )
    conta_corrente = _write_conta_corrente(
        tmp_path / "conta.xlsx",
        [_conta_row("APROVADO", "CLASSIFICAR", "2026-02-18", -3290, "PIX", "TOP LEVEL DIGITAL")],
    )
    output = tmp_path / "out.xlsx"

    result = main(
        [
            "gerar-lancamentos-banca-aprovados",
            "--comprovantes-xlsx",
            str(comprovantes),
            "--conta-corrente-xlsx",
            str(conta_corrente),
            "--saida",
            str(output),
            "--periodo-ini",
            "2026-02-01",
            "--periodo-fim",
            "2026-02-28",
        ]
    )

    assert result == 0
    assert output.exists()
    assert len(_lancamentos(output)) == 2


def test_parser_accepts_gerar_lancamentos_banca_aprovados_command():
    args = build_parser().parse_args(
        [
            "gerar-lancamentos-banca-aprovados",
            "--saida",
            "output/lancamentos_banca_aprovados.xlsx",
            "--periodo-ini",
            "2026-02-01",
            "--periodo-fim",
            "2026-02-28",
        ]
    )

    assert args.command == "gerar-lancamentos-banca-aprovados"
    assert args.saida == "output/lancamentos_banca_aprovados.xlsx"


# ---------------------------------------------------------------------------
# Testes: fontes CSV (FZ-GIORNALE-DRE-BRIDGE-BANCA-002B)
# ---------------------------------------------------------------------------


def test_bridge_le_comprovantes_csv_e_gera_lancamentos(tmp_path):
    """comprovantes_classificados.csv → Lancamento com origem COMPROVANTES_CSV."""
    comprovantes_csv = _write_comprovantes_csv(
        tmp_path / "comprovantes_classificados.csv",
        [
            _csv_comprovante_row(
                "FC CONTABILIDADE LTDA", "2026-02-05", "860.00",
                "6.01.001", "Honorarios Contabeis",
            ),
            _csv_comprovante_row(
                "FIBRANET BA TELECOM LTDA", "2026-02-16", "99.99",
                "4.01.014", "Internet e Telefonia",
            ),
        ],
    )
    comprovantes_xlsx, conta_corrente_xlsx = _empty_xlsx_pair(tmp_path)

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        comprovantes_csv=comprovantes_csv,
    )

    rows = _lancamentos(output)
    assert len(rows) == 2
    fornecedores = {r["fornecedor_ou_favorecido"] for r in rows}
    assert "FC CONTABILIDADE LTDA" in fornecedores
    assert "FIBRANET BA TELECOM LTDA" in fornecedores
    fc = next(r for r in rows if "FC CONTABILIDADE" in r["fornecedor_ou_favorecido"])
    assert fc["origem"] == "COMPROVANTES_CSV"
    assert fc["tipo_documento"] == "BANCA_COMPROVANTE_CSV"
    assert fc["conta_id"] == "6.01.001"
    assert fc["valor"] == 860.0

    manifest = json.loads((tmp_path / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert manifest["linhas_lidas_comprovantes_csv"] == 2
    assert manifest["linhas_aprovadas_comprovantes_csv"] == 2


def test_bridge_le_conta_corrente_csv_e_gera_lancamentos(tmp_path):
    """conta_corrente_classificados.csv → Lancamento com origem CONTA_CORRENTE_CSV."""
    conta_corrente_csv = _write_conta_corrente_csv(
        tmp_path / "conta_corrente_classificados.csv",
        [
            _csv_conta_corrente_row(
                "DA COELBA 7014914336", "", "2026-02-23", "-9164.48",
                "4.01.011", "Energia Eletrica",
                regra_id="REGRA_4_01_011_CONTA_CORRENTE_COELBA",
            ),
            _csv_conta_corrente_row(
                "EMBASA 142698660", "", "2026-02-13", "-351.23",
                "4.01.012", "Agua e Esgoto",
                regra_id="REGRA_4_01_012_CONTA_CORRENTE_EMBASA",
            ),
            _csv_conta_corrente_row(
                "DA VIVO-BA 20304230426", "", "2026-02-06", "-380.97",
                "4.01.014", "Internet e Telefonia",
                regra_id="REGRA_4_01_014_CONTA_CORRENTE_VIVO",
            ),
        ],
    )
    comprovantes_xlsx, conta_corrente_xlsx = _empty_xlsx_pair(tmp_path)

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        conta_corrente_csv=conta_corrente_csv,
    )

    rows = _lancamentos(output)
    assert len(rows) == 3
    lancamentos_desc = {r["fornecedor_ou_favorecido"] for r in rows}
    assert "DA COELBA 7014914336" in lancamentos_desc
    assert "EMBASA 142698660" in lancamentos_desc
    assert "DA VIVO-BA 20304230426" in lancamentos_desc
    coelba = next(r for r in rows if "COELBA" in r["fornecedor_ou_favorecido"])
    assert coelba["origem"] == "CONTA_CORRENTE_CSV"
    assert coelba["tipo_documento"] == "BANCA_CONTA_CORRENTE_CSV"
    assert coelba["conta_id"] == "4.01.011"
    assert coelba["valor"] == 9164.48

    manifest = json.loads((tmp_path / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert manifest["linhas_lidas_conta_corrente_csv"] == 3
    assert manifest["linhas_aprovadas_conta_corrente_csv"] == 3


def test_bridge_tre_fonti_insieme(tmp_path):
    """CSV comprovantes + CSV conta corrente + Excel APROVADO → tutte e tre le fonti."""
    comprovantes_csv = _write_comprovantes_csv(
        tmp_path / "comprovantes_classificados.csv",
        [
            _csv_comprovante_row(
                "FC CONTABILIDADE LTDA", "2026-02-05", "860.00",
                "6.01.001", "Honorarios Contabeis",
            ),
        ],
    )
    conta_corrente_csv = _write_conta_corrente_csv(
        tmp_path / "conta_corrente_classificados.csv",
        [
            _csv_conta_corrente_row(
                "DA COELBA 7014914336", "", "2026-02-23", "-9164.48",
                "4.01.011", "Energia Eletrica",
            ),
        ],
    )
    comprovantes_xlsx = _write_comprovantes(
        tmp_path / "Analisi_regole_registrazioni_banca_comprovanti.xlsx",
        [],
    )
    conta_corrente_xlsx = _write_conta_corrente(
        tmp_path / "Analisi_regole_registrazioni_banca_conto_corrente.xlsx",
        [
            _conta_row(
                "APROVADO", "CLASSIFICAR", "2026-02-18", -3290,
                "PIX ENVIADO", "TOP LEVEL DIGITAL",
                "9.01.001", "Publicidade e Propaganda",
            ),
        ],
    )

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        comprovantes_csv=comprovantes_csv,
        conta_corrente_csv=conta_corrente_csv,
    )

    rows = _lancamentos(output)
    assert len(rows) == 3
    origens = {r["origem"] for r in rows}
    assert "COMPROVANTES_CSV" in origens
    assert "CONTA_CORRENTE_CSV" in origens
    assert "CONTA_CORRENTE" in origens

    manifest = json.loads((tmp_path / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert manifest["linhas_aprovadas_comprovantes_csv"] == 1
    assert manifest["linhas_aprovadas_conta_corrente_csv"] == 1
    assert manifest["linhas_aprovadas_incluidas"] == 3


def test_bridge_csv_comprovantes_filtra_fora_periodo(tmp_path):
    """Righe del CSV con data fuori periodo vengono scartate."""
    comprovantes_csv = _write_comprovantes_csv(
        tmp_path / "comprovantes_classificados.csv",
        [
            _csv_comprovante_row(
                "FC CONTABILIDADE LTDA", "2026-02-05", "860.00",
                "6.01.001", "Honorarios Contabeis",
            ),
            _csv_comprovante_row(
                "FUORI PERIODO LTDA", "2026-03-01", "100.00",
                "6.01.001", "Honorarios Contabeis",
            ),
        ],
    )
    comprovantes_xlsx, conta_corrente_xlsx = _empty_xlsx_pair(tmp_path)

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        comprovantes_csv=comprovantes_csv,
    )

    rows = _lancamentos(output)
    assert len(rows) == 1
    assert rows[0]["fornecedor_ou_favorecido"] == "FC CONTABILIDADE LTDA"

    manifest = json.loads((tmp_path / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert manifest["linhas_lidas_comprovantes_csv"] == 2
    assert manifest["linhas_aprovadas_comprovantes_csv"] == 1
    assert manifest["linhas_fora_periodo"] == 1


def test_bridge_csv_comprovantes_filtra_nao_entra_dre(tmp_path):
    """Registros com entra_dre=false (comodato) sao ignorados."""
    comprovantes_csv = _write_comprovantes_csv(
        tmp_path / "comprovantes_classificados.csv",
        [
            _csv_comprovante_row(
                "FC CONTABILIDADE LTDA", "2026-02-05", "860.00",
                "6.01.001", "Honorarios Contabeis",
            ),
            _csv_comprovante_row(
                "COMODATO LTDA", "2026-02-10", "500.00",
                "8.01.001", "Comodato",
                entra_dre="false",
            ),
        ],
    )
    comprovantes_xlsx, conta_corrente_xlsx = _empty_xlsx_pair(tmp_path)

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        comprovantes_csv=comprovantes_csv,
    )

    rows = _lancamentos(output)
    assert len(rows) == 1
    assert rows[0]["fornecedor_ou_favorecido"] == "FC CONTABILIDADE LTDA"


def test_parser_aceita_argumentos_csv(tmp_path):
    """CLI aceita --comprovantes-csv e --conta-corrente-csv."""
    args = build_parser().parse_args(
        [
            "gerar-lancamentos-banca-aprovados",
            "--saida", "out.xlsx",
            "--periodo-ini", "2026-02-01",
            "--periodo-fim", "2026-02-28",
            "--comprovantes-csv", "reports/comprovantes_classificados.csv",
            "--conta-corrente-csv", "bank_raw/_decisoes_regras/conta_corrente_classificados.csv",
        ]
    )

    assert args.comprovantes_csv == "reports/comprovantes_classificados.csv"
    assert args.conta_corrente_csv == "bank_raw/_decisoes_regras/conta_corrente_classificados.csv"


# ---------------------------------------------------------------------------
# Testes: deduplicacao multifonte (FZ-GIORNALE-DRE-BRIDGE-BANCA-002C)
# ---------------------------------------------------------------------------


def test_dedup_conta_corrente_csv_contra_comprovantes_csv(tmp_path):
    comprovantes_csv = _write_comprovantes_csv(
        tmp_path / "comprovantes_classificados.csv",
        [_csv_comprovante_row("FC CONTABILIDADE LTDA", "2026-02-05", "860.00", "6.01.001", "Honorarios")],
    )
    conta_corrente_csv = _write_conta_corrente_csv(
        tmp_path / "conta_corrente_classificados.csv",
        [_csv_conta_corrente_row("PIX FC CONTABILIDADE", "FC CONTABILIDADE LTDA", "2026-02-05", "-860.00", "6.01.001", "Honorarios")],
    )
    comprovantes_xlsx, conta_corrente_xlsx = _empty_xlsx_pair(tmp_path)

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        comprovantes_csv=comprovantes_csv,
        conta_corrente_csv=conta_corrente_csv,
    )

    rows = _lancamentos(output)
    assert len(rows) == 1
    assert rows[0]["origem"] == "COMPROVANTES_CSV"
    rejeitados = _rejeitados(tmp_path)
    dup = next(row for row in rejeitados if row["motivo"] == "DUPLICADO_FONTE_PRIORITARIA")
    assert dup["fonte"] == "CONTA_CORRENTE_CSV"
    assert dup["fonte_prioritaria"] == "COMPROVANTES_CSV"


def test_dedup_excel_aprovado_contra_comprovantes_csv(tmp_path):
    comprovantes_csv = _write_comprovantes_csv(
        tmp_path / "comprovantes_classificados.csv",
        [_csv_comprovante_row("M A BRITO COMERCIO LTDA", "2026-02-05", "100.00", "6.01.001", "Honorarios")],
    )
    comprovantes_xlsx = _write_comprovantes(
        tmp_path / "comprovantes.xlsx",
        [_comprovante_row("APROVADO", "CLASSIFICAR", "2026-02-05", 100, "M A BRITO COMERCIO LTDA")],
    )
    conta_corrente_xlsx = _write_conta_corrente(tmp_path / "conta.xlsx", [])

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        comprovantes_csv=comprovantes_csv,
    )

    rows = _lancamentos(output)
    assert len(rows) == 1
    assert rows[0]["origem"] == "COMPROVANTES_CSV"
    dup = next(row for row in _rejeitados(tmp_path) if row["motivo"] == "DUPLICADO_FONTE_PRIORITARIA")
    assert dup["fonte"] == "COMPROVANTES_BANCA"
    assert dup["fonte_prioritaria"] == "COMPROVANTES_CSV"


def test_dedup_triplo_izaltino_respeita_precedencia(tmp_path):
    comprovantes_csv = _write_comprovantes_csv(
        tmp_path / "comprovantes_classificados.csv",
        [_csv_comprovante_row("IZALTINO JOSE ZANI JUNIOR", "2026-02-05", "2000.00", "6.01.002", "Honorarios")],
    )
    conta_corrente_csv = _write_conta_corrente_csv(
        tmp_path / "conta_corrente_classificados.csv",
        [_csv_conta_corrente_row("PIX IZALTINO", "IZALTINO JOSE ZANI JUNIOR", "2026-02-05", "-2000.00", "6.01.002", "Honorarios")],
    )
    comprovantes_xlsx = _write_comprovantes(
        tmp_path / "comprovantes.xlsx",
        [_comprovante_row("APROVADO", "CLASSIFICAR", "2026-02-05", 2000, "IZALTINO JOSE ZANI JUNIOR")],
    )
    conta_corrente_xlsx = _write_conta_corrente(tmp_path / "conta.xlsx", [])

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        comprovantes_csv=comprovantes_csv,
        conta_corrente_csv=conta_corrente_csv,
    )

    rows = _lancamentos(output)
    assert len(rows) == 1
    assert rows[0]["origem"] == "COMPROVANTES_CSV"
    duplicados = [row for row in _rejeitados(tmp_path) if row["motivo"] == "DUPLICADO_FONTE_PRIORITARIA"]
    assert {row["fonte"] for row in duplicados} == {"COMPROVANTES_BANCA", "CONTA_CORRENTE_CSV"}


def test_top_level_excel_aprovado_prevalece_sobre_conta_corrente_csv(tmp_path):
    comprovantes_xlsx = _write_comprovantes(tmp_path / "comprovantes.xlsx", [])
    conta_corrente_xlsx = _write_conta_corrente(
        tmp_path / "conta.xlsx",
        [_conta_row("APROVADO", "CLASSIFICAR", "2026-02-18", -3290, "PIX ENVIADO", "TOP LEVEL DIGITAL")],
    )
    conta_corrente_csv = _write_conta_corrente_csv(
        tmp_path / "conta_corrente_classificados.csv",
        [_csv_conta_corrente_row("PIX ENVIADO", "TOP LEVEL DIGITAL", "2026-02-18", "-3290.00", "9.01.001", "Publicidade")],
    )

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        conta_corrente_csv=conta_corrente_csv,
    )

    rows = _lancamentos(output)
    assert len(rows) == 1
    assert rows[0]["origem"] == "CONTA_CORRENTE"
    dup = next(row for row in _rejeitados(tmp_path) if row["motivo"] == "DUPLICADO_FONTE_PRIORITARIA")
    assert dup["fonte"] == "CONTA_CORRENTE_CSV"
    assert dup["fonte_prioritaria"] == "CONTA_CORRENTE"


def test_conta_corrente_csv_unica_mantem_coelba_embasa_vivo_da(tmp_path):
    conta_corrente_csv = _write_conta_corrente_csv(
        tmp_path / "conta_corrente_classificados.csv",
        [
            _csv_conta_corrente_row("DA COELBA 7014914336", "", "2026-02-23", "-9164.48", "4.01.011", "Energia"),
            _csv_conta_corrente_row("EMBASA 142698660", "", "2026-02-13", "-351.23", "4.01.012", "Agua"),
            _csv_conta_corrente_row("DA VIVO-BA 20304230426", "", "2026-02-06", "-380.97", "4.01.014", "Telefonia"),
        ],
    )
    comprovantes_xlsx, conta_corrente_xlsx = _empty_xlsx_pair(tmp_path)

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        conta_corrente_csv=conta_corrente_csv,
    )

    fornecedores = {row["fornecedor_ou_favorecido"] for row in _lancamentos(output)}
    assert fornecedores == {"DA COELBA 7014914336", "EMBASA 142698660", "DA VIVO-BA 20304230426"}


def test_data_divergente_um_dia_nao_deduplica_e_sinaliza(tmp_path):
    comprovantes_csv = _write_comprovantes_csv(
        tmp_path / "comprovantes_classificados.csv",
        [_csv_comprovante_row("M A BRITO COMERCIO LTDA", "2026-02-05", "100.00", "6.01.001", "Honorarios")],
    )
    conta_corrente_csv = _write_conta_corrente_csv(
        tmp_path / "conta_corrente_classificados.csv",
        [_csv_conta_corrente_row("PIX M A BRITO", "M A BRITO COMERCIO LTDA", "2026-02-06", "-100.00", "6.01.001", "Honorarios")],
    )
    comprovantes_xlsx, conta_corrente_xlsx = _empty_xlsx_pair(tmp_path)

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        comprovantes_csv=comprovantes_csv,
        conta_corrente_csv=conta_corrente_csv,
    )

    assert len(_lancamentos(output)) == 2
    possivel = next(row for row in _rejeitados(tmp_path) if row["motivo"] == "POSSIVEL_DUPLICADO_DATA_DIVERGENTE")
    assert possivel["incluido"] == "true"
    manifest = json.loads((tmp_path / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert manifest["deduplicacao"]["possiveis_duplicados_data_divergente"] == 1


def test_mesma_data_mesmo_valor_beneficiarios_diferentes_nao_deduplica(tmp_path):
    comprovantes_csv = _write_comprovantes_csv(
        tmp_path / "comprovantes_classificados.csv",
        [
            _csv_comprovante_row("FORNECEDOR A LTDA", "2026-02-05", "100.00", "6.01.001", "Honorarios"),
            _csv_comprovante_row("FORNECEDOR B LTDA", "2026-02-05", "100.00", "6.01.001", "Honorarios"),
        ],
    )
    comprovantes_xlsx, conta_corrente_xlsx = _empty_xlsx_pair(tmp_path)

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        comprovantes_csv=comprovantes_csv,
    )

    assert len(_lancamentos(output)) == 2


def test_conta_corrente_csv_entrada_positiva_nao_entra(tmp_path):
    conta_corrente_csv = _write_conta_corrente_csv(
        tmp_path / "conta_corrente_classificados.csv",
        [_csv_conta_corrente_row("PIX RECEBIDO", "CLIENTE LTDA", "2026-02-05", "100.00", "4.01.011", "Energia")],
    )
    comprovantes_xlsx, conta_corrente_xlsx = _empty_xlsx_pair(tmp_path)

    output = gerar_lancamentos_banca_aprovados(
        comprovantes_xlsx,
        conta_corrente_xlsx,
        tmp_path / "out.xlsx",
        date(2026, 2, 1),
        date(2026, 2, 28),
        conta_corrente_csv=conta_corrente_csv,
    )

    assert _lancamentos(output) == []
    rejeitado = _rejeitados(tmp_path)[0]
    assert rejeitado["motivo"] == "ACAO_ESCLUIDA"


def _write_comprovantes_csv(path: Path, rows: list[dict[str, object]]) -> Path:
    headers = [
        "cnpj", "mes", "data_pagamento", "beneficiario_normalizado", "referencia",
        "valor", "arquivo_origem", "status_classificacao", "conta_id", "conta_nome",
        "dre_linha_id", "entra_dre", "regra_id", "motivo_classificacao",
        "regras_candidatas", "payment_id", "beneficiary_name", "beneficiary_document",
        "payment_type", "status_pagamento", "source_file", "raw_hash",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
    return path


def _write_conta_corrente_csv(path: Path, rows: list[dict[str, object]]) -> Path:
    headers = [
        "cnpj", "mes", "data_movimento", "lancamento", "razao_social",
        "valor", "conta_id", "conta_nome", "regra_id", "chave_estavel",
        "hash_origem", "arquivo_origem", "tipo_movimento",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
    return path


def _csv_comprovante_row(
    beneficiario: str,
    data_pagamento: str,
    valor: str,
    conta_id: str,
    conta_nome: str,
    *,
    entra_dre: str = "true",
    regra_id: str = "REGRA_TEST",
) -> dict[str, object]:
    return {
        "cnpj": "08884180000123",
        "mes": data_pagamento[:7],
        "data_pagamento": data_pagamento,
        "beneficiario_normalizado": beneficiario,
        "referencia": "-",
        "valor": valor,
        "arquivo_origem": "ConsultaPagamentos.xls",
        "status_classificacao": "CLASSIFICADO_POR_REGRA",
        "conta_id": conta_id,
        "conta_nome": conta_nome,
        "dre_linha_id": "06.01",
        "entra_dre": entra_dre,
        "regra_id": regra_id,
        "motivo_classificacao": "REGRA_UNICA_EXPLICITA",
        "regras_candidatas": regra_id,
        "payment_id": f"pid_{beneficiario[:6]}",
        "beneficiary_name": beneficiario,
        "beneficiary_document": "",
        "payment_type": "PIX Transferencias",
        "status_pagamento": "Efetuado",
        "source_file": "ConsultaPagamentos.xls",
        "raw_hash": f"hash_{beneficiario[:6]}_{data_pagamento}",
    }


def _csv_conta_corrente_row(
    lancamento: str,
    razao_social: str,
    data_movimento: str,
    valor: str,
    conta_id: str,
    conta_nome: str,
    *,
    regra_id: str = "REGRA_TEST",
) -> dict[str, object]:
    return {
        "cnpj": "08884180000123",
        "mes": data_movimento[:7],
        "data_movimento": data_movimento,
        "lancamento": lancamento,
        "razao_social": razao_social,
        "valor": valor,
        "conta_id": conta_id,
        "conta_nome": conta_nome,
        "regra_id": regra_id,
        "chave_estavel": f"CONTA_CORRENTE|HISTORICO={lancamento}|DIRECAO=SAIDA",
        "hash_origem": f"hash_{lancamento[:8]}",
        "arquivo_origem": "Extrato_Lancamentos.xlsx",
        "tipo_movimento": "SAIDA_PAGAMENTO",
    }


def _empty_xlsx_pair(tmp_path: Path) -> tuple[Path, Path]:
    comprovantes = _write_comprovantes(
        tmp_path / "Analisi_regole_registrazioni_banca_comprovanti.xlsx", []
    )
    conta = _write_conta_corrente(
        tmp_path / "Analisi_regole_registrazioni_banca_conto_corrente.xlsx", []
    )
    return comprovantes, conta


def _write_comprovantes(path: Path, rows: list[dict[str, object]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = COMPROVANTES_SHEET
    sheet.append(COMPROVANTES_HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in COMPROVANTES_HEADERS])
    workbook.save(path)
    return path


def _write_conta_corrente(path: Path, rows: list[dict[str, object]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = CONTA_CORRENTE_SHEET
    sheet.append(CONTA_CORRENTE_HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in CONTA_CORRENTE_HEADERS])
    tecnico = workbook.create_sheet(TECNICO_CONTA_CORRENTE_SHEET)
    tecnico.append(TECNICO_CONTA_CORRENTE_HEADERS)
    for row in rows:
        tecnico.append(
            [
                _tecnico_value(row, header)
                for header in TECNICO_CONTA_CORRENTE_HEADERS
            ]
        )
    workbook.save(path)
    return path


def _comprovante_row(
    status: str,
    acao: str,
    data_pagamento: str,
    valor: object,
    beneficiario: str,
    *,
    risco: str = "OK",
) -> dict[str, object]:
    return {
        "STATUS_DECISAO": status,
        "ACAO_REGRA": acao,
        "PERIODO": "2026-02",
        "CNPJ": "08884180000123",
        "DATA_PAGAMENTO": data_pagamento,
        "VALOR": valor,
        "BENEFICIARIO_FAVORECIDO": beneficiario,
        "REFERENCIA_EMPRESA": "REFERENCIA",
        "TIPO_PAGAMENTO": "PIX",
        "CONTA_ID_APROVADA": "6.01.002",
        "CONTA_NOME_APROVADA": "Honorarios Contabeis",
        "CHAVE_REGRA_BENEFICIARIO": beneficiario,
        "CHAVE_REGRA_REFERENCIA": "REFERENCIA",
        "RISCO_DUPLICIDADE_XML": risco,
        "ESCOPO_REGRA": "COMPROVANTE",
    }


def _conta_row(
    status: str,
    acao: str,
    data_movimento: str,
    valor: object,
    lancamento: str,
    razao_social: str,
    conta_id: str = "9.01.001",
    conta_nome: str = "Publicidade e Propaganda",
    *,
    risco: str = "",
) -> dict[str, object]:
    return {
        "STATUS_DECISAO": status,
        "ACAO_REGRA": acao,
        "DATA": data_movimento,
        "VALOR": valor,
        "LANCAMENTO": lancamento,
        "RAZAO_SOCIAL": razao_social,
        "CONTA_ID_APROVADA": conta_id,
        "CONTA_NOME_APROVADA": conta_nome,
        "CHAVE_REGRA_LANCAMENTO": lancamento,
        "CHAVE_REGRA_RAZAO_SOCIAL": razao_social,
        "RISCO_DUPLICIDADE_XML": risco,
        "ESCOPO_REGRA": "CONTA_CORRENTE",
    }


def _tecnico_value(row: dict[str, object], header: str) -> object:
    if header == "CHAVE_ESTAVEL":
        return (
            f"CONTA_CORRENTE|HISTORICO={row['LANCAMENTO']}|"
            f"RAZAO={row['RAZAO_SOCIAL']}|TIPO=PIX_TRANSFERENCIA_SAIDA|"
            "DIRECAO=SAIDA"
        )
    if header == "CNPJ":
        return "08884180000123"
    return ""


def _lancamentos(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[LANCAMENTOS_SHEET]
    headers = [cell.value for cell in sheet[1]]
    return [
        {
            headers[index]: value
            for index, value in enumerate(row)
            if index < len(headers)
        }
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]


def _rejeitados(tmp_path: Path) -> list[dict[str, str]]:
    with (tmp_path / REJEITADOS_CSV_NAME).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))
