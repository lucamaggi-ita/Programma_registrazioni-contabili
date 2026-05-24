from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fantozzi.cli import main
from fantozzi.lancamentos_manuais_gerenciais import (
    FONTE_MANUAL_GERENCIAL,
    carregar_lancamentos_manuais_gerenciais,
    converter_lancamentos_manuais_para_lancamentos,
)
from fantozzi.lancamentos_manuais_template import (
    LancamentosManuaisTemplateError,
    gerar_template_lancamentos_manuais_gerenciais,
    validar_lancamentos_manuais_gerenciais_operador,
)


OPERATIVE_HEADERS = [
    "DATA",
    "HISTORICO",
    "DOCUMENTO",
    "FORNECEDOR_OU_CONTRAPARTE",
    "CATEGORIA_GERENCIAL",
    "VALOR",
    "CNPJ_ORIGEM",
    "CENTRO_CUSTO",
    "OBSERVACAO",
]

TECHNICAL_ACCOUNT_ID = "99.99.998"
TECHNICAL_ACCOUNT_NAME = "Contrapartida Tecnica Lancamentos Manuais"


def test_gera_template_com_abas_e_colunas_operacionais(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    template_path = tmp_path / "template.xlsx"

    result = gerar_template_lancamentos_manuais_gerenciais(
        template_path,
        plano_path,
        dre_path,
    )

    assert result.blockers == ()
    assert [item.categoria for item in result.categorias] == ["Manutencao Predial"]
    assert result.conta_credito_tecnica is not None
    assert result.conta_credito_tecnica.conta_id == TECHNICAL_ACCOUNT_ID
    workbook = load_workbook(template_path, data_only=False)
    assert workbook.sheetnames == [
        "lancamentos_manuais",
        "categorias",
        "plano_contas_operativo",
        "instrucoes",
    ]
    sheet = workbook["lancamentos_manuais"]
    assert [cell.value for cell in sheet[1]][:9] == OPERATIVE_HEADERS
    assert "FORMA_PAGAMENTO" not in [cell.value for cell in sheet[1]]
    assert sheet.column_dimensions["J"].hidden is True
    assert sheet.column_dimensions["P"].hidden is True


def test_template_usa_conta_tecnica_sem_cassa_banca_cartao(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    result = gerar_template_lancamentos_manuais_gerenciais(
        tmp_path / "template.xlsx",
        plano_path,
        dre_path,
    )

    assert result.blockers == ()
    assert result.conta_credito_tecnica is not None
    assert result.conta_credito_tecnica.conta_nome == TECHNICAL_ACCOUNT_NAME


def test_template_bloqueia_sem_conta_tecnica_neutra(tmp_path):
    plano_path, dre_path = _save_config(tmp_path, include_technical=False)
    result = gerar_template_lancamentos_manuais_gerenciais(
        tmp_path / "template.xlsx",
        plano_path,
        dre_path,
    )

    assert result.conta_credito_tecnica is None
    assert any("Mini-fase necessaria" in item for item in result.blockers)


def test_valida_xlsx_e_gera_csv_tecnico_compativel(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    template_path = tmp_path / "template.xlsx"
    output_csv = tmp_path / "out" / "manual_tecnico.csv"
    report_path = tmp_path / "reports" / "validacao.txt"
    gerar_template_lancamentos_manuais_gerenciais(template_path, plano_path, dre_path)
    _fill_template_row(template_path)

    result = validar_lancamentos_manuais_gerenciais_operador(
        template_path,
        output_csv,
        report_path,
        plano_path,
        dre_path,
    )

    assert result.linhas_lidas == 1
    assert result.linhas_validas == 1
    assert output_csv.read_bytes().startswith(b"\xef\xbb\xbf")
    with output_csv.open("r", encoding="utf-8-sig", newline="") as file:
        row = next(csv.DictReader(file, delimiter=";"))
    assert row["conta_debito_id"] == "8.01.001"
    assert row["conta_credito_id"] == TECHNICAL_ACCOUNT_ID
    assert row["conta_credito_nome"] == TECHNICAL_ACCOUNT_NAME
    loaded = carregar_lancamentos_manuais_gerenciais(output_csv, plano_path, dre_path)
    converted = converter_lancamentos_manuais_para_lancamentos(
        loaded,
        plano_path,
        output_csv,
    )
    assert len(converted) == 2
    assert {item.origem for item in converted} == {FONTE_MANUAL_GERENCIAL}
    assert converted[0].tipo_lancamento.endswith("_TOTALIZAVEL")
    assert not converted[1].tipo_lancamento.endswith("_TOTALIZAVEL")
    assert "nao garante que esta despesa ja nao exista em XML" in report_path.read_text(
        encoding="utf-8"
    )


def test_validador_bloqueia_categoria_inexistente_e_salva_relatorio(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    template_path = tmp_path / "template.xlsx"
    report_path = tmp_path / "reports" / "validacao.txt"
    gerar_template_lancamentos_manuais_gerenciais(template_path, plano_path, dre_path)
    _fill_template_row(template_path, CATEGORIA_GERENCIAL="Categoria Fantasma")

    with pytest.raises(LancamentosManuaisTemplateError, match="erro"):
        validar_lancamentos_manuais_gerenciais_operador(
            template_path,
            tmp_path / "manual_tecnico.csv",
            report_path,
            plano_path,
            dre_path,
        )

    assert "Categoria Fantasma" in report_path.read_text(encoding="utf-8")


def test_validador_bloqueia_valor_zero(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    template_path = tmp_path / "template.xlsx"
    gerar_template_lancamentos_manuais_gerenciais(template_path, plano_path, dre_path)
    _fill_template_row(template_path, VALOR="0")

    with pytest.raises(LancamentosManuaisTemplateError, match="erro"):
        validar_lancamentos_manuais_gerenciais_operador(
            template_path,
            tmp_path / "manual_tecnico.csv",
            None,
            plano_path,
            dre_path,
        )


def test_validador_bloqueia_linha_vazia_no_meio(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    template_path = tmp_path / "template.xlsx"
    gerar_template_lancamentos_manuais_gerenciais(template_path, plano_path, dre_path)
    _fill_template_row(template_path)
    _fill_template_row(template_path, row_index=4, DOCUMENTO="MAN-002")

    with pytest.raises(LancamentosManuaisTemplateError, match="erro"):
        validar_lancamentos_manuais_gerenciais_operador(
            template_path,
            tmp_path / "manual_tecnico.csv",
            None,
            plano_path,
            dre_path,
        )


def test_validador_csv_operacional_emite_avisos_sem_bloquear(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    csv_path = tmp_path / "operador.csv"
    output_csv = tmp_path / "manual_tecnico.csv"
    row = _operator_row(DOCUMENTO="", CNPJ_ORIGEM="", OBSERVACAO="")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=OPERATIVE_HEADERS, delimiter=";")
        writer.writeheader()
        writer.writerow(row)

    result = validar_lancamentos_manuais_gerenciais_operador(
        csv_path,
        output_csv,
        None,
        plano_path,
        dre_path,
    )

    assert output_csv.exists()
    assert len(result.avisos) >= 3


def test_gerar_template_retorna_exit_code_1_com_blockers(tmp_path):
    plano_path, dre_path = _save_config(tmp_path, include_technical=False)
    template_path = tmp_path / "template_blocker.xlsx"

    exit_code = main([
        "gerar-template-lancamentos-manuais-gerenciais",
        "--saida", str(template_path),
        "--plano-contas", str(plano_path),
        "--mapeamento-dre", str(dre_path),
    ])

    assert exit_code == 1
    assert template_path.exists()


def test_gerar_template_retorna_exit_code_0_sem_blockers(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    template_path = tmp_path / "template_ok.xlsx"

    exit_code = main([
        "gerar-template-lancamentos-manuais-gerenciais",
        "--saida", str(template_path),
        "--plano-contas", str(plano_path),
        "--mapeamento-dre", str(dre_path),
    ])

    assert exit_code == 0
    assert template_path.exists()


def test_validador_aviso_anti_xml_nao_multiplica_por_linha(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    template_path = tmp_path / "template.xlsx"
    output_csv = tmp_path / "manual_tecnico.csv"
    gerar_template_lancamentos_manuais_gerenciais(template_path, plano_path, dre_path)
    _fill_template_row(template_path, row_index=2)
    _fill_template_row(template_path, row_index=3, DOCUMENTO="MAN-002")

    result = validar_lancamentos_manuais_gerenciais_operador(
        template_path,
        output_csv,
        None,
        plano_path,
        dre_path,
    )

    assert result.linhas_validas == 2
    xml_avisos = [w for w in result.avisos if "XML" in w]
    assert xml_avisos == []


def _fill_template_row(path: Path, row_index: int = 2, **overrides: str) -> None:
    workbook = load_workbook(path)
    sheet = workbook["lancamentos_manuais"]
    row = _operator_row(**overrides)
    for column_index, header in enumerate(OPERATIVE_HEADERS, start=1):
        sheet.cell(row=row_index, column=column_index).value = row[header]
    workbook.save(path)


def _operator_row(**overrides: str) -> dict[str, str]:
    row = {
        "DATA": "2026-05-20",
        "HISTORICO": "Conserto pia cozinha",
        "DOCUMENTO": "MAN-001",
        "FORNECEDOR_OU_CONTRAPARTE": "Fornecedor Teste",
        "CATEGORIA_GERENCIAL": "Manutencao Predial",
        "VALOR": "10,50",
        "CNPJ_ORIGEM": "08884180000123",
        "CENTRO_CUSTO": "OPERACAO",
        "OBSERVACAO": "despesa manual de teste",
    }
    row.update(overrides)
    return row


def _save_config(tmp_path: Path, include_technical: bool = True) -> tuple[Path, Path]:
    plano_path = tmp_path / "plano_contas.yaml"
    dre_path = tmp_path / "mapeamento_dre.yaml"
    technical_line = (
        f'  - {{id: "{TECHNICAL_ACCOUNT_ID}", nome: "{TECHNICAL_ACCOUNT_NAME}", grupo: "Controles Tecnicos", tipo: tecnico}}\n'
        if include_technical
        else ""
    )
    plano_path.write_text(
        (
            "contas:\n"
            '  - {id: "8.01.001", nome: "Manutencao Predial", grupo: "Despesas Operacionais", tipo: despesa}\n'
            f"{technical_line}"
            '  - {id: "13.01.001", nome: "A Revisar", grupo: "Revisao", tipo: revisao}\n'
            '  - {id: "13.01.002", nome: "Ambiguo", grupo: "Revisao", tipo: revisao}\n'
            '  - {id: "13.01.006", nome: "Valor Inconsistente", grupo: "Revisao", tipo: revisao}\n'
            '  - {id: "13.01.009", nome: "Possivel Investimento", grupo: "Revisao", tipo: revisao}\n'
            '  - {id: "13.01.010", nome: "Pessoal a Regularizar", grupo: "Revisao", tipo: revisao}\n'
            '  - {id: "13.01.999", nome: "Outros", grupo: "Revisao", tipo: revisao}\n'
        ),
        encoding="utf-8",
    )
    dre_path.write_text(
        """
linhas_dre:
  - {id: "08.01", nome: "Ocupacao", grupo: "Despesas", sinal: "-", ordem: 1}
mapeamento:
  "8.01.001": {dre_linha_id: "08.01", entra_dre: true}
  "99.99.998": {dre_linha_id: "", entra_dre: false}
""".strip()
        + "\n",
        encoding="utf-8",
    )
    return plano_path, dre_path
