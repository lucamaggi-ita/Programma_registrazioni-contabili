from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from fantozzi.cli import main
from fantozzi.dre_pdf_gerencial import (
    DrePdfGerencialError,
    DrePdfResult,
    _find_dre_xlsx,
    _fmt_brl,
    _load_dre_geral,
    _load_periodos,
    gerar_pdf_dre,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_dre_xlsx(tmp_path: Path, rows: list[tuple], mensal_rows: list[tuple] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "DRE_Geral"
    ws.append(["ordem", "grupo_dre", "linha_dre", "valor", "observacao"])
    for row in rows:
        ws.append(list(row))

    ws2 = wb.create_sheet("DRE_Mensal")
    ws2.append(["ano_mes", "ordem", "grupo_dre", "linha_dre", "valor"])
    for mrow in (mensal_rows or []):
        ws2.append(list(mrow))

    p = tmp_path / "dre_gerencial_ok_provisoria.xlsx"
    wb.save(str(p))
    return p


_FIXTURE_ROWS = [
    (1, "Custos", "CMV - Alimentos", 1000.50, "PROVISORIO"),
    (2, "Custos", "CMV - Bebidas", 500.00, "PROVISORIO"),
    (3, "Despesas Operacionais", "Pessoal CLT", 2000.00, "PROVISORIO"),
    (4, "Despesas Operacionais", "Aluguel", 1500.00, "PROVISORIO"),
]

_FIXTURE_MENSAL = [
    ("2026-01", 1, "Custos", "CMV - Alimentos", 1000.50),
    ("2026-01", 3, "Despesas Operacionais", "Pessoal CLT", 2000.00),
    ("2026-02", 1, "Custos", "CMV - Alimentos", 800.00),
]


# ---------------------------------------------------------------------------
# Testes _fmt_brl
# ---------------------------------------------------------------------------

def test_fmt_brl_positivo():
    assert _fmt_brl(1234.56) == "R$ 1.234,56"


def test_fmt_brl_zero():
    assert _fmt_brl(0) == "R$ 0,00"


def test_fmt_brl_negativo():
    assert _fmt_brl(-500.0) == "R$ -500,00"


def test_fmt_brl_decimal():
    assert _fmt_brl(Decimal("999999.99")) == "R$ 999.999,99"


def test_fmt_brl_grande():
    assert _fmt_brl(1000000.00) == "R$ 1.000.000,00"


def test_fmt_brl_none():
    result = _fmt_brl(None)
    assert result == "R$ 0,00"


# ---------------------------------------------------------------------------
# Testes _load_dre_geral
# ---------------------------------------------------------------------------

def test_load_dre_geral_retorna_linhas(tmp_path):
    xlsx = _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    rows = _load_dre_geral(xlsx)
    assert len(rows) == 4
    assert rows[0]["grupo_dre"] == "Custos"
    assert rows[0]["linha_dre"] == "CMV - Alimentos"
    assert float(rows[0]["valor"]) == pytest.approx(1000.50)


def test_load_dre_geral_aba_ausente_levanta_erro(tmp_path):
    wb = Workbook()
    wb.active.title = "OutraAba"
    p = tmp_path / "sem_dre.xlsx"
    wb.save(str(p))
    with pytest.raises(DrePdfGerencialError, match="DRE_Geral"):
        _load_dre_geral(p)


def test_load_dre_geral_arquivo_inexistente_levanta_erro(tmp_path):
    with pytest.raises(DrePdfGerencialError):
        _load_dre_geral(tmp_path / "nao_existe.xlsx")


# ---------------------------------------------------------------------------
# Testes _load_periodos
# ---------------------------------------------------------------------------

def test_load_periodos_retorna_ordenados(tmp_path):
    xlsx = _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL)
    periodos = _load_periodos(xlsx)
    assert periodos == ["2026-01", "2026-02"]


def test_load_periodos_sem_aba_mensal(tmp_path):
    wb = Workbook()
    wb.active.title = "DRE_Geral"
    p = tmp_path / "sem_mensal.xlsx"
    wb.save(str(p))
    assert _load_periodos(p) == []


# ---------------------------------------------------------------------------
# Testes _find_dre_xlsx
# ---------------------------------------------------------------------------

def test_find_dre_xlsx_encontra_provisoria(tmp_path):
    xlsx = _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    found = _find_dre_xlsx(tmp_path)
    assert found == xlsx


def test_find_dre_xlsx_encontra_pacote(tmp_path):
    pacote = tmp_path / "pacote_contabil_gerencial_provisorio"
    pacote.mkdir()
    target = pacote / "01_DRE_GERENCIAL_CONSOLIDADO.xlsx"
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)  # cria o fallback
    # copia para o pacote
    import shutil
    shutil.copy(tmp_path / "dre_gerencial_ok_provisoria.xlsx", target)
    # pacote deve ter prioridade
    found = _find_dre_xlsx(tmp_path)
    assert found == target


def test_find_dre_xlsx_ausente_levanta_erro(tmp_path):
    with pytest.raises(DrePdfGerencialError, match="Nenhum arquivo DRE"):
        _find_dre_xlsx(tmp_path)


# ---------------------------------------------------------------------------
# Testes gerar_pdf_dre
# ---------------------------------------------------------------------------

def test_gerar_pdf_dre_cria_arquivo_com_nome_correto(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL)
    result = gerar_pdf_dre(tmp_path, tmp_path / "out")
    assert result.pdf_path.name == "01_DRE_GERENCIAL.pdf"
    assert result.pdf_path.exists()


def test_gerar_pdf_dre_header_pdf_valido(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_pdf_dre(tmp_path, tmp_path / "out")
    header = result.pdf_path.read_bytes()[:4]
    assert header == b"%PDF"


def test_gerar_pdf_dre_tamanho_minimo(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL)
    result = gerar_pdf_dre(tmp_path, tmp_path / "out")
    assert result.pdf_path.stat().st_size > 5000


def test_gerar_pdf_dre_status_teste_tecnico(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_pdf_dre(tmp_path, tmp_path / "out")
    assert result.status == "TESTE_TECNICO_NAO_OFICIAL"


def test_gerar_pdf_dre_total_linhas_correto(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_pdf_dre(tmp_path, tmp_path / "out")
    assert result.total_linhas == 4


def test_gerar_pdf_dre_periodos_corretos(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL)
    result = gerar_pdf_dre(tmp_path, tmp_path / "out")
    assert result.periodos == ("2026-01", "2026-02")


def test_gerar_pdf_dre_output_dir_default(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_pdf_dre(tmp_path)
    assert result.pdf_path.parent == tmp_path / "output_visual"


def test_gerar_pdf_dre_run_dir_inexistente_levanta_erro(tmp_path):
    with pytest.raises(DrePdfGerencialError, match="run-dir nao encontrado"):
        gerar_pdf_dre(tmp_path / "nao_existe")


def test_gerar_pdf_dre_sem_dre_xlsx_levanta_erro(tmp_path):
    with pytest.raises(DrePdfGerencialError):
        gerar_pdf_dre(tmp_path, tmp_path / "out")


def test_gerar_pdf_dre_le_run_id_do_manifest(tmp_path):
    import json
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    manifest = {"run_id": "2026-05-29_001", "status_run": "PROVISORIO_CONCLUIDO"}
    (tmp_path / "manifest_run.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    result = gerar_pdf_dre(tmp_path, tmp_path / "out")
    assert result.run_id == "2026-05-29_001"


# ---------------------------------------------------------------------------
# Testes CLI
# ---------------------------------------------------------------------------

def test_cli_gerar_pdf_dre_ok(tmp_path, capsys):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL)
    out_dir = tmp_path / "out"
    result = main(
        ["gerar-pdf-dre", "--run-dir", str(tmp_path), "--output-dir", str(out_dir)]
    )
    captured = capsys.readouterr()
    assert result == 0
    assert "[OK] PDF DRE gerencial gerado." in captured.out
    assert "TESTE_TECNICO_NAO_OFICIAL" in captured.out
    assert (out_dir / "01_DRE_GERENCIAL.pdf").exists()


def test_cli_gerar_pdf_dre_run_dir_inexistente(tmp_path, capsys):
    result = main(
        ["gerar-pdf-dre", "--run-dir", str(tmp_path / "nao_existe")]
    )
    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO]" in captured.out


def test_cli_gerar_pdf_dre_sem_dre(tmp_path, capsys):
    result = main(["gerar-pdf-dre", "--run-dir", str(tmp_path)])
    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO]" in captured.out


def test_cli_gerar_pdf_dre_output_dir_default(tmp_path, capsys):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = main(["gerar-pdf-dre", "--run-dir", str(tmp_path)])
    assert result == 0
    assert (tmp_path / "output_visual" / "01_DRE_GERENCIAL.pdf").exists()
