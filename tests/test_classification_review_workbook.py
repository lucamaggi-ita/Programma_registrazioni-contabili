from pathlib import Path

from openpyxl import load_workbook

from fantozzi.classification_review_export import CSV_HEADERS
from fantozzi.classification_review_workbook import (
    ACOES_VALIDAS_SHEET,
    EXEMPLOS_SHEET,
    GUIA_ALIAS_REGRAS_SHEET,
    INSTRUCOES_SHEET,
    PLANO_CONTAS_SHEET,
    REVISAO_SHEET,
    gerar_planilha_revisao_classificacao,
)


def test_gera_workbook_a_partir_de_csv_valido(tmp_path):
    csv_path = _save_review_csv(tmp_path)
    plano_path = _save_plano_contas(tmp_path)
    output_path = tmp_path / "reports" / "revisao_assistida.xlsx"

    result = gerar_planilha_revisao_classificacao(csv_path, output_path, plano_path)

    assert result == output_path
    assert output_path.exists()


def test_workbook_contem_abas_obrigatorias(tmp_path):
    output_path = _generate_workbook(tmp_path)

    workbook = load_workbook(output_path)

    assert workbook.sheetnames == [
        REVISAO_SHEET,
        PLANO_CONTAS_SHEET,
        ACOES_VALIDAS_SHEET,
        GUIA_ALIAS_REGRAS_SHEET,
        EXEMPLOS_SHEET,
        INSTRUCOES_SHEET,
    ]


def test_aba_revisao_mantem_cabecalhos_contratados_do_csv(tmp_path):
    output_path = _generate_workbook(tmp_path)
    sheet = load_workbook(output_path)[REVISAO_SHEET]

    headers = [cell.value for cell in sheet[1]]

    for header in CSV_HEADERS:
        assert header in headers


def test_colunas_manuais_esperadas_existem_na_aba_revisao(tmp_path):
    output_path = _generate_workbook(tmp_path)
    sheet = load_workbook(output_path)[REVISAO_SHEET]
    headers = [cell.value for cell in sheet[1]]

    assert "acao_manual" in headers
    assert "conta_escolhida" in headers
    assert "conta_id_sugerida" in headers


def test_lista_de_acoes_validas_e_criada(tmp_path):
    output_path = _generate_workbook(tmp_path)
    workbook = load_workbook(output_path)
    acoes_sheet = workbook[ACOES_VALIDAS_SHEET]
    revisao_sheet = workbook[REVISAO_SHEET]
    actions = [row[0].value for row in acoes_sheet.iter_rows(min_row=2)]
    validations = list(revisao_sheet.data_validations.dataValidation)

    assert actions == [
        "ANALISAR_COM_CONTADOR",
        "CLASSIFICAR",
        "CRIAR_ALIAS",
        "CRIAR_REGRA",
        "IGNORAR",
        "MANTER_REVISAO",
    ]
    assert validations
    assert validations[0].type == "list"


def test_plano_de_contas_e_carregado_do_yaml(tmp_path):
    output_path = _generate_workbook(tmp_path)
    sheet = load_workbook(output_path)[PLANO_CONTAS_SHEET]

    assert [cell.value for cell in sheet[1]] == [
        "conta_id",
        "conta_nome",
        "grupo",
        "nivel",
        "conta_formatada",
        "ativo",
        "observacao",
    ]
    assert sheet["A2"].value == "2.01.004"
    assert sheet["B2"].value == "Refrigerantes"
    assert sheet["C2"].value == "Custos de Mercadorias e Insumos"
    assert sheet["E2"].value == "2.01.004 - Refrigerantes"


def test_csv_original_nao_e_modificado(tmp_path):
    csv_path = _save_review_csv(tmp_path)
    plano_path = _save_plano_contas(tmp_path)
    original_csv = csv_path.read_text(encoding="utf-8-sig")

    gerar_planilha_revisao_classificacao(
        csv_path, tmp_path / "revisao_assistida.xlsx", plano_path
    )

    assert csv_path.read_text(encoding="utf-8-sig") == original_csv


def test_yaml_original_nao_e_modificado(tmp_path):
    csv_path = _save_review_csv(tmp_path)
    plano_path = _save_plano_contas(tmp_path)
    original_yaml = plano_path.read_text(encoding="utf-8")

    gerar_planilha_revisao_classificacao(
        csv_path, tmp_path / "revisao_assistida.xlsx", plano_path
    )

    assert plano_path.read_text(encoding="utf-8") == original_yaml


def test_aba_revisao_tem_filtro_automatico(tmp_path):
    output_path = _generate_workbook(tmp_path)
    sheet = load_workbook(output_path)[REVISAO_SHEET]

    assert sheet.auto_filter.ref == sheet.dimensions


def test_aba_revisao_congela_primeira_linha(tmp_path):
    output_path = _generate_workbook(tmp_path)
    sheet = load_workbook(output_path)[REVISAO_SHEET]

    assert sheet.freeze_panes == "A2"


def test_planilha_melhorada_tem_dropdown_conta_escolhida(tmp_path):
    output_path = _generate_workbook(tmp_path)
    sheet = load_workbook(output_path)[REVISAO_SHEET]
    headers = [cell.value for cell in sheet[1]]
    conta_col = headers.index("conta_escolhida") + 1
    conta_letter = sheet.cell(row=1, column=conta_col).column_letter
    validations = list(sheet.data_validations.dataValidation)

    assert any(
        validation.type == "list"
        and "Plano_de_Contas!$E$2:$E$3" in validation.formula1
        and f"{conta_letter}2:{conta_letter}3" in str(validation.sqref)
        for validation in validations
    )


def test_planilha_melhorada_tem_guia_alias_e_regras(tmp_path):
    output_path = _generate_workbook(tmp_path)
    sheet = load_workbook(output_path)[GUIA_ALIAS_REGRAS_SHEET]
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]

    assert "ALIAS" in values
    assert "REGRA" in values
    assert any("prefira CRIAR_ALIAS" in value for value in values if isinstance(value, str))


def test_planilha_melhorada_tem_exemplos_alias_e_regras(tmp_path):
    output_path = _generate_workbook(tmp_path)
    sheet = load_workbook(output_path)[EXEMPLOS_SHEET]
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]

    assert "CRIAR_ALIAS" in values
    assert "CRIAR_REGRA" in values
    assert any("Evite palavras genericas" in value for value in values)


def test_planilha_melhorada_nao_exige_digitacao_manual_conta_id(tmp_path):
    output_path = _generate_workbook(tmp_path)
    sheet = load_workbook(output_path)[REVISAO_SHEET]
    headers = [cell.value for cell in sheet[1]]
    conta_escolhida_col = headers.index("conta_escolhida") + 1
    conta_id_col = headers.index("conta_id_sugerida") + 1

    assert not sheet.cell(row=2, column=conta_escolhida_col).protection.locked
    assert sheet.cell(row=2, column=conta_id_col).protection.locked
    assert sheet.protection.sheet


def _generate_workbook(tmp_path: Path) -> Path:
    csv_path = _save_review_csv(tmp_path)
    plano_path = _save_plano_contas(tmp_path)
    output_path = tmp_path / "revisao_assistida.xlsx"
    gerar_planilha_revisao_classificacao(csv_path, output_path, plano_path)
    return output_path


def _save_review_csv(tmp_path: Path) -> Path:
    csv_path = tmp_path / "revisao_classificacao.csv"
    csv_path.write_text(
        ";".join(CSV_HEADERS)
        + "\n"
        + ";".join(
            [
                "CLASSIFICACAO_PRODUTO",
                "agua mineral",
                "AGUA MINERAL 500ML",
                "3",
                "Fornecedor A",
                "123,124",
                "45.90",
                "",
                "",
                "REVISAR",
                "",
                "MANTER_REVISAO",
                "",
                "",
                "",
                "",
            ]
        )
        + "\n"
        + ";".join(
            [
                "CLASSIFICACAO_PRODUTO",
                "freezer comodato",
                "FREEZER 300L",
                "1",
                "Fornecedor B",
                "999",
                "0.00",
                "",
                "",
                "REVISAR",
                "POSSIVEL_COMODATO",
                "ANALISAR_COM_CONTADOR",
                "",
                "",
                "",
                "",
            ]
        )
        + "\n",
        encoding="utf-8-sig",
    )
    return csv_path


def _save_plano_contas(tmp_path: Path) -> Path:
    plano_path = tmp_path / "plano_contas_gerencial.yaml"
    plano_path.write_text(
        """\
contas:
  - id: "2.01.004"
    nome: "Refrigerantes"
    grupo: "Custos de Mercadorias e Insumos"
    tipo: "custo"
  - id: "13.01.001"
    nome: "Revisar - Produto Nao Classificado"
    grupo: "Contas de Revisao e Classificacao Pendente"
    tipo: "revisao"
""",
        encoding="utf-8",
    )
    return plano_path
