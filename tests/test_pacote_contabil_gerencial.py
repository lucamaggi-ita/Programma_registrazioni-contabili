from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from fantozzi.pacote_contabil_gerencial import (
    BONIFICACAO_AUDIT_FILE,
    COMODATO_AUDIT_FILE,
    DRE_TXT_FILE,
    DRE_XLSX_FILE,
    INDICE_FILE,
    MANIFEST_FILE,
    MANUAIS_AUDIT_FILE,
    PACOTE_DIR_NAME,
    PRE_LIVRO_XLSX_FILE,
    RELATORIOS_XLSX_FILE,
    gerar_pacote_contabil_gerencial_provisorio,
)


def test_gera_indice_manifest_e_copia_arquivos_presentes(tmp_path):
    reports_run = _make_run(tmp_path, manual_executado=True)
    _save_core_files(reports_run)
    _save_dre_workbook(reports_run / "dre_gerencial_ok_provisoria.xlsx")
    _save_pre_livro_workbook(reports_run / "pre_livro_diario_gerencial_teste.xlsx")

    result = gerar_pacote_contabil_gerencial_provisorio(reports_run)

    pacote_dir = reports_run / PACOTE_DIR_NAME
    manifest = json.loads((pacote_dir / MANIFEST_FILE).read_text(encoding="utf-8"))
    indice = (pacote_dir / INDICE_FILE).read_text(encoding="utf-8")
    assert result.status == "OK"
    assert (pacote_dir / DRE_XLSX_FILE).exists()
    assert (pacote_dir / DRE_TXT_FILE).exists()
    assert (pacote_dir / PRE_LIVRO_XLSX_FILE).exists()
    assert (pacote_dir / RELATORIOS_XLSX_FILE).exists()
    assert manifest["tipo_pacote"] == "pacote_contabil_gerencial_provisorio"
    assert manifest["fase"] == "9O"
    assert manifest["run_id"] == "2026-05-16_999"
    assert manifest["cnpjs_consolidados"] == ["08884180000123", "55746722000181"]
    assert manifest["comodato_entra_dre"] is False
    assert manifest["bonificacao_recebida_entra_dre"] is True
    assert manifest["fase9_executada"] is False
    assert manifest["arquivos"][INDICE_FILE]["status"] == "OK"
    assert manifest["arquivos"][MANIFEST_FILE]["status"] == "OK"
    assert "COMODATO fora DRE: OK" in indice
    assert "Extrato bancario" in indice


def test_arquivos_ausentes_sao_marcados_como_nao_gerado(tmp_path):
    reports_run = _make_run(tmp_path, manual_executado=False)
    _save_dre_workbook(reports_run / "dre_gerencial_ok_provisoria.xlsx")

    result = gerar_pacote_contabil_gerencial_provisorio(reports_run)

    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert result.status == "PARCIAL"
    assert manifest["arquivos"][DRE_XLSX_FILE]["status"] == "OK"
    assert manifest["arquivos"][DRE_TXT_FILE]["status"] == "NAO_GERADO"
    assert manifest["arquivos"][PRE_LIVRO_XLSX_FILE]["status"] == "NAO_GERADO"


def test_audit_comodato_confirma_fora_dre(tmp_path):
    reports_run = _make_run(tmp_path, manual_executado=False)
    _save_dre_workbook(reports_run / "dre_gerencial_ok_provisoria.xlsx")

    gerar_pacote_contabil_gerencial_provisorio(reports_run)

    text = (
        reports_run
        / PACOTE_DIR_NAME
        / COMODATO_AUDIT_FILE
    ).read_text(encoding="utf-8")
    assert "COMODATO_BASE_DRE: 0" in text
    assert "COMODATO fora DRE: OK" in text


def test_audit_bonificacao_confirma_205003_reducao_custo(tmp_path):
    reports_run = _make_run(tmp_path, manual_executado=False)
    _save_dre_workbook(reports_run / "dre_gerencial_ok_provisoria.xlsx")

    gerar_pacote_contabil_gerencial_provisorio(reports_run)

    text = (
        reports_run
        / PACOTE_DIR_NAME
        / BONIFICACAO_AUDIT_FILE
    ).read_text(encoding="utf-8")
    assert "conta_id: 2.05.003" in text
    assert "tratamento: REDUCAO_CUSTO" in text
    assert "entra_dre: true" in text
    assert "quantidade_lancamentos: 1" in text


def test_audit_manuais_nao_aplicavel_quando_executado_false(tmp_path):
    reports_run = _make_run(tmp_path, manual_executado=False)

    gerar_pacote_contabil_gerencial_provisorio(reports_run)

    text = (
        reports_run
        / PACOTE_DIR_NAME
        / MANUAIS_AUDIT_FILE
    ).read_text(encoding="utf-8")
    assert "executado: false" in text
    assert "NAO_APLICAVEL" in text


def test_audit_manuais_executado_true_confirma_partida_dobrada(tmp_path):
    reports_run = _make_run(tmp_path, manual_executado=True)
    _save_pre_livro_workbook(reports_run / "pre_livro_diario_gerencial_teste.xlsx")

    gerar_pacote_contabil_gerencial_provisorio(reports_run)

    text = (
        reports_run
        / PACOTE_DIR_NAME
        / MANUAIS_AUDIT_FILE
    ).read_text(encoding="utf-8")
    assert "executado: true" in text
    assert "lancamentos_aceitos: 1" in text
    assert "linhas_tecnicas_geradas: 2" in text
    assert "MANUAL_GERENCIAL_DEBITO_TOTALIZAVEL" in text
    assert "MANUAL_GERENCIAL_CREDITO" in text
    assert "pre_livro_partida_dobrada: OK" in text
    assert "dre_impacta_uma_vez: OK" in text


def _make_run(tmp_path: Path, *, manual_executado: bool) -> Path:
    reports_run = tmp_path / "reports" / "importacao_real" / "runs" / "2026-05-16_999"
    reports_run.mkdir(parents=True)
    manifest = {
        "run_id": "2026-05-16_999",
        "classificacao": {"fase9_executada": False},
        "fontes_xml": [
            {"cnpj_real": "08884180000123"},
            {"cnpj_real": "55746722000181"},
        ],
        "lancamentos_manuais_gerenciais": (
            {
                "executado": True,
                "input_csv": "manual.csv",
                "lancamentos_aceitos": 1,
                "linhas_tecnicas_geradas": 2,
                "total_manual": "10.50",
            }
            if manual_executado
            else {"executado": False}
        ),
    }
    (reports_run / "manifest_run.json").write_text(
        json.dumps(manifest),
        encoding="utf-8",
    )
    return reports_run


def _save_core_files(reports_run: Path) -> None:
    for name in (
        "dre_gerencial_ok_provisoria.txt",
        "auditoria_pre_livro_diario_gerencial_teste.txt",
        "relatorios_gerenciais_ok_provisorios.txt",
    ):
        (reports_run / name).write_text(name, encoding="utf-8")
    _empty_xlsx(reports_run / "relatorios_gerenciais_ok_provisorios.xlsx")


def _save_dre_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Base_OK_DRE"
    sheet.append(
        [
            "documento",
            "conta_id",
            "conta_nome",
            "valor",
            "tipo_operacao_especial",
            "tratamento_operacao_especial",
        ]
    )
    sheet.append(
        [
            "BON-001",
            "2.05.003",
            "Reducao de Custos - Bonificacoes Recebidas",
            -10.5,
            "BONIFICACAO_RECEBIDA",
            "REDUCAO_CUSTO",
        ]
    )
    workbook.save(path)


def _save_pre_livro_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lancamentos_OK"
    sheet.append(["origem", "tipo_lancamento", "valor"])
    sheet.append(["MANUAL_GERENCIAL", "MANUAL_GERENCIAL_DEBITO_TOTALIZAVEL", 10.5])
    sheet.append(["MANUAL_GERENCIAL", "MANUAL_GERENCIAL_CREDITO", 10.5])
    workbook.save(path)


def _empty_xlsx(path: Path) -> None:
    workbook = Workbook()
    workbook.save(path)
