from __future__ import annotations

import csv
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fantozzi.banking.payment_classification_run import (
    execute_payment_classification_run,
)
from fantozzi.banking.payment_review_package import (
    FORA_ESCOPO_SISPAG,
    REVISAO_SHEET,
    RESUMO_BENEFICIARIO_SHEET,
    SISPAG_SHEET,
    gerar_pacote_revisao_comprovantes,
)
from fantozzi.banking.payment_rules import (
    STATUS_CLASSIFICADO,
    STATUS_EXCLUIDO,
    STATUS_MULTIPLAS_REGRAS,
    STATUS_SEM_REGRA,
    PaymentRulesError,
    audit_payment_rules_file,
    classify_payment_receipt,
    classify_payment_receipts_file,
    load_payment_rules,
    load_payment_exclusion_rules,
    normalize_match_text,
    payment_rules_audit_report_text,
    validate_payment_rules,
)
from fantozzi.cli import build_parser, main


def test_match_by_beneficiary_and_reference(tmp_path):
    rules_path = _write_rules(
        tmp_path,
        [
            _rule(
                "REGRA_EMBASA",
                "4.01.012",
                "Agua e Esgoto",
                "06.02",
                beneficiary=["EMBASA"],
                reference=["AGUA LOJA"],
            )
        ],
    )
    rules = load_payment_rules(rules_path)

    result = classify_payment_receipt(
        {
            "beneficiary_name": "Embasa SA",
            "reference_id": "Agua Loja Centro",
        },
        rules,
    )

    assert result.status_classificacao == STATUS_CLASSIFICADO
    assert result.conta_id == "4.01.012"
    assert result.regra_id == "REGRA_EMBASA"


def test_match_by_reference_only(tmp_path):
    rules_path = _write_rules(
        tmp_path,
        [
            _rule(
                "REGRA_REF",
                "6.01.001",
                "Honorarios Contabeis",
                "08.01",
                reference=["CONTABILIDADE"],
            )
        ],
    )
    rules = load_payment_rules(rules_path)

    result = classify_payment_receipt(
        {
            "beneficiary_name": "Fornecedor Generico",
            "reference_id": "Mensalidade Contabilidade",
        },
        rules,
    )

    assert result.status_classificacao == STATUS_CLASSIFICADO
    assert result.conta_id == "6.01.001"


def test_without_rule_goes_to_review(tmp_path):
    rules = load_payment_rules(_write_rules(tmp_path, []))

    result = classify_payment_receipt(
        {"beneficiary_name": "Sem Regra", "reference_id": "Outro"},
        rules,
    )

    assert result.status_classificacao == STATUS_SEM_REGRA
    assert result.motivo == "SEM_REGRA"
    assert result.conta_id == ""


def test_exclusion_rule_marks_receipt_without_account_or_dre(tmp_path):
    rules_path = _write_rules_with_exclusions(
        tmp_path,
        positive_rules=[],
        exclusions=[
            _exclusion_rule(
                "EXCLUSAO_XML_FORNECEDOR",
                "XML",
                beneficiary=["FORNECEDOR XML"],
            )
        ],
    )
    rules = load_payment_rules(rules_path)
    exclusions = load_payment_exclusion_rules(rules_path)

    result = classify_payment_receipt(
        {
            "beneficiary_name": "Fornecedor XML",
            "reference_id": "-",
        },
        rules,
        exclusions,
    )

    assert result.status_classificacao == STATUS_EXCLUIDO
    assert result.acao == "XML"
    assert result.conta_id == ""
    assert result.conta_nome == ""
    assert result.dre_linha_id == ""
    assert result.entra_dre is False


def test_inss_reference_exclusion_does_not_match_empty_reference(tmp_path):
    rules_path = _write_rules_with_exclusions(
        tmp_path,
        positive_rules=[],
        exclusions=[
            _exclusion_rule(
                "EXCLUSAO_ESCLUDERE_REFERENCIA_INSS",
                "ESCLUDERE",
                reference=["INSS"],
            )
        ],
    )
    rules = load_payment_rules(rules_path)
    exclusions = load_payment_exclusion_rules(rules_path)

    inss_result = classify_payment_receipt(
        {
            "beneficiary_name": "RECEITA FED DARF NUMERADO CB",
            "reference_id": "INSS e IRRF MARCO 20",
        },
        rules,
        exclusions,
    )
    caixa_result = classify_payment_receipt(
        {
            "beneficiary_name": "CAIXA ECONOMICA FEDERAL",
            "reference_id": "-",
        },
        rules,
        exclusions,
    )

    assert inss_result.status_classificacao == STATUS_EXCLUIDO
    assert inss_result.acao == "ESCLUDERE"
    assert caixa_result.status_classificacao == STATUS_SEM_REGRA


def test_multiple_matching_rules_go_to_review(tmp_path):
    rules_path = _write_rules(
        tmp_path,
        [
            _rule(
                "REGRA_A",
                "7.01.001",
                "Tarifas Bancarias",
                "09.01",
                reference=["TARIFA PIX"],
            ),
            _rule(
                "REGRA_B",
                "7.01.003",
                "Tarifas PIX",
                "09.01",
                reference=["TARIFA PIX"],
            ),
        ],
    )
    rules = load_payment_rules(rules_path)

    result = classify_payment_receipt(
        {"beneficiary_name": "Banco", "reference_id": "Tarifa Pix Pagamento"},
        rules,
    )

    assert result.status_classificacao == STATUS_MULTIPLAS_REGRAS
    assert result.regras_candidatas == ("REGRA_A", "REGRA_B")


def test_multiple_matching_rules_same_account_are_classified(tmp_path):
    rules_path = _write_rules(
        tmp_path,
        [
            _rule(
                "REGRA_PARCSN",
                "5.01.003",
                "Parcelamento Federal / Receita Federal / PGFN / Simples",
                "07.01",
                reference=["PARCSN"],
            ),
            _rule(
                "REGRA_PERTSN",
                "5.01.003",
                "Parcelamento Federal / Receita Federal / PGFN / Simples",
                "07.01",
                reference=["PERTSN"],
            ),
        ],
    )
    rules = load_payment_rules(rules_path)

    result = classify_payment_receipt(
        {
            "beneficiary_name": "DAS-SIMPLES NACIONAL-COD BARRA",
            "reference_id": "PARCSN 2 RATA 28 PERTSN 9101 RATA 94",
        },
        rules,
    )

    assert result.status_classificacao == STATUS_CLASSIFICADO
    assert result.conta_id == "5.01.003"
    assert result.regra_id == "REGRA_PARCSN"
    assert result.regras_candidatas == ("REGRA_PARCSN", "REGRA_PERTSN")


def test_normalization_removes_accents_case_and_punctuation():
    assert normalize_match_text("  Companhia de ÁGUA/esgoto!!! ") == (
        "COMPANHIA DE AGUA ESGOTO"
    )


def test_conta_corrente_recurring_rules_match_direct_fields(tmp_path):
    rules_path = tmp_path / "regras.yaml"
    rules_path.write_text(
        """regras:
- regra_id: REGRA_COELBA
  conta_id: "4.01.011"
  conta_nome: Energia Eletrica
  dre_linha_id: "06.02"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - COELBA
- regra_id: REGRA_EMBASA
  conta_id: "4.01.012"
  conta_nome: Agua e Esgoto
  dre_linha_id: "06.02"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - EMBASA
- regra_id: REGRA_VIVO
  conta_id: "4.01.014"
  conta_nome: Internet e Telefonia
  dre_linha_id: "06.02"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - VIVO
- regra_id: REGRA_DEBITO_SEGURO
  conta_id: "7.01.001"
  conta_nome: Tarifas Bancarias
  dre_linha_id: "09.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - DEBITO SEGURO
- regra_id: REGRA_SIMPLES_NACIONAL
  conta_id: "5.01.001"
  conta_nome: Simples Nacional - DAS Mensal - Anexo I
  dre_linha_id: "02.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
    - SIMPLES NACIONAL
    referencia_empresa_contem_algum: []
exclusoes:
- regra_id: EXCLUSAO_SISPAG_SALARIOS
  acao: ESCLUDERE
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum:
    - SISPAG SALARIOS
    tipo_pagamento_contem_algum:
    - SISPAG_AGREGADO
""",
        encoding="utf-8",
    )
    rules = load_payment_rules(rules_path)
    exclusions = load_payment_exclusion_rules(rules_path)

    cases = [
        ("DA COELBA 9999999999", "", "SAIDA_PAGAMENTO", "4.01.011"),
        ("EMBASA 987654321", "", "SAIDA_PAGAMENTO", "4.01.012"),
        ("DA VIVO-BA 20304230426", "", "SAIDA_PAGAMENTO", "4.01.014"),
        ("DEBITO SEGURO 12/57", "", "SAIDA_PAGAMENTO", "7.01.001"),
        (
            "PAGAMENTOS TRIB COD BARRAS",
            "SIMPLES NACIONAL",
            "SAIDA_PAGAMENTO",
            "5.01.001",
        ),
    ]
    for lancamento, razao_social, tipo_movimento, conta_id in cases:
        result = classify_payment_receipt(
            {
                "beneficiary_name": razao_social,
                "reference_id": lancamento,
                "payment_type": tipo_movimento,
                "description_receipt": lancamento,
            },
            rules,
            exclusions,
        )

        assert result.status_classificacao == STATUS_CLASSIFICADO
        assert result.conta_id == conta_id

    generic_simples = classify_payment_receipt(
        {
            "beneficiary_name": "",
            "reference_id": "PAGAMENTOS TRIB COD BARRAS",
            "payment_type": "SAIDA_PAGAMENTO",
            "description_receipt": "PAGAMENTOS TRIB COD BARRAS",
        },
        rules,
        exclusions,
    )
    sispag = classify_payment_receipt(
        {
            "beneficiary_name": "",
            "reference_id": "SISPAG SALARIOS",
            "payment_type": "SISPAG_AGREGADO",
            "description_receipt": "SISPAG SALARIOS",
        },
        rules,
        exclusions,
    )

    assert generic_simples.status_classificacao == STATUS_SEM_REGRA
    assert sispag.status_classificacao == STATUS_EXCLUIDO
    assert sispag.acao == "ESCLUDERE"
    assert sispag.entra_dre is False


def test_rule_with_only_origin_metadata_and_empty_match_is_rejected(tmp_path):
    rules_path = tmp_path / "regras.yaml"
    rules_path.write_text(
        """regras:
- regra_id: COMPROVANTE_LINHA_002
  origem_linha_planilha: 2
  conta_id: "9.01.001"
  conta_nome: Publicidade e Propaganda
  dre_linha_id: "11.01"
  entra_dre: true
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum: []
    referencia_empresa_contem_algum: []
""",
        encoding="utf-8",
    )

    with pytest.raises(PaymentRulesError, match="ao menos uma condicao"):
        load_payment_rules(rules_path)


def test_rule_with_beneficiary_and_reference_does_not_match_empty_reference(tmp_path):
    rules = load_payment_rules(
        _write_rules(
            tmp_path,
            [
                _rule(
                    "REGRA_EUDES_COCO",
                    "2.01.009",
                    "Coco verde",
                    "03.02",
                    beneficiary=["EUDES BRAGA NASCIMENTO"],
                    reference=["COCO VERDE"],
                    entra_dre=False,
                )
            ],
        )
    )

    result = classify_payment_receipt(
        {
            "beneficiary_name": "EUDES BRAGA NASCIMENTO",
            "reference_id": "",
        },
        rules,
    )

    assert result.status_classificacao == STATUS_SEM_REGRA


def test_rule_with_beneficiary_only_matches_even_when_reference_is_empty(tmp_path):
    rules = load_payment_rules(
        _write_rules(
            tmp_path,
            [
                _rule(
                    "REGRA_EUDES",
                    "2.01.009",
                    "Coco verde",
                    "03.02",
                    beneficiary=["EUDES BRAGA NASCIMENTO"],
                    entra_dre=False,
                )
            ],
        )
    )

    result = classify_payment_receipt(
        {
            "beneficiary_name": "EUDES BRAGA NASCIMENTO",
            "reference_id": "",
        },
        rules,
    )

    assert result.status_classificacao == STATUS_CLASSIFICADO
    assert result.regra_id == "REGRA_EUDES"


def test_pjbank_rule_matches_beneficiary_with_dash_reference(tmp_path):
    rules = load_payment_rules(
        _write_rules(
            tmp_path,
            [
                _rule(
                    "COMPROVANTE_LINHA_002",
                    "9.01.001",
                    "Publicidade e Propaganda",
                    "11.01",
                    beneficiary=["PJBANK PAGAMENTOS S A"],
                )
            ],
        )
    )

    result = classify_payment_receipt(
        {
            "beneficiary_name": "PJBANK PAGAMENTOS S A",
            "reference_id": "-",
        },
        rules,
    )

    assert result.status_classificacao == STATUS_CLASSIFICADO
    assert result.conta_id == "9.01.001"
    assert result.conta_nome == "Publicidade e Propaganda"
    assert result.regra_id == "COMPROVANTE_LINHA_002"


def test_sispar_pgfn_rule_matches_real_reference_variants(tmp_path):
    rules = load_payment_rules(
        _write_rules(
            tmp_path,
            [
                _rule(
                    "REGRA_5_01_003_SISPAR_PGFN_8916087",
                    "5.01.003",
                    "Parcelamento Federal / Receita Federal / PGFN / Simples",
                    "07.01",
                    reference=["sispar", "PGFN 8916087", "SISPAR 8916087", "pgfn"],
                )
            ],
        )
    )

    for reference in (
        "PGFN-SISPAR 8916087",
        "SISPAR 8916087 RT29",
        "SISPAR 8916087 R.30",
    ):
        result = classify_payment_receipt(
            {
                "beneficiary_name": "DAS-SIMPLES NACIONAL-COD BARRA",
                "reference_id": reference,
            },
            rules,
        )

        assert result.status_classificacao == STATUS_CLASSIFICADO
        assert result.conta_id == "5.01.003"
        assert result.regra_id == "REGRA_5_01_003_SISPAR_PGFN_8916087"


def test_audit_payment_rules_flags_planilha_identity_and_conflict(tmp_path):
    rules_path = _write_rules(
        tmp_path,
        [
            _rule(
                "COMPROVANTE_LINHA_002",
                "7.01.001",
                "Tarifas Bancarias",
                "09.01",
                reference=["TARIFA PIX"],
                origem_linha=2,
            ),
            _rule(
                "REGRA_TARIFA_PIX_DIVERGENTE",
                "7.01.003",
                "Tarifas PIX",
                "09.01",
                reference=["TARIFA PIX"],
            ),
        ],
    )

    result = audit_payment_rules_file(rules_path)
    report = payment_rules_audit_report_text(result)

    assert result.total_regras == 2
    assert result.total_erros == 1
    assert "CAMPO_ORIGEM_PLANILHA" in report
    assert "REGRA_ID_BASEADO_EM_LINHA" in report
    assert "CONFLITO_CRITERIO_CONTA" in report


def test_current_payment_rules_yaml_is_autonomous_from_planilha():
    rules_path = Path(__file__).resolve().parents[1] / "config" / (
        "regras_classificacao_comprovantes.yaml"
    )

    result = audit_payment_rules_file(rules_path)
    categories = {issue.categoria for issue in result.issues}

    assert "REFERENCIA_PLANILHA" not in categories
    assert "CAMPO_ORIGEM_PLANILHA" not in categories
    assert "REGRA_ID_BASEADO_EM_LINHA" not in categories


def test_current_payment_rules_apply_13y_confirmed_review_decisions_only():
    rules_path = Path(__file__).resolve().parents[1] / "config" / (
        "regras_classificacao_comprovantes.yaml"
    )
    rules = load_payment_rules(rules_path)
    exclusions = load_payment_exclusion_rules(rules_path)

    total = classify_payment_receipt(
        {"beneficiary_name": "TOTAL", "reference_id": "", "payment_type": ""},
        rules,
        exclusions,
    )
    assert total.status_classificacao == STATUS_EXCLUIDO
    assert total.acao == "ESCLUDERE"
    assert total.conta_id == ""

    for reference, expected_rule in (
        ("PARCSN 2 RATA 29", "REGRA_5_01_003_PARCSN"),
        ("PERTSN 9101 RATA 90", "REGRA_5_01_003_PERTSN"),
        ("PGFN-SISPAR 8916087", "REGRA_5_01_003_SISPAR_PGFN_8916087"),
    ):
        result = classify_payment_receipt(
            {
                "beneficiary_name": "DAS-SIMPLES NACIONAL-COD BARRA",
                "reference_id": reference,
            },
            rules,
            exclusions,
        )
        assert result.status_classificacao == STATUS_CLASSIFICADO
        assert result.conta_id == "5.01.003"
        assert result.dre_linha_id == "07.01"
        assert result.regra_id == expected_rule

    simples_ordinario = classify_payment_receipt(
        {
            "beneficiary_name": "PAGAMENTOS TRIB COD BARRAS",
            "reference_id": "SIMPLES NACIONAL",
        },
        rules,
        exclusions,
    )
    assert simples_ordinario.status_classificacao == STATUS_CLASSIFICADO
    assert simples_ordinario.conta_id == "5.01.001"
    assert simples_ordinario.dre_linha_id == "02.01"

    jose = classify_payment_receipt(
        {"beneficiary_name": "JOSE RUCAS ACHY NETO", "reference_id": "-"},
        rules,
        exclusions,
    )
    assert jose.status_classificacao == STATUS_CLASSIFICADO
    assert jose.conta_id == "2.02.001"
    assert jose.dre_linha_id == "03.01"
    assert jose.regra_id == "REGRA_2_02_001_JOSE_RUCAS_ACHY_NETO"

    pataxo = classify_payment_receipt(
        {"beneficiary_name": "PATAXO MATERIAIS DE CONSTRUCAO", "reference_id": "-"},
        rules,
        exclusions,
    )
    assert pataxo.status_classificacao == STATUS_EXCLUIDO
    assert pataxo.acao == "XML"
    assert pataxo.conta_id == ""

    himeros = classify_payment_receipt(
        {"beneficiary_name": "HIMEROS COMERCIO DE PESCADOS L", "reference_id": "-"},
        rules,
        exclusions,
    )
    assert himeros.status_classificacao == STATUS_SEM_REGRA


def test_rule_with_missing_account_fails_validation(tmp_path):
    rules = load_payment_rules(
        _write_rules(
            tmp_path,
            [
                _rule(
                    "REGRA_MISSING",
                    "9.99.999",
                    "Conta Inexistente",
                    "99.01",
                    reference=["TESTE"],
                )
            ],
        )
    )
    plano_path, mapeamento_path = _write_plan_and_mapping(tmp_path)

    with pytest.raises(PaymentRulesError, match="conta_id inexistente"):
        validate_payment_rules(
            rules,
            plano_contas_path=plano_path,
            mapeamento_dre_path=mapeamento_path,
        )


def test_rule_with_divergent_dre_mapping_fails_validation(tmp_path):
    rules = load_payment_rules(
        _write_rules(
            tmp_path,
            [
                _rule(
                    "REGRA_DRE",
                    "4.01.012",
                    "Agua e Esgoto",
                    "99.01",
                    reference=["EMBASA"],
                )
            ],
        )
    )
    plano_path, mapeamento_path = _write_plan_and_mapping(tmp_path)

    with pytest.raises(PaymentRulesError, match="dre_linha_id"):
        validate_payment_rules(
            rules,
            plano_contas_path=plano_path,
            mapeamento_dre_path=mapeamento_path,
        )


def test_cli_classifies_and_writes_expected_csvs(tmp_path):
    plano_path, mapeamento_path = _write_plan_and_mapping(tmp_path)
    rules_path = _write_rules(
        tmp_path,
        [
            _rule(
                "REGRA_EMBASA",
                "4.01.012",
                "Agua e Esgoto",
                "06.02",
                beneficiary=["EMBASA"],
                reference=["AGUA"],
            )
        ],
    )
    comprovantes_path = _write_receipts_csv(
        tmp_path,
        [
            _receipt("RC-1", "Embasa", "Agua mensal"),
            _receipt("RC-2", "Fornecedor", "Sem regra"),
        ],
    )

    result = main(
        [
            "classificar-comprovantes-pagamento",
            "--comprovantes",
            str(comprovantes_path),
            "--regras",
            str(rules_path),
            "--plano-contas",
            str(plano_path),
            "--mapeamento-dre",
            str(mapeamento_path),
            "--saida",
            str(tmp_path / "saida"),
        ]
    )

    assert result == 0
    classified = _read_csv(tmp_path / "saida" / "comprovantes_classificados.csv")
    review = _read_csv(tmp_path / "saida" / "comprovantes_a_revisar.csv")
    manifest = json.loads(
        (tmp_path / "saida" / "manifest_classificacao_comprovantes.json").read_text(
            encoding="utf-8"
        )
    )
    assert classified[0]["payment_id"] == "RC-1"
    assert classified[0]["status_classificacao"] == STATUS_CLASSIFICADO
    assert review[0]["payment_id"] == "RC-2"
    assert review[0]["status_classificacao"] == STATUS_SEM_REGRA
    assert manifest["classification_executed"] is True
    assert manifest["nao_inferir_conta"] is True


def test_file_api_validates_rules_before_output(tmp_path):
    plano_path, mapeamento_path = _write_plan_and_mapping(tmp_path)
    rules_path = _write_rules(
        tmp_path,
        [
            _rule(
                "REGRA_OK",
                "4.01.012",
                "Agua e Esgoto",
                "06.02",
                reference=["AGUA"],
            )
        ],
    )
    comprovantes_path = _write_receipts_csv(
        tmp_path,
        [_receipt("RC-1", "Embasa", "Agua mensal")],
    )

    manifest = classify_payment_receipts_file(
        comprovantes_path=comprovantes_path,
        regras_path=rules_path,
        plano_contas_path=plano_path,
        mapeamento_dre_path=mapeamento_path,
        output_dir=tmp_path / "out",
    )

    assert manifest.classified_count == 1
    assert (tmp_path / "out" / "comprovantes_classificados.csv").exists()


def test_parser_accepts_classificar_comprovantes_pagamento_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "classificar-comprovantes-pagamento",
            "--comprovantes",
            "payment_receipts_normalized.csv",
            "--saida",
            "output/fase_13h",
        ]
    )

    assert args.command == "classificar-comprovantes-pagamento"
    assert args.comprovantes == "payment_receipts_normalized.csv"
    assert args.saida == "output/fase_13h"


def test_parser_accepts_auditar_regras_comprovantes_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "auditar-regras-comprovantes",
            "--regras",
            "config/regras_classificacao_comprovantes.yaml",
            "--relatorio",
            "reports/auditoria_regras_comprovantes.txt",
        ]
    )

    assert args.command == "auditar-regras-comprovantes"
    assert args.regras == "config/regras_classificacao_comprovantes.yaml"
    assert args.relatorio == "reports/auditoria_regras_comprovantes.txt"


def test_cli_auditar_regras_comprovantes_writes_report(tmp_path):
    rules_path = _write_rules(
        tmp_path,
        [
            _rule(
                "COMPROVANTE_LINHA_002",
                "7.01.001",
                "Tarifas Bancarias",
                "09.01",
                reference=["TARIFA PIX"],
                origem_linha=2,
            )
        ],
    )
    report_path = tmp_path / "auditoria.txt"

    result = main(
        [
            "auditar-regras-comprovantes",
            "--regras",
            str(rules_path),
            "--relatorio",
            str(report_path),
        ]
    )

    assert result == 0
    report = report_path.read_text(encoding="utf-8")
    assert "Auditoria de regras de comprovantes" in report
    assert "REGRA_ID_BASEADO_EM_LINHA" in report


def test_execute_payment_classification_run_generates_13s_package(tmp_path):
    plano_path, mapeamento_path = _write_plan_and_mapping(tmp_path)
    rules_path = _write_rules(
        tmp_path,
        [
            _rule(
                "REGRA_EMBASA",
                "4.01.012",
                "Agua e Esgoto",
                "06.02",
                beneficiary=["EMBASA"],
                reference=["AGUA"],
            )
        ],
    )
    input_dir = (
        tmp_path
        / "bank_raw"
        / "_audit"
        / "fase_13c"
        / "2026-02"
        / "08884180000123"
        / "13a_import"
        / "08884180000123"
    )
    input_dir.mkdir(parents=True)
    _write_receipts_csv_to(
        input_dir / "payment_receipts_normalized.csv",
        [
            _receipt("RC-1", "Embasa", "Agua mensal"),
            _receipt("RC-2", "Fornecedor", "Sem regra"),
        ],
    )

    manifest = execute_payment_classification_run(
        entrada_raiz=tmp_path / "bank_raw" / "_audit" / "fase_13c",
        saida_raiz=tmp_path / "reports" / "comprovantes_classificacao",
        regras_path=rules_path,
        plano_contas_path=plano_path,
        mapeamento_dre_path=mapeamento_path,
        today=date(2026, 5, 18),
    )

    output_dir = tmp_path / "reports" / "comprovantes_classificacao" / "runs" / "2026-05-18_001"
    classified = _read_csv(output_dir / "comprovantes_classificados.csv")
    review = _read_csv(output_dir / "comprovantes_a_revisar.csv")
    summary_account = _read_csv(output_dir / "resumo_por_conta.csv")
    summary_cnpj_month = _read_csv(output_dir / "resumo_por_cnpj_mes.csv")
    manifest_json = json.loads(
        (output_dir / "manifest_classificacao_comprovantes.json").read_text(
            encoding="utf-8"
        )
    )

    assert manifest.run_id == "2026-05-18_001"
    assert manifest.total_comprovantes == 2
    assert manifest.total_classificados == 1
    assert classified[0]["conta_id"] == "4.01.012"
    assert review[0]["status_classificacao"] == STATUS_SEM_REGRA
    assert review[0]["DECISAO_OPERADOR"] == ""
    assert summary_account[0]["quantidade"] == "1"
    assert summary_account[0]["valor_total"] == "20.00"
    assert summary_cnpj_month[0]["total_comprovantes"] == "2"
    assert manifest_json["fase"] == "13S"
    assert manifest_json["modo"] == "provisorio"
    assert manifest_json["matching_automatico_executado"] is False
    assert manifest_json["nao_inferir_conta"] is True


def test_execute_payment_classification_run_keeps_exclusions_out_of_account_and_dre(
    tmp_path,
):
    plano_path, mapeamento_path = _write_plan_and_mapping(tmp_path)
    rules_path = _write_rules_with_exclusions(
        tmp_path,
        positive_rules=[
            _rule(
                "REGRA_EMBASA",
                "4.01.012",
                "Agua e Esgoto",
                "06.02",
                beneficiary=["EMBASA"],
            )
        ],
        exclusions=[
            _exclusion_rule(
                "EXCLUSAO_XML_FORNECEDOR",
                "XML",
                beneficiary=["FORNECEDOR XML"],
            )
        ],
    )
    input_dir = (
        tmp_path
        / "bank_raw"
        / "_audit"
        / "fase_13c"
        / "2026-02"
        / "08884180000123"
        / "13a_import"
        / "08884180000123"
    )
    input_dir.mkdir(parents=True)
    _write_receipts_csv_to(
        input_dir / "payment_receipts_normalized.csv",
        [
            _receipt("RC-1", "Embasa", "Agua mensal"),
            _receipt("RC-2", "Fornecedor XML", "-"),
            _receipt("RC-3", "Fornecedor sem regra", "-"),
        ],
    )

    manifest = execute_payment_classification_run(
        entrada_raiz=tmp_path / "bank_raw" / "_audit" / "fase_13c",
        saida_raiz=tmp_path / "reports" / "comprovantes_classificacao",
        regras_path=rules_path,
        plano_contas_path=plano_path,
        mapeamento_dre_path=mapeamento_path,
        today=date(2026, 5, 18),
    )

    output_dir = tmp_path / "reports" / "comprovantes_classificacao" / "runs" / "2026-05-18_001"
    classified = _read_csv(output_dir / "comprovantes_classificados.csv")
    excluded = _read_csv(output_dir / "comprovantes_excluidos.csv")
    review = _read_csv(output_dir / "comprovantes_a_revisar.csv")
    summary_account = _read_csv(output_dir / "resumo_por_conta.csv")
    summary_dre = _read_csv(output_dir / "resumo_por_dre.csv")

    assert manifest.total_classificados == 1
    assert manifest.total_excluidos_xml == 1
    assert manifest.total_excluidos_escludere == 0
    assert manifest.total_a_revisar_sem_regra == 1
    assert classified[0]["payment_id"] == "RC-1"
    assert excluded[0]["payment_id"] == "RC-2"
    assert excluded[0]["status_classificacao"] == STATUS_EXCLUIDO
    assert excluded[0]["acao"] == "XML"
    assert "conta_id" not in excluded[0]
    assert review[0]["payment_id"] == "RC-3"
    assert sum(int(row["quantidade"]) for row in summary_account) == 1
    assert sum(int(row["quantidade"]) for row in summary_dre) == 1


def test_gerar_pacote_revisao_comprovantes_splits_sispag_and_guides_review(tmp_path):
    run_dir = tmp_path / "reports" / "comprovantes_classificacao" / "runs" / "2026-05-18_001"
    run_dir.mkdir(parents=True)
    csv_path = run_dir / "comprovantes_a_revisar.csv"
    _write_review_csv(
        csv_path,
        [
            _review_row("RC-1", "FORNECEDOR TESTE", "MANUTENCAO", "100.00"),
            _review_row("RC-2", "FORNECEDOR TESTE", "MANUTENCAO EXTRA", "50.00"),
            _review_row("RC-3", "SISPAG FORNECEDORES", "SISPAG PAGAMENTOS", "200.00"),
        ],
    )
    yaml_path = tmp_path / "regras_classificacao_comprovantes.yaml"
    yaml_path.write_text("regras: []\n", encoding="utf-8")
    original_yaml = yaml_path.read_text(encoding="utf-8")

    result = gerar_pacote_revisao_comprovantes(run_dir)

    workbook = load_workbook(result.xlsx_path)
    assert REVISAO_SHEET in workbook.sheetnames
    assert RESUMO_BENEFICIARIO_SHEET in workbook.sheetnames
    assert SISPAG_SHEET in workbook.sheetnames

    revisao = workbook[REVISAO_SHEET]
    headers = [cell.value for cell in revisao[1]]
    assert "DECISAO_OPERADOR" in headers
    assert "CRIAR_REGRA_FUTURA" in headers
    assert "TIPO_REGRA_SUGERIDA" in headers
    assert revisao.max_row == 3
    assert all(
        "SISPAG" not in str(revisao.cell(row=row, column=4).value)
        for row in range(2, revisao.max_row + 1)
    )

    resumo = workbook[RESUMO_BENEFICIARIO_SHEET]
    resumo_rows = list(resumo.iter_rows(min_row=2, values_only=True))
    assert resumo_rows[0][0] == "FORNECEDOR TESTE"
    assert resumo_rows[0][1] == 2
    assert resumo_rows[0][2] == 150

    sispag = workbook[SISPAG_SHEET]
    sispag_headers = [cell.value for cell in sispag[1]]
    status_col = sispag_headers.index("status_classificacao") + 1
    criar_regra_col = sispag_headers.index("CRIAR_REGRA_FUTURA") + 1
    conta_col = sispag_headers.index("CONTA_ID_APROVADA") + 1
    assert sispag.max_row == 2
    assert sispag.cell(row=2, column=status_col).value == FORA_ESCOPO_SISPAG
    assert sispag.cell(row=2, column=criar_regra_col).value == "nao"
    assert sispag.cell(row=2, column=conta_col).value is None

    instrucoes_text = "\n".join(
        str(cell.value or "") for cell in workbook["INSTRUCOES_OPERADOR"]["A"]
    )
    assert "CRIAR_REGRA_FUTURA" in instrucoes_text
    assert "Alias vs keyword" in instrucoes_text
    assert "matching automatico bancario nao executado" in instrucoes_text

    report_text = result.report_path.read_text(encoding="utf-8")
    assert "nenhuma regra foi criada automaticamente" in report_text
    assert "matching automatico bancario nao executado" in report_text
    assert result.matching_automatico_executado is False
    assert result.yaml_modificado is False
    assert yaml_path.read_text(encoding="utf-8") == original_yaml


def _rule(
    regra_id: str,
    conta_id: str,
    conta_nome: str,
    dre_linha_id: str,
    *,
    beneficiary: list[str] | None = None,
    reference: list[str] | None = None,
    entra_dre: bool = True,
    origem_linha: int | None = None,
) -> str:
    beneficiary_items = _yaml_list(beneficiary or [])
    reference_items = _yaml_list(reference or [])
    origem_text = (
        f"  origem_linha_planilha: {origem_linha}\n"
        if origem_linha is not None
        else ""
    )
    return f"""- regra_id: {regra_id}
{origem_text}  conta_id: "{conta_id}"
  conta_nome: {conta_nome}
  dre_linha_id: "{dre_linha_id}"
  entra_dre: {str(entra_dre).lower()}
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
{beneficiary_items}
    referencia_empresa_contem_algum:
{reference_items}
"""


def _yaml_list(values: list[str]) -> str:
    if not values:
        return "      []"
    return "\n".join(f"      - {value}" for value in values)


def _write_rules(tmp_path: Path, rules: list[str]) -> Path:
    path = tmp_path / "regras.yaml"
    if not rules:
        path.write_text("regras: []\n", encoding="utf-8")
        return path
    rules_text = "\n".join(rules)
    path.write_text(f"regras:\n{rules_text}\n", encoding="utf-8")
    return path


def _write_rules_with_exclusions(
    tmp_path: Path,
    *,
    positive_rules: list[str],
    exclusions: list[str],
) -> Path:
    path = tmp_path / "regras.yaml"
    if positive_rules:
        text = "regras:\n" + "\n".join(positive_rules) + "\n"
    else:
        text = "regras: []\n"
    if exclusions:
        text += "exclusoes:\n" + "\n".join(exclusions) + "\n"
    else:
        text += "exclusoes: []\n"
    path.write_text(text, encoding="utf-8")
    return path


def _exclusion_rule(
    regra_id: str,
    acao: str,
    *,
    beneficiary: list[str] | None = None,
    reference: list[str] | None = None,
) -> str:
    beneficiary_items = _yaml_list(beneficiary or [])
    reference_items = _yaml_list(reference or [])
    return f"""- regra_id: {regra_id}
  acao: {acao}
  match:
    modo: TODAS_AS_CONDICOES_NAO_VAZIAS
    beneficiario_favorecido_contem_algum:
{beneficiary_items}
    referencia_empresa_contem_algum:
{reference_items}
"""


def _write_plan_and_mapping(tmp_path: Path) -> tuple[Path, Path]:
    plano_path = tmp_path / "plano.yaml"
    plano_path.write_text(
        """contas:
- id: "4.01.012"
  nome: Agua e Esgoto
  grupo: Ocupacao
  tipo: despesa
- id: "6.01.001"
  nome: Honorarios Contabeis
  grupo: Administrativo
  tipo: despesa
- id: "7.01.001"
  nome: Tarifas Bancarias
  grupo: Financeiro
  tipo: despesa
- id: "7.01.003"
  nome: Tarifas PIX
  grupo: Financeiro
  tipo: despesa
""",
        encoding="utf-8",
    )
    mapeamento_path = tmp_path / "mapeamento.yaml"
    mapeamento_path.write_text(
        """mapeamento:
  "4.01.012":
    dre_linha_id: "06.02"
    entra_dre: true
  "6.01.001":
    dre_linha_id: "08.01"
    entra_dre: true
  "7.01.001":
    dre_linha_id: "09.01"
    entra_dre: true
  "7.01.003":
    dre_linha_id: "09.01"
    entra_dre: true
""",
        encoding="utf-8",
    )
    return plano_path, mapeamento_path


def _receipt(payment_id: str, beneficiary: str, reference: str) -> dict[str, str]:
    return {
        "payment_id": payment_id,
        "cnpj": "08884180000123",
        "payment_date": "2026-02-05",
        "amount": "20.00",
        "beneficiary_name": beneficiary,
        "beneficiary_document": "12.345.678/0001-90",
        "payment_type": "PIX",
        "reference_id": reference,
        "status": "Efetuado",
        "source_file": "ConsultaPagamentos.xls",
        "raw_hash": payment_id,
    }


def _write_receipts_csv(tmp_path: Path, rows: list[dict[str, str]]) -> Path:
    path = tmp_path / "payment_receipts_normalized.csv"
    _write_receipts_csv_to(path, rows)
    return path


def _write_receipts_csv_to(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(rows[0].keys()), delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def _review_row(
    payment_id: str,
    beneficiary: str,
    reference: str,
    amount: str,
) -> dict[str, str]:
    return {
        "cnpj": "08884180000123",
        "mes": "2026-02",
        "data_pagamento": "2026-02-05",
        "beneficiario_normalizado": beneficiary,
        "referencia": reference,
        "valor": amount,
        "arquivo_origem": "payment_receipts_normalized.csv",
        "status_classificacao": STATUS_SEM_REGRA,
        "motivo_revisao": "SEM_REGRA",
        "regras_candidatas": "",
        "contas_candidatas": "",
        "DECISAO_OPERADOR": "",
        "CONTA_ID_APROVADA": "",
        "CRIAR_REGRA_FUTURA": "",
        "ALIAS_O_KEYWORD_APROVADO": "",
        "OBSERVACOES": "",
        "payment_id": payment_id,
        "beneficiary_name": beneficiary,
        "beneficiary_document": "",
        "payment_type": "PIX",
        "status_pagamento": "Efetuado",
        "source_file": "ConsultaPagamentos.xls",
        "raw_hash": payment_id,
    }


def _write_review_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(rows[0].keys()), delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file, delimiter=";"))
