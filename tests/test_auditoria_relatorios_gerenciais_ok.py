from datetime import date
from decimal import Decimal
from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook

from fantozzi.auditoria_relatorios_gerenciais_ok import (
    AMOSTRA_EXCLUIDOS_SHEET,
    AUDITORIA_GERAL_SHEET,
    INCONSISTENCIAS_SHEET,
    INSTRUCOES_SHEET,
    RECONCILIACAO_OPERACOES_SHEET,
    RECONCILIACAO_STATUS_SHEET,
    SOBREPOSICOES_SHEET,
    TOTAIS_ORIGEM_SHEET,
    auditar_relatorios_gerenciais_ok_provisorios,
)
from fantozzi.excel_writer import write_lancamentos_to_excel
from fantozzi.models import Duplicata, Lancamento
from fantozzi.nfe_importer import parse_nfe


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_XML = ROOT / "samples" / "nfe_exemplo" / "nfe_oesa_12566_20260505.xml"


def test_calcula_status_ok_revisar_e_reconciliacao(tmp_path):
    excel_path = tmp_path / "55746722000181" / "lancamentos.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento("ok-1"),
            _lancamento("revisar-1", status_revisao="REVISAR"),
        ],
    )

    result = auditar_relatorios_gerenciais_ok_provisorios(
        [excel_path], tmp_path / "auditoria.xlsx"
    )

    assert result.total_lancamentos_origem == 2
    assert result.total_status_ok == 1
    assert result.total_status_revisar == 1
    assert result.total_incluidos_ok_nao_especiais == 1
    assert result.resultado_auditoria == "ATENCAO"


def test_calcula_ok_excluido_por_comodato(tmp_path):
    xml_dir = _xml_dir_with_special_operation(tmp_path, natureza="Comodato", cfop="5908")
    nfe = parse_nfe(next(xml_dir.glob("*.xml")))
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(excel_path, [_lancamento("ok-comodato", chave_origem=nfe.chave)])

    result = auditar_relatorios_gerenciais_ok_provisorios(
        [excel_path], tmp_path / "auditoria.xlsx", pastas_operacoes_especiais=[xml_dir]
    )

    assert result.total_status_ok == 1
    assert result.total_incluidos_ok_nao_especiais == 0
    assert result.total_ok_excluidos_por_operacao_especial == 1
    assert result.total_operacoes_especiais == 1


def test_calcula_ok_excluido_por_bonificacao(tmp_path):
    xml_dir = _xml_dir_with_special_operation(
        tmp_path, natureza="Bonificacao recebida", cfop="5910"
    )
    nfe = parse_nfe(next(xml_dir.glob("*.xml")))
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(excel_path, [_lancamento("ok-bonif", chave_origem=nfe.chave)])

    result = auditar_relatorios_gerenciais_ok_provisorios(
        [excel_path], tmp_path / "auditoria.xlsx", pastas_operacoes_especiais=[xml_dir]
    )

    assert result.total_ok_excluidos_por_operacao_especial == 1
    assert result.valor_total_incluido_ok_nao_especial == Decimal("0.00")


def test_calcula_revisar_com_operacao_especial(tmp_path):
    xml_dir = _xml_dir_with_special_operation(tmp_path, natureza="Comodato", cfop="5908")
    nfe = parse_nfe(next(xml_dir.glob("*.xml")))
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [_lancamento("revisar-comodato", status_revisao="REVISAR", chave_origem=nfe.chave)],
    )

    result = auditar_relatorios_gerenciais_ok_provisorios(
        [excel_path], tmp_path / "auditoria.xlsx", pastas_operacoes_especiais=[xml_dir]
    )

    assert result.total_revisar_com_operacao_especial == 1
    assert result.total_ok_excluidos_por_operacao_especial == 0


def test_gera_xlsx_com_abas_obrigatorias(tmp_path):
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(excel_path, [_lancamento("ok-1")])
    xlsx_path = tmp_path / "auditoria.xlsx"

    auditar_relatorios_gerenciais_ok_provisorios([excel_path], xlsx_path)

    workbook = load_workbook(xlsx_path, data_only=True)
    assert workbook.sheetnames == [
        AUDITORIA_GERAL_SHEET,
        RECONCILIACAO_STATUS_SHEET,
        RECONCILIACAO_OPERACOES_SHEET,
        SOBREPOSICOES_SHEET,
        TOTAIS_ORIGEM_SHEET,
        INCONSISTENCIAS_SHEET,
        "Amostra_Incluidos_OK",
        AMOSTRA_EXCLUIDOS_SHEET,
        INSTRUCOES_SHEET,
    ]


def test_gera_txt_com_resultado_auditoria(tmp_path):
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(excel_path, [_lancamento("ok-1")])
    report_path = tmp_path / "auditoria.txt"

    auditar_relatorios_gerenciais_ok_provisorios(
        [excel_path], tmp_path / "auditoria.xlsx", report_path
    )

    text = report_path.read_text(encoding="utf-8")
    assert "AUDITORIA DOS RELATORIOS GERENCIAIS OK PROVISORIOS" in text
    assert "Resultado auditoria:\nOK" in text


def _xml_dir_with_special_operation(tmp_path: Path, *, natureza: str, cfop: str) -> Path:
    xml_dir = tmp_path / f"xml_{cfop}"
    xml_dir.mkdir()
    copyfile(SAMPLE_XML, xml_dir / "especial.xml")
    xml_path = xml_dir / "especial.xml"
    text = xml_path.read_text(encoding="utf-8")
    text = text.replace("<natOp>VENDA DE MERCADORIAS</natOp>", f"<natOp>{natureza}</natOp>")
    text = text.replace("<CFOP>5102</CFOP>", f"<CFOP>{cfop}</CFOP>", 1)
    xml_path.write_text(text, encoding="utf-8")
    return xml_dir


def _lancamento(
    hash_registro: str,
    *,
    valor: Decimal = Decimal("10.00"),
    status_revisao: str = "OK",
    chave_origem: str = "NFE123",
) -> Lancamento:
    return Lancamento(
        data_documento=date(2026, 5, 5),
        competencia="2026-05",
        cnpj_empresa="55746722000181",
        nome_empresa="TRATTORIA DO SOSSEGO RESTAURANTE LTDA",
        fornecedor_ou_favorecido="Fornecedor teste",
        cnpj_cpf_favorecido="12345678000195",
        documento="12566",
        tipo_documento="XML_NFE",
        conta_id="2.01.004",
        conta_nome="Refrigerantes",
        grupo_conta="Custos de Mercadorias e Insumos",
        descricao_itens="COCA-COLA KS 290ML",
        valor=valor,
        duplicatas=[Duplicata(numero="001", vencimento=date(2026, 5, 20), valor=valor)],
        origem="XML_NFE",
        arquivo_origem="nfe.xml",
        chave_origem=chave_origem,
        hash_registro=hash_registro,
        status_revisao=status_revisao,
        alertas="" if status_revisao == "OK" else "CLASSIFICACAO_NAO_ENCONTRADA",
        tipo_lancamento="custo",
    )
