from datetime import date
from decimal import Decimal
import json

from openpyxl import Workbook, load_workbook

from fantozzi.excel_writer import HEADERS, write_lancamentos_to_excel
from fantozzi.models import Duplicata, Lancamento


def test_writes_new_excel_file_with_lancamentos_sheet_and_headers(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"

    write_lancamentos_to_excel(excel_path, [_lancamento("hash-001")])

    workbook = load_workbook(excel_path)
    assert excel_path.exists()
    assert "Lancamentos" in workbook.sheetnames
    assert [cell.value for cell in workbook["Lancamentos"][1]] == HEADERS


def test_does_not_duplicate_same_lancamentos_when_written_twice(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"
    lancamentos = [_lancamento("hash-001"), _lancamento("hash-002", conta_id="2.01.002")]

    write_lancamentos_to_excel(excel_path, lancamentos)
    write_lancamentos_to_excel(excel_path, lancamentos)

    sheet = load_workbook(excel_path)["Lancamentos"]
    assert sheet.max_row == 3


def test_preserves_existing_manual_row_in_lancamentos_sheet(tmp_path):
    excel_path = tmp_path / "manual.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lancamentos"
    sheet.append(HEADERS)
    manual_row = _manual_row()
    sheet.append(manual_row)
    workbook.save(excel_path)

    write_lancamentos_to_excel(excel_path, [_lancamento("hash-001")])

    sheet = load_workbook(excel_path)["Lancamentos"]
    rows = list(sheet.iter_rows(values_only=True))
    _assert_manual_row_preserved(rows[1], manual_row)
    assert sheet.max_row == 3


def test_revisao_sheet_receives_review_lancamentos(tmp_path):
    excel_path = tmp_path / "revisao.xlsx"
    lancamento = _lancamento(
        "hash-revisar", status_revisao="REVISAR", alertas="NAO_CLASSIFICADO"
    )

    write_lancamentos_to_excel(excel_path, [lancamento])

    sheet = load_workbook(excel_path)["Revisao"]
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(sheet.iter_rows(min_row=2, values_only=True))))

    assert sheet.max_row == 2
    assert row["hash_registro"] == "hash-revisar"
    assert row["status_revisao"] == "REVISAR"
    assert row["alertas"] == "NAO_CLASSIFICADO"


def test_preserves_manual_row_in_revisao_and_avoids_review_duplicates(tmp_path):
    excel_path = tmp_path / "revisao_manual.xlsx"
    workbook = Workbook()
    lancamentos_sheet = workbook.active
    lancamentos_sheet.title = "Lancamentos"
    lancamentos_sheet.append(HEADERS)
    revisao_sheet = workbook.create_sheet("Revisao")
    revisao_sheet.append(HEADERS)
    manual_row = _manual_row(status_revisao="REVISAR", alertas="MANUAL")
    revisao_sheet.append(manual_row)
    workbook.save(excel_path)

    lancamento = _lancamento(
        "hash-revisar", status_revisao="REVISAR", alertas="NAO_CLASSIFICADO"
    )
    write_lancamentos_to_excel(excel_path, [lancamento])
    write_lancamentos_to_excel(excel_path, [lancamento])

    sheet = load_workbook(excel_path)["Revisao"]
    rows = list(sheet.iter_rows(values_only=True))
    _assert_manual_row_preserved(rows[1], manual_row)
    assert sheet.max_row == 3


def test_uses_stable_fallback_key_when_hash_is_empty(tmp_path):
    excel_path = tmp_path / "fallback.xlsx"
    lancamento = _lancamento("")

    write_lancamentos_to_excel(excel_path, [lancamento])
    write_lancamentos_to_excel(excel_path, [lancamento])

    sheet = load_workbook(excel_path)["Lancamentos"]
    assert sheet.max_row == 2


def test_writes_duplicatas_json(tmp_path):
    excel_path = tmp_path / "duplicatas.xlsx"

    write_lancamentos_to_excel(excel_path, [_lancamento("hash-001")])

    sheet = load_workbook(excel_path)["Lancamentos"]
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(sheet.iter_rows(min_row=2, values_only=True))))
    duplicatas = json.loads(row["duplicatas_json"])

    assert duplicatas == [
        {"numero": "001", "valor": "5.00", "vencimento": "2026-05-20"}
    ]


def _lancamento(
    hash_registro: str,
    *,
    conta_id: str = "2.01.004",
    status_revisao: str = "OK",
    alertas: str = "",
) -> Lancamento:
    return Lancamento(
        data_documento=date(2026, 5, 5),
        competencia="2026-05",
        cnpj_empresa="55746722000181",
        nome_empresa="TRATTORIA DO SOSSEGO RESTAURANTE LTDA",
        fornecedor_ou_favorecido="OESA COMERCIO E REPRESENTACOES S/A",
        cnpj_cpf_favorecido="12345678000195",
        documento="12566",
        tipo_documento="XML_NFE",
        conta_id=conta_id,
        conta_nome="Refrigerantes",
        grupo_conta="Custos de Mercadorias e Insumos",
        descricao_itens="COCA-COLA KS 290ML",
        valor=Decimal("10.00"),
        duplicatas=[
            Duplicata(
                numero="001",
                vencimento=date(2026, 5, 20),
                valor=Decimal("5.00"),
            )
        ],
        origem="XML_NFE",
        arquivo_origem="nfe.xml",
        chave_origem="NFE123",
        hash_registro=hash_registro,
        status_revisao=status_revisao,
        alertas=alertas,
        tipo_lancamento="custo",
    )


def _manual_row(
    *, status_revisao: str = "OK", alertas: str = ""
) -> list[object]:
    return [
        date(2026, 5, 1),
        "2026-05",
        "manual-cnpj",
        "Manual Empresa",
        "Manual Favorecido",
        "manual-doc",
        "MAN-001",
        "MANUAL",
        "manual-conta",
        "Manual Conta",
        "Manual Grupo",
        "Manual item",
        Decimal("1.00"),
        "MANUAL",
        "",
        "",
        "",
        status_revisao,
        alertas,
        "linha manual preservada",
        "",
        "manual",
        "",
        "",
        None,
        "[]",
    ]


def _assert_manual_row_preserved(saved_row, expected_row):
    assert saved_row[0].date() == expected_row[0]
    assert saved_row[1:12] == tuple(expected_row[1:12])
    assert Decimal(str(saved_row[12])).quantize(Decimal("0.01")) == expected_row[12]
    assert tuple(_blank_to_empty(value) for value in saved_row[13:]) == tuple(
        _blank_to_empty(value) for value in expected_row[13:]
    )


def _blank_to_empty(value):
    return "" if value is None else value
