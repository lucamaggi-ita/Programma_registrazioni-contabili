from __future__ import annotations

import json
import shutil
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from fantozzi.cli import main
from fantozzi.lancamentos_manuais_template import (
    gerar_template_lancamentos_manuais_gerenciais,
)


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PLANO_REAL = ROOT / "config" / "plano_contas_gerencial.yaml"
CONFIG_DRE_REAL = ROOT / "config" / "mapeamento_dre.yaml"
CONFIG_PRODUTOS_REAL = ROOT / "config" / "regras_classificacao_produtos.yaml"
SAMPLE_XML = ROOT / "samples" / "nfe_exemplo" / "nfe_oesa_12566_20260505.xml"

CONTA_MANUAL_ID = "8.01.001"
CONTA_MANUAL_NOME = "Manutencao Predial"
CONTA_TECNICA_ID = "99.99.998"
VALOR_MANUAL = Decimal("321.45")
DESCRICAO_MANUAL = "Conserto emergencial forno - teste e2e"

def test_fluxo_manuais_gerenciais_cli_end_to_end(tmp_path):
    project_root = _preparar_project_root_com_config_real(tmp_path)
    xml_base = _preparar_xml_base(project_root)
    template_path = _preparar_template_preenchido(project_root)
    saida_csv = (
        project_root
        / "Database_Dev"
        / "Registrazioni_manuali"
        / "registrazioni_manuali_tecnico.csv"
    )
    relatorio_validacao = project_root / "reports" / "validacao_manuais.txt"
    saida_output = project_root / "output" / "importacao_real"
    saida_reports = project_root / "reports" / "importacao_real"

    result = main(
        [
            "executar-lancamentos-manuais-gerenciais",
            "--project-root",
            str(project_root),
            "--template",
            str(template_path),
            "--saida-csv",
            str(saida_csv),
            "--xml-base",
            str(xml_base),
            "--saida-output",
            str(saida_output),
            "--saida-reports",
            str(saida_reports),
            "--plano-contas",
            str(project_root / "config" / "plano_contas_gerencial.yaml"),
            "--mapeamento-dre",
            str(project_root / "config" / "mapeamento_dre.yaml"),
            "--relatorio-validacao",
            str(relatorio_validacao),
        ]
    )

    assert result == 0
    assert saida_csv.exists()
    assert relatorio_validacao.exists()

    run_id = (saida_reports / "current_run.txt").read_text(encoding="utf-8").strip()
    reports_run_dir = saida_reports / "runs" / run_id
    manifest_path = reports_run_dir / "manifest_run.json"

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["status_run"] == "PROVISORIO_CONCLUIDO"
    manuais = manifest["lancamentos_manuais_gerenciais"]
    assert manuais["executado"] is True
    assert manuais["lancamentos_aceitos"] == 1
    assert manuais["fonte"] == "MANUAL_GERENCIAL"
    assert manifest["pacote_contabil_gerencial"]["executado"] is True

    for path_text in manifest["relatorios_gerados"]:
        path = Path(path_text)
        assert path.resolve().is_relative_to(project_root.resolve())

    pre_livro_path = reports_run_dir / "pre_livro_diario_gerencial_teste.xlsx"
    assert pre_livro_path.exists()
    manuais_rows = _linhas_manuais_do_pre_livro(pre_livro_path)
    assert {row["conta_id"] for row in manuais_rows} == {CONTA_MANUAL_ID, CONTA_TECNICA_ID}
    for row in manuais_rows:
        assert Decimal(str(row["valor"])) == VALOR_MANUAL
        assert row["status_revisao"] == "OK"

    dre_path = reports_run_dir / "dre_gerencial_ok_provisoria.xlsx"
    assert dre_path.exists()
    dre_workbook = load_workbook(dre_path)
    dre_geral = dre_workbook["DRE_Geral"]
    dre_rows = list(dre_geral.iter_rows(min_row=2, values_only=True))
    dre_grupos = {row[1] for row in dre_rows}
    dre_by_linha = {row[2]: row[3] for row in dre_rows}

    assert "Controles Tecnicos" not in dre_grupos
    assert "Manutencao e Reparos" in dre_by_linha
    assert Decimal(str(dre_by_linha["Manutencao e Reparos"])) == VALOR_MANUAL


def _preparar_project_root_com_config_real(tmp_path: Path) -> Path:
    project_root = tmp_path / "project_root"
    config_dir = project_root / "config"
    config_dir.mkdir(parents=True)
    shutil.copy2(CONFIG_PLANO_REAL, config_dir / "plano_contas_gerencial.yaml")
    shutil.copy2(CONFIG_DRE_REAL, config_dir / "mapeamento_dre.yaml")
    shutil.copy2(CONFIG_PRODUTOS_REAL, config_dir / "regras_classificacao_produtos.yaml")
    return project_root


def _preparar_xml_base(tmp_path: Path) -> Path:
    xml_base = tmp_path / "xml_base"
    xml_dir = xml_base / "08884180000123" / "xml"
    xml_dir.mkdir(parents=True)
    shutil.copy2(SAMPLE_XML, xml_dir / SAMPLE_XML.name)
    return xml_base


def _preparar_template_preenchido(project_root: Path) -> Path:
    plano_path = project_root / "config" / "plano_contas_gerencial.yaml"
    dre_path = project_root / "config" / "mapeamento_dre.yaml"
    template_path = (
        project_root
        / "Database_Dev"
        / "Registrazioni_manuali"
        / "registrazioni_manuali_base.xlsx"
    )
    gerar_template_lancamentos_manuais_gerenciais(template_path, plano_path, dre_path)

    workbook = load_workbook(template_path)
    sheet = workbook["lancamentos_manuais"]
    sheet["A2"] = "2026-05-20"
    sheet["B2"] = f"{CONTA_MANUAL_ID} - {CONTA_MANUAL_NOME}"
    sheet["C2"] = CONTA_MANUAL_ID
    sheet["D2"] = CONTA_MANUAL_NOME
    sheet["E2"] = "321,45"
    sheet["F2"] = DESCRICAO_MANUAL
    sheet["G2"] = "lancamento sintetico de teste e2e"
    workbook.save(template_path)
    return template_path


def _linhas_manuais_do_pre_livro(pre_livro_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(pre_livro_path)
    sheet = workbook["Lancamentos_OK"]
    headers = [cell.value for cell in sheet[1]]
    origem_index = headers.index("origem")
    rows = []
    for raw_row in sheet.iter_rows(min_row=2, values_only=True):
        if raw_row[origem_index] == "MANUAL_GERENCIAL":
            rows.append(dict(zip(headers, raw_row)))
    assert len(rows) == 2, f"esperadas 2 linhas MANUAL_GERENCIAL, encontrada(s) {len(rows)}"
    return rows
