from __future__ import annotations

import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fantozzi.excel_writer import write_lancamentos_to_excel
from fantozzi.models import Lancamento
from fantozzi.runs_incrementais import (
    BLOCCO_BASELINE_AUSENTE,
    BLOCCO_BASELINE_INSUFICIENTE,
    BLOCCO_DELTA_JA_APLICADO,
    BLOCCO_STRUCTURAL_CHANGE,
    RunsIncrementaisError,
    executar_run_incremental_manuais,
    registrar_baseline_aprovada,
)


def test_incremental_manuais_nao_reprocessa_xml_banca_e_atualiza_dre(tmp_path):
    project_root, baseline_path, manuali_csv, manuali_manifest = _fixture_incremental(
        tmp_path
    )

    result = executar_run_incremental_manuais(
        baseline_path,
        manuali_csv,
        manuali_manifest,
        tmp_path / "incremental",
        project_root=project_root,
    )

    manifest = _load_json(result.manifest_path)
    assert manifest["full_run_executed"] is False
    assert manifest["fontes_reprocessadas"]["xml"] is False
    assert manifest["fontes_reprocessadas"]["banca"] is False
    assert manifest["fontes_reprocessadas"]["manuali"] is True
    assert manifest["fontes_reutilizadas"]["xml"] is True
    assert manifest["fontes_reutilizadas"]["banca"] is True
    assert manifest["manuali"]["linhas"] == 2
    assert manifest["manuali"]["valor_total"] == "125.30"
    assert manifest["manuali"]["competencias_afetadas"] == ["2026-04"]
    assert manifest["dre_status"] == "OK_PROVISORIO"

    baseline_atualizada = _load_json(baseline_path)
    assert baseline_atualizada["source_run_path"] == str(result.output_run_dir)
    assert baseline_atualizada["dre_path"] == str(Path(manifest["outputs"]["dre"]))
    assert baseline_atualizada["periodo_fim"] == "2026-04-30"

    base_rows = _dre_base_rows(Path(manifest["outputs"]["dre"]))
    conta_values = [
        Decimal(str(row["valor"]))
        for row in base_rows
        if row["conta_id"] in {"2.02.008", "8.03.001"}
    ]
    assert abs(sum(conta_values, Decimal("0.00"))) == Decimal("325.30")
    assert "99.99.998" not in _all_workbook_values(Path(manifest["outputs"]["dre"]))


def test_incremental_manuais_retroativos_acumulam_nova_versao_baseline(tmp_path):
    project_root, baseline_path, manuali_csv, manuali_manifest = _fixture_incremental(
        tmp_path
    )
    first = executar_run_incremental_manuais(
        baseline_path,
        manuali_csv,
        manuali_manifest,
        tmp_path / "incremental_1",
        project_root=project_root,
    )
    _write_manuali(
        project_root,
        rows=[
            (
                "2026-04-12",
                "Manual retroativo adicional",
                "2.02.008",
                "Hortifruti",
                "10.00",
            )
        ],
    )

    second = executar_run_incremental_manuais(
        baseline_path,
        manuali_csv,
        manuali_manifest,
        tmp_path / "incremental_2",
        project_root=project_root,
    )

    assert first.output_run_dir != second.output_run_dir
    baseline_atualizada = _load_json(baseline_path)
    assert baseline_atualizada["source_run_path"] == str(second.output_run_dir)
    assert len(baseline_atualizada["applied_manuali_hashes"]) == 2

    manifest = _load_json(second.manifest_path)
    base_rows = _dre_base_rows(Path(manifest["outputs"]["dre"]))
    manual_values = [
        Decimal(str(row["valor"]))
        for row in base_rows
        if "MANUAL_GERENCIAL_INCREMENTAL" in str(row.get("origem_arquivo") or "")
    ]
    assert len(manual_values) == 3
    assert abs(sum(manual_values, Decimal("0.00"))) == Decimal("135.30")


def test_incremental_manuais_retroativos_dentro_baseline_sao_aceitos(tmp_path):
    project_root, baseline_path, manuali_csv, manuali_manifest = _fixture_incremental(
        tmp_path
    )
    baseline = _load_json(baseline_path)
    baseline["periodo_ini"] = "2026-01-01"
    baseline["periodo_fim"] = "2026-04-30"
    baseline_path.write_text(json.dumps(baseline), encoding="utf-8")
    _write_manuali(
        project_root,
        rows=[
            (
                "2026-02-15",
                "Manual retroativo dentro baseline",
                "2.02.008",
                "Hortifruti",
                "42.00",
            )
        ],
    )

    result = executar_run_incremental_manuais(
        baseline_path,
        manuali_csv,
        manuali_manifest,
        tmp_path / "incremental_retro",
        project_root=project_root,
    )

    manifest = _load_json(result.manifest_path)
    assert manifest["full_run_executed"] is False
    assert manifest["fontes_reprocessadas"]["xml"] is False
    assert manifest["fontes_reprocessadas"]["banca"] is False
    assert manifest["fontes_reprocessadas"]["comprovantes"] is False
    assert manifest["fontes_reprocessadas"]["conta_corrente"] is False
    assert manifest["manuali"]["periodo_ini"] == "2026-02-15"
    assert manifest["manuali"]["periodo_fim"] == "2026-02-15"
    assert manifest["manuali"]["competencias_afetadas"] == ["2026-02"]
    assert manifest["dre_status"] == "OK_PROVISORIO"


def test_incremental_manuais_duas_competencias_retroativas_ordem_estavel(tmp_path):
    project_root, baseline_path, manuali_csv, manuali_manifest = _fixture_incremental(
        tmp_path
    )
    baseline = _load_json(baseline_path)
    baseline["periodo_ini"] = "2026-01-01"
    baseline["periodo_fim"] = "2026-04-30"
    baseline_path.write_text(json.dumps(baseline), encoding="utf-8")
    _write_manuali(
        project_root,
        rows=[
            (
                "2026-04-03",
                "Manual abril dentro baseline",
                "2.02.008",
                "Hortifruti",
                "25.30",
            ),
            (
                "2026-02-15",
                "Manual fevereiro dentro baseline",
                "8.03.001",
                "Combustivel",
                "100.00",
            ),
        ],
    )

    result = executar_run_incremental_manuais(
        baseline_path,
        manuali_csv,
        manuali_manifest,
        tmp_path / "incremental_duas_competencias",
        project_root=project_root,
    )

    manifest = _load_json(result.manifest_path)
    assert manifest["full_run_executed"] is False
    assert manifest["fontes_reprocessadas"]["manuali"] is True
    assert manifest["fontes_reprocessadas"]["xml"] is False
    assert manifest["fontes_reprocessadas"]["banca"] is False
    assert manifest["fontes_reprocessadas"]["comprovantes"] is False
    assert manifest["fontes_reprocessadas"]["conta_corrente"] is False
    assert manifest["manuali"]["competencias_afetadas"] == ["2026-02", "2026-04"]
    assert manifest["dre_status"] == "OK_PROVISORIO"


def test_incremental_manuais_data_antes_da_baseline_bloqueia_sem_promover(tmp_path):
    project_root, baseline_path, manuali_csv, manuali_manifest = _fixture_incremental(
        tmp_path
    )
    baseline = _load_json(baseline_path)
    baseline["periodo_ini"] = "2026-01-01"
    baseline["periodo_fim"] = "2026-04-30"
    baseline_path.write_text(json.dumps(baseline), encoding="utf-8")
    _write_manuali(
        project_root,
        rows=[
            (
                "2025-12-20",
                "Manual antes da baseline",
                "2.02.008",
                "Hortifruti",
                "10.00",
            )
        ],
    )

    with pytest.raises(RunsIncrementaisError) as exc_info:
        executar_run_incremental_manuais(
            baseline_path,
            manuali_csv,
            manuali_manifest,
            tmp_path / "incremental_antes",
            project_root=project_root,
        )

    assert exc_info.value.code == BLOCCO_BASELINE_INSUFICIENTE
    assert "fora do periodo coberto pela baseline" in exc_info.value.message
    assert _load_json(baseline_path) == baseline
    assert not (tmp_path / "incremental_antes").exists()


def test_incremental_manuais_data_depois_da_baseline_bloqueia_sem_full_run(tmp_path):
    project_root, baseline_path, manuali_csv, manuali_manifest = _fixture_incremental(
        tmp_path
    )
    baseline = _load_json(baseline_path)
    baseline["periodo_ini"] = "2026-01-01"
    baseline["periodo_fim"] = "2026-04-30"
    baseline_path.write_text(json.dumps(baseline), encoding="utf-8")
    _write_manuali(
        project_root,
        rows=[
            (
                "2026-05-10",
                "Manual depois da baseline",
                "2.02.008",
                "Hortifruti",
                "10.00",
            )
        ],
    )

    with pytest.raises(RunsIncrementaisError) as exc_info:
        executar_run_incremental_manuais(
            baseline_path,
            manuali_csv,
            manuali_manifest,
            tmp_path / "incremental_depois",
            project_root=project_root,
        )

    assert exc_info.value.code == BLOCCO_BASELINE_INSUFICIENTE
    assert "fora do periodo coberto pela baseline" in exc_info.value.message
    assert _load_json(baseline_path) == baseline
    assert not (tmp_path / "incremental_depois").exists()


def test_incremental_manuais_delta_ja_aplicado_nao_duplica(tmp_path):
    project_root, baseline_path, manuali_csv, manuali_manifest = _fixture_incremental(
        tmp_path
    )
    executar_run_incremental_manuais(
        baseline_path,
        manuali_csv,
        manuali_manifest,
        tmp_path / "incremental_1",
        project_root=project_root,
    )

    with pytest.raises(RunsIncrementaisError) as exc_info:
        executar_run_incremental_manuais(
            baseline_path,
            manuali_csv,
            manuali_manifest,
            tmp_path / "incremental_2",
            project_root=project_root,
        )

    assert exc_info.value.code == BLOCCO_DELTA_JA_APLICADO


def test_incremental_baseline_ausente_bloqueia_sem_full_run(tmp_path):
    project_root = _write_project_config(tmp_path)
    manuali_csv, manuali_manifest = _write_manuali(project_root)

    with pytest.raises(RunsIncrementaisError) as exc_info:
        executar_run_incremental_manuais(
            tmp_path / "missing_baseline.json",
            manuali_csv,
            manuali_manifest,
            tmp_path / "incremental",
            project_root=project_root,
        )

    assert exc_info.value.code == BLOCCO_BASELINE_AUSENTE


def test_incremental_structural_hash_mudou_bloqueia_sem_full_run(tmp_path):
    project_root, baseline_path, manuali_csv, manuali_manifest = _fixture_incremental(
        tmp_path
    )
    mapping_path = project_root / "config" / "mapeamento_dre.yaml"
    mapping_path.write_text(mapping_path.read_text(encoding="utf-8") + "\n# mudou\n")

    with pytest.raises(RunsIncrementaisError) as exc_info:
        executar_run_incremental_manuais(
            baseline_path,
            manuali_csv,
            manuali_manifest,
            tmp_path / "incremental",
            project_root=project_root,
        )

    assert exc_info.value.code == BLOCCO_STRUCTURAL_CHANGE


def test_registrar_baseline_aprovada_cria_manifest_incrementalizavel(tmp_path):
    project_root = _write_project_config(tmp_path)
    manifest_run = _write_baseline_run(project_root)
    baseline_path = tmp_path / "baseline_atual.json"

    result = registrar_baseline_aprovada(
        manifest_run,
        baseline_path,
        project_root=project_root,
    )

    baseline = _load_json(result.baseline_path)
    assert baseline["status"] == "APROVADA"
    assert baseline["run_id"] == "2026-05-28_001"
    assert baseline["periodo_ini"] == "2026-04-01"
    assert baseline["periodo_fim"] == "2026-04-30"
    assert baseline["inputs"]["xml"]["count"] == 1
    assert baseline["structural_hash"]
    assert baseline["applied_manuali_hashes"] == []


def _fixture_incremental(tmp_path: Path) -> tuple[Path, Path, Path, Path]:
    project_root = _write_project_config(tmp_path)
    manifest_run = _write_baseline_run(project_root)
    baseline_path = project_root / "output" / "baselines" / "baseline_atual.json"
    registrar_baseline_aprovada(manifest_run, baseline_path, project_root=project_root)
    manuali_csv, manuali_manifest = _write_manuali(project_root)
    return project_root, baseline_path, manuali_csv, manuali_manifest


def _write_project_config(tmp_path: Path) -> Path:
    project_root = tmp_path
    config = project_root / "config"
    config.mkdir(parents=True)
    review_accounts = "\n".join(
        f"- id: {account_id}\n  nome: Revisao {account_id}\n  grupo: Revisao\n  tipo: revisao"
        for account_id in (
            "13.01.001",
            "13.01.002",
            "13.01.006",
            "13.01.009",
            "13.01.010",
            "13.01.999",
        )
    )
    (config / "plano_contas_gerencial.yaml").write_text(
        "\n".join(
            [
                "contas:",
                "- id: 2.02.008",
                "  nome: Hortifruti",
                "  grupo: Custos",
                "  tipo: custo",
                "- id: 8.03.001",
                "  nome: Combustivel",
                "  grupo: Despesas",
                "  tipo: despesa",
                "- id: 99.99.998",
                "  nome: Contrapartida Tecnica Lancamentos Manuais",
                "  grupo: Tecnico",
                "  tipo: tecnico",
                review_accounts,
                "",
            ]
        ),
        encoding="utf-8",
    )
    (config / "mapeamento_dre.yaml").write_text(
        "\n".join(
            [
                "linhas_dre:",
                "- id: '03.01'",
                "  nome: CMV - Alimentos",
                "  grupo: Custos",
                "  sinal: '-'",
                "  ordem: 301",
                "- id: '06.01'",
                "  nome: Despesas com Veiculos",
                "  grupo: Despesas",
                "  sinal: '-'",
                "  ordem: 601",
                "mapeamento:",
                "  2.02.008:",
                "    dre_linha_id: '03.01'",
                "    entra_dre: true",
                "  8.03.001:",
                "    dre_linha_id: '06.01'",
                "    entra_dre: true",
                "  99.99.998:",
                "    dre_linha_id: ''",
                "    entra_dre: false",
                "",
            ]
        ),
        encoding="utf-8",
    )
    (config / "regras_classificacao_produtos.yaml").write_text(
        "regras: []\n",
        encoding="utf-8",
    )
    (config / "regras_classificacao_comprovantes.yaml").write_text(
        "regras: []\n",
        encoding="utf-8",
    )
    return project_root


def _write_baseline_run(project_root: Path) -> Path:
    run_id = "2026-05-28_001"
    output_root = project_root / "output" / "baseline_output"
    reports_root = project_root / "output" / "baseline_reports"
    source_run = output_root / "runs" / run_id / "08884180000123"
    source_run.mkdir(parents=True)
    write_lancamentos_to_excel(
        source_run / "lancamentos_fantozzi_real.xlsx",
        [
            Lancamento(
                data_documento=date(2026, 4, 1),
                competencia="2026-04",
                cnpj_empresa="08884180000123",
                nome_empresa="Empresa Teste",
                fornecedor_ou_favorecido="Fornecedor Baseline",
                cnpj_cpf_favorecido="00000000000100",
                documento="NF-1",
                tipo_documento="NF-e",
                conta_id="2.02.008",
                conta_nome="Hortifruti",
                grupo_conta="Custos",
                descricao_itens="Compra baseline",
                valor=Decimal("200.00"),
                origem="XML_NFE",
                status_revisao="OK",
            ),
            Lancamento(
                data_documento=date(2026, 4, 30),
                competencia="2026-04",
                cnpj_empresa="08884180000123",
                nome_empresa="Empresa Teste",
                fornecedor_ou_favorecido="Fornecedor Revisar",
                cnpj_cpf_favorecido="00000000000199",
                documento="NF-2",
                tipo_documento="NF-e",
                conta_id="2.02.008",
                conta_nome="Hortifruti",
                grupo_conta="Custos",
                descricao_itens="Limite periodo baseline",
                valor=Decimal("1.00"),
                origem="XML_NFE",
                status_revisao="REVISAR",
            )
        ],
    )
    reports_run = reports_root / "runs" / run_id
    reports_run.mkdir(parents=True)
    (reports_run / "pre_livro_diario_gerencial_teste.xlsx").write_bytes(b"placeholder")
    (reports_run / "dre_gerencial_ok_provisoria.xlsx").write_bytes(b"placeholder")
    (reports_run / "pacote_contabil_gerencial_provisorio").mkdir()
    manifest = {
        "run_id": run_id,
        "status_run": "PROVISORIO_CONCLUIDO",
        "parametros": {
            "xml_base": str(project_root / "xml"),
            "saida_output": str(output_root),
            "saida_reports": str(reports_root),
        },
        "totais": {"xml_processados": 1},
        "lancamentos_banca_aprovados": {
            "incluido": True,
            "input_xlsx": "",
            "linhas_lidas": 0,
        },
        "lancamentos_manuais_gerenciais": {"executado": False},
    }
    manifest_path = reports_run / "manifest_run.json"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    return manifest_path


def _write_manuali(
    project_root: Path,
    rows: list[tuple[str, str, str, str, str]] | None = None,
) -> tuple[Path, Path]:
    manuali_dir = project_root / "Database_Dev" / "Registrazioni_manuali"
    manuali_dir.mkdir(parents=True, exist_ok=True)
    csv_path = manuali_dir / "registrazioni_manuali_tecnico.csv"
    rows = rows or [
        (
            "2026-04-10",
            "Compra manual",
            "2.02.008",
            "Hortifruti",
            "100.00",
        ),
        (
            "2026-04-11",
            "Combustivel manual",
            "8.03.001",
            "Combustivel",
            "25.30",
        ),
    ]
    csv_rows = [
        "data;descricao_da_despesa;conta_debito_id;conta_debito_nome;conta_credito_id;conta_credito_nome;valor;observacao"
    ]
    csv_rows.extend(
        (
            f"{data};{descricao};{conta_id};{conta_nome};"
            "99.99.998;Contrapartida Tecnica Lancamentos Manuais;"
            f"{valor};obs"
        )
        for data, descricao, conta_id, conta_nome, valor in rows
    )
    csv_path.write_text(
        "\n".join([*csv_rows, ""]),
        encoding="utf-8",
    )
    manifest_path = manuali_dir / "registrazioni_manuali_validazione_manifest.json"
    manifest_path.write_text(
        json.dumps({"status_validacao": "APROVADO", "linhas_validas": len(rows)}),
        encoding="utf-8",
    )
    return csv_path, manifest_path


def _dre_base_rows(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["Base_OK_DRE"]
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if any(value is not None for value in values):
            rows.append(dict(zip(headers, values)))
    return rows


def _all_workbook_values(path: Path) -> set[str]:
    workbook = load_workbook(path, data_only=True)
    values = set()
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if value is not None:
                    values.add(str(value))
    return values


def _load_json(path: Path) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))
