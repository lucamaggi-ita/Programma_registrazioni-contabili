import csv
from pathlib import Path

from openpyxl import load_workbook
import pytest

from fantozzi.classification_review_export import CSV_HEADERS
from fantozzi.classification_review_package import (
    BASE_COMPLETA_SHEET,
    ClassificationReviewPackageError,
    gerar_pacote_revisao_contabil,
)


EXPECTED_SHEETS = [
    "Resumo",
    "Prioridades",
    "Top_Valor",
    "Top_Ocorrencias",
    "Por_Fornecedor",
    "Possiveis_Especiais",
    "Para_Contador",
    "Base_Completa",
    "Instrucoes",
]


def test_le_dois_csvs_e_gera_xlsx_com_abas_obrigatorias(tmp_path):
    csv_a = _save_review_csv(tmp_path / "empresa_a" / "revisao.csv", [_row()])
    csv_b = _save_review_csv(
        tmp_path / "empresa_b" / "revisao.csv",
        [_row(descricao_normalizada="cerveja lata", valor_total="R$ 120,00")],
    )
    xlsx_path = tmp_path / "reports" / "pacote.xlsx"

    result = gerar_pacote_revisao_contabil([csv_a, csv_b], xlsx_path)

    assert result.total_candidatos == 2
    assert xlsx_path.exists()
    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == EXPECTED_SHEETS


def test_gera_txt_com_resumo(tmp_path):
    csv_path = _save_review_csv(tmp_path / "empresa" / "revisao.csv", [_row()])
    xlsx_path = tmp_path / "pacote.xlsx"
    report_path = tmp_path / "pacote.txt"

    gerar_pacote_revisao_contabil([csv_path], xlsx_path, report_path)

    text = report_path.read_text(encoding="utf-8")
    assert "PACOTE ASSISTIDO PARA REVISAO CONTABIL HUMANA" in text
    assert "Total de candidatos: 1" in text
    assert "Resultado: OK" in text


def test_calcula_prioridade_alta_por_valor_alto(tmp_path):
    row = _package_row(tmp_path, _row(valor_total="R$ 5.000,00"))

    assert row["prioridade_revisao"] == "ALTA"
    assert "Valor alto" in row["motivo_prioridade"]


def test_calcula_prioridade_alta_por_muitas_ocorrencias(tmp_path):
    row = _package_row(tmp_path, _row(ocorrencias="10"))

    assert row["prioridade_revisao"] == "ALTA"
    assert "Muitas ocorrencias" in row["motivo_prioridade"]


def test_calcula_prioridade_baixa_para_baixo_impacto(tmp_path):
    row = _package_row(
        tmp_path,
        _row(
            ocorrencias="1",
            valor_total="R$ 10,00",
            conta_atual_id="2.01.004",
            conta_atual_nome="Refrigerantes",
            status_revisao="OK",
            alertas="",
        ),
    )

    assert row["prioridade_revisao"] == "BAIXA"
    assert row["motivo_prioridade"] == "Baixo impacto"


def test_identifica_requer_contador_por_acao_manual(tmp_path):
    row = _package_row(tmp_path, _row(acao_manual="ANALISAR_COM_CONTADOR"))

    assert row["requer_contador"] == "SIM"


def test_identifica_possivel_bonificacao_por_alertas(tmp_path):
    row = _package_row(tmp_path, _row(alertas="POSSIVEL_BONIFICACAO_RECEBIDA"))

    assert row["possivel_operacao_especial"] == "BONIFICACAO_RECEBIDA"
    assert row["requer_contador"] == "SIM"


def test_identifica_possivel_comodato_por_texto(tmp_path):
    row = _package_row(
        tmp_path,
        _row(
            descricao_normalizada="freezer comodato",
            descricao_exemplo="FREEZER EM COMODATO",
        ),
    )

    assert row["possivel_operacao_especial"] == "COMODATO"
    assert row["sugestao_tratamento"] == "Revisar com contador"


def test_falha_claramente_se_coluna_obrigatoria_ausente(tmp_path):
    csv_path = tmp_path / "empresa" / "revisao.csv"
    csv_path.parent.mkdir(parents=True)
    headers = [header for header in CSV_HEADERS if header != "acao_manual"]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers, delimiter=";")
        writer.writeheader()

    with pytest.raises(ClassificationReviewPackageError) as exc_info:
        gerar_pacote_revisao_contabil([csv_path], tmp_path / "pacote.xlsx")

    assert "Colunas obrigatorias ausentes" in str(exc_info.value)
    assert "acao_manual" in str(exc_info.value)


def _package_row(tmp_path: Path, row: dict[str, str]) -> dict[str, object]:
    csv_path = _save_review_csv(tmp_path / "empresa" / "revisao.csv", [row])
    xlsx_path = tmp_path / "pacote.xlsx"
    gerar_pacote_revisao_contabil([csv_path], xlsx_path)
    sheet = load_workbook(xlsx_path)[BASE_COMPLETA_SHEET]
    headers = [cell.value for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]]
    return dict(zip(headers, values))


def _save_review_csv(csv_path: Path, rows: list[dict[str, str]]) -> Path:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=CSV_HEADERS, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
    return csv_path


def _row(**overrides: str) -> dict[str, str]:
    row = {
        "tipo_candidato": "CLASSIFICACAO_PRODUTO",
        "descricao_normalizada": "agua mineral",
        "descricao_exemplo": "AGUA MINERAL 500ML",
        "ocorrencias": "1",
        "fornecedores": "Fornecedor A",
        "documentos": "NF-1",
        "valor_total": "R$ 10,00",
        "conta_atual_id": "2.01.004",
        "conta_atual_nome": "Refrigerantes",
        "status_revisao": "REVISAR",
        "alertas": "",
        "acao_manual": "MANTER_REVISAO",
        "conta_id_sugerida": "",
        "conta_nome_sugerida": "",
        "palavras_chave_sugeridas": "",
        "observacao_manual": "",
    }
    row.update(overrides)
    return row
