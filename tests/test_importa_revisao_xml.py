from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from fantozzi.importa_revisao_xml import (
    DECISOES_SHEET,
    EXCEL_NAME,
    MANIFEST_NAME,
    ImportaRevisaoXmlError,
    importar_revisao_xml,
    resolve_run_reports_dir,
)
from fantozzi.cli import main


CSV_HEADERS = [
    "tipo_candidato",
    "descricao_normalizada",
    "descricao_exemplo",
    "ocorrencias",
    "fornecedores",
    "documentos",
    "valor_total",
    "conta_atual_id",
    "conta_atual_nome",
    "status_revisao",
    "alertas",
    "cfop",
    "natureza_operacao",
    "tipo_operacao_especial",
    "confianca_operacao_especial",
    "tratamento_operacao_especial",
    "acao_manual",
    "conta_id_sugerida",
    "conta_nome_sugerida",
    "palavras_chave_sugeridas",
    "observacao_manual",
]


def test_resolve_run_reports_dir_from_run_id(tmp_path):
    path = resolve_run_reports_dir(
        run_id="2026-05-23_001",
        reports_runs_root=tmp_path / "reports" / "importacao_real" / "runs",
    )

    assert path == tmp_path / "reports" / "importacao_real" / "runs" / "2026-05-23_001"


def test_stale_hash_blocks_by_default(tmp_path):
    env = _make_env(tmp_path, manifest_hash="old-hash")
    _write_csv(env.run_dir, [_row()])

    try:
        importar_revisao_xml(
            run_reports_dir=env.run_dir,
            output_dir=env.output_dir,
            regras_path=env.regras_path,
            plano_contas_path=env.plano_path,
        )
    except ImportaRevisaoXmlError as exc:
        message = str(exc)
    else:
        raise AssertionError("stale run should block")

    assert "BLOCCO: la run 2026-05-23_001" in message
    assert "Hash run:     old-hash" in message
    assert "Hash attuale:" in message
    assert not (env.output_dir / EXCEL_NAME).exists()


def test_missing_yaml_hash_blocks_by_default(tmp_path):
    env = _make_env(tmp_path, include_hash=False)
    _write_csv(env.run_dir, [_row()])

    try:
        importar_revisao_xml(
            run_reports_dir=env.run_dir,
            output_dir=env.output_dir,
            regras_path=env.regras_path,
            plano_contas_path=env.plano_path,
        )
    except ImportaRevisaoXmlError as exc:
        message = str(exc)
    else:
        raise AssertionError("missing run hash should block")

    assert "non contiene l'hash del YAML" in message
    assert "Per audit/debug su dati storici usare --allow-stale-run." in message
    assert not (env.output_dir / EXCEL_NAME).exists()


def test_missing_yaml_hash_with_allow_stale_generates_audit_manifest(tmp_path):
    env = _make_env(tmp_path, include_hash=False)
    _write_csv(env.run_dir, [_row()])

    result = importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
        allow_stale_run=True,
    )

    manifest = json.loads((env.output_dir / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert result.hash_yaml_run == ""
    assert result.hash_yaml_run_assente is True
    assert result.stale_run is True
    assert manifest["hash_yaml_run"] == ""
    assert manifest["hash_yaml_run_assente"] is True
    assert manifest["run_verificavel"] is False
    assert manifest["allow_stale_run"] is True


def test_manifest_without_config_usado_blocks_like_missing_hash(tmp_path):
    env = _make_env(tmp_path, include_config=False)
    _write_csv(env.run_dir, [_row()])

    try:
        importar_revisao_xml(
            run_reports_dir=env.run_dir,
            output_dir=env.output_dir,
            regras_path=env.regras_path,
            plano_contas_path=env.plano_path,
        )
    except ImportaRevisaoXmlError as exc:
        message = str(exc)
    else:
        raise AssertionError("manifest without config_usado should block")

    assert "non contiene l'hash del YAML" in message


def test_manifest_without_config_usado_allows_only_explicit_audit_mode(tmp_path):
    env = _make_env(tmp_path, include_config=False)
    _write_csv(env.run_dir, [_row()])

    result = importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
        allow_stale_run=True,
    )

    manifest = json.loads((env.output_dir / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert result.hash_yaml_run_assente is True
    assert manifest["hash_yaml_run_assente"] is True
    assert manifest["allow_stale_run"] is True


def test_missing_plano_contas_raises_import_error(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(env.run_dir, [_row()])
    env.plano_path.unlink()

    try:
        importar_revisao_xml(
            run_reports_dir=env.run_dir,
            output_dir=env.output_dir,
            regras_path=env.regras_path,
            plano_contas_path=env.plano_path,
        )
    except ImportaRevisaoXmlError as exc:
        message = str(exc)
    else:
        raise AssertionError("missing plano contas should raise ImportaRevisaoXmlError")

    assert "Arquivo de plano de contas nao encontrado" in message
    assert str(env.plano_path) in message


def test_allow_stale_run_generates_manifest_with_flag(tmp_path):
    env = _make_env(tmp_path, manifest_hash="old-hash")
    _write_csv(env.run_dir, [_row()])

    result = importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
        allow_stale_run=True,
    )

    manifest = json.loads((env.output_dir / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert result.stale_run is True
    assert manifest["stale_run"] is True
    assert manifest["allow_stale_run"] is True
    assert manifest["hash_yaml_run"] == "old-hash"


def test_strong_comodato_and_bonificacao_are_excluded_and_counted(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(
        env.run_dir,
        [
            _row(
                descricao_normalizada="freezer comodato",
                tipo_operacao_especial="COMODATO",
                confianca_operacao_especial="FORTE",
            ),
            _row(
                descricao_normalizada="cerveja bonificacao",
                tipo_operacao_especial="BONIFICACAO_RECEBIDA",
                confianca_operacao_especial="FORTE",
            ),
            _row(descricao_normalizada="produto revisar"),
        ],
    )

    result = importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    assert result.comodato_forte_esclusi == 1
    assert result.bonificacao_forte_esclusa == 1
    assert result.linhas_operacionais_geradas == 1
    rows = _decision_rows(env.output_dir / EXCEL_NAME)
    assert [row["DESCRICAO_NORMALIZADA"] for row in rows] == ["produto revisar"]


def test_special_operations_not_strong_enter_as_blocked(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(
        env.run_dir,
        [
            _row(
                descricao_normalizada="comodato provavel",
                tipo_operacao_especial="COMODATO",
                confianca_operacao_especial="PROVAVEL",
            ),
            _row(
                descricao_normalizada="bonificacao provavel",
                tipo_operacao_especial="BONIFICACAO_RECEBIDA",
                confianca_operacao_especial="PROVAVEL",
            ),
        ],
    )

    result = importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    rows = _decision_rows(env.output_dir / EXCEL_NAME)
    assert result.candidatos_bloqueados_operacao_especial == 2
    assert {row["STATUS_DECISAO"] for row in rows} == {"BLOQUEADO"}
    assert {row["NIVEL_RISCO_REGRA"] for row in rows} == {"ALTO"}


def test_blocked_rows_have_restricted_action_validation(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(
        env.run_dir,
        [
            _row(
                descricao_normalizada="comodato provavel",
                tipo_operacao_especial="COMODATO",
                confianca_operacao_especial="PROVAVEL",
            ),
            _row(descricao_normalizada="produto normal"),
        ],
    )

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    workbook = load_workbook(env.output_dir / EXCEL_NAME, data_only=False)
    sheet = workbook[DECISOES_SHEET]
    rows = _decision_rows(env.output_dir / EXCEL_NAME)
    blocked_index = next(
        index
        for index, row in enumerate(rows, start=2)
        if row["STATUS_DECISAO"] == "BLOQUEADO"
    )
    headers = [cell.value for cell in sheet[1]]
    action_column = chr(ord("A") + headers.index("ACAO_REGRA"))
    blocked_cell = f"{action_column}{blocked_index}"

    blocked_validations = [
        validation
        for validation in sheet.data_validations.dataValidation
        if validation.formula1 == '"IGNORAR,BLOQUEADO,REVISAR"'
    ]
    assert blocked_validations
    assert any(blocked_cell in str(validation.sqref) for validation in blocked_validations)
    assert "CRIAR_ALIAS" not in blocked_validations[0].formula1
    assert "CRIAR_REGRA" not in blocked_validations[0].formula1


def test_output_sheet_and_required_columns_are_compatible(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(env.run_dir, [_row()])

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    workbook = load_workbook(env.output_dir / EXCEL_NAME, data_only=False)
    assert workbook.sheetnames == [
        "decisoes_xml",
        "tecnico_xml",
        "plano_contas",
        "listas_validacao",
        "instrucoes",
    ]
    headers = [cell.value for cell in workbook[DECISOES_SHEET][1]]
    assert "CHAVE_ESTAVEL" in headers
    assert "FORNECEDORES" in headers
    assert "CNPJS_ORIGEM" in headers
    assert "CONTA_ESCOLHIDA" in headers
    assert "CONTA_ID_APROVADA" in headers
    assert "CONTA_NOME_APROVADA" in headers


def test_manual_decisions_are_preserved_and_account_is_derived_from_choice(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(env.run_dir, [_row(descricao_normalizada="queijo mussarela")])
    key = "XML_PRODUTO|FORNECEDOR=FORNECEDOR TESTE|DESCRICAO=QUEIJO MUSSARELA"
    _write_existing_excel(env.output_dir / EXCEL_NAME, key)

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    row = _decision_rows(env.output_dir / EXCEL_NAME)[0]
    assert row["STATUS_DECISAO"] == "APROVADO"
    assert row["ACAO_REGRA"] == "CRIAR_REGRA"
    assert row["CONTA_ESCOLHIDA"] == "2.02.006 - Laticinios e Queijos"
    assert row["CONTA_ID_APROVADA"] == "2.02.006"
    assert row["CONTA_NOME_APROVADA"] == "Laticinios e Queijos"
    assert row["CONTA_ID_APROVADA"] != "=LEFT(H99,7)"
    assert row["OBSERVACAO_OPERADOR"] == "ok operador"


def test_importar_revisao_xml_idempotente(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(
        env.run_dir,
        [
            _row(descricao_normalizada="queijo mussarela"),
            _row(descricao_normalizada="cerveja lata"),
        ],
    )

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )
    first_rows = _decision_rows(env.output_dir / EXCEL_NAME)

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )
    second_rows = _decision_rows(env.output_dir / EXCEL_NAME)

    assert second_rows == first_rows


def test_importar_revisao_xml_aggiunge_nova_chave_sem_perder_decisoes(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(env.run_dir, [_row(descricao_normalizada="queijo mussarela")])

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )
    _set_cell_by_header(
        env.output_dir / EXCEL_NAME,
        row=2,
        values={
            "STATUS_DECISAO": "APROVADO",
            "ACAO_REGRA": "CRIAR_ALIAS",
            "CONTA_ESCOLHIDA": "2.02.006 - Laticinios e Queijos",
            "PALAVRAS_CHAVE_SUGERIDAS": "queijo; mussarela",
            "OBSERVACAO_OPERADOR": "decisao preservada",
        },
    )
    _write_csv(
        env.run_dir,
        [
            _row(descricao_normalizada="queijo mussarela"),
            _row(descricao_normalizada="cerveja lata"),
        ],
    )

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    rows = _decision_rows(env.output_dir / EXCEL_NAME)
    assert len(rows) == 2
    queijo = next(row for row in rows if row["DESCRICAO_NORMALIZADA"] == "queijo mussarela")
    cerveja = next(row for row in rows if row["DESCRICAO_NORMALIZADA"] == "cerveja lata")
    assert queijo["STATUS_DECISAO"] == "APROVADO"
    assert queijo["ACAO_REGRA"] == "CRIAR_ALIAS"
    assert queijo["CONTA_ESCOLHIDA"] == "2.02.006 - Laticinios e Queijos"
    assert queijo["CONTA_ID_APROVADA"] == "2.02.006"
    assert queijo["CONTA_NOME_APROVADA"] == "Laticinios e Queijos"
    assert queijo["PALAVRAS_CHAVE_SUGERIDAS"] == "queijo; mussarela"
    assert queijo["OBSERVACAO_OPERADOR"] == "decisao preservada"
    assert cerveja["STATUS_DECISAO"] == "PENDENTE"


def test_importar_revisao_xml_remove_chave_desaparecida_do_operacional(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(
        env.run_dir,
        [
            _row(descricao_normalizada="queijo mussarela"),
            _row(descricao_normalizada="cerveja lata"),
        ],
    )

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )
    _write_csv(env.run_dir, [_row(descricao_normalizada="queijo mussarela")])

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    rows = _decision_rows(env.output_dir / EXCEL_NAME)
    assert [row["DESCRICAO_NORMALIZADA"] for row in rows] == ["queijo mussarela"]


def test_status_ok_is_excluded_from_operational_sheet(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(env.run_dir, [_row(status_revisao="OK")])

    result = importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    assert result.linhas_operacionais_geradas == 0
    assert _decision_rows(env.output_dir / EXCEL_NAME) == []


def test_same_stable_key_across_multiple_cnpjs_groups_into_one_decision(tmp_path):
    env = _make_env(tmp_path)
    row = _row(
        descricao_normalizada="queijo mussarela",
        descricao_exemplo="QUEIJO MUSSARELA",
        ocorrencias="1",
        valor_total="R$ 10,00",
        fornecedores="Fornecedor Teste",
    )
    _write_csv(env.run_dir, [row], cnpj="08884180000123")
    _write_csv(env.run_dir, [row], cnpj="55746722000181")

    result = importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    rows = _decision_rows(env.output_dir / EXCEL_NAME)
    assert result.linhas_operacionais_geradas == 1
    assert rows[0]["QTD_OCORRENCIAS"] == 2
    assert rows[0]["VALOR_TOTAL"] == 20
    assert rows[0]["CNPJS_ORIGEM"] == "08884180000123; 55746722000181"

    workbook = load_workbook(env.output_dir / EXCEL_NAME, data_only=False)
    tecnico = workbook["tecnico_xml"]
    tecnico_rows = list(tecnico.iter_rows(values_only=True))
    headers = [str(value or "") for value in tecnico_rows[0]]
    cnpj_col = headers.index("CNPJ")
    assert {row[cnpj_col] for row in tecnico_rows[1:]} == {
        "08884180000123",
        "55746722000181",
    }


def test_manifest_keeps_critical_fields(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(env.run_dir, [_row()])

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    manifest = json.loads((env.output_dir / MANIFEST_NAME).read_text(encoding="utf-8"))
    for field in [
        "hash_yaml_run",
        "hash_yaml_atual",
        "hash_yaml_run_assente",
        "stale_run",
        "allow_stale_run",
        "comodato_forte_esclusi",
        "bonificacao_forte_esclusa",
        "candidatos_bloqueados_operacao_especial",
    ]:
        assert field in manifest


def test_importar_revisao_xml_escreve_manifest_historico_por_run(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(env.run_dir, [_row()])

    result = importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    current = env.output_dir / MANIFEST_NAME
    historic = env.output_dir / "manifest_importa_revisao_xml_2026-05-23_001.json"
    assert result.manifest_path == current
    assert result.manifest_historico_path == historic
    assert current.exists()
    assert historic.exists()
    current_payload = json.loads(current.read_text(encoding="utf-8"))
    historic_payload = json.loads(historic.read_text(encoding="utf-8"))
    assert current_payload == historic_payload
    assert current_payload["run_id_origem"] == "2026-05-23_001"
    assert current_payload["manifest_corrente"] == str(current)
    assert current_payload["manifest_historico"] == str(historic)


def test_importar_revisao_xml_preserva_decisoes_e_atualiza_cnpjs_origem_entre_runs(
    tmp_path,
):
    env = _make_env(tmp_path)
    row = _row(descricao_normalizada="queijo mussarela", fornecedores="Fornecedor Teste")
    _write_csv(env.run_dir, [row], cnpj="08884180000123")

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )
    _set_cell_by_header(
        env.output_dir / EXCEL_NAME,
        row=2,
        values={
            "STATUS_DECISAO": "APROVADO",
            "ACAO_REGRA": "CRIAR_ALIAS",
            "CONTA_ESCOLHIDA": "2.02.006 - Laticinios e Queijos",
            "OBSERVACAO_OPERADOR": "preservar entre cnpjs",
        },
    )
    _write_csv(env.run_dir, [row], cnpj="55746722000181")

    importar_revisao_xml(
        run_reports_dir=env.run_dir,
        output_dir=env.output_dir,
        regras_path=env.regras_path,
        plano_contas_path=env.plano_path,
    )

    rows = _decision_rows(env.output_dir / EXCEL_NAME)
    assert len(rows) == 1
    assert rows[0]["STATUS_DECISAO"] == "APROVADO"
    assert rows[0]["ACAO_REGRA"] == "CRIAR_ALIAS"
    assert rows[0]["CONTA_ID_APROVADA"] == "2.02.006"
    assert rows[0]["OBSERVACAO_OPERADOR"] == "preservar entre cnpjs"
    assert rows[0]["CNPJS_ORIGEM"] == "08884180000123; 55746722000181"


def test_cli_run_id_imports_xml_review(tmp_path):
    env = _make_env(tmp_path)
    _write_csv(env.run_dir, [_row()])

    exit_code = main(
        [
            "importar-revisao-xml",
            "--run-id",
            env.run_dir.name,
            "--reports-runs-root",
            str(env.run_dir.parent),
            "--saida-dir",
            str(env.output_dir),
            "--regras",
            str(env.regras_path),
            "--plano-contas",
            str(env.plano_path),
        ]
    )

    assert exit_code == 0
    assert (env.output_dir / EXCEL_NAME).exists()


class _Env:
    def __init__(
        self,
        *,
        run_dir: Path,
        output_dir: Path,
        regras_path: Path,
        plano_path: Path,
    ):
        self.run_dir = run_dir
        self.output_dir = output_dir
        self.regras_path = regras_path
        self.plano_path = plano_path


def _make_env(
    tmp_path,
    *,
    manifest_hash: str | None = None,
    include_hash: bool = True,
    include_config: bool = True,
) -> _Env:
    regras_path = tmp_path / "config" / "regras_classificacao_produtos.yaml"
    plano_path = tmp_path / "config" / "plano_contas_gerencial.yaml"
    regras_path.parent.mkdir(parents=True)
    regras_path.write_text("regras: []\n", encoding="utf-8")
    plano_path.write_text(
        "\n".join(
            [
                "contas:",
                "- id: 2.01.001",
                "  nome: Cervejas",
                "- id: 2.02.006",
                "  nome: Laticinios e Queijos",
            ]
        ),
        encoding="utf-8",
    )
    run_dir = tmp_path / "reports" / "importacao_real" / "runs" / "2026-05-23_001"
    run_dir.mkdir(parents=True)
    hash_value = manifest_hash or _hash_file(regras_path)
    payload = {"run_id": "2026-05-23_001"}
    if include_config:
        payload["config_usado"] = {"regras_classificacao_produtos": {}}
        if include_hash:
            payload["config_usado"]["regras_classificacao_produtos"][
                "hash_sha256"
            ] = hash_value
    (run_dir / "manifest_run.json").write_text(
        json.dumps(payload),
        encoding="utf-8",
    )
    return _Env(
        run_dir=run_dir,
        output_dir=tmp_path / "Database_Dev" / "_decisoes_regras" / "xml",
        regras_path=regras_path,
        plano_path=plano_path,
    )


def _write_csv(
    run_dir: Path,
    rows: list[dict[str, str]],
    *,
    cnpj: str = "08884180000123",
) -> Path:
    csv_dir = run_dir / cnpj
    csv_dir.mkdir(parents=True, exist_ok=True)
    csv_path = csv_dir / "revisao_classificacao_real.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=CSV_HEADERS, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
    return csv_path


def _row(**overrides) -> dict[str, str]:
    row = {
        "tipo_candidato": "CLASSIFICACAO_PRODUTO",
        "descricao_normalizada": "produto revisar",
        "descricao_exemplo": "PRODUTO REVISAR",
        "ocorrencias": "2",
        "fornecedores": "Fornecedor Teste",
        "documentos": "123",
        "valor_total": "R$ 10,50",
        "conta_atual_id": "13.01.001",
        "conta_atual_nome": "Revisar - Produto Nao Classificado",
        "status_revisao": "REVISAR",
        "alertas": "",
        "cfop": "5102",
        "natureza_operacao": "VENDA",
        "tipo_operacao_especial": "",
        "confianca_operacao_especial": "",
        "tratamento_operacao_especial": "",
        "acao_manual": "",
        "conta_id_sugerida": "",
        "conta_nome_sugerida": "",
        "palavras_chave_sugeridas": "",
        "observacao_manual": "",
    }
    row.update({key: str(value) for key, value in overrides.items()})
    return row


def _write_existing_excel(path: Path, key: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DECISOES_SHEET
    headers = [
        "CHAVE_ESTAVEL",
        "STATUS_DECISAO",
        "ACAO_REGRA",
        "CONTA_ESCOLHIDA",
        "CONTA_ID_APROVADA",
        "CONTA_NOME_APROVADA",
        "PALAVRAS_CHAVE_SUGERIDAS",
        "OBSERVACAO_OPERADOR",
    ]
    sheet.append(headers)
    sheet.append(
        [
            key,
            "APROVADO",
            "CRIAR_REGRA",
            "2.02.006 - Laticinios e Queijos",
            "=LEFT(H99,7)",
            "=MID(H99,10,99)",
            "queijo",
            "ok operador",
        ]
    )
    workbook.save(path)


def _set_cell_by_header(
    path: Path,
    *,
    row: int,
    values: dict[str, object],
) -> None:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[DECISOES_SHEET]
    headers = [cell.value for cell in sheet[1]]
    for header, value in values.items():
        sheet.cell(row=row, column=headers.index(header) + 1).value = value
    workbook.save(path)


def _decision_rows(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[DECISOES_SHEET]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "") for value in rows[0]]
    return [
        {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
        }
        for values in rows[1:]
        if any(value not in (None, "") for value in values)
    ]


def _hash_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
