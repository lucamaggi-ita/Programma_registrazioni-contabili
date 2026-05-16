from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fantozzi.excel_writer import LANCAMENTOS_SHEET
from fantozzi.dre_gerencial_provisoria import (
    DRE_GERAL_SHEET,
    gerar_dre_gerencial_ok_provisoria,
)
from fantozzi.lancamentos_manuais_gerenciais import (
    FONTE_MANUAL_GERENCIAL,
    LancamentosManuaisGerenciaisError,
    carregar_lancamentos_manuais_gerenciais,
    converter_lancamentos_manuais_para_lancamentos,
    processar_lancamentos_manuais_gerenciais,
)
from fantozzi.pre_livro_diario import LANCAMENTOS_OK_SHEET, gerar_pre_livro_diario_teste


CSV_HEADERS = [
    "data",
    "historico",
    "documento",
    "fornecedor_ou_contraparte",
    "conta_debito_id",
    "conta_debito_nome",
    "conta_credito_id",
    "conta_credito_nome",
    "valor",
    "cnpj_origem_opcional",
    "centro_custo_opcional",
    "observacao",
]


def test_le_csv_valido_com_decimal_brasileiro(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = _save_manual_csv(tmp_path, valor="1.234,56")

    result = carregar_lancamentos_manuais_gerenciais(csv_path, *paths)

    assert len(result) == 1
    assert result[0].valor == Decimal("1234.56")
    assert result[0].linha_origem_csv == 2
    assert result[0].conta_dre_mapeada_count == 1


def test_aceita_decimal_com_ponto(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = _save_manual_csv(tmp_path, valor="1234.56")

    result = carregar_lancamentos_manuais_gerenciais(csv_path, *paths)

    assert result[0].valor == Decimal("1234.56")


def test_rejeita_coluna_obrigatoria_ausente(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = tmp_path / "manual.csv"
    headers = [header for header in CSV_HEADERS if header != "historico"]
    _write_csv(csv_path, headers, [_row()])

    with pytest.raises(LancamentosManuaisGerenciaisError, match="historico"):
        carregar_lancamentos_manuais_gerenciais(csv_path, *paths)


def test_rejeita_coluna_duplicada(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = tmp_path / "manual.csv"
    csv_path.write_text(
        "data;data;historico;documento;fornecedor_ou_contraparte;"
        "conta_debito_id;conta_debito_nome;conta_credito_id;"
        "conta_credito_nome;valor\n"
        "2026-05-16;2026-05-16;Ajuste;DOC;Fornecedor;2.02.001;"
        "Carnes Bovinas;10.01.001;Equipamentos de Cozinha;10,00\n",
        encoding="utf-8",
    )

    with pytest.raises(LancamentosManuaisGerenciaisError, match="duplicadas"):
        carregar_lancamentos_manuais_gerenciais(csv_path, *paths)


def test_rejeita_data_invalida(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = _save_manual_csv(tmp_path, data="16/05/2026")

    with pytest.raises(LancamentosManuaisGerenciaisError, match="data invalida"):
        carregar_lancamentos_manuais_gerenciais(csv_path, *paths)


@pytest.mark.parametrize("valor", ["0", "-10,00"])
def test_rejeita_valor_zero_ou_negativo(tmp_path, valor):
    paths = _save_config(tmp_path)
    csv_path = _save_manual_csv(tmp_path, valor=valor)

    with pytest.raises(LancamentosManuaisGerenciaisError, match="valor"):
        carregar_lancamentos_manuais_gerenciais(csv_path, *paths)


def test_rejeita_debito_igual_credito(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = _save_manual_csv(
        tmp_path,
        conta_credito_id="2.02.001",
        conta_credito_nome="Carnes Bovinas",
    )

    with pytest.raises(LancamentosManuaisGerenciaisError, match="igual"):
        carregar_lancamentos_manuais_gerenciais(csv_path, *paths)


def test_rejeita_conta_inexistente(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = _save_manual_csv(tmp_path, conta_debito_id="99.99.999")

    with pytest.raises(LancamentosManuaisGerenciaisError, match="inexistente"):
        carregar_lancamentos_manuais_gerenciais(csv_path, *paths)


def test_rejeita_nome_conta_incompativel(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = _save_manual_csv(tmp_path, conta_debito_nome="Nome errado")

    with pytest.raises(LancamentosManuaisGerenciaisError, match="incompativel"):
        carregar_lancamentos_manuais_gerenciais(csv_path, *paths)


def test_rejeita_duplicata_no_proprio_csv(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = tmp_path / "manual.csv"
    row = _row()
    _write_csv(csv_path, CSV_HEADERS, [row, row])

    with pytest.raises(LancamentosManuaisGerenciaisError, match="Duplicata"):
        carregar_lancamentos_manuais_gerenciais(csv_path, *paths)


def test_guardrail_dre_permite_zero_contas_mapeadas(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = _save_manual_csv(
        tmp_path,
        conta_debito_id="10.01.001",
        conta_debito_nome="Equipamentos de Cozinha",
        conta_credito_id="10.01.003",
        conta_credito_nome="Obras e Reformas",
    )

    result = carregar_lancamentos_manuais_gerenciais(csv_path, *paths)

    assert result[0].conta_dre_mapeada_count == 0


def test_guardrail_dre_bloqueia_duas_contas_mapeadas(tmp_path):
    paths = _save_config(tmp_path)
    csv_path = _save_manual_csv(
        tmp_path,
        conta_credito_id="1.01.001",
        conta_credito_nome="Receita de Vendas no Salao",
    )

    with pytest.raises(LancamentosManuaisGerenciaisError, match="duas contas|debito e credito"):
        carregar_lancamentos_manuais_gerenciais(csv_path, *paths)


def test_chave_deduplicacao_estavel_e_conversao_em_duas_linhas(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    csv_path = _save_manual_csv(tmp_path)

    first = carregar_lancamentos_manuais_gerenciais(csv_path, plano_path, dre_path)
    second = carregar_lancamentos_manuais_gerenciais(csv_path, plano_path, dre_path)
    tecnicos = converter_lancamentos_manuais_para_lancamentos(
        first,
        plano_path,
        csv_path,
    )

    assert first[0].chave_deduplicacao == second[0].chave_deduplicacao
    assert [item.tipo_lancamento for item in tecnicos] == [
        "MANUAL_GERENCIAL_DEBITO",
        "MANUAL_GERENCIAL_CREDITO",
    ]
    assert {item.origem for item in tecnicos} == {FONTE_MANUAL_GERENCIAL}
    assert tecnicos[0].chave_origem.endswith(":D")
    assert tecnicos[1].chave_origem.endswith(":C")


def test_processa_e_gera_xlsx_csv_auditoria(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    csv_path = _save_manual_csv(tmp_path)
    output_xlsx = tmp_path / "out" / "lancamentos_manuais.xlsx"
    output_csv = tmp_path / "reports" / "lancamentos_manuais.csv"
    auditoria_txt = tmp_path / "reports" / "auditoria.txt"

    result = processar_lancamentos_manuais_gerenciais(
        csv_path,
        plano_path,
        dre_path,
        output_xlsx,
        output_csv,
        auditoria_txt,
    )

    assert result.lancamentos_aceitos == 1
    assert result.linhas_tecnicas_geradas == 2
    assert result.total_manual == Decimal("10.50")
    assert output_xlsx.exists()
    assert output_csv.exists()
    assert auditoria_txt.exists()
    workbook = load_workbook(output_xlsx, data_only=True)
    sheet = workbook[LANCAMENTOS_SHEET]
    headers = [cell.value for cell in sheet[1]]
    origem_index = headers.index("origem") + 1
    assert sheet.max_row == 3
    assert sheet.cell(row=2, column=origem_index).value == FONTE_MANUAL_GERENCIAL
    assert "lancamentos_aceitos: 1" in auditoria_txt.read_text(encoding="utf-8")


def test_manual_com_uma_conta_dre_entra_uma_vez_no_pre_livro_e_dre(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    csv_path = _save_manual_csv(tmp_path)
    output_dir = tmp_path / "output"
    output_xlsx = output_dir / "MANUAL_GERENCIAL" / "lancamentos_manuais.xlsx"
    processar_lancamentos_manuais_gerenciais(
        csv_path,
        plano_path,
        dre_path,
        output_xlsx,
        tmp_path / "reports" / "manual.csv",
        tmp_path / "reports" / "manual.txt",
    )

    pre_livro_path = tmp_path / "reports" / "pre_livro.xlsx"
    dre_path_out = tmp_path / "reports" / "dre.xlsx"
    gerar_pre_livro_diario_teste([output_xlsx], pre_livro_path)
    gerar_dre_gerencial_ok_provisoria([output_dir], dre_path, dre_path_out)

    pre_livro = load_workbook(pre_livro_path, data_only=True)
    dre = load_workbook(dre_path_out, data_only=True)
    assert pre_livro[LANCAMENTOS_OK_SHEET].max_row == 3
    assert dre[DRE_GERAL_SHEET].max_row == 2
    assert dre[DRE_GERAL_SHEET]["C2"].value == "CMV"
    assert dre[DRE_GERAL_SHEET]["D2"].value == 10.5


def test_manual_com_zero_contas_dre_nao_gera_linha_dre(tmp_path):
    plano_path, dre_path = _save_config(tmp_path)
    csv_path = _save_manual_csv(
        tmp_path,
        conta_debito_id="10.01.001",
        conta_debito_nome="Equipamentos de Cozinha",
        conta_credito_id="10.01.003",
        conta_credito_nome="Obras e Reformas",
    )
    output_dir = tmp_path / "output"
    output_xlsx = output_dir / "MANUAL_GERENCIAL" / "lancamentos_manuais.xlsx"
    processar_lancamentos_manuais_gerenciais(
        csv_path,
        plano_path,
        dre_path,
        output_xlsx,
        tmp_path / "reports" / "manual.csv",
        tmp_path / "reports" / "manual.txt",
    )

    pre_livro_path = tmp_path / "reports" / "pre_livro.xlsx"
    dre_path_out = tmp_path / "reports" / "dre.xlsx"
    gerar_pre_livro_diario_teste([output_xlsx], pre_livro_path)
    gerar_dre_gerencial_ok_provisoria([output_dir], dre_path, dre_path_out)

    pre_livro = load_workbook(pre_livro_path, data_only=True)
    dre = load_workbook(dre_path_out, data_only=True)
    assert pre_livro[LANCAMENTOS_OK_SHEET].max_row == 3
    assert dre[DRE_GERAL_SHEET].max_row == 1


def _save_config(tmp_path: Path) -> tuple[Path, Path]:
    plano_path = tmp_path / "plano_contas.yaml"
    dre_path = tmp_path / "mapeamento_dre.yaml"
    plano_path.write_text(
        """
contas:
  - {id: "1.01.001", nome: "Receita de Vendas no Salao", grupo: "Receitas", tipo: receita}
  - {id: "2.02.001", nome: "Carnes Bovinas", grupo: "Custos", tipo: custo}
  - {id: "10.01.001", nome: "Equipamentos de Cozinha", grupo: "Investimentos", tipo: investimento}
  - {id: "10.01.003", nome: "Obras e Reformas", grupo: "Investimentos", tipo: investimento}
  - {id: "13.01.001", nome: "A Revisar", grupo: "Revisao", tipo: revisao}
  - {id: "13.01.002", nome: "Ambiguo", grupo: "Revisao", tipo: revisao}
  - {id: "13.01.006", nome: "Valor Inconsistente", grupo: "Revisao", tipo: revisao}
  - {id: "13.01.009", nome: "Possivel Investimento", grupo: "Revisao", tipo: revisao}
  - {id: "13.01.010", nome: "Pessoal a Regularizar", grupo: "Revisao", tipo: revisao}
  - {id: "13.01.999", nome: "Outros", grupo: "Revisao", tipo: revisao}
""".strip()
        + "\n",
        encoding="utf-8",
    )
    dre_path.write_text(
        """
linhas_dre:
  - {id: "01.01", nome: "Receita", grupo: "Receita", sinal: "+", ordem: 1}
  - {id: "03.01", nome: "CMV", grupo: "Custos", sinal: "-", ordem: 2}
  - {id: "12.01", nome: "Investimentos", grupo: "Investimentos", sinal: "-", ordem: 3}
mapeamento:
  "1.01.001": {dre_linha_id: "01.01", entra_dre: true}
  "2.02.001": {dre_linha_id: "03.01", entra_dre: true}
  "10.01.001": {dre_linha_id: "12.01", entra_dre: false}
""".strip()
        + "\n",
        encoding="utf-8",
    )
    return plano_path, dre_path


def _save_manual_csv(tmp_path: Path, **overrides) -> Path:
    csv_path = tmp_path / "manual.csv"
    _write_csv(csv_path, CSV_HEADERS, [_row(**overrides)])
    return csv_path


def _write_csv(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=headers,
            delimiter=";",
            extrasaction="ignore",
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _row(**overrides) -> dict[str, str]:
    row = {
        "data": "2026-05-16",
        "historico": "Ajuste manual gerencial",
        "documento": "MAN-001",
        "fornecedor_ou_contraparte": "Fornecedor Teste",
        "conta_debito_id": "2.02.001",
        "conta_debito_nome": "Carnes Bovinas",
        "conta_credito_id": "10.01.001",
        "conta_credito_nome": "Equipamentos de Cozinha",
        "valor": "10,50",
        "cnpj_origem_opcional": "",
        "centro_custo_opcional": "CC-001",
        "observacao": "fixture sintetica",
    }
    row.update({key: str(value) for key, value in overrides.items()})
    return row
