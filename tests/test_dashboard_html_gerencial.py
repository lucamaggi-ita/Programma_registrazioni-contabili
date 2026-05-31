from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from fantozzi.cli import main
from fantozzi.dashboard_html_gerencial import (
    DashboardHtmlGerencialError,
    DashboardHtmlResult,
    _build_cards_html,
    _build_table_html,
    _load_dre_mensal,
    gerar_dashboard_html,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_dre_xlsx(
    tmp_path: Path,
    rows: list[tuple],
    mensal_rows: list[tuple] | None = None,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "DRE_Geral"
    ws.append(["ordem", "grupo_dre", "linha_dre", "valor", "observacao"])
    for row in rows:
        ws.append(list(row))

    ws2 = wb.create_sheet("DRE_Mensal")
    ws2.append(["ano_mes", "ordem", "grupo_dre", "linha_dre", "valor"])
    for mrow in (mensal_rows or []):
        ws2.append(list(mrow))

    p = tmp_path / "dre_gerencial_ok_provisoria.xlsx"
    wb.save(str(p))
    return p


_FIXTURE_ROWS = [
    (1, "Custos", "CMV - Alimentos", 1000.50, "PROVISORIO"),
    (2, "Custos", "CMV - Bebidas", 500.00, "PROVISORIO"),
    (3, "Despesas Operacionais", "Pessoal CLT", 2000.00, "PROVISORIO"),
    (4, "Despesas Operacionais", "Aluguel", 1500.00, "PROVISORIO"),
]

_FIXTURE_MENSAL_1 = [
    ("2026-01", 1, "Custos", "CMV - Alimentos", 1000.50),
    ("2026-01", 3, "Despesas Operacionais", "Pessoal CLT", 2000.00),
]

_FIXTURE_MENSAL_2 = [
    ("2026-01", 1, "Custos", "CMV - Alimentos", 1000.50),
    ("2026-01", 3, "Despesas Operacionais", "Pessoal CLT", 2000.00),
    ("2026-02", 1, "Custos", "CMV - Alimentos", 800.00),
    ("2026-02", 3, "Despesas Operacionais", "Pessoal CLT", 1800.00),
]


# ---------------------------------------------------------------------------
# Testes _load_dre_mensal
# ---------------------------------------------------------------------------

def test_load_dre_mensal_retorna_linhas(tmp_path):
    xlsx = _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL_2)
    rows = _load_dre_mensal(xlsx)
    assert len(rows) == 4
    assert rows[0]["ano_mes"] == "2026-01"
    assert rows[0]["grupo_dre"] == "Custos"


def test_load_dre_mensal_sem_aba_mensal(tmp_path):
    wb = Workbook()
    wb.active.title = "DRE_Geral"
    p = tmp_path / "sem_mensal.xlsx"
    wb.save(str(p))
    assert _load_dre_mensal(p) == []


def test_load_dre_mensal_arquivo_inexistente(tmp_path):
    assert _load_dre_mensal(tmp_path / "nao_existe.xlsx") == []


def test_load_dre_mensal_aba_vazia(tmp_path):
    wb = Workbook()
    wb.active.title = "DRE_Geral"
    ws2 = wb.create_sheet("DRE_Mensal")
    ws2.append(["ano_mes", "grupo_dre", "linha_dre", "valor"])
    p = tmp_path / "mensal_vazia.xlsx"
    wb.save(str(p))
    assert _load_dre_mensal(p) == []


# ---------------------------------------------------------------------------
# Testes _build_cards_html
# ---------------------------------------------------------------------------

def test_build_cards_html_contem_total_geral():
    grupos = {"Custos": Decimal("1500.50"), "Despesas": Decimal("3500.00")}
    html = _build_cards_html(grupos, Decimal("5000.50"), 2, "2026-01 · 2026-02")
    assert "Total Geral" in html
    assert "R$ 5.000,50" in html


def test_build_cards_html_contem_todos_grupos():
    grupos = {"Custos": Decimal("1500"), "Despesas": Decimal("3500")}
    html = _build_cards_html(grupos, Decimal("5000"), 1, "2026-01")
    assert "Custos" in html
    assert "Despesas" in html


def test_build_cards_html_trunca_nome_longo():
    grupos = {"A" * 40: Decimal("100")}
    html = _build_cards_html(grupos, Decimal("100"), 1, "2026-01")
    assert "…" in html


# ---------------------------------------------------------------------------
# Testes _build_table_html
# ---------------------------------------------------------------------------

def test_build_table_html_contem_linhas_dre():
    rows = [
        {"grupo_dre": "Custos", "linha_dre": "CMV - Alimentos", "valor": 1000},
        {"grupo_dre": "Custos", "linha_dre": "CMV - Bebidas", "valor": 500},
    ]
    html = _build_table_html(rows)
    assert "CMV - Alimentos" in html
    assert "CMV - Bebidas" in html
    assert "TOTAL GERAL" in html


def test_build_table_html_contem_subtotal():
    rows = [
        {"grupo_dre": "Custos", "linha_dre": "CMV - Alimentos", "valor": 1000},
    ]
    html = _build_table_html(rows)
    assert "Sub-total Custos" in html


# ---------------------------------------------------------------------------
# Testes gerar_dashboard_html
# ---------------------------------------------------------------------------

def test_gerar_dashboard_html_cria_arquivo_com_nome_correto(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL_2)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    assert result.html_path.name == "02_DASHBOARD_GERENCIAL.html"
    assert result.html_path.exists()


def test_gerar_dashboard_html_tamanho_minimo(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL_2)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    # plotly.js embedded inline — file deve essere molto più grande di 10 KB
    assert result.html_path.stat().st_size > 10_000


def test_gerar_dashboard_html_status_resultado(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    assert result.status == "TESTE_TECNICO_NAO_OFICIAL"


def test_gerar_dashboard_html_contem_status_badge(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    content = result.html_path.read_text(encoding="utf-8")
    assert "TESTE_TECNICO_NAO_OFICIAL" in content or "TESTE TÉCNICO" in content


def test_gerar_dashboard_html_contem_plotly(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    content = result.html_path.read_text(encoding="utf-8")
    assert "Plotly" in content or "plotly" in content.lower()


def test_gerar_dashboard_html_e_html_valido(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    content = result.html_path.read_text(encoding="utf-8")
    assert content.strip().startswith("<!DOCTYPE html>")
    assert "</html>" in content


def test_gerar_dashboard_html_total_linhas(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    assert result.total_linhas == 4


def test_gerar_dashboard_html_periodos_corretos(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL_2)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    assert result.periodos == ("2026-01", "2026-02")


def test_gerar_dashboard_html_output_dir_default(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_dashboard_html(tmp_path)
    assert result.html_path.parent == tmp_path / "output_visual"


def test_gerar_dashboard_html_run_dir_inexistente_levanta_erro(tmp_path):
    with pytest.raises(DashboardHtmlGerencialError, match="run-dir nao encontrado"):
        gerar_dashboard_html(tmp_path / "nao_existe")


def test_gerar_dashboard_html_sem_dre_xlsx_levanta_erro(tmp_path):
    with pytest.raises(DashboardHtmlGerencialError):
        gerar_dashboard_html(tmp_path, tmp_path / "out")


def test_gerar_dashboard_html_le_run_id_do_manifest(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    manifest = {"run_id": "2026-05-29_001", "status_run": "PROVISORIO_CONCLUIDO"}
    (tmp_path / "manifest_run.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    assert result.run_id == "2026-05-29_001"


def test_gerar_dashboard_html_sem_mensal_sem_linha_chart(tmp_path):
    # Apenas 1 período — sem linha chart
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL_1)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    content = result.html_path.read_text(encoding="utf-8")
    assert "Andamento Mensal" not in content


def test_gerar_dashboard_html_com_mensal_contem_linha_chart(tmp_path):
    # 2 periodos — linha chart deve aparecer
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL_2)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    content = result.html_path.read_text(encoding="utf-8")
    assert "Andamento Mensal" in content


def test_gerar_dashboard_html_contem_subtotal_por_grupo(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    content = result.html_path.read_text(encoding="utf-8")
    assert "Subtotal por Grupo DRE" in content


def test_gerar_dashboard_html_contem_composicao_por_grupo(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    content = result.html_path.read_text(encoding="utf-8")
    assert "Composicao por Grupo" in content


def test_gerar_dashboard_html_run_id_none_sem_manifest(tmp_path):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    assert result.run_id is None


def test_gerar_dashboard_html_periodos_vazios_sem_mensal(tmp_path):
    # Sem DRE_Mensal → periodos vazio
    wb = Workbook()
    ws = wb.active
    ws.title = "DRE_Geral"
    ws.append(["ordem", "grupo_dre", "linha_dre", "valor", "observacao"])
    ws.append([1, "Custos", "CMV", 100.0, "PROVISORIO"])
    p = tmp_path / "dre_gerencial_ok_provisoria.xlsx"
    wb.save(str(p))
    result = gerar_dashboard_html(tmp_path, tmp_path / "out")
    assert result.periodos == ()


# ---------------------------------------------------------------------------
# Testes CLI
# ---------------------------------------------------------------------------

def test_cli_gerar_dashboard_html_ok(tmp_path, capsys):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS, _FIXTURE_MENSAL_2)
    out_dir = tmp_path / "out"
    result = main(
        ["gerar-dashboard-html", "--run-dir", str(tmp_path), "--output-dir", str(out_dir)]
    )
    captured = capsys.readouterr()
    assert result == 0
    assert "[OK] Dashboard HTML gerencial gerado." in captured.out
    assert "TESTE_TECNICO_NAO_OFICIAL" in captured.out
    assert (out_dir / "02_DASHBOARD_GERENCIAL.html").exists()


def test_cli_gerar_dashboard_html_run_dir_inexistente(tmp_path, capsys):
    result = main(
        ["gerar-dashboard-html", "--run-dir", str(tmp_path / "nao_existe")]
    )
    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO]" in captured.out


def test_cli_gerar_dashboard_html_sem_dre(tmp_path, capsys):
    result = main(["gerar-dashboard-html", "--run-dir", str(tmp_path)])
    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO]" in captured.out


def test_cli_gerar_dashboard_html_output_dir_default(tmp_path, capsys):
    _make_dre_xlsx(tmp_path, _FIXTURE_ROWS)
    result = main(["gerar-dashboard-html", "--run-dir", str(tmp_path)])
    assert result == 0
    assert (tmp_path / "output_visual" / "02_DASHBOARD_GERENCIAL.html").exists()
