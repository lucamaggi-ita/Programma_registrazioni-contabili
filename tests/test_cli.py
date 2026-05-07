from pathlib import Path
from shutil import copyfile

from openpyxl import Workbook, load_workbook
import pytest

from fantozzi.cli import build_parser, main
from fantozzi.classification_review_export import CSV_HEADERS
from fantozzi.excel_writer import HEADERS, LANCAMENTOS_SHEET, REVISAO_SHEET


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_XML = ROOT / "samples" / "nfe_exemplo" / "nfe_oesa_12566_20260505.xml"


def test_parser_accepts_importar_nfe_command():
    parser = build_parser()

    args = parser.parse_args(
        ["importar-nfe", "--xml", "sample.xml", "--excel", "output.xlsx"]
    )

    assert args.command == "importar-nfe"
    assert args.xml == "sample.xml"
    assert args.excel == "output.xlsx"


def test_parser_accepts_importar_pasta_nfe_command():
    parser = build_parser()

    args = parser.parse_args(
        ["importar-pasta-nfe", "--pasta", "xmls", "--excel", "output.xlsx"]
    )

    assert args.command == "importar-pasta-nfe"
    assert args.pasta == "xmls"
    assert args.excel == "output.xlsx"


def test_parser_accepts_diagnosticar_excel_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "diagnosticar-excel",
            "--excel",
            "output.xlsx",
            "--relatorio",
            "reports/diagnostico.txt",
        ]
    )

    assert args.command == "diagnosticar-excel"
    assert args.excel == "output.xlsx"
    assert args.relatorio == "reports/diagnostico.txt"


def test_parser_accepts_relatorio_revisao_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "relatorio-revisao",
            "--excel",
            "output.xlsx",
            "--saida",
            "reports/revisao.txt",
        ]
    )

    assert args.command == "relatorio-revisao"
    assert args.excel == "output.xlsx"
    assert args.saida == "reports/revisao.txt"


def test_parser_accepts_relatorio_candidatos_classificacao_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "relatorio-candidatos-classificacao",
            "--excel",
            "output.xlsx",
            "--saida",
            "reports/candidatos.txt",
        ]
    )

    assert args.command == "relatorio-candidatos-classificacao"
    assert args.excel == "output.xlsx"
    assert args.saida == "reports/candidatos.txt"


def test_parser_accepts_exportar_revisao_classificacao_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "exportar-revisao-classificacao",
            "--excel",
            "output.xlsx",
            "--saida",
            "reports/revisao_classificacao.csv",
        ]
    )

    assert args.command == "exportar-revisao-classificacao"
    assert args.excel == "output.xlsx"
    assert args.saida == "reports/revisao_classificacao.csv"


def test_parser_accepts_validar_revisao_classificacao_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "validar-revisao-classificacao",
            "--csv",
            "reports/revisao_classificacao.csv",
            "--relatorio",
            "reports/validacao_revisao_classificacao.txt",
        ]
    )

    assert args.command == "validar-revisao-classificacao"
    assert args.csv == "reports/revisao_classificacao.csv"
    assert args.relatorio == "reports/validacao_revisao_classificacao.txt"


def test_parser_accepts_gerar_proposta_classificacao_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "gerar-proposta-classificacao",
            "--csv",
            "reports/revisao_classificacao.csv",
            "--saida-yaml",
            "reports/proposta_regras_classificacao.yaml",
            "--relatorio",
            "reports/proposta_regras_classificacao.txt",
        ]
    )

    assert args.command == "gerar-proposta-classificacao"
    assert args.csv == "reports/revisao_classificacao.csv"
    assert args.saida_yaml == "reports/proposta_regras_classificacao.yaml"
    assert args.relatorio == "reports/proposta_regras_classificacao.txt"


def test_parser_accepts_gerar_proposta_regras_classificacao_real_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "gerar-proposta-regras-classificacao-real",
            "--csv",
            "reports/importacao_real/55746722000181/revisao_classificacao_revisada.csv",
            "--csv",
            "reports/importacao_real/8884180000123/revisao_classificacao_revisada.csv",
            "--saida",
            "reports/importacao_real/proposta_real.yaml",
            "--relatorio",
            "reports/importacao_real/resumo_proposta_real.txt",
        ]
    )

    assert args.command == "gerar-proposta-regras-classificacao-real"
    assert args.csv == [
        "reports/importacao_real/55746722000181/revisao_classificacao_revisada.csv",
        "reports/importacao_real/8884180000123/revisao_classificacao_revisada.csv",
    ]
    assert args.saida == "reports/importacao_real/proposta_real.yaml"
    assert args.relatorio == "reports/importacao_real/resumo_proposta_real.txt"


def test_parser_accepts_auditar_operacoes_especiais_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "auditar-operacoes-especiais",
            "--pasta",
            "xmls",
            "--saida",
            "reports/operacoes_especiais.txt",
        ]
    )

    assert args.command == "auditar-operacoes-especiais"
    assert args.pasta == "xmls"
    assert args.saida == "reports/operacoes_especiais.txt"


def test_parser_accepts_gerar_planilha_revisao_classificacao_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "gerar-planilha-revisao-classificacao",
            "--csv",
            "reports/revisao_classificacao.csv",
            "--saida",
            "reports/revisao_classificacao_assistida.xlsx",
        ]
    )

    assert args.command == "gerar-planilha-revisao-classificacao"
    assert args.csv == "reports/revisao_classificacao.csv"
    assert args.saida == "reports/revisao_classificacao_assistida.xlsx"


def test_parser_accepts_importar_planilha_revisao_classificacao_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "importar-planilha-revisao-classificacao",
            "--xlsx",
            "reports/revisao_classificacao_assistida.xlsx",
            "--saida",
            "reports/revisao_classificacao_revisada.csv",
            "--relatorio",
            "reports/validacao_planilha.txt",
        ]
    )

    assert args.command == "importar-planilha-revisao-classificacao"
    assert args.xlsx == "reports/revisao_classificacao_assistida.xlsx"
    assert args.saida == "reports/revisao_classificacao_revisada.csv"
    assert args.relatorio == "reports/validacao_planilha.txt"


def test_parser_accepts_gerar_pre_livro_diario_teste_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "gerar-pre-livro-diario-teste",
            "--excel",
            "output/empresa_1.xlsx",
            "--excel",
            "output/empresa_2.xlsx",
            "--pasta-operacoes-especiais",
            "XML_fatture_acquisto/55746722000181/xml",
            "--relatorio-auditoria",
            "reports/auditoria_pre_livro.txt",
            "--saida",
            "reports/pre_livro_diario_teste.xlsx",
        ]
    )

    assert args.command == "gerar-pre-livro-diario-teste"
    assert args.excel == ["output/empresa_1.xlsx", "output/empresa_2.xlsx"]
    assert args.pasta_operacoes_especiais == [
        "XML_fatture_acquisto/55746722000181/xml"
    ]
    assert args.relatorio_auditoria == "reports/auditoria_pre_livro.txt"
    assert args.saida == "reports/pre_livro_diario_teste.xlsx"


def test_parser_accepts_gerar_pacote_revisao_contabil_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "gerar-pacote-revisao-contabil",
            "--csv",
            "reports/importacao_real/55746722000181/revisao_classificacao_revisada.csv",
            "--csv",
            "reports/importacao_real/8884180000123/revisao_classificacao_revisada.csv",
            "--saida-xlsx",
            "reports/importacao_real/pacote_revisao_contabil.xlsx",
            "--relatorio",
            "reports/importacao_real/pacote_revisao_contabil.txt",
            "--limite-top",
            "25",
        ]
    )

    assert args.command == "gerar-pacote-revisao-contabil"
    assert args.csv == [
        "reports/importacao_real/55746722000181/revisao_classificacao_revisada.csv",
        "reports/importacao_real/8884180000123/revisao_classificacao_revisada.csv",
    ]
    assert args.saida_xlsx == "reports/importacao_real/pacote_revisao_contabil.xlsx"
    assert args.relatorio == "reports/importacao_real/pacote_revisao_contabil.txt"
    assert args.limite_top == 25


def test_parser_accepts_gerar_relatorios_gerenciais_ok_provisorios_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "gerar-relatorios-gerenciais-ok-provisorios",
            "--entrada",
            "output/importacao_real",
            "--saida-xlsx",
            "reports/importacao_real/relatorios_gerenciais_ok_provisorios.xlsx",
            "--relatorio",
            "reports/importacao_real/relatorios_gerenciais_ok_provisorios.txt",
            "--pasta-operacoes-especiais",
            "XML_fatture_acquisto/55746722000181/xml",
        ]
    )

    assert args.command == "gerar-relatorios-gerenciais-ok-provisorios"
    assert args.entrada == ["output/importacao_real"]
    assert (
        args.saida_xlsx
        == "reports/importacao_real/relatorios_gerenciais_ok_provisorios.xlsx"
    )
    assert (
        args.relatorio
        == "reports/importacao_real/relatorios_gerenciais_ok_provisorios.txt"
    )
    assert args.pasta_operacoes_especiais == [
        "XML_fatture_acquisto/55746722000181/xml"
    ]


def test_parser_accepts_auditar_relatorios_gerenciais_ok_provisorios_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "auditar-relatorios-gerenciais-ok-provisorios",
            "--entrada",
            "output/importacao_real",
            "--saida-xlsx",
            "reports/importacao_real/auditoria_relatorios_ok.xlsx",
            "--relatorio",
            "reports/importacao_real/auditoria_relatorios_ok.txt",
            "--pasta-operacoes-especiais",
            "XML_fatture_acquisto/55746722000181/xml",
        ]
    )

    assert args.command == "auditar-relatorios-gerenciais-ok-provisorios"
    assert args.entrada == ["output/importacao_real"]
    assert args.saida_xlsx == "reports/importacao_real/auditoria_relatorios_ok.xlsx"
    assert args.relatorio == "reports/importacao_real/auditoria_relatorios_ok.txt"
    assert args.pasta_operacoes_especiais == [
        "XML_fatture_acquisto/55746722000181/xml"
    ]


def test_parser_accepts_gerar_dre_gerencial_ok_provisoria_command():
    parser = build_parser()

    args = parser.parse_args(
        [
            "gerar-dre-gerencial-ok-provisoria",
            "--entrada",
            "output/importacao_real",
            "--mapeamento-dre",
            "config/mapeamento_dre.yaml",
            "--saida-xlsx",
            "reports/importacao_real/dre_gerencial_ok_provisoria.xlsx",
            "--relatorio",
            "reports/importacao_real/dre_gerencial_ok_provisoria.txt",
        ]
    )

    assert args.command == "gerar-dre-gerencial-ok-provisoria"
    assert args.entrada == ["output/importacao_real"]
    assert args.mapeamento_dre == "config/mapeamento_dre.yaml"
    assert args.saida_xlsx == "reports/importacao_real/dre_gerencial_ok_provisoria.xlsx"
    assert args.relatorio == "reports/importacao_real/dre_gerencial_ok_provisoria.txt"


def test_general_help_returns_zero_and_lists_commands(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "Fantozzi" in captured.out
    assert "importar-nfe" in captured.out
    assert "importar-pasta-nfe" in captured.out
    assert "diagnosticar-excel" in captured.out
    assert "relatorio-revisao" in captured.out
    assert "relatorio-candidatos-classificacao" in captured.out
    assert "exportar-revisao-classificacao" in captured.out
    assert "validar-revisao-classificacao" in captured.out
    assert "gerar-proposta-classificacao" in captured.out
    assert "gerar-proposta-regras-classificacao-real" in captured.out
    assert "auditar-operacoes-especiais" in captured.out
    assert "gerar-planilha-revisao-classificacao" in captured.out
    assert "importar-planilha-revisao-classificacao" in captured.out
    assert "gerar-pre-livro-diario-teste" in captured.out
    assert "gerar-pacote-revisao-contabil" in captured.out
    assert "gerar-relatorios-gerenciais-ok-provisorios" in captured.out
    assert "auditar-relatorios-gerenciais-ok-provisorios" in captured.out
    assert "gerar-dre-gerencial-ok-provisoria" in captured.out
    assert "reprocessar-importacao-real" in captured.out


def test_importar_nfe_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["importar-nfe", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--xml" in captured.out
    assert "--excel" in captured.out
    assert "--plano-contas" in captured.out
    assert "--regras-produtos" in captured.out


def test_importar_pasta_nfe_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["importar-pasta-nfe", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--pasta" in captured.out
    assert "--excel" in captured.out
    assert "--plano-contas" in captured.out
    assert "--regras-produtos" in captured.out


def test_diagnosticar_excel_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["diagnosticar-excel", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--excel" in captured.out
    assert "--relatorio" in captured.out


def test_relatorio_revisao_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["relatorio-revisao", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--excel" in captured.out
    assert "--saida" in captured.out


def test_relatorio_candidatos_classificacao_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["relatorio-candidatos-classificacao", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--excel" in captured.out
    assert "--saida" in captured.out


def test_exportar_revisao_classificacao_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["exportar-revisao-classificacao", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--excel" in captured.out
    assert "--saida" in captured.out


def test_validar_revisao_classificacao_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["validar-revisao-classificacao", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--csv" in captured.out
    assert "--relatorio" in captured.out


def test_gerar_proposta_classificacao_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["gerar-proposta-classificacao", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--csv" in captured.out
    assert "--saida-yaml" in captured.out
    assert "--relatorio" in captured.out
    assert "--plano-contas" in captured.out
    assert "--regras-produtos" in captured.out


def test_gerar_proposta_regras_classificacao_real_help_returns_zero_and_lists_options(
    capsys,
):
    with pytest.raises(SystemExit) as exc_info:
        main(["gerar-proposta-regras-classificacao-real", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--csv" in captured.out
    assert "--saida" in captured.out
    assert "--relatorio" in captured.out
    assert "--plano-contas" in captured.out


def test_auditar_operacoes_especiais_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["auditar-operacoes-especiais", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--pasta" in captured.out
    assert "--saida" in captured.out


def test_gerar_planilha_revisao_classificacao_help_returns_zero_and_lists_options(
    capsys,
):
    with pytest.raises(SystemExit) as exc_info:
        main(["gerar-planilha-revisao-classificacao", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--csv" in captured.out
    assert "--saida" in captured.out
    assert "--plano-contas" in captured.out


def test_importar_planilha_revisao_classificacao_help_returns_zero_and_lists_options(
    capsys,
):
    with pytest.raises(SystemExit) as exc_info:
        main(["importar-planilha-revisao-classificacao", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--xlsx" in captured.out
    assert "--saida" in captured.out
    assert "--relatorio" in captured.out
    assert "--plano-contas" in captured.out


def test_gerar_pre_livro_diario_teste_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["gerar-pre-livro-diario-teste", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--excel" in captured.out
    assert "--saida" in captured.out
    assert "--pasta-operacoes-especiais" in captured.out
    assert "--relatorio-auditoria" in captured.out


def test_gerar_pacote_revisao_contabil_help_returns_zero_and_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["gerar-pacote-revisao-contabil", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--csv" in captured.out
    assert "--saida-xlsx" in captured.out
    assert "--relatorio" in captured.out
    assert "--limite-top" in captured.out
    assert "--separador" in captured.out


def test_gerar_relatorios_gerenciais_ok_provisorios_help_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["gerar-relatorios-gerenciais-ok-provisorios", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--entrada" in captured.out
    assert "--saida-xlsx" in captured.out
    assert "--relatorio" in captured.out
    assert "--pasta-operacoes-especiais" in captured.out


def test_auditar_relatorios_gerenciais_ok_provisorios_help_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["auditar-relatorios-gerenciais-ok-provisorios", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--entrada" in captured.out
    assert "--saida-xlsx" in captured.out
    assert "--relatorio" in captured.out
    assert "--pasta-operacoes-especiais" in captured.out


def test_gerar_dre_gerencial_ok_provisoria_help_lists_options(capsys):
    with pytest.raises(SystemExit) as exc_info:
        main(["gerar-dre-gerencial-ok-provisoria", "--help"])

    captured = capsys.readouterr()
    assert exc_info.value.code == 0
    assert "--entrada" in captured.out
    assert "--mapeamento-dre" in captured.out
    assert "--saida-xlsx" in captured.out
    assert "--relatorio" in captured.out


def test_importar_nfe_command_writes_excel_with_sample_xml(tmp_path):
    excel_path = tmp_path / "lancamentos_cli.xlsx"

    result = main(
        ["importar-nfe", "--xml", str(SAMPLE_XML), "--excel", str(excel_path)]
    )

    assert result == 0
    assert excel_path.exists()
    workbook = load_workbook(excel_path)
    assert "Lancamentos" in workbook.sheetnames
    assert workbook["Lancamentos"].max_row >= 2


def test_importar_nfe_command_is_idempotent(tmp_path):
    excel_path = tmp_path / "lancamentos_cli.xlsx"
    argv = ["importar-nfe", "--xml", str(SAMPLE_XML), "--excel", str(excel_path)]

    first_result = main(argv)
    first_rows = load_workbook(excel_path)["Lancamentos"].max_row
    second_result = main(argv)
    second_rows = load_workbook(excel_path)["Lancamentos"].max_row

    assert first_result == 0
    assert second_result == 0
    assert second_rows == first_rows


def test_importar_nfe_command_returns_error_for_missing_xml(tmp_path):
    excel_path = tmp_path / "missing.xlsx"
    missing_xml = tmp_path / "arquivo_inexistente.xml"

    result = main(
        ["importar-nfe", "--xml", str(missing_xml), "--excel", str(excel_path)]
    )

    assert result != 0
    assert not excel_path.exists()


def test_importar_pasta_nfe_command_writes_single_excel(tmp_path):
    xml_dir = _xml_dir_with_sample(tmp_path)
    excel_path = tmp_path / "lancamentos_pasta.xlsx"

    result = main(
        ["importar-pasta-nfe", "--pasta", str(xml_dir), "--excel", str(excel_path)]
    )

    assert result == 0
    assert excel_path.exists()
    workbook = load_workbook(excel_path)
    assert "Lancamentos" in workbook.sheetnames
    assert workbook["Lancamentos"].max_row >= 2


def test_importar_pasta_nfe_command_is_idempotent(tmp_path):
    xml_dir = _xml_dir_with_sample(tmp_path)
    excel_path = tmp_path / "lancamentos_pasta.xlsx"
    argv = ["importar-pasta-nfe", "--pasta", str(xml_dir), "--excel", str(excel_path)]

    first_result = main(argv)
    first_rows = load_workbook(excel_path)["Lancamentos"].max_row
    second_result = main(argv)
    second_rows = load_workbook(excel_path)["Lancamentos"].max_row

    assert first_result == 0
    assert second_result == 0
    assert second_rows == first_rows


def test_importar_pasta_nfe_command_returns_error_for_missing_folder(tmp_path):
    excel_path = tmp_path / "missing_folder.xlsx"
    missing_folder = tmp_path / "pasta_inexistente"

    result = main(
        [
            "importar-pasta-nfe",
            "--pasta",
            str(missing_folder),
            "--excel",
            str(excel_path),
        ]
    )

    assert result != 0
    assert not excel_path.exists()


def test_importar_pasta_nfe_command_returns_error_when_folder_has_no_xml(tmp_path):
    xml_dir = tmp_path / "sem_xml"
    xml_dir.mkdir()
    (xml_dir / "nota.txt").write_text("nao e xml", encoding="utf-8")
    excel_path = tmp_path / "sem_xml.xlsx"

    result = main(
        ["importar-pasta-nfe", "--pasta", str(xml_dir), "--excel", str(excel_path)]
    )

    assert result != 0
    assert not excel_path.exists()


def test_importar_pasta_nfe_command_ignores_invalid_xml_and_imports_valid(
    tmp_path, capsys
):
    xml_dir = _xml_dir_with_sample(tmp_path, filename="valid_sample.XML")
    (xml_dir / "invalid.xml").write_text("<root></root>", encoding="utf-8")
    excel_path = tmp_path / "mixed_xml.xlsx"

    result = main(
        ["importar-pasta-nfe", "--pasta", str(xml_dir), "--excel", str(excel_path)]
    )

    assert result == 0
    assert excel_path.exists()
    assert load_workbook(excel_path)["Lancamentos"].max_row >= 2
    captured = capsys.readouterr()
    assert "XML ignorado" in captured.out


def test_relatorio_revisao_command_with_saida_returns_success(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"
    report_path = tmp_path / "reports" / "revisao.txt"
    _save_revision_workbook(excel_path)

    result = main(
        [
            "relatorio-revisao",
            "--excel",
            str(excel_path),
            "--saida",
            str(report_path),
        ]
    )

    assert result == 0
    assert report_path.exists()


def test_relatorio_revisao_command_returns_error_for_missing_excel(tmp_path):
    missing_excel = tmp_path / "inexistente.xlsx"

    result = main(["relatorio-revisao", "--excel", str(missing_excel)])

    assert result == 1


def test_relatorio_candidatos_command_with_saida_returns_success(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"
    report_path = tmp_path / "reports" / "candidatos.txt"
    _save_candidates_workbook(excel_path)

    result = main(
        [
            "relatorio-candidatos-classificacao",
            "--excel",
            str(excel_path),
            "--saida",
            str(report_path),
        ]
    )

    assert result == 0
    assert report_path.exists()


def test_relatorio_candidatos_command_without_saida_prints_report(tmp_path, capsys):
    excel_path = tmp_path / "lancamentos.xlsx"
    _save_candidates_workbook(excel_path)

    result = main(["relatorio-candidatos-classificacao", "--excel", str(excel_path)])

    captured = capsys.readouterr()
    assert result == 0
    assert "RELATORIO DE CANDIDATOS PARA CLASSIFICACAO" in captured.out
    assert "PRODUTO DESCONHECIDO" in captured.out


def test_relatorio_candidatos_command_returns_error_for_missing_excel(tmp_path):
    missing_excel = tmp_path / "inexistente.xlsx"

    result = main(
        ["relatorio-candidatos-classificacao", "--excel", str(missing_excel)]
    )

    assert result == 1


def test_exportar_revisao_classificacao_command_returns_success(tmp_path):
    excel_path = tmp_path / "lancamentos.xlsx"
    csv_path = tmp_path / "reports" / "revisao_classificacao.csv"
    _save_candidates_workbook(excel_path)

    result = main(
        [
            "exportar-revisao-classificacao",
            "--excel",
            str(excel_path),
            "--saida",
            str(csv_path),
        ]
    )

    assert result == 0
    assert csv_path.exists()


def test_exportar_revisao_classificacao_command_returns_error_for_missing_excel(
    tmp_path,
):
    missing_excel = tmp_path / "inexistente.xlsx"
    csv_path = tmp_path / "reports" / "revisao_classificacao.csv"

    result = main(
        [
            "exportar-revisao-classificacao",
            "--excel",
            str(missing_excel),
            "--saida",
            str(csv_path),
        ]
    )

    assert result == 1
    assert not csv_path.exists()


def test_validar_revisao_classificacao_command_returns_success(tmp_path, capsys):
    csv_path = tmp_path / "revisao_classificacao.csv"
    _save_review_csv(csv_path)

    result = main(["validar-revisao-classificacao", "--csv", str(csv_path)])

    captured = capsys.readouterr()
    assert result == 0
    assert "VALIDACAO DA REVISAO MANUAL DE CLASSIFICACAO" in captured.out
    assert "Total de linhas: 1" in captured.out


def test_validar_revisao_classificacao_command_saves_report(tmp_path):
    csv_path = tmp_path / "revisao_classificacao.csv"
    report_path = tmp_path / "reports" / "validacao.txt"
    _save_review_csv(csv_path)

    result = main(
        [
            "validar-revisao-classificacao",
            "--csv",
            str(csv_path),
            "--relatorio",
            str(report_path),
        ]
    )

    assert result == 0
    assert report_path.exists()


def test_validar_revisao_classificacao_command_returns_error_for_missing_csv(
    tmp_path,
):
    csv_path = tmp_path / "inexistente.csv"

    result = main(["validar-revisao-classificacao", "--csv", str(csv_path)])

    assert result == 1


def test_gerar_proposta_classificacao_command_with_yaml_returns_success(tmp_path):
    csv_path = tmp_path / "revisao_classificacao.csv"
    yaml_path = tmp_path / "reports" / "proposta.yaml"
    plano_path = _save_proposal_plano_contas(tmp_path)
    regras_path = _save_proposal_regras_produtos(tmp_path)
    _save_review_csv(csv_path)

    result = main(
        [
            "gerar-proposta-classificacao",
            "--csv",
            str(csv_path),
            "--saida-yaml",
            str(yaml_path),
            "--plano-contas",
            str(plano_path),
            "--regras-produtos",
            str(regras_path),
        ]
    )

    assert result == 0
    assert yaml_path.exists()
    assert "propostas:" in yaml_path.read_text(encoding="utf-8")


def test_gerar_proposta_classificacao_command_with_report_returns_success(tmp_path):
    csv_path = tmp_path / "revisao_classificacao.csv"
    report_path = tmp_path / "reports" / "proposta.txt"
    plano_path = _save_proposal_plano_contas(tmp_path)
    regras_path = _save_proposal_regras_produtos(tmp_path)
    _save_review_csv(csv_path)

    result = main(
        [
            "gerar-proposta-classificacao",
            "--csv",
            str(csv_path),
            "--relatorio",
            str(report_path),
            "--plano-contas",
            str(plano_path),
            "--regras-produtos",
            str(regras_path),
        ]
    )

    assert result == 0
    assert report_path.exists()
    assert "PROPOSTA DE REGRAS DE CLASSIFICACAO" in report_path.read_text(
        encoding="utf-8"
    )


def test_gerar_proposta_classificacao_command_with_yaml_and_report_returns_success(
    tmp_path,
):
    csv_path = tmp_path / "revisao_classificacao.csv"
    yaml_path = tmp_path / "reports" / "proposta.yaml"
    report_path = tmp_path / "reports" / "proposta.txt"
    plano_path = _save_proposal_plano_contas(tmp_path)
    regras_path = _save_proposal_regras_produtos(tmp_path)
    original_plan = plano_path.read_text(encoding="utf-8")
    original_rules = regras_path.read_text(encoding="utf-8")
    _save_review_csv(csv_path)

    result = main(
        [
            "gerar-proposta-classificacao",
            "--csv",
            str(csv_path),
            "--saida-yaml",
            str(yaml_path),
            "--relatorio",
            str(report_path),
            "--plano-contas",
            str(plano_path),
            "--regras-produtos",
            str(regras_path),
        ]
    )

    assert result == 0
    assert yaml_path.exists()
    assert report_path.exists()
    assert plano_path.read_text(encoding="utf-8") == original_plan
    assert regras_path.read_text(encoding="utf-8") == original_rules


def test_gerar_proposta_classificacao_command_requires_output(tmp_path, capsys):
    csv_path = tmp_path / "revisao_classificacao.csv"
    _save_review_csv(csv_path)

    result = main(["gerar-proposta-classificacao", "--csv", str(csv_path)])

    captured = capsys.readouterr()
    assert result == 1
    assert "Informe --saida-yaml e/ou --relatorio" in captured.out


def test_gerar_proposta_classificacao_command_returns_error_for_missing_csv(
    tmp_path,
):
    csv_path = tmp_path / "inexistente.csv"
    yaml_path = tmp_path / "reports" / "proposta.yaml"

    result = main(
        [
            "gerar-proposta-classificacao",
            "--csv",
            str(csv_path),
            "--saida-yaml",
            str(yaml_path),
        ]
    )

    assert result == 1
    assert not yaml_path.exists()


def test_gerar_proposta_regras_classificacao_real_command_returns_success(tmp_path):
    csv_path = tmp_path / "revisao_classificacao_revisada.csv"
    yaml_path = tmp_path / "reports" / "proposta_real.yaml"
    report_path = tmp_path / "reports" / "proposta_real.txt"
    plano_path = _save_proposal_plano_contas(tmp_path)
    original_plan = plano_path.read_text(encoding="utf-8")
    _save_review_csv(csv_path)

    result = main(
        [
            "gerar-proposta-regras-classificacao-real",
            "--csv",
            str(csv_path),
            "--saida",
            str(yaml_path),
            "--relatorio",
            str(report_path),
            "--plano-contas",
            str(plano_path),
        ]
    )

    assert result == 0
    assert yaml_path.exists()
    assert report_path.exists()
    assert "metadata:" in yaml_path.read_text(encoding="utf-8")
    assert "propostas_alias:" in yaml_path.read_text(encoding="utf-8")
    assert plano_path.read_text(encoding="utf-8") == original_plan


def test_gerar_proposta_regras_classificacao_real_command_returns_error_for_missing_csv(
    tmp_path,
):
    csv_path = tmp_path / "inexistente.csv"
    yaml_path = tmp_path / "reports" / "proposta_real.yaml"
    plano_path = _save_proposal_plano_contas(tmp_path)

    result = main(
        [
            "gerar-proposta-regras-classificacao-real",
            "--csv",
            str(csv_path),
            "--saida",
            str(yaml_path),
            "--plano-contas",
            str(plano_path),
        ]
    )

    assert result == 1
    assert not yaml_path.exists()


def test_auditar_operacoes_especiais_command_generates_report(tmp_path):
    xml_dir = _xml_dir_with_special_operation(tmp_path)
    report_path = tmp_path / "reports" / "operacoes_especiais.txt"

    result = main(
        [
            "auditar-operacoes-especiais",
            "--pasta",
            str(xml_dir),
            "--saida",
            str(report_path),
        ]
    )

    assert result == 0
    assert report_path.exists()
    text = report_path.read_text(encoding="utf-8")
    assert "FANTOZZI - OPERACOES ESPECIAIS" in text
    assert "COMODATO" in text
    assert "ANALISAR_COM_CONTADOR" in text


def test_auditar_operacoes_especiais_command_returns_error_for_missing_folder(
    tmp_path,
):
    missing_folder = tmp_path / "xmls_inexistentes"
    report_path = tmp_path / "reports" / "operacoes_especiais.txt"

    result = main(
        [
            "auditar-operacoes-especiais",
            "--pasta",
            str(missing_folder),
            "--saida",
            str(report_path),
        ]
    )

    assert result == 1
    assert not report_path.exists()


def test_gerar_planilha_revisao_classificacao_command_gera_xlsx(tmp_path, capsys):
    csv_path = tmp_path / "revisao_classificacao.csv"
    output_path = tmp_path / "reports" / "revisao_classificacao_assistida.xlsx"
    plano_path = _save_proposal_plano_contas(tmp_path)
    _save_review_csv(csv_path)
    original_csv = csv_path.read_text(encoding="utf-8-sig")

    result = main(
        [
            "gerar-planilha-revisao-classificacao",
            "--csv",
            str(csv_path),
            "--saida",
            str(output_path),
            "--plano-contas",
            str(plano_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 0
    assert output_path.exists()
    assert "[OK] Planilha assistida gerada:" in captured.out
    assert load_workbook(output_path)["Revisao"].max_row == 2
    assert csv_path.read_text(encoding="utf-8-sig") == original_csv


def test_gerar_planilha_revisao_classificacao_command_retorna_erro_csv_inexistente(
    tmp_path,
    capsys,
):
    missing_csv = tmp_path / "inexistente.csv"
    output_path = tmp_path / "reports" / "revisao_classificacao_assistida.xlsx"
    plano_path = _save_proposal_plano_contas(tmp_path)

    result = main(
        [
            "gerar-planilha-revisao-classificacao",
            "--csv",
            str(missing_csv),
            "--saida",
            str(output_path),
            "--plano-contas",
            str(plano_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO] CSV de revisao nao encontrado:" in captured.out
    assert not output_path.exists()


def test_importar_planilha_revisao_classificacao_command_gera_csv(tmp_path, capsys):
    xlsx_path = _save_assisted_review_workbook(tmp_path)
    csv_path = tmp_path / "reports" / "revisao_classificacao_revisada.csv"
    plano_path = _save_proposal_plano_contas(tmp_path)
    original_xlsx = xlsx_path.read_bytes()

    result = main(
        [
            "importar-planilha-revisao-classificacao",
            "--xlsx",
            str(xlsx_path),
            "--saida",
            str(csv_path),
            "--plano-contas",
            str(plano_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 0
    assert csv_path.exists()
    assert "[OK] Planilha importada com sucesso." in captured.out
    assert "[OK] CSV revisado gerado:" in captured.out
    assert xlsx_path.read_bytes() == original_xlsx


def test_importar_planilha_revisao_classificacao_command_aceita_relatorio(
    tmp_path,
    capsys,
):
    xlsx_path = _save_assisted_review_workbook(tmp_path)
    csv_path = tmp_path / "reports" / "revisao_classificacao_revisada.csv"
    report_path = tmp_path / "reports" / "validacao_planilha.txt"
    plano_path = _save_proposal_plano_contas(tmp_path)

    result = main(
        [
            "importar-planilha-revisao-classificacao",
            "--xlsx",
            str(xlsx_path),
            "--saida",
            str(csv_path),
            "--relatorio",
            str(report_path),
            "--plano-contas",
            str(plano_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 0
    assert report_path.exists()
    assert "[OK] Relatorio salvo em:" in captured.out


def test_importar_planilha_revisao_classificacao_command_retorna_erro_xlsx_inexistente(
    tmp_path,
    capsys,
):
    missing_xlsx = tmp_path / "inexistente.xlsx"
    csv_path = tmp_path / "reports" / "revisao_classificacao_revisada.csv"
    plano_path = _save_proposal_plano_contas(tmp_path)

    result = main(
        [
            "importar-planilha-revisao-classificacao",
            "--xlsx",
            str(missing_xlsx),
            "--saida",
            str(csv_path),
            "--plano-contas",
            str(plano_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO] Planilha assistida nao encontrada:" in captured.out
    assert not csv_path.exists()


def test_importar_planilha_revisao_classificacao_command_retorna_erro_acao_invalida(
    tmp_path,
    capsys,
):
    xlsx_path = _save_assisted_review_workbook(tmp_path, acao_manual="APROVAR")
    csv_path = tmp_path / "reports" / "revisao_classificacao_revisada.csv"
    plano_path = _save_proposal_plano_contas(tmp_path)

    result = main(
        [
            "importar-planilha-revisao-classificacao",
            "--xlsx",
            str(xlsx_path),
            "--saida",
            str(csv_path),
            "--plano-contas",
            str(plano_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO] Linha 2: acao_manual invalida: APROVAR" in captured.out
    assert not csv_path.exists()


def test_gerar_pre_livro_diario_teste_command_gera_xlsx(tmp_path, capsys):
    excel_path = tmp_path / "lancamentos.xlsx"
    output_path = tmp_path / "reports" / "pre_livro.xlsx"
    _save_revision_workbook(excel_path)

    result = main(
        [
            "gerar-pre-livro-diario-teste",
            "--excel",
            str(excel_path),
            "--saida",
            str(output_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 0
    assert output_path.exists()
    assert "[OK] Pre-livro diario gerencial de teste gerado." in captured.out


def test_gerar_pre_livro_diario_teste_command_aceita_relatorio_auditoria(
    tmp_path, capsys
):
    excel_path = tmp_path / "lancamentos.xlsx"
    output_path = tmp_path / "reports" / "pre_livro.xlsx"
    report_path = tmp_path / "reports" / "auditoria.txt"
    _save_revision_workbook(excel_path)

    result = main(
        [
            "gerar-pre-livro-diario-teste",
            "--excel",
            str(excel_path),
            "--saida",
            str(output_path),
            "--relatorio-auditoria",
            str(report_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 0
    assert report_path.exists()
    assert "[OK] Relatorio de auditoria salvo em:" in captured.out


def test_gerar_pre_livro_diario_teste_command_retorna_erro_excel_inexistente(
    tmp_path,
):
    missing_excel = tmp_path / "inexistente.xlsx"
    output_path = tmp_path / "reports" / "pre_livro.xlsx"

    result = main(
        [
            "gerar-pre-livro-diario-teste",
            "--excel",
            str(missing_excel),
            "--saida",
            str(output_path),
        ]
    )

    assert result == 1
    assert not output_path.exists()


def test_gerar_pacote_revisao_contabil_command_gera_xlsx_e_txt(tmp_path, capsys):
    csv_path = tmp_path / "55746722000181" / "revisao_classificacao_revisada.csv"
    csv_path.parent.mkdir(parents=True)
    _save_review_csv(csv_path)
    output_path = tmp_path / "reports" / "pacote_revisao_contabil.xlsx"
    report_path = tmp_path / "reports" / "pacote_revisao_contabil.txt"

    result = main(
        [
            "gerar-pacote-revisao-contabil",
            "--csv",
            str(csv_path),
            "--saida-xlsx",
            str(output_path),
            "--relatorio",
            str(report_path),
            "--limite-top",
            "10",
        ]
    )

    captured = capsys.readouterr()
    assert result == 0
    assert output_path.exists()
    assert report_path.exists()
    assert "[OK] Pacote assistido para revisao contabil gerado." in captured.out


def test_gerar_pacote_revisao_contabil_command_retorna_erro_csv_inexistente(
    tmp_path,
):
    missing_csv = tmp_path / "inexistente.csv"
    output_path = tmp_path / "reports" / "pacote.xlsx"

    result = main(
        [
            "gerar-pacote-revisao-contabil",
            "--csv",
            str(missing_csv),
            "--saida-xlsx",
            str(output_path),
        ]
    )

    assert result == 1
    assert not output_path.exists()


def test_gerar_relatorios_gerenciais_ok_provisorios_command_gera_xlsx_e_txt(
    tmp_path,
    capsys,
):
    excel_path = tmp_path / "55746722000181" / "lancamentos.xlsx"
    excel_path.parent.mkdir(parents=True)
    _save_revision_workbook(excel_path)
    output_path = tmp_path / "reports" / "relatorios_ok.xlsx"
    report_path = tmp_path / "reports" / "relatorios_ok.txt"

    result = main(
        [
            "gerar-relatorios-gerenciais-ok-provisorios",
            "--entrada",
            str(excel_path.parent),
            "--saida-xlsx",
            str(output_path),
            "--relatorio",
            str(report_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 0
    assert output_path.exists()
    assert report_path.exists()
    assert "[OK] Relatorios gerenciais OK provisorios gerados." in captured.out


def test_gerar_relatorios_gerenciais_ok_provisorios_command_retorna_erro_entrada_inexistente(
    tmp_path,
):
    output_path = tmp_path / "reports" / "relatorios_ok.xlsx"

    result = main(
        [
            "gerar-relatorios-gerenciais-ok-provisorios",
            "--entrada",
            str(tmp_path / "inexistente"),
            "--saida-xlsx",
            str(output_path),
        ]
    )

    assert result == 1
    assert not output_path.exists()


def test_auditar_relatorios_gerenciais_ok_provisorios_command_gera_xlsx_e_txt(
    tmp_path,
    capsys,
):
    excel_path = tmp_path / "55746722000181" / "lancamentos.xlsx"
    excel_path.parent.mkdir(parents=True)
    _save_revision_workbook(excel_path)
    output_path = tmp_path / "reports" / "auditoria.xlsx"
    report_path = tmp_path / "reports" / "auditoria.txt"

    result = main(
        [
            "auditar-relatorios-gerenciais-ok-provisorios",
            "--entrada",
            str(excel_path.parent),
            "--saida-xlsx",
            str(output_path),
            "--relatorio",
            str(report_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 0
    assert output_path.exists()
    assert report_path.exists()
    assert "[OK] Auditoria dos relatorios gerenciais OK provisorios gerada." in captured.out


def test_gerar_dre_gerencial_ok_provisoria_command_gera_xlsx_e_txt(
    tmp_path,
    capsys,
):
    excel_path = tmp_path / "55746722000181" / "lancamentos.xlsx"
    excel_path.parent.mkdir(parents=True)
    _save_revision_workbook(excel_path)
    mapping_path = tmp_path / "mapeamento_dre.yaml"
    mapping_path.write_text(
        """\
linhas_dre:
  - id: "03.02"
    nome: "CMV - Bebidas"
    grupo: "Custos"
    sinal: "-"
    ordem: 302
mapeamento:
  "2.01.004": { dre_linha_id: "03.02", entra_dre: true }
""",
        encoding="utf-8",
    )
    output_path = tmp_path / "reports" / "dre.xlsx"
    report_path = tmp_path / "reports" / "dre.txt"

    result = main(
        [
            "gerar-dre-gerencial-ok-provisoria",
            "--entrada",
            str(excel_path.parent),
            "--mapeamento-dre",
            str(mapping_path),
            "--saida-xlsx",
            str(output_path),
            "--relatorio",
            str(report_path),
        ]
    )

    captured = capsys.readouterr()
    assert result == 0
    assert output_path.exists()
    assert report_path.exists()
    assert "[OK] DRE gerencial OK provisoria gerada." in captured.out


def _xml_dir_with_sample(tmp_path: Path, *, filename: str = "sample.xml") -> Path:
    xml_dir = tmp_path / "nfe_xml"
    xml_dir.mkdir()
    copyfile(SAMPLE_XML, xml_dir / filename)
    return xml_dir


def _xml_dir_with_special_operation(tmp_path: Path) -> Path:
    xml_dir = tmp_path / "nfe_xml_especial"
    xml_dir.mkdir()
    text = SAMPLE_XML.read_text(encoding="utf-8")
    text = text.replace("<natOp>VENDA DE MERCADORIAS</natOp>", "<natOp>Comodato</natOp>")
    text = text.replace("<CFOP>5102</CFOP>", "<CFOP>5908</CFOP>", 1)
    (xml_dir / "comodato.xml").write_text(text, encoding="utf-8")
    return xml_dir


def _save_revision_workbook(excel_path: Path) -> None:
    workbook = Workbook()
    lancamentos_sheet = workbook.active
    lancamentos_sheet.title = LANCAMENTOS_SHEET
    revisao_sheet = workbook.create_sheet(REVISAO_SHEET)
    lancamentos_sheet.append(HEADERS)
    revisao_sheet.append(HEADERS)
    lancamentos_sheet.append(
        [
            {
                "data_documento": "2026-05-05",
                "fornecedor_ou_favorecido": "Fornecedor teste",
                "documento": "NF-1",
                "conta_id": "2.01.004",
                "conta_nome": "Refrigerantes",
                "descricao_itens": "COCA COLA",
                "valor": 10,
                "origem": "XML_NFE",
                "chave_origem": "chave-1",
                "status_revisao": "OK",
                "hash_registro": "hash-1",
                "duplicatas_json": "[]",
            }.get(header)
            for header in HEADERS
        ]
    )
    workbook.save(excel_path)


def _save_candidates_workbook(excel_path: Path) -> None:
    workbook = Workbook()
    lancamentos_sheet = workbook.active
    lancamentos_sheet.title = LANCAMENTOS_SHEET
    revisao_sheet = workbook.create_sheet(REVISAO_SHEET)
    lancamentos_sheet.append(HEADERS)
    revisao_sheet.append(HEADERS)
    revisao_sheet.append(
        [
            {
                "data_documento": "2026-05-05",
                "fornecedor_ou_favorecido": "Fornecedor teste",
                "documento": "NF-1",
                "conta_id": "13.01.001",
                "conta_nome": "Revisar - Produto Nao Classificado",
                "descricao_itens": "PRODUTO DESCONHECIDO",
                "valor": 10,
                "origem": "XML_NFE",
                "chave_origem": "chave-1",
                "status_revisao": "REVISAR",
                "hash_registro": "hash-1",
                "alertas": "CLASSIFICACAO_NAO_ENCONTRADA",
                "duplicatas_json": "[]",
            }.get(header)
            for header in HEADERS
        ]
    )
    workbook.save(excel_path)


def _save_review_csv(csv_path: Path) -> None:
    csv_path.write_text(
        ";".join(CSV_HEADERS)
        + "\n"
        + ";".join(
            [
                "CLASSIFICACAO_PRODUTO",
                "produto desconhecido",
                "PRODUTO DESCONHECIDO",
                "1",
                "Fornecedor teste",
                "NF-1",
                "R$ 10,00",
                "13.01.001",
                "Revisar - Produto Nao Classificado",
                "REVISAR",
                "CLASSIFICACAO_NAO_ENCONTRADA",
                "CRIAR_ALIAS",
                "2.01.004",
                "Refrigerantes",
                "produto desconhecido",
                "",
            ]
        )
        + "\n",
        encoding="utf-8-sig",
    )


def _save_assisted_review_workbook(
    tmp_path: Path,
    *,
    acao_manual: str = "CLASSIFICAR",
) -> Path:
    xlsx_path = tmp_path / "revisao_classificacao_assistida.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Revisao"
    sheet.append(CSV_HEADERS)
    row = {
        "tipo_candidato": "CLASSIFICACAO_PRODUTO",
        "descricao_normalizada": "produto desconhecido",
        "descricao_exemplo": "PRODUTO DESCONHECIDO",
        "ocorrencias": "1",
        "fornecedores": "Fornecedor teste",
        "documentos": "NF-1",
        "valor_total": "R$ 10,00",
        "conta_atual_id": "13.01.001",
        "conta_atual_nome": "Revisar - Produto Nao Classificado",
        "status_revisao": "REVISAR",
        "alertas": "CLASSIFICACAO_NAO_ENCONTRADA",
        "acao_manual": acao_manual,
        "conta_id_sugerida": "2.01.004",
        "conta_nome_sugerida": "Refrigerantes",
        "palavras_chave_sugeridas": "produto desconhecido",
        "observacao_manual": "",
    }
    sheet.append([row.get(header, "") for header in CSV_HEADERS])
    workbook.save(xlsx_path)
    return xlsx_path


def _save_proposal_plano_contas(tmp_path: Path) -> Path:
    plano_path = tmp_path / "plano_contas_gerencial.yaml"
    plano_path.write_text(
        """\
contas:
  - id: "2.01.004"
    nome: "Refrigerantes"
    grupo: "Custos de Mercadorias e Insumos"
    tipo: "custo"
  - id: "13.01.001"
    nome: "Revisar - Produto Nao Classificado"
    grupo: "Contas de Revisao e Classificacao Pendente"
    tipo: "revisao"
  - id: "13.01.002"
    nome: "Revisar - Classificacao Ambigua"
    grupo: "Contas de Revisao e Classificacao Pendente"
    tipo: "revisao"
  - id: "13.01.006"
    nome: "Revisar - Valor Inconsistente"
    grupo: "Contas de Revisao e Classificacao Pendente"
    tipo: "revisao"
  - id: "13.01.009"
    nome: "Revisar - Possivel Investimento"
    grupo: "Contas de Revisao e Classificacao Pendente"
    tipo: "revisao"
  - id: "13.01.010"
    nome: "Revisar - Pagamento de Pessoal a Regularizar"
    grupo: "Contas de Revisao e Classificacao Pendente"
    tipo: "revisao"
  - id: "13.01.999"
    nome: "Revisar - Outros Casos"
    grupo: "Contas de Revisao e Classificacao Pendente"
    tipo: "revisao"
""",
        encoding="utf-8",
    )
    return plano_path


def _save_proposal_regras_produtos(tmp_path: Path) -> Path:
    regras_path = tmp_path / "regras_classificacao_produtos.yaml"
    regras_path.write_text(
        """\
versao: "1.0"
regras:
  - conta_id: "2.01.004"
    conta_nome: "Refrigerantes"
    aliases:
      - coca cola
regras_por_fornecedor: []
""",
        encoding="utf-8",
    )
    return regras_path


# ---------------------------------------------------------------------------
# preparar-run-reprocessamento
# ---------------------------------------------------------------------------

def test_parser_aceita_preparar_run_reprocessamento():
    parser = build_parser()

    args = parser.parse_args(
        [
            "preparar-run-reprocessamento",
            "--xml-base", "XML_fatture_acquisto",
            "--saida-output", "output/importacao_real",
            "--saida-reports", "reports/importacao_real",
        ]
    )

    assert args.command == "preparar-run-reprocessamento"
    assert args.xml_base == "XML_fatture_acquisto"
    assert args.saida_output == "output/importacao_real"
    assert args.saida_reports == "reports/importacao_real"
    assert args.dry_run is False


def test_parser_aceita_reprocessar_importacao_real():
    parser = build_parser()

    args = parser.parse_args(
        [
            "reprocessar-importacao-real",
            "--xml-base", "XML_fatture_acquisto",
            "--saida-output", "output/importacao_real",
            "--saida-reports", "reports/importacao_real",
        ]
    )

    assert args.command == "reprocessar-importacao-real"
    assert args.xml_base == "XML_fatture_acquisto"
    assert args.saida_output == "output/importacao_real"
    assert args.saida_reports == "reports/importacao_real"
    assert args.modo == "provisorio"
    assert args.dry_run is False


def test_cli_preparar_run_cria_pastas_manifesto_current_run(tmp_path, capsys):
    import json

    xml_base = tmp_path / "xml_base"
    xml_base.mkdir()
    saida_output = tmp_path / "output"
    saida_reports = tmp_path / "reports"

    from fantozzi import run_manager as _rm
    original_root = _rm.PROJECT_ROOT
    _rm.PROJECT_ROOT = tmp_path
    try:
        result = main(
            [
                "preparar-run-reprocessamento",
                "--xml-base", str(xml_base),
                "--saida-output", str(saida_output),
                "--saida-reports", str(saida_reports),
            ]
        )
    finally:
        _rm.PROJECT_ROOT = original_root

    captured = capsys.readouterr()
    assert result == 0
    assert "[OK] Run preparada:" in captured.out
    assert "[OK] Manifesto criado:" in captured.out
    assert "[OK] current_run.txt atualizado:" in captured.out

    # Verifica arquivos criados
    import re
    run_id_match = re.search(r"\d{4}-\d{2}-\d{2}_\d{3}", captured.out)
    assert run_id_match, "run_id nao encontrado no output"
    run_id = run_id_match.group()

    manifest_path = saida_reports / "runs" / run_id / "manifest_run.json"
    assert manifest_path.exists()
    data = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert data["run_id"] == run_id
    assert data["status_run"] == "INICIADA"

    assert (saida_output / "runs" / run_id).is_dir()
    assert (saida_output / "current_run.txt").read_text(encoding="utf-8").strip() == run_id
    assert (saida_reports / "current_run.txt").read_text(encoding="utf-8").strip() == run_id

    # Manifesto NAO em output
    assert not (saida_output / "runs" / run_id / "manifest_run.json").exists()


def test_cli_preparar_run_dry_run_nao_escreve_nada(tmp_path, capsys):
    xml_base = tmp_path / "xml_base"
    xml_base.mkdir()
    saida_output = tmp_path / "output"
    saida_reports = tmp_path / "reports"

    from fantozzi import run_manager as _rm
    original_root = _rm.PROJECT_ROOT
    _rm.PROJECT_ROOT = tmp_path
    try:
        result = main(
            [
                "preparar-run-reprocessamento",
                "--xml-base", str(xml_base),
                "--saida-output", str(saida_output),
                "--saida-reports", str(saida_reports),
                "--dry-run",
            ]
        )
    finally:
        _rm.PROJECT_ROOT = original_root

    captured = capsys.readouterr()
    assert result == 0
    assert "[DRY-RUN]" in captured.out
    assert "Nenhum arquivo foi escrito" in captured.out
    # Nenhum arquivo ou pasta criado
    assert not saida_output.exists()
    assert not saida_reports.exists()


def test_cli_preparar_run_rejeita_xml_base_dentro_saida(tmp_path, capsys):
    saida_output = tmp_path / "output"
    saida_output.mkdir()
    xml_dentro_saida = saida_output / "xml_base"
    xml_dentro_saida.mkdir()

    from fantozzi import run_manager as _rm
    original_root = _rm.PROJECT_ROOT
    _rm.PROJECT_ROOT = tmp_path
    try:
        result = main(
            [
                "preparar-run-reprocessamento",
                "--xml-base", str(xml_dentro_saida),
                "--saida-output", str(saida_output),
                "--saida-reports", str(tmp_path / "reports"),
            ]
        )
    finally:
        _rm.PROJECT_ROOT = original_root

    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO]" in captured.out


def test_cli_reprocessar_importacao_real_dry_run_nao_escreve(tmp_path, capsys):
    xml_base = tmp_path / "XML_fatture_acquisto"
    (xml_base / "55746722000181" / "xml").mkdir(parents=True)
    saida_output = tmp_path / "output"
    saida_reports = tmp_path / "reports"

    from fantozzi import run_manager as _rm
    original_root = _rm.PROJECT_ROOT
    _rm.PROJECT_ROOT = tmp_path
    try:
        result = main(
            [
                "reprocessar-importacao-real",
                "--xml-base", str(xml_base),
                "--saida-output", str(saida_output),
                "--saida-reports", str(saida_reports),
                "--dry-run",
            ]
        )
    finally:
        _rm.PROJECT_ROOT = original_root

    captured = capsys.readouterr()
    assert result == 0
    assert "[DRY-RUN] Nenhum arquivo foi criado." in captured.out
    assert "[DRY-RUN] Nenhum XML foi importado." in captured.out
    assert "[DRY-RUN] Nenhum relatorio foi gerado." in captured.out
    assert not (saida_output / "runs").exists()
    assert not (saida_reports / "runs").exists()
    assert not (saida_output / "current_run.txt").exists()
    assert not (saida_reports / "current_run.txt").exists()


def test_cli_reprocessar_importacao_real_bloqueia_modo_definitivo(tmp_path, capsys):
    xml_base = tmp_path / "XML_fatture_acquisto"
    xml_base.mkdir()

    from fantozzi import run_manager as _rm
    original_root = _rm.PROJECT_ROOT
    _rm.PROJECT_ROOT = tmp_path
    try:
        result = main(
            [
                "reprocessar-importacao-real",
                "--xml-base", str(xml_base),
                "--saida-output", str(tmp_path / "output"),
                "--saida-reports", str(tmp_path / "reports"),
                "--modo", "definitivo",
            ]
        )
    finally:
        _rm.PROJECT_ROOT = original_root

    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO] Modo definitivo bloqueado: Fase 9 ainda nao foi aprovada." in captured.out


def test_cli_retorna_erro_quando_pipeline_falha(monkeypatch, tmp_path, capsys):
    from fantozzi import reprocessamento_pipeline

    def _falha(**kwargs):
        raise reprocessamento_pipeline.ReprocessamentoPipelineError("falha controlada")

    monkeypatch.setattr(
        reprocessamento_pipeline,
        "reprocessar_importacao_real_provisoria",
        _falha,
    )

    result = main(
        [
            "reprocessar-importacao-real",
            "--xml-base", str(tmp_path / "xml"),
            "--saida-output", str(tmp_path / "output"),
            "--saida-reports", str(tmp_path / "reports"),
        ]
    )

    captured = capsys.readouterr()
    assert result == 1
    assert "[ERRO] falha controlada" in captured.out
