from datetime import date
from decimal import Decimal
from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook

from fantozzi.dre_gerencial_provisoria import (
    AUDITORIA_DRE_SHEET,
    BASE_OK_DRE_SHEET,
    CONTAS_SEM_MAPEAMENTO_SHEET,
    DRE_CNPJ_SHEET,
    DRE_GERAL_SHEET,
    DRE_MENSAL_SHEET,
    INSTRUCOES_SHEET,
    gerar_dre_gerencial_ok_provisoria,
)
from fantozzi.excel_writer import write_lancamentos_to_excel
from fantozzi.models import Duplicata, Lancamento
from fantozzi.nfe_importer import parse_nfe


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_XML = ROOT / "samples" / "nfe_exemplo" / "nfe_oesa_12566_20260505.xml"


def test_gera_dre_apenas_com_ok_nao_especiais(tmp_path):
    excel_path = tmp_path / "55746722000181" / "lancamentos.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento("ok-1", conta_id="2.01.004", valor=Decimal("10.00")),
            _lancamento("revisar-1", status_revisao="REVISAR", valor=Decimal("7.00")),
        ],
    )
    mapping_path = _save_mapping(tmp_path)

    result = gerar_dre_gerencial_ok_provisoria(
        [excel_path], mapping_path, tmp_path / "dre.xlsx"
    )

    assert result.total_lancamentos_origem == 2
    assert result.lancamentos_base_ok == 1
    assert result.lancamentos_revisar_excluidos == 1
    assert result.valor_total_base_ok == Decimal("10.00")


def test_exclui_comodato_e_bonificacao(tmp_path):
    xml_comodato = _xml_dir_with_special_operation(tmp_path, natureza="Comodato", cfop="5908")
    xml_bonif = _xml_dir_with_special_operation(
        tmp_path, natureza="Bonificacao recebida", cfop="5910"
    )
    nfe_comodato = parse_nfe(next(xml_comodato.glob("*.xml")))
    nfe_bonif = parse_nfe(next(xml_bonif.glob("*.xml")))
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento("ok-normal", chave_origem="NORMAL"),
            _lancamento("ok-comodato", chave_origem=nfe_comodato.chave),
            _lancamento("ok-bonif", chave_origem=nfe_bonif.chave),
        ],
    )

    result = gerar_dre_gerencial_ok_provisoria(
        [excel_path],
        _save_mapping(tmp_path),
        tmp_path / "dre.xlsx",
        pastas_operacoes_especiais=[xml_comodato, xml_bonif],
    )

    assert result.lancamentos_base_ok == 1
    assert result.operacoes_especiais_excluidas == 2


def test_separa_contas_sem_mapeamento(tmp_path):
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(excel_path, [_lancamento("ok-sem", conta_id="9.99.999")])
    xlsx_path = tmp_path / "dre.xlsx"

    result = gerar_dre_gerencial_ok_provisoria([excel_path], _save_mapping(tmp_path), xlsx_path)

    sheet = load_workbook(xlsx_path, data_only=True)[CONTAS_SEM_MAPEAMENTO_SHEET]
    assert result.status_dre == "ATENCAO_MAPEAMENTO_INCOMPLETO"
    assert sheet.max_row == 2
    assert sheet["A2"].value == "9.99.999"


def test_calcula_dre_geral_e_mensal(tmp_path):
    excel_path = tmp_path / "55746722000181" / "lancamentos.xlsx"
    write_lancamentos_to_excel(
        excel_path,
        [
            _lancamento("receita", conta_id="1.01.001", valor=Decimal("100.00")),
            _lancamento("custo", conta_id="2.01.004", valor=Decimal("25.00")),
        ],
    )
    xlsx_path = tmp_path / "dre.xlsx"

    gerar_dre_gerencial_ok_provisoria([excel_path], _save_mapping(tmp_path), xlsx_path)

    workbook = load_workbook(xlsx_path, data_only=True)
    dre_geral = _sheet_rows(workbook, DRE_GERAL_SHEET)
    dre_mensal = _sheet_rows(workbook, DRE_MENSAL_SHEET)
    assert {row["linha_dre"] for row in dre_geral} == {
        "Receita Bruta de Vendas",
        "CMV - Bebidas",
    }
    assert len(dre_mensal) == 2


def test_gera_xlsx_com_abas_obrigatorias_e_txt_provisorio(tmp_path):
    excel_path = tmp_path / "empresa" / "lancamentos.xlsx"
    write_lancamentos_to_excel(excel_path, [_lancamento("ok-1")])
    xlsx_path = tmp_path / "dre.xlsx"
    report_path = tmp_path / "dre.txt"

    gerar_dre_gerencial_ok_provisoria(
        [excel_path], _save_mapping(tmp_path), xlsx_path, report_path
    )

    workbook = load_workbook(xlsx_path, data_only=True)
    assert workbook.sheetnames == [
        DRE_GERAL_SHEET,
        DRE_MENSAL_SHEET,
        DRE_CNPJ_SHEET,
        BASE_OK_DRE_SHEET,
        CONTAS_SEM_MAPEAMENTO_SHEET,
        AUDITORIA_DRE_SHEET,
        INSTRUCOES_SHEET,
    ]
    text = report_path.read_text(encoding="utf-8")
    assert "DRE GERENCIAL PROVISORIA" in text
    assert "Status: PROVISORIO" in text
    assert "Nao e DRE definitiva" in text


def _save_mapping(tmp_path: Path) -> Path:
    path = tmp_path / "mapeamento_dre.yaml"
    path.write_text(
        """\
linhas_dre:
  - id: "01.01"
    nome: "Receita Bruta de Vendas"
    grupo: "Receita Bruta"
    sinal: "+"
    ordem: 101
    calculada: false
  - id: "03.02"
    nome: "CMV - Bebidas"
    grupo: "Custos"
    sinal: "-"
    ordem: 302
    calculada: false
mapeamento:
  "1.01.001": { dre_linha_id: "01.01", entra_dre: true }
  "2.01.004": { dre_linha_id: "03.02", entra_dre: true }
""",
        encoding="utf-8",
    )
    return path


def _xml_dir_with_special_operation(tmp_path: Path, *, natureza: str, cfop: str) -> Path:
    xml_dir = tmp_path / f"xml_{cfop}"
    xml_dir.mkdir()
    copyfile(SAMPLE_XML, xml_dir / "especial.xml")
    xml_path = xml_dir / "especial.xml"
    text = xml_path.read_text(encoding="utf-8")
    text = text.replace("<natOp>VENDA DE MERCADORIAS</natOp>", f"<natOp>{natureza}</natOp>")
    text = text.replace("<CFOP>5102</CFOP>", f"<CFOP>{cfop}</CFOP>", 1)
    xml_path.write_text(text, encoding="utf-8")
    return xml_dir


def _sheet_rows(workbook, sheet_name: str) -> list[dict[str, object]]:
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    return [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in values)
    ]


def _lancamento(
    hash_registro: str,
    *,
    conta_id: str = "2.01.004",
    valor: Decimal = Decimal("10.00"),
    status_revisao: str = "OK",
    chave_origem: str = "NFE123",
) -> Lancamento:
    return Lancamento(
        data_documento=date(2026, 5, 5),
        competencia="2026-05",
        cnpj_empresa="55746722000181",
        nome_empresa="TRATTORIA DO SOSSEGO RESTAURANTE LTDA",
        fornecedor_ou_favorecido="Fornecedor teste",
        cnpj_cpf_favorecido="12345678000195",
        documento="12566",
        tipo_documento="XML_NFE",
        conta_id=conta_id,
        conta_nome="Conta teste",
        grupo_conta="Grupo teste",
        descricao_itens="ITEM TESTE",
        valor=valor,
        duplicatas=[Duplicata(numero="001", vencimento=date(2026, 5, 20), valor=valor)],
        origem="XML_NFE",
        arquivo_origem="nfe.xml",
        chave_origem=chave_origem,
        hash_registro=hash_registro,
        status_revisao=status_revisao,
        alertas="" if status_revisao == "OK" else "CLASSIFICACAO_NAO_ENCONTRADA",
        tipo_lancamento="custo",
    )
